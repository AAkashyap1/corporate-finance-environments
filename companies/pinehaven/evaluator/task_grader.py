from __future__ import annotations

import json
import math
import os
import posixpath
import re
import shutil
import sqlite3
import stat
import tempfile
import zipfile
from collections import Counter
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable, Mapping
from xml.etree import ElementTree

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook
from openpyxl.formula import Tokenizer
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.utils.datetime import from_excel
from pypdf import PdfReader
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR

from evaluator.native_artifacts import (
    NativePackageSafetyError,
    inspect_native_package,
    pptx_axis_id_issues,
    run_sandboxed_libreoffice,
    validated_native_package_copy,
)


GOLD_PATH = Path(__file__).resolve().parent / "reference" / "tasks.json"
FORMULA_ERRORS = {"#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!"}
_SOFFICE_RECALC_DISABLED = False
_WORKBOOK_NUMERIC_CLIP_OVERFLOW_RATIO = 1.25
_WORKBOOK_TEXT_CLIP_OVERFLOW_RATIO = 1.15
_WORKBOOK_MAX_CELL_VISUAL_ISSUES = 50
_WORKBOOK_MAX_CLIP_SCAN_CELLS = 10_000
_MONTH_NUMBERS = {
    name: number
    for number, names in enumerate(
        (
            (),
            ("january", "jan"),
            ("february", "feb"),
            ("march", "mar"),
            ("april", "apr"),
            ("may",),
            ("june", "jun"),
            ("july", "jul"),
            ("august", "aug"),
            ("september", "sep", "sept"),
            ("october", "oct"),
            ("november", "nov"),
            ("december", "dec"),
        )
    )
    for name in names
}
_MONTH_NAMES = (
    "",
    "january",
    "february",
    "march",
    "april",
    "may",
    "june",
    "july",
    "august",
    "september",
    "october",
    "november",
    "december",
)


@dataclass(frozen=True)
class Criterion:
    id: str
    description: str
    met: bool
    evidence: str
    category: str = "core_finance"
    weight: int = 10
    semantic: bool = False
    failure_cap: float | None = None

    def as_dict(self) -> dict[str, Any]:
        return {
            "id": self.id,
            "description": self.description,
            "value": int(self.met),
            "evidence": self.evidence,
            "category": self.category,
            "weight": self.weight,
            "semantic": self.semantic,
            **(
                {"failure_cap": self.failure_cap}
                if self.failure_cap is not None
                else {}
            ),
        }


def load_gold(task_id: str | None = None) -> dict[str, Any]:
    payload = json.loads(GOLD_PATH.read_text(encoding="utf-8"))
    return payload[task_id] if task_id else payload


def _normalize(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(value or "").casefold()).strip()


def _year_month(value: Any) -> tuple[int, int] | None:
    if isinstance(value, (date, datetime)):
        return value.year, value.month
    parts = _normalize(value).split()
    if len(parts) != 2:
        return None
    if re.fullmatch(r"20\d{2}", parts[0]):
        year = int(parts[0])
        month = (
            int(parts[1])
            if parts[1].isdigit()
            else _MONTH_NUMBERS.get(parts[1])
        )
    elif re.fullmatch(r"20\d{2}", parts[1]):
        year = int(parts[1])
        month = _MONTH_NUMBERS.get(parts[0])
    else:
        return None
    return (year, month) if month and 1 <= month <= 12 else None


def _number(value: Any) -> float | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)) and math.isfinite(float(value)):
        return float(value)
    if not isinstance(value, str):
        return None
    text = (
        value.strip()
        .translate(
            str.maketrans(
                {
                    "\N{MINUS SIGN}": "-",
                    "\N{FIGURE DASH}": "-",
                    "\N{EN DASH}": "-",
                    "\N{EM DASH}": "-",
                }
            )
        )
        .replace("$", "")
        .strip()
    )
    # Favorable/unfavorable suffixes are presentation classifications, not
    # part of the numeric token and not an instruction to change its sign.
    # Keep the displayed accounting sign authoritative while accepting the
    # common ``$58,928.32 U`` and ``($1,250.50) F`` forms.
    text = re.sub(
        r"\s+(?:f|u|favorable|unfavorable)\s*$",
        "",
        text,
        flags=re.I,
    ).rstrip()
    word_suffix = re.search(
        r"\b(thousand|million|billion)\s*$",
        text,
        flags=re.I,
    )
    word_multiplier = {
        "thousand": 1_000.0,
        "million": 1_000_000.0,
        "billion": 1_000_000_000.0,
    }.get(
        word_suffix.group(1).casefold() if word_suffix else "",
        1.0,
    )
    if word_suffix:
        text = text[: word_suffix.start()].rstrip()
    # Accounting exports may put a magnitude/percent suffix either inside or
    # outside the parentheses: ($1.2M), $(1.2)M, and (12.5)% are all negative.
    accounting = re.fullmatch(r"\((.*)\)\s*([%kmb]?)", text, flags=re.I)
    negative = accounting is not None
    if accounting:
        text = f"{accounting.group(1)}{accounting.group(2)}"
    elif text.startswith("(") and not text.endswith(")"):
        # A prose parenthetical such as ``(24.14% effective rate)`` yields a
        # number token with the opening parenthesis but not the later prose
        # close. It is a positive displayed rate, not accounting notation.
        text = text[1:].lstrip()
    percent = text.rstrip().endswith("%")
    text = text.replace(",", "")
    suffix = re.search(r"([kmb])\s*$", text, flags=re.I)
    multiplier = {
        "k": 1_000.0,
        "m": 1_000_000.0,
        "b": 1_000_000_000.0,
    }.get(suffix.group(1).casefold() if suffix else "", 1.0)
    if suffix:
        text = text[: suffix.start()]
    text = text.strip("% x×")
    try:
        result = float(text) * multiplier * word_multiplier
    except ValueError:
        return None
    if negative:
        result = -abs(result)
    return result / 100 if percent else result


_RATE_METRIC_WORDS = {
    "rate",
    "ratio",
    "margin",
    "percent",
    "coverage",
    "utilization",
    "yield",
    "growth",
    "concentration",
    "roic",
    "leverage",
    "headroom",
    "score",
    "irr",
    "index",
}
_SCENARIO_RATE_METRICS = {
    "volume change",
    "price change",
    "cost change",
    "material cost change",
    "required hours change",
    "net available hours change",
}
_QUANTITY_METRIC_WORDS = {
    "quantity",
    "units",
    "hours",
    "payback",
}


def _is_rate_metric(metric_key: str | None) -> bool:
    """Classify decimal-rate metrics without substring false positives."""

    key = _normalize(metric_key)
    if key == "labor rate variance":
        return False
    return (
        bool(set(key.split()).intersection(_RATE_METRIC_WORDS))
        or any(
            f" {phrase} " in f" {key} "
            for phrase in _SCENARIO_RATE_METRICS
        )
    )


def _allows_percent_display(metric_key: str | None) -> bool:
    """Return whether a decimal metric may be rendered with a percent sign."""

    key = _normalize(metric_key)
    if not _is_rate_metric(key):
        return False
    return not set(key.split()).intersection(
        {"index", "leverage", "headroom"}
    )


def _is_quantity_metric(metric_key: str | None) -> bool:
    key = _normalize(metric_key)
    return bool(set(key.split()).intersection(_QUANTITY_METRIC_WORDS))


def _metric_precision(metric_key: str | None, expected: float) -> float:
    """Return the largest acceptable display-rounding delta for a metric.

    Task prompts require cents for money, four decimals for quantities, and
    decimal ratios. Relative tolerances on large dollar values silently
    accepted materially wrong answers, so tolerances follow the declared
    output precision instead of the magnitude.
    """

    key = _normalize(metric_key)
    if set(key.split()).intersection({"score", "irr", "index"}):
        return 0.0000005
    if _is_rate_metric(key) or (not key and abs(expected) <= 2):
        return 0.00005
    if _is_quantity_metric(key):
        return 0.00005
    # Money and ordinary two-decimal business measures are exact to the
    # nearest cent. A half-cent window accepts the same rounded value without
    # accepting a one-cent error.
    return 0.005


def _close(
    actual: Any,
    expected: float,
    metric_key: str | None = None,
    *,
    integer: bool = False,
) -> bool:
    normalized_key = _normalize(metric_key)
    ratio_metric = _is_rate_metric(normalized_key)
    if (
        isinstance(actual, str)
        and "%" in actual
        and normalized_key
        and not ratio_metric
    ):
        return False
    value = _number(actual)
    if value is None:
        return False
    tolerance = 0.0 if integer else _metric_precision(metric_key, expected)
    # Include a scale-aware machine epsilon so a value exactly on the
    # documented tolerance boundary is not rejected by binary-float noise.
    epsilon = max(1.0, abs(value), abs(expected)) * 1e-12
    if abs(value - expected) <= tolerance + epsilon:
        return True
    # Rates may be professionally displayed as percentage points.
    return (
        not integer
        and (
            _allows_percent_display(normalized_key)
            or (not normalized_key and abs(expected) <= 2)
        )
        and abs(value - expected * 100) <= 0.005 + epsilon
    )


def _calendar_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        try:
            converted = from_excel(value)
        except (TypeError, ValueError, OverflowError):
            return None
        return converted.date() if isinstance(converted, datetime) else converted
    text = str(value or "").strip()
    for pattern in (
        "%Y-%m-%d",
        "%Y/%m/%d",
        "%m/%d/%Y",
        "%B %d, %Y",
        "%b %d, %Y",
        "%d %B %Y",
        "%d %b %Y",
    ):
        try:
            return datetime.strptime(text, pattern).date()
        except ValueError:
            continue
    return None


def _value_matches(
    actual: Any,
    expected: Any,
    metric_key: str | None = None,
) -> bool:
    if isinstance(expected, bool):
        if isinstance(actual, bool):
            return actual is expected
        aliases = (
            {"true", "yes", "compliant"}
            if expected
            else {"false", "no", "not compliant"}
        )
        if "compliant" in _normalize(metric_key):
            aliases = aliases | ({"pass"} if expected else {"fail", "failed"})
        return _normalize(actual) in aliases
    if isinstance(expected, (int, float)) and not isinstance(expected, bool):
        return _close(
            actual,
            float(expected),
            metric_key,
            integer=isinstance(expected, int),
        )
    expected_date = _calendar_date(expected)
    if expected_date is not None:
        actual_date = _calendar_date(actual)
        if actual_date is not None:
            return actual_date == expected_date
    actual_text = _normalize(actual)
    expected_text = _normalize(expected)
    if actual_text == expected_text:
        return True
    if "country" in _normalize(metric_key):
        united_states_aliases = {
            "us",
            "u s",
            "usa",
            "u s a",
            "united states",
            "united states of america",
        }
        if (
            expected_text in united_states_aliases
            and actual_text in united_states_aliases
        ):
            return True
    if (
        {"no", "recovery", "annual", "savings"} <= set(expected_text.split())
        and (
            "non positive" in expected_text
            or "nonpositive" in expected_text
        )
        and {"no", "recovery", "annual", "savings"} <= set(actual_text.split())
        and (
            "non positive" in actual_text
            or "nonpositive" in actual_text
        )
    ):
        return True
    actual_period = _year_month(actual)
    expected_period = _year_month(expected)
    if actual_period is not None and actual_period == expected_period:
        return True
    expected_raw = str(expected).strip()
    actual_raw = str(actual).strip()
    if re.fullmatch(
        r"[A-Za-z][A-Za-z0-9]*(?:-[A-Za-z0-9]+)+", expected_raw
    ) and re.fullmatch(
        rf"{re.escape(expected_raw)}"
        r"(?:\s*[-–—|:/]\s*.+|\s+\(.+\))",
        actual_raw,
        flags=re.I,
    ):
        return True
    preclose_aliases = {"pre close", "open pre close"}
    return actual_text in preclose_aliases and expected_text in preclose_aliases


def _strict_json(answer: Any) -> tuple[dict[str, Any], bool, str]:
    if isinstance(answer, dict):
        return answer, True, "answer was a mapping"
    text = str(answer or "").strip()
    try:
        parsed = json.loads(text)
        if isinstance(parsed, dict):
            return parsed, True, "entire answer parsed as one JSON object"
    except json.JSONDecodeError:
        pass
    fenced = re.fullmatch(r"```(?:json)?\s*(\{.*\})\s*```", text, re.I | re.S)
    if fenced:
        try:
            parsed = json.loads(fenced.group(1))
            if isinstance(parsed, dict):
                return parsed, False, "JSON parsed only after removing a code fence"
        except json.JSONDecodeError:
            pass
    brace = re.search(r"\{.*\}", text, re.S)
    if brace:
        try:
            parsed = json.loads(brace.group(0))
            if isinstance(parsed, dict):
                return parsed, False, "JSON parsed only after removing surrounding text"
        except json.JSONDecodeError:
            pass
    return {}, False, "no valid JSON object was found"


def _criterion(
    task_id: str,
    suffix: str,
    description: str,
    met: bool,
    evidence: str,
    *,
    category: str = "core_finance",
    weight: int = 10,
    semantic: bool = False,
    failure_cap: float | None = None,
) -> Criterion:
    return Criterion(
        f"{task_id}__{suffix}",
        description,
        met,
        evidence,
        category,
        weight,
        semantic,
        failure_cap,
    )


def _console_criteria(
    task_id: str, answer: Any, gold: dict[str, Any]
) -> list[Criterion]:
    mapping, exact_json, parse_evidence = _strict_json(answer)
    expected = gold["values"]
    keys = gold["value_keys"]
    criteria = [
        _criterion(
            task_id,
            "json_parse",
            "The response contains a valid JSON object",
            bool(mapping),
            parse_evidence,
            category="structure",
            weight=3,
        ),
        _criterion(
            task_id,
            "json_only",
            "The entire response is the requested JSON object with no wrapper prose",
            exact_json,
            parse_evidence,
            category="structure",
            weight=1,
        ),
        _criterion(
            task_id,
            "exact_key_order",
            "The response uses exactly the required top-level keys in order",
            list(mapping) == keys,
            f"actual={list(mapping)!r}; expected={keys!r}",
            category="decision",
            weight=10,
            failure_cap=0.49,
        ),
        _criterion(
            task_id,
            "no_null_values",
            "No required field is null",
            all(mapping.get(key) is not None for key in keys),
            "checked every required field",
            category="structure",
            weight=3,
        ),
    ]
    for key in keys:
        present = key in mapping
        actual = mapping.get(key)
        target = expected[key]
        if isinstance(target, bool):
            valid_type = isinstance(actual, bool)
            matched = present and valid_type and actual is target
        elif isinstance(target, int):
            valid_type = (
                isinstance(actual, int)
                and not isinstance(actual, bool)
            )
            matched = present and valid_type and actual == target
        elif isinstance(target, float):
            valid_type = (
                isinstance(actual, (int, float))
                and not isinstance(actual, bool)
                and math.isfinite(float(actual))
            )
            tolerance = _metric_precision(key, target)
            epsilon = max(
                1.0,
                abs(float(actual)) if valid_type else 0.0,
                abs(target),
            ) * 1e-12
            matched = (
                present
                and valid_type
                and abs(float(actual) - target) <= tolerance + epsilon
            )
        else:
            valid_type = isinstance(actual, str) and bool(actual.strip())
            matched = present and valid_type and _value_matches(
                actual,
                target,
                key,
            )
            matched = matched or (
                present
                and valid_type
                and any(
                    _value_matches(actual, alias, key)
                    for alias in gold.get("value_aliases", {}).get(key, [])
                )
            )
        slug = _normalize(key).replace(" ", "_")
        criteria.extend(
            [
                _criterion(
                    task_id,
                    f"{slug}__present_valid",
                    f"`{key}` is present with a finite, usable value of the required type",
                    present and valid_type,
                    f"actual={actual!r}; expected type={type(target).__name__}",
                    category="auditability",
                    weight=5,
                    failure_cap=0.69,
                ),
                _criterion(
                    task_id,
                    f"{slug}__correct",
                    f"`{key}` agrees with independently derived ERP/source truth",
                    matched,
                    (
                        "value matched"
                        if matched
                        else f"actual={actual!r}; expected={target!r}"
                    ),
                    weight=10,
                ),
            ]
        )
    return criteria


@dataclass(frozen=True)
class _NativeDocxParagraph:
    text: str
    style_name: str
    table_depth: int


@dataclass(frozen=True)
class _NativeDocxTable:
    rows: tuple[tuple[str, ...], ...]
    table_depth: int


@dataclass(frozen=True)
class _NativeDocxSnapshot:
    paragraphs: tuple[_NativeDocxParagraph, ...]
    tables: tuple[_NativeDocxTable, ...]
    blocks: tuple[_NativeDocxParagraph | _NativeDocxTable, ...]


_DOCX_PARAGRAPH_TAG = qn("w:p")
_DOCX_TABLE_TAG = qn("w:tbl")
_DOCX_SDT_TAG = qn("w:sdt")
_DOCX_SDT_CONTENT_TAG = qn("w:sdtContent")
_DOCX_VERTICAL_MERGE_TAG = qn("w:vMerge")
_DOCX_VALUE_ATTRIBUTE = qn("w:val")


def _docx_block_elements(element: Any) -> Iterable[Any]:
    """Yield native paragraphs and tables without flattening table cells."""

    for child in element.iterchildren():
        if child.tag in {_DOCX_PARAGRAPH_TAG, _DOCX_TABLE_TAG}:
            yield child
        elif child.tag == _DOCX_SDT_TAG:
            content = child.find(_DOCX_SDT_CONTENT_TAG)
            if content is not None:
                yield from _docx_block_elements(content)


def _docx_vmerge_continuation(table_cell: Any) -> bool:
    """Return True only for a non-visible vertical-merge continuation."""

    properties = table_cell.tcPr
    vertical_merge = properties.find(_DOCX_VERTICAL_MERGE_TAG)
    if vertical_merge is None:
        return False
    value = vertical_merge.get(_DOCX_VALUE_ATTRIBUTE)
    return value is None or value.casefold() in {"", "continue"}


def _docx_paragraph_record(
    element: Any,
    parent: Any,
    *,
    table_depth: int,
) -> _NativeDocxParagraph:
    paragraph = Paragraph(element, parent)
    return _NativeDocxParagraph(
        text=paragraph.text,
        style_name=str(getattr(paragraph.style, "name", "") or ""),
        table_depth=table_depth,
    )


def _docx_direct_cell_text(cell: _Cell, *, continuation: bool) -> str:
    """Read visible direct cell paragraphs, excluding nested table text."""

    if continuation:
        return ""
    values = [
        _docx_paragraph_record(
            element,
            cell,
            table_depth=1,
        ).text.strip()
        for element in _docx_block_elements(cell._tc)
        if element.tag == _DOCX_PARAGRAPH_TAG
    ]
    return "\n".join(value for value in values if value)


def _native_docx_snapshot(container: Any) -> _NativeDocxSnapshot:
    """Extract body-order Word evidence from physical OOXML cells.

    Horizontal grid spans are represented by their single physical ``w:tc``.
    Vertical-merge continuations retain an empty positional cell but never
    duplicate the restart cell's text. Nested tables are recursively emitted
    immediately after their containing table and use the same rules.
    """

    paragraphs: list[_NativeDocxParagraph] = []
    tables: list[_NativeDocxTable] = []
    blocks: list[_NativeDocxParagraph | _NativeDocxTable] = []

    def visit_table(table: Table, *, table_depth: int) -> None:
        rows: list[tuple[str, ...]] = []
        physical_cells: list[tuple[_Cell, bool]] = []
        for table_row in table._tbl.tr_lst:
            row_values: list[str] = []
            for table_cell in table_row.tc_lst:
                continuation = _docx_vmerge_continuation(table_cell)
                cell = _Cell(table_cell, table)
                physical_cells.append((cell, continuation))
                row_values.append(
                    _docx_direct_cell_text(
                        cell,
                        continuation=continuation,
                    )
                )
            rows.append(tuple(row_values))
        native_table = _NativeDocxTable(
            rows=tuple(rows),
            table_depth=table_depth,
        )
        tables.append(native_table)
        blocks.append(native_table)

        for cell, continuation in physical_cells:
            if continuation:
                continue
            for element in _docx_block_elements(cell._tc):
                if element.tag == _DOCX_PARAGRAPH_TAG:
                    paragraphs.append(
                        _docx_paragraph_record(
                            element,
                            cell,
                            table_depth=table_depth + 1,
                        )
                    )
                elif element.tag == _DOCX_TABLE_TAG:
                    visit_table(
                        Table(element, cell),
                        table_depth=table_depth + 1,
                    )

    document_element = getattr(container, "element", None)
    if document_element is not None and hasattr(document_element, "body"):
        root = document_element.body
    else:
        root = container._element
    for element in _docx_block_elements(root):
        if element.tag == _DOCX_PARAGRAPH_TAG:
            paragraph = _docx_paragraph_record(
                element,
                container,
                table_depth=0,
            )
            paragraphs.append(paragraph)
            blocks.append(paragraph)
        elif element.tag == _DOCX_TABLE_TAG:
            visit_table(Table(element, container), table_depth=0)

    return _NativeDocxSnapshot(
        paragraphs=tuple(paragraphs),
        tables=tuple(tables),
        blocks=tuple(blocks),
    )


def _native_docx_visible_lines(
    snapshot: _NativeDocxSnapshot,
) -> list[str]:
    lines: list[str] = []
    for block in snapshot.blocks:
        if isinstance(block, _NativeDocxParagraph):
            if block.text.strip():
                lines.append(block.text.strip())
            continue
        lines.extend(
            " | ".join(value.strip() for value in row if value.strip())
            for row in block.rows
            if any(value.strip() for value in row)
        )
    return lines


def _flatten_docx(document: Document) -> tuple[str, list[list[str]]]:
    snapshot = _native_docx_snapshot(document)
    visible_lines = _native_docx_visible_lines(snapshot)
    # Top-level paragraphs are native, editable document text blocks. A
    # clearly written "metric: value" paragraph is an auditable pairing even
    # when the required decision table is used for actions rather than KPIs.
    paired_rows: list[list[str]] = [
        [block.text]
        for block in snapshot.blocks
        if isinstance(block, _NativeDocxParagraph) and block.text.strip()
    ]
    paired_rows.extend(
        list(row)
        for table in snapshot.tables
        for row in table.rows
        if any(value.strip() for value in row)
    )
    return "\n".join(visible_lines), paired_rows


_DOCX_GENERIC_SOURCE_PATTERNS = (
    r"\bauthoritative\s+extract\b",
    r"\berp\s+source\b",
    r"\bsource\s*(?:input|\d+)\b",
    r"\bnamed\s+source\b",
)
_DOCX_GENERIC_CALCULATION_PATTERNS = (
    r"^\s*round\s*\(\s*source(?:\s+\d+)?\s+input\b",
    r"^\s*(?:use|copy|report)\s+(?:the\s+)?(?:source|input|provided value)\b",
    r"^\s*(?:evaluate|calculate|derive)\s+(?:the\s+)?(?:metric|result|value)\b",
    r"^\s*(?:see|refer to)\s+(?:the\s+)?source\b",
)
_DOCX_CALCULATION_METHOD_TERMS = {
    "add",
    "aggregate",
    "average",
    "classify",
    "compare",
    "count",
    "discount",
    "divide",
    "group",
    "less",
    "multiply",
    "rank",
    "reconcile",
    "select",
    "subtract",
    "sum",
    "test",
}


def _docx_source_is_management_assumption(source_text: str) -> bool:
    normalized = _normalize(source_text)
    return any(
        phrase in normalized
        for phrase in (
            "management assumption",
            "management planning assumption",
            "management valuation assumption",
            "management approved assumption",
        )
    )


def _docx_source_lineage_valid(
    source_text: str,
    *,
    metric_label: str,
    sources: Iterable[str],
) -> bool:
    """Require a named authoritative source or a metric-specific assumption."""

    normalized = _normalize(source_text)
    if not normalized or any(
        re.search(pattern, normalized, flags=re.I)
        for pattern in _DOCX_GENERIC_SOURCE_PATTERNS
    ):
        return False
    if _docx_source_is_management_assumption(source_text):
        metric_terms = {
            term
            for term in _normalize(metric_label).split()
            if len(term) > 2
        }
        visible_terms = set(normalized.split())
        return bool(metric_terms.intersection(visible_terms))
    return any(_source_present(source_text, source) for source in sources)


def _docx_calculation_lineage_valid(
    calculation_text: str,
    *,
    metric_label: str,
    source_is_assumption: bool,
) -> bool:
    """Reject template prose that does not disclose a usable derivation."""

    normalized = _normalize(calculation_text)
    if len(normalized.split()) < 4 or any(
        re.search(pattern, calculation_text, flags=re.I)
        for pattern in _DOCX_GENERIC_CALCULATION_PATTERNS
    ):
        return False
    if source_is_assumption:
        metric_terms = {
            term
            for term in _normalize(metric_label).split()
            if len(term) > 2
        }
        visible_terms = {
            term for term in normalized.split() if len(term) > 2
        }
        metric_specific = any(
            metric_term[:4] == visible_term[:4]
            for metric_term in metric_terms
            for visible_term in visible_terms
        )
        return (
            any(
                token in normalized
                for token in (
                    "assumption",
                    "executed",
                    "approved",
                )
            )
            and metric_specific
        )
    calculation_terms = set(normalized.split())
    return bool(
        _DOCX_CALCULATION_METHOD_TERMS.intersection(calculation_terms)
        or any(
            term.startswith(
                (
                    "aggregat",
                    "averag",
                    "classif",
                    "compar",
                    "discount",
                    "multipli",
                    "reconcil",
                    "subtract",
                )
            )
            for term in calculation_terms
        )
        or re.search(r"(?:[+\-*/=()]|\bper\b)", calculation_text)
    )


def _docx_native_structure(
    document: Document,
    *,
    title: str,
    metric_keys: Iterable[str],
    sources: Iterable[str] = (),
) -> dict[str, Any]:
    """Collect contract evidence from native Word styles and table structure."""

    snapshot = _native_docx_snapshot(document)
    title_candidates = [
        paragraph.text.strip()
        for paragraph in snapshot.paragraphs
        if paragraph.text.strip()
        and _normalize(paragraph.style_name) in {
            "title",
            "heading 0",
        }
    ]
    headings = [
        paragraph.text.strip()
        for paragraph in snapshot.paragraphs
        if paragraph.text.strip()
        and _normalize(paragraph.style_name).startswith(
            "heading "
        )
    ]
    visible_blocks = _native_docx_visible_lines(snapshot)
    visible_text = "\n".join(visible_blocks)
    date_present = bool(
        re.search(
            (
                r"\b(?:2026-06-30|June\s+30,?\s+2026|"
                r"30\s+June\s+2026)\b"
            ),
            visible_text,
            flags=re.I,
        )
    )
    audience_present = bool(
        re.search(r"\baudience\s*:\s*[^\s|;]+", visible_text, flags=re.I)
    )

    source_contract = tuple(str(source) for source in sources)
    calculation_rows: dict[str, list[str]] = {}
    calculation_lineage_rows: dict[str, list[str]] = {}
    calculation_lineage_issues: list[str] = []
    calculation_candidates: dict[str, list[list[str]]] = {}
    for table in snapshot.tables:
        rows = [
            [value.strip() for value in row]
            for row in table.rows
        ]
        if len(rows) < 2:
            continue
        headers = [_normalize(value) for value in rows[0]]

        def header_index(*tokens: str) -> int | None:
            return next(
                (
                    index
                    for index, header in enumerate(headers)
                    if any(token in header for token in tokens)
                ),
                None,
            )

        metric_column = header_index("metric", "kpi")
        source_column = header_index("source", "input")
        calculation_column = header_index(
            "calculation",
            "formula",
            "logic",
            "derivation",
        )
        result_column = header_index("result", "value", "output")
        combined_lineage_column = next(
            (
                index
                for index, header in enumerate(headers)
                if any(
                    token in header
                    for token in (
                        "calculation",
                        "formula",
                        "logic",
                        "derivation",
                    )
                )
                and any(
                    token in header
                    for token in (
                        "source",
                        "input",
                        "authoritative",
                        "lineage",
                        "chain",
                    )
                )
            ),
            None,
        )
        combined_lineage = (
            source_column is None
            and combined_lineage_column is not None
            and calculation_column == combined_lineage_column
        )
        columns = (
            metric_column,
            calculation_column,
            result_column,
        )
        if None in columns or len(set(columns)) != 3:
            continue
        if not combined_lineage and (
            source_column is None
            or source_column in columns
        ):
            continue
        for row in rows[1:]:
            physical_columns = (
                metric_column,
                calculation_column,
                result_column,
            ) if combined_lineage else (
                metric_column,
                source_column,
                calculation_column,
                result_column,
            )
            if max(int(column) for column in physical_columns) >= len(row):
                continue
            if combined_lineage:
                lineage = row[int(combined_lineage_column)].strip()
                selected = [
                    row[int(metric_column)].strip(),
                    lineage,
                    lineage,
                    row[int(result_column)].strip(),
                ]
            else:
                selected = [
                    row[int(metric_column)].strip(),
                    row[int(source_column)].strip(),
                    row[int(calculation_column)].strip(),
                    row[int(result_column)].strip(),
                ]
            if (
                not all(selected)
                or any(
                    _normalize(value)
                    in {"na", "n a", "none", "tbd", "todo", "placeholder"}
                    for value in selected
                )
            ):
                continue
            normalized_metric = _normalize(
                selected[0].replace("_", " ")
            )
            calculation_candidates.setdefault(
                normalized_metric,
                [],
            ).append(selected)

    for normalized_metric, candidates in calculation_candidates.items():
        unique_candidates = {
            tuple(_normalize(value) for value in candidate): candidate
            for candidate in candidates
        }
        if len(unique_candidates) != 1:
            if len(calculation_lineage_issues) < 20:
                calculation_lineage_issues.append(
                    f"{normalized_metric}: ambiguous native calculation "
                    f"rows={len(unique_candidates)}"
                )
            continue
        selected = next(iter(unique_candidates.values()))
        calculation_rows[normalized_metric] = selected
        source_is_assumption = _docx_source_is_management_assumption(
            selected[1]
        )
        source_valid = _docx_source_lineage_valid(
            selected[1],
            metric_label=normalized_metric,
            sources=source_contract,
        )
        calculation_valid = _docx_calculation_lineage_valid(
            selected[2],
            metric_label=normalized_metric,
            source_is_assumption=source_is_assumption,
        )
        # A deliberately combined "calculation / authoritative chain" column
        # is a valid native design when the row itself contains concrete
        # lineage or a direct derivation. Reject generic Source/Input N
        # placeholders so the combined form cannot weaken the contract.
        combined_text = _normalize(selected[2])
        combined_placeholder = re.search(
            r"\b(?:source|input)\s*\d+\b",
            combined_text,
        ) is not None
        combined_authority = bool(
            re.search(
                (
                    r"\b(?:specified|as of|erp|general ledger|gl|p l|"
                    r"subledger|report|binder|policy|instruction|account|"
                    r"commitment|cutoff|window)\b"
                ),
                combined_text,
            )
        )
        combined_derivation = bool(
            re.search(r"\d", selected[2])
            and re.search(r"(?:[+\-*/=÷×−]|\bper\b)", selected[2])
        )
        if (
            selected[1] == selected[2]
            and not combined_placeholder
            and (calculation_valid or combined_authority or combined_derivation)
            and (source_valid or combined_authority or combined_derivation)
        ):
            source_valid = True
            calculation_valid = True
        if source_valid and calculation_valid:
            calculation_lineage_rows[normalized_metric] = selected
        elif len(calculation_lineage_issues) < 20:
            calculation_lineage_issues.append(
                f"{normalized_metric}: "
                f"source_valid={source_valid}; "
                f"calculation_valid={calculation_valid}"
            )

    expected_metric_labels = {
        _normalize(key.replace("_", " ")) for key in metric_keys
    }
    calculation_coverage = expected_metric_labels.intersection(
        calculation_rows
    )
    calculation_lineage_coverage = expected_metric_labels.intersection(
        calculation_lineage_rows
    )

    owner_timing_blocks: list[str] = []
    for block in visible_blocks:
        owner_match = re.search(
            r"\bowner\s*[:\-]\s*([^|;\n.]+)",
            block,
            flags=re.I,
        )
        timing_match = re.search(
            r"\b(?:timing|due(?:\s+date)?)\s*[:\-]\s*([^|;\n.]+)",
            block,
            flags=re.I,
        )
        if owner_match is None or timing_match is None:
            continue
        owner = _normalize(owner_match.group(1))
        timing = _normalize(timing_match.group(1))
        if owner in {
            "",
            "functional owner",
            "named functional owner",
            "owner",
            "team",
            "tbd",
        }:
            continue
        if timing in {
            "",
            "decision gate",
            "next close or decision gate",
            "tbd",
        }:
            continue
        owner_timing_blocks.append(block)

    for table in snapshot.tables:
        rows = [[value.strip() for value in row] for row in table.rows]
        if len(rows) < 2:
            continue
        headers = [_normalize(value) for value in rows[0]]
        owner_column = next(
            (
                index
                for index, header in enumerate(headers)
                if "owner" in header
            ),
            None,
        )
        timing_column = next(
            (
                index
                for index, header in enumerate(headers)
                if "timing" in header or "due" in header
            ),
            None,
        )
        if (
            owner_column is None
            or timing_column is None
            or owner_column == timing_column
        ):
            continue
        for row in rows[1:]:
            if max(owner_column, timing_column) >= len(row):
                continue
            owner = _normalize(row[owner_column])
            timing = _normalize(row[timing_column])
            if owner in {
                "",
                "functional owner",
                "named functional owner",
                "owner",
                "team",
                "tbd",
            } or timing in {
                "",
                "decision gate",
                "next close or decision gate",
                "tbd",
            }:
                continue
            owner_timing_blocks.append(" | ".join(row))

    return {
        "title_candidates": title_candidates,
        "title_present": any(
            _normalize(candidate)
            in {
                _normalize(title),
                _normalize(
                    re.sub(
                        r"^(?:Build|Write|Create)\s+the\s+",
                        "",
                        title.strip(),
                        flags=re.I,
                    )
                ),
            }
            for candidate in title_candidates
        ),
        "headings": headings,
        "date_present": date_present,
        "audience_present": audience_present,
        "calculation_rows": calculation_rows,
        "calculation_coverage": calculation_coverage,
        "calculation_lineage_rows": calculation_lineage_rows,
        "calculation_lineage_coverage": calculation_lineage_coverage,
        "calculation_lineage_issues": calculation_lineage_issues,
        "expected_metric_labels": expected_metric_labels,
        "owner_timing_blocks": owner_timing_blocks,
    }


_DECISION_BLOCK_LABELS = {
    "approval requested",
    "board resolution",
    "central decision",
    "close recommendation",
    "council resolution requested",
    "decision",
    "decision requested",
    "executive conclusion",
    "executive decision",
    "primary decision",
    "recommendation",
    "recommended action",
}


def _labeled_decision_block(values: Iterable[Any]) -> str | None:
    """Return one native row only when it is explicitly labeled as a decision."""

    cells = [str(value).strip() for value in values if str(value).strip()]
    if not cells:
        return None
    first = _normalize(cells[0])
    if first in _DECISION_BLOCK_LABELS:
        return " | ".join(cells)
    if len(cells) == 1 and any(
        first.startswith(f"{label} ")
        for label in _DECISION_BLOCK_LABELS
    ):
        return cells[0]
    return None


def _labeled_decision_blocks(
    rows: Iterable[Iterable[Any]],
) -> list[str]:
    blocks: list[str] = []
    seen: set[str] = set()
    for row in rows:
        block = _labeled_decision_block(row)
        if block and block not in seen:
            seen.add(block)
            blocks.append(block)
    return blocks


def _docx_decision_blocks(document: Document) -> list[str]:
    """Extract native decision paragraphs/tables without using arbitrary prose."""

    snapshot = _native_docx_snapshot(document)
    rows: list[list[str]] = [
        [block.text.strip()]
        for block in snapshot.blocks
        if isinstance(block, _NativeDocxParagraph)
        and block.text.strip()
    ]
    for index, block in enumerate(snapshot.blocks[:-1]):
        next_block = snapshot.blocks[index + 1]
        if (
            isinstance(block, _NativeDocxParagraph)
            and isinstance(next_block, _NativeDocxParagraph)
            and _normalize(block.text) == "executive conclusion"
            and next_block.text.strip()
        ):
            rows.append(
                ["Executive conclusion", next_block.text.strip()]
            )
    rows.extend(
        list(row)
        for table in snapshot.tables
        for row in table.rows
    )
    return _labeled_decision_blocks(rows)


def _docx_decision_action_rows(document: Document) -> list[str]:
    """Extract each data row from an explicitly labeled decision/action table."""

    blocks: list[str] = []
    for table in _native_docx_snapshot(document).tables:
        rows = [
            [value.strip() for value in row]
            for row in table.rows
        ]
        if (
            not rows
            or not rows[0]
            or _normalize(rows[0][0]) != "decision action"
        ):
            continue
        blocks.extend(
            "Decision | " + " | ".join(
                value for value in row if value
            )
            for row in rows[1:]
            if any(row)
        )
    return blocks


def _decision_rate_semantic_pair(
    text: str,
    *,
    expected: float,
    metric_key: str,
    own_concept: str,
    competing_concept: str,
) -> bool:
    """Pair a decision rate to its closest preservation objective."""

    numeric_pattern = re.compile(
        (
            r"-?\$?\s*\(?-?\$?\d[\d,]*(?:\.\d+)?"
            r"\)?%?(?:[ \t]*[KMB](?![A-Za-z]))?\)?"
        ),
        flags=re.I,
    )
    target_positions = [
        (match.start() + match.end()) // 2
        for match in numeric_pattern.finditer(text)
        if _close(match.group(0), expected, metric_key)
    ]
    if not target_positions:
        return False
    preservation = (
        r"(?:preserv\w*|maintain\w*|hold\w*|protect\w*|restore\w*)"
    )
    semantic_text = text.translate(
        str.maketrans(
            {
                "-": " ",
                "\N{MINUS SIGN}": " ",
                "\N{FIGURE DASH}": " ",
                "\N{EN DASH}": " ",
                "\N{EM DASH}": " ",
            }
        )
    )

    def concept_positions(concept: str) -> list[int]:
        pattern = re.compile(
            (
                rf"(?:{preservation}.{{0,55}}{concept}|"
                rf"{concept}.{{0,55}}{preservation})"
            ),
            flags=re.I,
        )
        return [
            (match.start() + match.end()) // 2
            for match in pattern.finditer(semantic_text)
        ]

    own_positions = concept_positions(own_concept)
    competing_positions = concept_positions(competing_concept)
    if not own_positions:
        return False
    for target in target_positions:
        own_distance = min(abs(target - position) for position in own_positions)
        competing_distance = (
            min(
                abs(target - position)
                for position in competing_positions
            )
            if competing_positions
            else math.inf
        )
        if own_distance <= 160 and own_distance + 5 < competing_distance:
            return True
    return False


def _flatten_pptx(
    presentation: Presentation,
) -> tuple[str, int, int, list[list[str]]]:
    parts: list[str] = []
    chart_count = 0
    table_count = 0
    title_candidates: list[list[str]] = []
    for slide in presentation.slides:
        candidates: list[str] = []
        if slide.shapes.title and slide.shapes.title.text.strip():
            candidates.append(slide.shapes.title.text)
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                parts.append(shape.text)
                # Professionally authored decks often use a positioned text box
                # instead of a layout title placeholder. Treat an exact,
                # standalone text frame in the top quarter as a title candidate.
                if (
                    shape.text.strip()
                    and "\n" not in shape.text.strip()
                    and int(getattr(shape, "top", presentation.slide_height))
                    <= int(presentation.slide_height * 0.25)
                    and shape.text not in candidates
                ):
                    candidates.append(shape.text)
            if getattr(shape, "has_chart", False):
                chart_count += 1
            if getattr(shape, "has_table", False):
                table_count += 1
                for row in shape.table.rows:
                    parts.extend(cell.text for cell in row.cells)
        title_candidates.append(candidates)
    return "\n".join(parts), chart_count, table_count, title_candidates


def _pptx_chart_evidence(
    presentation: Presentation,
) -> list[dict[str, Any]]:
    """Extract native chart categories and series values for exact checks."""

    charts: list[dict[str, Any]] = []
    for slide in presentation.slides:
        for shape in slide.shapes:
            if not getattr(shape, "has_chart", False):
                continue
            chart = shape.chart
            title = ""
            if chart.has_title:
                title = chart.chart_title.text_frame.text
            categories: list[str] = []
            if chart.plots:
                try:
                    categories = [
                        str(category)
                        for category in chart.plots[0].categories
                    ]
                except (AttributeError, TypeError, ValueError):
                    categories = []
            series = []
            for item in chart.series:
                try:
                    values = list(item.values)
                except (AttributeError, TypeError, ValueError):
                    values = []
                series.append({"name": str(item.name), "values": values})
            charts.append(
                {
                    "title": title,
                    "categories": categories,
                    "series": series,
                }
            )
    return charts


def _pptx_decision_blocks(presentation: Presentation) -> list[str]:
    """Extract text from native ``Decision`` and owned ``Actions`` slides."""

    blocks: list[str] = []
    seen: set[str] = set()
    for slide in presentation.slides:
        title = (
            slide.shapes.title.text.strip()
            if slide.shapes.title and slide.shapes.title.text.strip()
            else ""
        )
        if _normalize(title) not in {"decision", "actions"}:
            title = next(
                (
                    shape.text.strip()
                    for shape in slide.shapes
                    if getattr(shape, "has_text_frame", False)
                    and _normalize(shape.text) in {"decision", "actions"}
                    and int(getattr(shape, "top", presentation.slide_height))
                    <= int(presentation.slide_height * 0.25)
                ),
                "",
            )
        if _normalize(title) not in {"decision", "actions"}:
            continue
        parts: list[str] = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.text.strip():
                parts.append(shape.text.strip())
            if getattr(shape, "has_table", False):
                parts.extend(
                    " | ".join(cell.text.strip() for cell in row.cells)
                    for row in shape.table.rows
                )
        block = "\n".join(part for part in parts if part)
        if block and block not in seen:
            seen.add(block)
            blocks.append(block)
    return blocks


def _pptx_paired_rows(
    presentation: Presentation,
    detail_specification: dict[str, Any] | None = None,
) -> list[list[str]]:
    """Return conservative native label/value groups from a presentation.

    Accept native table rows and columns, individual text frames, and text
    frames contained by the same bounded card/background shape. Do not treat
    arbitrary text elsewhere on a slide as paired.
    """
    paired_rows: list[list[str]] = []
    seen: set[tuple[str, ...]] = set()

    def add(values: Iterable[Any]) -> None:
        row = [str(value).strip() for value in values if str(value).strip()]
        signature = tuple(row)
        if row and signature not in seen:
            seen.add(signature)
            paired_rows.append(row)

    detail_headers = (
        [
            _normalize(column)
            for column in detail_specification.get("columns", [])
        ]
        if isinstance(detail_specification, dict)
        else []
    )
    slide_area = max(
        1, int(presentation.slide_width) * int(presentation.slide_height)
    )
    for slide in presentation.slides:
        text_shapes = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.text.strip():
                text_shapes.append(shape)
                add([shape.text])
            if getattr(shape, "has_table", False):
                table_rows = [
                    [cell.text for cell in row.cells]
                    for row in shape.table.rows
                ]
                for row in table_rows:
                    add(row)
                is_detail_table = bool(
                    detail_headers
                    and table_rows
                    and [
                        _normalize(value)
                        for value in table_rows[0]
                    ]
                    == detail_headers
                )
                # Column grouping is valid for compact column-oriented metric
                # tables. Do not apply it to an authoritative detail schedule:
                # otherwise a scoped field such as ``plant_yield_rate`` can
                # lend its values to the aggregate ``yield_rate`` criterion.
                if table_rows and not is_detail_table:
                    for column_index in range(len(table_rows[0])):
                        add(
                            row[column_index]
                            for row in table_rows
                            if column_index < len(row)
                        )

        # Cards are commonly built from separate native label and value text
        # boxes over one bounded rectangle. Group only text whose center falls
        # inside a reasonably small shared container; full-slide backgrounds
        # are deliberately excluded.
        for container in slide.shapes:
            if (
                getattr(container, "has_table", False)
                or getattr(container, "has_chart", False)
            ):
                continue
            area = int(container.width) * int(container.height)
            if not (slide_area * 0.002 <= area <= slide_area * 0.25):
                continue
            left = int(container.left)
            top = int(container.top)
            right = left + int(container.width)
            bottom = top + int(container.height)
            members: list[str] = []
            if (
                getattr(container, "has_text_frame", False)
                and container.text.strip()
            ):
                members.append(container.text)
            for shape in text_shapes:
                if shape is container:
                    continue
                center_x = int(shape.left) + int(shape.width) // 2
                center_y = int(shape.top) + int(shape.height) // 2
                if left <= center_x <= right and top <= center_y <= bottom:
                    members.append(shape.text)
            if len(members) >= 2:
                add(members)
    return paired_rows


def _displayed_claim_matches(
    token: str,
    expected: float,
    *,
    absolute: bool = False,
) -> bool:
    value = _number(token)
    if value is None:
        return False
    if absolute:
        value = abs(value)
        expected = abs(expected)
    normalized = token.strip().upper().replace(",", "")
    numeric = re.search(r"\d+(?:\.(\d+))?", normalized)
    decimals = len(numeric.group(1) or "") if numeric else 0
    if "%" in normalized:
        tolerance = 0.5 * 10 ** (-(decimals + 2))
    else:
        scale = (
            1_000_000_000
            if re.search(r"(?:B\b|BILLION\b)", normalized)
            else 1_000_000
            if re.search(r"(?:M\b|MILLION\b)", normalized)
            else 1_000
            if re.search(r"(?:K\b|THOUSAND\b)", normalized)
            else 1
        )
        tolerance = 0.5 * scale * 10 ** (-decimals)
    epsilon = max(1.0, abs(value), abs(expected)) * 1e-12
    return abs(value - expected) <= tolerance + epsilon


def _scaled_currency_claim_matches(token: str, expected: float) -> bool:
    """Match an otherwise unit-ambiguous ``$92.550``-style executive value."""

    if "$" not in token or re.search(r"[KMB]\s*$", token, flags=re.I):
        return False
    value = _number(token)
    if value is None or abs(expected) < 1_000_000 or abs(value) >= 1_000:
        return False
    numeric = re.search(r"\d+(?:\.(\d+))?", token.replace(",", ""))
    decimals = len(numeric.group(1) or "") if numeric else 0
    for scale in (1_000, 1_000_000, 1_000_000_000):
        tolerance = 0.5 * scale * 10 ** (-decimals)
        scaled = value * scale
        epsilon = max(1.0, abs(scaled), abs(expected)) * 1e-12
        if abs(scaled - expected) <= tolerance + epsilon:
            return True
    return False


def _sop_past_due_claim_consistency(
    rows: Iterable[Iterable[Any]],
    supporting_values: dict[str, Any],
) -> tuple[bool, list[dict[str, Any]]]:
    expected_value = float(supporting_values["past_due_backlog_value"])
    expected_units = float(supporting_values["past_due_backlog_units"])
    expected_percent = float(
        supporting_values["past_due_backlog_percent"]
    )
    checks: list[dict[str, Any]] = []
    for row in rows:
        cells = [str(value) for value in row]
        normalized_cells = [_normalize(cell) for cell in cells]
        context = any(
            re.search(r"\b(?:past[\s-]+due|overdue)\b", cell, flags=re.I)
            for cell in cells
        )
        if not context:
            continue
        for cell_index, cell in enumerate(cells):
            normalized = _normalize(cell)
            relevant = bool(
                re.search(
                    r"\b(?:past[\s-]+due|overdue)\b",
                    cell,
                    flags=re.I,
                )
                or "promised through" in normalized
                or "current backlog" in normalized
            )
            if not relevant:
                continue
            for token in re.findall(
                r"\$\s*\d[\d,]*(?:\.\d+)?\s*[KMB]?",
                cell,
                flags=re.I,
            ):
                token_match = re.search(re.escape(token), cell, flags=re.I)
                following_text = (
                    cell[token_match.end() :]
                    if token_match is not None
                    else ""
                )
                explicitly_total_backlog = re.search(
                    r"\b(?:current|total)\s+backlog\b",
                    following_text,
                    flags=re.I,
                )
                adjacent_total_backlog_label = (
                    _normalize(cell) == _normalize(token)
                    and any(
                        abs(other_index - cell_index) == 1
                        and normalized_cells[other_index]
                        in {
                            "current backlog",
                            "total backlog",
                            "backlog total",
                        }
                        for other_index in range(len(cells))
                    )
                )
                if (
                    explicitly_total_backlog
                    or adjacent_total_backlog_label
                ):
                    continue
                checks.append(
                    {
                        "kind": "value",
                        "token": token,
                        "matched": _displayed_claim_matches(
                            token,
                            expected_value,
                        ),
                    }
                )
            for token in re.findall(
                r"\d[\d,]*(?:\.\d+)?(?=\s*units?\b)",
                cell,
                flags=re.I,
            ):
                checks.append(
                    {
                        "kind": "units",
                        "token": token,
                        "matched": _displayed_claim_matches(
                            token,
                            expected_units,
                        ),
                    }
                )
            for token in re.findall(r"\d+(?:\.\d+)?%", cell):
                checks.append(
                    {
                        "kind": "percent",
                        "token": token,
                        "matched": _displayed_claim_matches(
                            token,
                            expected_percent,
                        ),
                    }
                )
    return all(check["matched"] for check in checks), checks


def _optional_labeled_claim_consistency(
    rows: Iterable[Iterable[Any]],
    *,
    label: str,
    expected: Any,
    metric_key: str,
    label_aliases: Iterable[str] = (),
) -> tuple[bool, list[dict[str, Any]]]:
    """Validate every optional, explicitly labeled quantitative claim."""

    normalized_label = _normalize(label)
    normalized_aliases = {
        _normalize(alias) for alias in label_aliases if _normalize(alias)
    }
    checks: list[dict[str, Any]] = []
    for row in rows:
        row_values = list(row)
        raw_row = " ".join(map(str, row_values))
        normalized_row = _normalize(raw_row)
        exact_label_match = normalized_label in normalized_row
        alias_match = any(
            alias in normalized_row for alias in normalized_aliases
        )
        if not exact_label_match and not alias_match:
            continue
        if alias_match and not exact_label_match:
            if "count" in _normalize(metric_key):
                quantified = bool(
                    re.search(
                        r"\b\d[\d,]*\s+(?:negative\s+)?"
                        r"(?:wip\s+)?(?:orders?|exceptions?|residuals?)\b",
                        raw_row,
                        flags=re.I,
                    )
                )
            else:
                quantified = bool(
                    re.search(
                        r"(?:total(?:ing)?|value)\s+"
                        r"\(?\s*\$?\s*-?\d[\d,]*(?:\.\d+)?",
                        raw_row,
                        flags=re.I,
                    )
                )
            if not quantified:
                continue
        # ``paired_rows`` contains both the formula workbook and its
        # recalculated value workbook. A formula row is lineage, not a
        # separately displayed quantitative claim; its cached-value twin is
        # validated below.
        if any(
            isinstance(value, str) and value.startswith("=")
            for value in row_values
        ):
            continue
        matched = any(
            _value_matches(value, expected, metric_key)
            or _text_contains(str(value), expected, metric_key)
            for value in row_values
        )
        checks.append(
            {
                "label": label,
                "row": row_values,
                "expected": expected,
                "matched": matched,
            }
        )
    return all(check["matched"] for check in checks), checks


def _ppt_text_width(text: str, font_size: float) -> float:
    return sum(
        (
            0.28
            if character in " il.,'`:;|!"
            else 0.78
            if character in "MW@%&#"
            else 0.52
        )
        * font_size
        for character in text
    )


def _ppt_paragraph_layout(
    paragraph: Any,
    available_width_points: float,
    *,
    wrap: bool,
) -> tuple[float, float] | None:
    """Return (widest explicit line, required height) in points."""

    line_width = 0.0
    line_height = 0.0
    explicit_lines: list[tuple[float, float]] = []
    runs = [run for run in paragraph.runs if run.text]
    if not runs:
        return (0.0, 0.0)
    if "\v" in paragraph.text:
        sizes = [run.font.size or paragraph.font.size for run in runs]
        if any(size is None for size in sizes):
            return None
        font_size = max(float(size.pt) for size in sizes if size is not None)
        explicit_lines = [
            (_ppt_text_width(line, font_size), font_size)
            for line in paragraph.text.split("\v")
        ]
        widest = max((width for width, _ in explicit_lines), default=0.0)
        required = sum(
            (
                max(1, math.ceil(width / max(1.0, available_width_points)))
                if wrap
                else 1
            )
            * height
            * 1.15
            for width, height in explicit_lines
        )
        return widest, required
    for run in runs:
        size = run.font.size or paragraph.font.size
        if size is None:
            return None
        font_size = float(size.pt)
        pieces = re.split(r"([\v\n])", run.text)
        for piece in pieces:
            if piece in {"\v", "\n"}:
                explicit_lines.append((line_width, line_height or font_size))
                line_width = 0.0
                line_height = 0.0
                continue
            line_width += _ppt_text_width(piece, font_size)
            line_height = max(line_height, font_size)
    explicit_lines.append((line_width, line_height or 11.0))

    widest = max((width for width, _ in explicit_lines), default=0.0)
    required = 0.0
    for width, height in explicit_lines:
        wrapped_lines = (
            max(1, math.ceil(width / max(1.0, available_width_points)))
            if wrap
            else 1
        )
        required += wrapped_lines * height * 1.15
    for spacing in (paragraph.space_before, paragraph.space_after):
        if spacing is not None and hasattr(spacing, "pt"):
            required += float(spacing.pt)
    return widest, required


def _ppt_explicit_text_style(shape: Any) -> tuple[bool, bool]:
    """Return whether visible text has any explicit style and a full label style."""

    has_any = False
    has_size = False
    has_color = False
    for paragraph in shape.text_frame.paragraphs:
        fonts = [paragraph.font, *(run.font for run in paragraph.runs if run.text)]
        for font in fonts:
            size = font.size
            bold = font.bold
            name = font.name
            try:
                color = font.color.rgb
            except (AttributeError, TypeError, ValueError):
                color = None
            has_any = has_any or any(
                value is not None for value in (size, bold, name, color)
            )
            has_size = has_size or size is not None
            has_color = has_color or color is not None
    return has_any, has_size and has_color


def _ppt_dark_solid_fill(shape: Any) -> bool:
    try:
        rgb = shape.fill.fore_color.rgb
    except (AttributeError, TypeError, ValueError):
        return False
    if rgb is None:
        return False
    value = str(rgb)
    if not re.fullmatch(r"[0-9A-Fa-f]{6}", value):
        return False
    red, green, blue = (
        int(value[index : index + 2], 16) for index in (0, 2, 4)
    )
    luminance = 0.2126 * red + 0.7152 * green + 0.0722 * blue
    return luminance <= 96


def _ppt_opaque_solid_fill(shape: Any) -> bool:
    try:
        rgb = shape.fill.fore_color.rgb
    except (AttributeError, TypeError, ValueError):
        return False
    return bool(
        rgb is not None
        and re.fullmatch(r"[0-9A-Fa-f]{6}", str(rgb))
    )


def _pptx_visual_issues(presentation: Presentation) -> list[str]:
    """Return conservative, native presentation-layout defects.

    A manual line break is stored as a separate run. If a generator formats
    only the first run, text after the break falls back to the theme default
    and can visibly spill out even though ``python-pptx`` parses the package.
    Flag only that unambiguous mixed explicit/inherited case, short labels
    missing the explicit style shared by two geometrically equivalent peers on
    a dark panel, plus shapes that extend beyond the slide canvas.
    """

    issues: list[str] = []
    tolerance = 12_700  # one point
    slide_width = int(presentation.slide_width)
    slide_height = int(presentation.slide_height)
    for slide_index, slide in enumerate(presentation.slides, start=1):
        slide_shapes = list(slide.shapes)
        for shape_index, shape in enumerate(slide_shapes, start=1):
            left = int(getattr(shape, "left", 0))
            top = int(getattr(shape, "top", 0))
            right = left + int(getattr(shape, "width", 0))
            bottom = top + int(getattr(shape, "height", 0))
            if (
                left < -tolerance
                or top < -tolerance
                or right > slide_width + tolerance
                or bottom > slide_height + tolerance
            ):
                issues.append(
                    f"slide {slide_index} shape {shape_index} extends "
                    "outside the slide canvas"
                )

            if not getattr(shape, "has_text_frame", False):
                continue
            short_label = shape.text.strip()
            if not short_label:
                continue
            has_explicit_style, _ = _ppt_explicit_text_style(shape)
            if (
                short_label
                and len(short_label) <= 12
                and "\n" not in short_label
                and re.fullmatch(
                    r"[A-Z][A-Z0-9]*(?:[ /&-][A-Z0-9]+)*",
                    short_label,
                )
                and not has_explicit_style
            ):
                containing_panels = []
                for container in slide_shapes:
                    if (
                        container is shape
                        or getattr(container, "has_chart", False)
                        or getattr(container, "has_table", False)
                        or (
                            getattr(container, "has_text_frame", False)
                            and container.text.strip()
                        )
                        or not _ppt_dark_solid_fill(container)
                    ):
                        continue
                    container_left = int(container.left)
                    container_top = int(container.top)
                    container_right = container_left + int(container.width)
                    container_bottom = container_top + int(container.height)
                    container_area = int(container.width) * int(container.height)
                    if (
                        container_area <= slide_width * slide_height * 0.25
                        and container_left <= left
                        and right <= container_right
                        and container_top <= top
                        and bottom <= container_bottom
                    ):
                        containing_panels.append((container_area, container))
                if containing_panels:
                    _, panel = min(containing_panels, key=lambda row: row[0])
                    panel_left = int(panel.left)
                    panel_top = int(panel.top)
                    panel_right = panel_left + int(panel.width)
                    panel_bottom = panel_top + int(panel.height)
                    peer_count = 0
                    for peer in slide_shapes:
                        if (
                            peer is shape
                            or not getattr(peer, "has_text_frame", False)
                        ):
                            continue
                        peer_text = peer.text.strip()
                        if (
                            not peer_text
                            or len(peer_text) > 12
                            or "\n" in peer_text
                            or re.fullmatch(
                                r"[A-Z][A-Z0-9]*(?:[ /&-][A-Z0-9]+)*",
                                peer_text,
                            )
                            is None
                        ):
                            continue
                        peer_left = int(peer.left)
                        peer_top = int(peer.top)
                        peer_right = peer_left + int(peer.width)
                        peer_bottom = peer_top + int(peer.height)
                        _, peer_has_full_style = _ppt_explicit_text_style(peer)
                        if (
                            panel_left <= peer_left
                            and peer_right <= panel_right
                            and panel_top <= peer_top
                            and peer_bottom <= panel_bottom
                            and abs(peer_left - left) <= 91_440
                            and abs(int(peer.width) - int(shape.width))
                            <= max(45_720, int(shape.width) * 0.15)
                            and abs(int(peer.height) - int(shape.height))
                            <= max(45_720, int(shape.height) * 0.15)
                            and peer_has_full_style
                        ):
                            peer_count += 1
                    if peer_count >= 2:
                        issues.append(
                            f"slide {slide_index} shape {shape_index} label "
                            f"{short_label!r} inherits theme-default styling "
                            "while equivalent peer labels on its dark panel "
                            "use explicit formatting"
                        )
            for paragraph_index, paragraph in enumerate(
                shape.text_frame.paragraphs, start=1
            ):
                runs = [run for run in paragraph.runs if run.text]
                if (
                    "\v" in paragraph.text
                    and runs
                    and any(run.font.size is None for run in runs)
                    and any(run.font.size is not None for run in runs)
                ):
                    issues.append(
                        f"slide {slide_index} shape {shape_index} paragraph "
                        f"{paragraph_index} has an unformatted run after a "
                        "manual line break"
                    )

            text_frame = shape.text_frame
            available_width = max(
                1.0,
                (
                    int(shape.width)
                    - int(text_frame.margin_left)
                    - int(text_frame.margin_right)
                )
                / 12_700,
            )
            available_height = max(
                1.0,
                (
                    int(shape.height)
                    - int(text_frame.margin_top)
                    - int(text_frame.margin_bottom)
                )
                / 12_700,
            )
            wrap = text_frame.word_wrap is not False
            layouts = [
                _ppt_paragraph_layout(
                    paragraph,
                    available_width,
                    wrap=wrap,
                )
                for paragraph in text_frame.paragraphs
            ]
            known_layouts = [layout for layout in layouts if layout is not None]
            if known_layouts and len(known_layouts) == len(layouts):
                widest = max(layout[0] for layout in known_layouts)
                required_height = sum(layout[1] for layout in known_layouts)
                auto_size = str(text_frame.auto_size or "")
                shrinks_to_fit = "TEXT_TO_FIT_SHAPE" in auto_size
                grows_to_fit = "SHAPE_TO_FIT_TEXT" in auto_size
                predicted_right = (
                    left
                    + int(text_frame.margin_left)
                    + int(text_frame.margin_right)
                    + int(widest * 12_700)
                )
                overlapping_right = any(
                    int(other.left) >= right - tolerance
                    and int(other.left) < predicted_right + 27_432
                    and int(other.top) < bottom
                    and int(other.top) + int(other.height) > top
                    for other in slide_shapes
                    if other is not shape
                )
                obvious_horizontal_growth = (
                    not wrap
                    and widest > available_width * 1.10
                    and not shrinks_to_fit
                    and (
                        predicted_right > slide_width + tolerance
                        or overlapping_right
                    )
                )
                explicit_font_sizes = [
                    float(size.pt)
                    for paragraph in text_frame.paragraphs
                    for size in (
                        paragraph.font.size,
                        *(
                            run.font.size
                            for run in paragraph.runs
                            if run.text
                        ),
                    )
                    if size is not None
                ]
                edge_title_growth = (
                    not wrap
                    and grows_to_fit
                    and not shrinks_to_fit
                    and widest > available_width * 1.01
                    and max(explicit_font_sizes, default=0.0) >= 18.0
                    and top <= slide_height * 0.30
                    and right >= slide_width - 685_800  # 0.75 inch
                    and right <= slide_width + tolerance
                    and (
                        predicted_right > slide_width + tolerance
                        or overlapping_right
                    )
                )
                if obvious_horizontal_growth:
                    issues.append(
                        f"slide {slide_index} shape {shape_index} text "
                        "requires horizontal growth and can collide with "
                        "adjacent content"
                    )
                elif edge_title_growth:
                    issues.append(
                        f"slide {slide_index} shape {shape_index} large "
                        "no-wrap title uses shape autofit near the slide edge "
                        "and can be clipped"
                    )
                needs_reflow = (
                    any(
                        layout[0] > available_width * 1.03
                        for layout in known_layouts
                    )
                    or len(text_frame.paragraphs) > 1
                    or any(
                        "\v" in paragraph.text
                        for paragraph in text_frame.paragraphs
                    )
                )
                vertical_growth = max(
                    0.0,
                    required_height - available_height,
                )
                vertical_anchor = text_frame.vertical_anchor
                if vertical_anchor == MSO_ANCHOR.MIDDLE:
                    growth_above = vertical_growth / 2
                    growth_below = vertical_growth / 2
                elif vertical_anchor == MSO_ANCHOR.BOTTOM:
                    growth_above = vertical_growth
                    growth_below = 0.0
                else:
                    # PowerPoint's inherited/default vertical anchor is top.
                    growth_above = 0.0
                    growth_below = vertical_growth
                predicted_top = top - int(growth_above * 12_700)
                predicted_bottom = bottom + int(growth_below * 12_700)
                overlapping_above = growth_above > 0 and any(
                    int(other.top) + int(other.height) <= top + tolerance
                    and int(other.top) + int(other.height)
                    > predicted_top - 27_432
                    and int(other.left) < right
                    and int(other.left) + int(other.width) > left
                    for other in slide_shapes
                    if other is not shape
                )
                overlapping_below = growth_below > 0 and any(
                    int(other.top) >= bottom - tolerance
                    and int(other.top) < predicted_bottom + 27_432
                    and int(other.left) < right
                    and int(other.left) + int(other.width) > left
                    for other in slide_shapes
                    if other is not shape
                )
                if (
                    wrap
                    and required_height > available_height * 1.15
                    and not shrinks_to_fit
                    and needs_reflow
                    and (
                        predicted_top < -tolerance
                        or overlapping_above
                        or predicted_bottom > slide_height + tolerance
                        or overlapping_below
                    )
                ):
                    issues.append(
                        f"slide {slide_index} shape {shape_index} text "
                        "requires vertical growth beyond its native bounds"
                    )

            # Detect a text box placed partly outside the small card/panel it
            # is clearly intended to occupy. Full-slide backgrounds and
            # decorative rules are excluded.
            text_area = max(1, int(shape.width) * int(shape.height))
            candidates: list[tuple[int, Any]] = []
            for container_index, container in enumerate(
                slide_shapes[: shape_index - 1],
                start=1,
            ):
                if shape.shape_type != MSO_SHAPE_TYPE.TEXT_BOX:
                    break
                if (
                    getattr(container, "has_chart", False)
                    or getattr(container, "has_table", False)
                    or (
                        getattr(container, "has_text_frame", False)
                        and container.text.strip()
                    )
                ):
                    continue
                container_area = int(container.width) * int(container.height)
                if not (
                    text_area * 2 <= container_area <= slide_width * slide_height * 0.25
                ):
                    continue
                container_left = int(container.left)
                container_top = int(container.top)
                container_right = container_left + int(container.width)
                container_bottom = container_top + int(container.height)
                if (
                    container_left - tolerance <= left
                    and right <= container_right + tolerance
                    and container_top - tolerance <= top <= container_bottom
                ):
                    candidates.append((container_area, container))
            if candidates:
                _, container = min(candidates, key=lambda row: row[0])
                container_bottom = int(container.top) + int(container.height)
                if bottom > container_bottom + 45_720:  # 0.05 inch
                    issues.append(
                        f"slide {slide_index} shape {shape_index} extends "
                        "below its containing panel"
                    )

            # Z-order is significant: a later opaque panel can cover an
            # earlier text box even though both shapes individually remain
            # inside the slide. Flag only material overlap by a later,
            # text-free solid panel to avoid penalizing intentional labels,
            # icons, or transparent decoration.
            shape_area = max(1, int(shape.width) * int(shape.height))
            for overlay_index, overlay in enumerate(
                slide_shapes[shape_index:],
                start=shape_index + 1,
            ):
                if (
                    getattr(overlay, "has_chart", False)
                    or getattr(overlay, "has_table", False)
                    or (
                        getattr(overlay, "has_text_frame", False)
                        and overlay.text.strip()
                    )
                    or not _ppt_opaque_solid_fill(overlay)
                ):
                    continue
                overlay_left = int(overlay.left)
                overlay_top = int(overlay.top)
                overlay_right = overlay_left + int(overlay.width)
                overlay_bottom = overlay_top + int(overlay.height)
                overlap_width = max(
                    0,
                    min(right, overlay_right) - max(left, overlay_left),
                )
                overlap_height = max(
                    0,
                    min(bottom, overlay_bottom) - max(top, overlay_top),
                )
                if overlap_width * overlap_height >= shape_area * 0.20:
                    issues.append(
                        f"slide {slide_index} shape {shape_index} text is "
                        f"materially occluded by later opaque shape "
                        f"{overlay_index}"
                    )
                    break
    return issues


def _recalculate_workbook(
    path: Path,
    workspace_root: Path | None = None,
) -> Path:
    global _SOFFICE_RECALC_DISABLED
    inspection_root = workspace_root or path.parent
    inspect_native_package(
        path,
        workspace_root=inspection_root,
    )
    if _SOFFICE_RECALC_DISABLED:
        return path
    temporary = Path(tempfile.mkdtemp(prefix="pinehaven-grade-xlsx-"))
    input_dir = temporary / "input"
    output_dir = temporary / "output"
    profile_dir = temporary / "profile"
    for directory in (
        input_dir,
        output_dir,
        profile_dir,
        temporary / "cache",
        temporary / "config",
        temporary / "home",
    ):
        directory.mkdir(mode=0o700)
    copied = input_dir / path.name
    try:
        with validated_native_package_copy(
            path,
            workspace_root=inspection_root,
        ) as (snapshot, _inspection):
            shutil.copy2(snapshot, copied)
        inspect_native_package(copied, workspace_root=temporary)
    except Exception:
        shutil.rmtree(temporary, ignore_errors=True)
        raise
    try:
        completed = run_sandboxed_libreoffice(
            sandbox_root=temporary,
            input_path=copied,
            output_dir=output_dir,
            profile_dir=profile_dir,
            convert_to="xlsx",
        )
    except Exception:
        shutil.rmtree(temporary, ignore_errors=True)
        raise
    if completed is None:
        # Never fall back to direct host execution. Non-release development
        # without a working sandbox uses the original validated workbook.
        _SOFFICE_RECALC_DISABLED = True
        return copied
    recalculated = output_dir / path.name
    output_exists = recalculated.is_file()
    if completed.returncode != 0 or not output_exists:
        shutil.rmtree(temporary, ignore_errors=True)
        raise NativePackageSafetyError(
            "sandboxed LibreOffice recalculation failed closed: "
            f"returncode={completed.returncode}; "
            f"output_exists={output_exists}"
        )
    try:
        inspect_native_package(recalculated, workspace_root=temporary)
    except Exception:
        shutil.rmtree(temporary, ignore_errors=True)
        raise
    return recalculated


def _estimated_text_width(text: str, font_size: float = 11.0) -> float:
    """Approximate text width in Excel's character-width units."""

    return (
        sum(
            0.42
            if character in " il.,'`:;|!"
            else 1.38
            if character in "MW@%&#"
            else 1.0
            for character in text
        )
        * font_size
        / 11.0
    )


def _formatted_cell_text(value: Any, number_format: str) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        calendar_value = value.date() if isinstance(value, datetime) else value
        date_code = str(number_format or "General").casefold()
        weekday = (
            f"{calendar_value.strftime('%A')}, "
            if "dddd" in date_code
            else (
                f"{calendar_value.strftime('%a')}, "
                if "ddd" in date_code
                else ""
            )
        )
        if "mmmm" in date_code:
            return (
                f"{weekday}{calendar_value.strftime('%B')} "
                f"{calendar_value.day}, {calendar_value.year}"
            )
        if "mmm" in date_code:
            return (
                f"{weekday}{calendar_value.strftime('%b')} "
                f"{calendar_value.day}, {calendar_value.year}"
            )
        # ISO text is a conservative approximation for ordinary numeric date
        # formats and, unlike datetime.isoformat(), does not invent a midnight
        # time that Excel does not render.
        return f"{weekday}{calendar_value.isoformat()}"
    if not isinstance(value, (int, float)) or isinstance(value, bool):
        return str(value)
    if not math.isfinite(float(value)):
        return str(value)
    code = str(number_format or "General")
    if code.casefold() == "general":
        return str(value)
    section = code.split(";")[1 if value < 0 and ";" in code else 0]
    percent = "%" in section
    scaled = float(value) * (100 if percent else 1)
    # A comma immediately after the numeric placeholder scales by 1,000.
    scale_commas = len(re.findall(r"(?<=[0#]),(?=[^0#]*$)", section))
    if scale_commas:
        scaled /= 1_000 ** scale_commas
    decimal_match = re.search(r"\.([0#]+)", section)
    decimals = len(decimal_match.group(1)) if decimal_match else 0
    grouping = "#,##" in section or "0,00" in section
    absolute = abs(scaled)
    rendered = (
        f"{absolute:,.{decimals}f}"
        if grouping
        else f"{absolute:.{decimals}f}"
    )
    if "$" in section:
        rendered = f"${rendered}"
    if percent:
        rendered = f"{rendered}%"
    placeholder_positions = [
        index
        for index, character in enumerate(section)
        if character in "0#"
    ]
    suffix = (
        section[max(placeholder_positions) + 1 :]
        if placeholder_positions
        else ""
    )
    literal_suffixes = "".join(re.findall(r'"([^"]+)"', suffix))
    rendered = f"{rendered}{literal_suffixes}"
    if value < 0:
        rendered = f"({rendered})" if "(" in section else f"-{rendered}"
    return rendered


def _workbook_cell_visual_issues(
    sheet: Any,
    value_sheet: Any,
    *,
    limit: int,
) -> list[str]:
    """Return bounded, high-confidence cell clipping findings.

    Column-width arithmetic is only an approximation of Excel's renderer. It
    is useful for concise report tabs, but generates noisy false positives on
    dense raw-data tabs whose display also depends on zoom and renderer font
    metrics. Those large tabs still receive formula, error, chart, header, and
    native-artifact checks elsewhere in the workbook grader.
    """

    if (
        limit <= 0
        or sheet.max_row * sheet.max_column > _WORKBOOK_MAX_CLIP_SCAN_CELLS
    ):
        return []

    # A merged cell's value lives at its top-left anchor. Index those anchors
    # once so lookup is O(1), rather than scanning every merged range for
    # every populated cell.
    merged_anchors = {
        (merged_range.min_row, merged_range.min_col): merged_range
        for merged_range in sheet.merged_cells.ranges
    }
    issues: list[str] = []
    for row in sheet.iter_rows():
        for cell in row:
            source_value = cell.value
            value = (
                value_sheet[cell.coordinate].value
                if isinstance(source_value, str)
                and source_value.startswith("=")
                else source_value
            )
            if (
                value in (None, "")
                or cell.alignment.wrap_text
                or cell.alignment.shrink_to_fit
            ):
                continue
            merged = merged_anchors.get((cell.row, cell.column))
            if merged is not None:
                available_width = sum(
                    sheet.column_dimensions[
                        get_column_letter(column)
                    ].width
                    or 13.0
                    for column in range(
                        merged.min_col,
                        merged.max_col + 1,
                    )
                )
                right_column = merged.max_col + 1
            else:
                available_width = (
                    sheet.column_dimensions[
                        get_column_letter(cell.column)
                    ].width
                    or 13.0
                )
                right_column = cell.column + 1
            right_value = (
                value_sheet.cell(
                    row=cell.row,
                    column=right_column,
                ).value
                if right_column <= sheet.max_column
                else None
            )
            rendered = _formatted_cell_text(
                value,
                cell.number_format,
            )
            text_width = _estimated_text_width(
                rendered,
                float(cell.font.sz or 11.0),
            )
            if cell.font.bold:
                text_width *= 1.10
            is_numeric = (
                isinstance(value, (int, float))
                and not isinstance(value, bool)
                and str(cell.number_format or "General").casefold()
                != "general"
            )
            if (
                not is_numeric
                and not (
                    isinstance(source_value, str)
                    and source_value.startswith("=")
                )
            ):
                probe_column = right_column
                right_value = None
                while probe_column <= sheet.max_column:
                    candidate = value_sheet.cell(
                        row=cell.row,
                        column=probe_column,
                    ).value
                    if candidate not in (None, ""):
                        right_value = candidate
                        break
                    available_width += (
                        sheet.column_dimensions[
                            get_column_letter(probe_column)
                        ].width
                        or 13.0
                    )
                    probe_column += 1
            text_blocked = (
                not is_numeric
                and not (
                    isinstance(source_value, str)
                    and source_value.startswith("=")
                )
                and right_value not in (None, "")
            )
            overflow_ratio = (
                _WORKBOOK_NUMERIC_CLIP_OVERFLOW_RATIO
                if is_numeric
                else _WORKBOOK_TEXT_CLIP_OVERFLOW_RATIO
            )
            if not (
                "###" in rendered
                or (
                    (is_numeric or text_blocked)
                    and text_width
                    > available_width * overflow_ratio
                )
            ):
                continue
            kind = "numeric display" if is_numeric else "text"
            issues.append(
                f"{sheet.title}!{cell.coordinate}: {kind} is "
                "likely clipped"
                + (
                    " before a populated adjacent cell"
                    if text_blocked
                    else ""
                )
                + " ("
                f"estimated width {text_width:.1f}, "
                f"available {available_width:.1f})"
            )
            if len(issues) >= limit:
                return issues
    return issues


def _chart_data_count(source: Any) -> int | None:
    if source is None:
        return None
    for reference_name, literal_name in (
        ("numRef", "numLit"),
        ("strRef", "strLit"),
        ("multiLvlStrRef", None),
    ):
        reference = getattr(source, reference_name, None)
        cache = (
            getattr(reference, "numCache", None)
            or getattr(reference, "strCache", None)
            or getattr(reference, "multiLvlStrCache", None)
            if reference is not None
            else None
        )
        if cache is not None:
            count = getattr(cache, "ptCount", None)
            if count is not None:
                return int(count)
            points = getattr(cache, "pt", None)
            if points is not None:
                return len(points)
        if literal_name:
            literal = getattr(source, literal_name, None)
            if literal is not None:
                count = getattr(literal, "ptCount", None)
                if count is not None:
                    return int(count)
                points = getattr(literal, "pt", None)
                if points is not None:
                    return len(points)
    return None


def _chart_numeric_values(series: Any) -> list[float]:
    source = getattr(series, "val", None) or getattr(series, "yVal", None)
    reference = getattr(source, "numRef", None) if source is not None else None
    cache = getattr(reference, "numCache", None) if reference is not None else None
    result: list[float] = []
    for point in getattr(cache, "pt", ()) if cache is not None else ():
        try:
            result.append(float(point.v))
        except (TypeError, ValueError):
            continue
    return result


def _mixed_sign_zero_axis_label_collision(
    chart: Any,
    numeric_values: Iterable[float],
) -> bool:
    """Detect horizontal mixed-sign bars whose category labels sit at zero."""

    values = list(numeric_values)
    if (
        getattr(chart, "type", None) != "bar"
        or not values
        or min(values) >= 0
        or max(values) <= 0
    ):
        return False
    category_axis = next(
        (
            axis
            for axis in (
                getattr(chart, "x_axis", None),
                getattr(chart, "y_axis", None),
            )
            if type(axis).__name__ == "TextAxis"
        ),
        None,
    )
    numeric_axis = next(
        (
            axis
            for axis in (
                getattr(chart, "x_axis", None),
                getattr(chart, "y_axis", None),
            )
            if type(axis).__name__ == "NumericAxis"
        ),
        None,
    )
    return (
        category_axis is not None
        and numeric_axis is not None
        and getattr(category_axis, "tickLblPos", None) == "nextTo"
        and getattr(numeric_axis, "crosses", None) == "autoZero"
    )


def _defined_name_references(workbook: Any) -> dict[str, str]:
    """Resolve bounded workbook names to native single-cell/range references."""

    references: dict[str, str] = {}
    for name, definition in workbook.defined_names.items():
        if getattr(definition, "type", None) != "RANGE":
            continue
        try:
            destinations = list(definition.destinations)
        except (AttributeError, TypeError, ValueError):
            continue
        if len(destinations) != 1:
            continue
        sheet_name, address = destinations[0]
        if not isinstance(sheet_name, str) or not isinstance(address, str):
            continue
        cleaned_address = address.replace("$", "")
        if not re.fullmatch(
            r"[A-Za-z]{1,3}[1-9]\d*(?::[A-Za-z]{1,3}[1-9]\d*)?",
            cleaned_address,
        ):
            continue
        escaped_sheet = sheet_name.replace("'", "''")
        references[name.casefold()] = f"'{escaped_sheet}'!{address}"
    return references


def _expand_formula_defined_names(
    formula: str,
    defined_name_references: Mapping[str, str],
) -> str:
    """Expand native workbook names without touching strings or functions."""

    try:
        tokens = Tokenizer(formula).items
    except Exception:
        return formula
    changed = False
    pieces: list[str] = []
    for token in tokens:
        replacement = (
            defined_name_references.get(str(token.value).casefold())
            if token.type == "OPERAND" and token.subtype == "RANGE"
            else None
        )
        if replacement is not None:
            pieces.append(replacement)
            changed = True
        else:
            pieces.append(str(token.value))
    if not changed:
        return formula
    return ("=" if formula.startswith("=") else "") + "".join(pieces)


def _native_formula_label(sheet: Any, cell: Any, row: Iterable[Any]) -> str:
    """Return an exact adjacent label for row or compact card layouts."""

    normalized_sheet = _normalize(sheet.title)
    # Reconciliation tables conventionally put the metric key in column A and
    # put calculated values, authority descriptions, deltas, and statuses to
    # its right.  Prefer that row key on explicitly named control/check sheets;
    # otherwise a nearby source-description literal can incorrectly become
    # the formula label and make a complete control appear unrelated.
    if normalized_sheet in {"control", "controls", "check", "checks"}:
        row_key = sheet.cell(row=cell.row, column=1).value
        if (
            cell.column > 1
            and isinstance(row_key, str)
            and row_key.strip()
            and not row_key.startswith("=")
        ):
            return row_key.strip()
    if cell.column > 1:
        left = sheet.cell(row=cell.row, column=cell.column - 1).value
        if (
            isinstance(left, str)
            and left.strip()
            and not left.startswith("=")
            and _year_month(left) is None
        ):
            return left.strip()
    # Native reconciliation tables conventionally keep the metric name in
    # the first populated cell and place calculated, authority, difference,
    # tolerance, and status formulas across the rest of that row.  Requiring
    # the label to be immediately adjacent makes the status formula look
    # unlabeled and falsely rejects a complete independent control.  Restrict
    # the wider row lookup to sheets explicitly presented as controls/checks
    # so detail schedules cannot borrow an unrelated distant row label.
    if normalized_sheet in {"control", "controls", "check", "checks"}:
        for prior_column in range(cell.column - 2, 0, -1):
            candidate = sheet.cell(
                row=cell.row,
                column=prior_column,
            ).value
            if (
                isinstance(candidate, str)
                and candidate.strip()
                and not candidate.startswith("=")
            ):
                return candidate.strip()
    row_has_literal = any(
        candidate.value not in (None, "")
        and not (
            isinstance(candidate.value, str)
            and candidate.value.startswith("=")
        )
        for candidate in row
    )
    if not row_has_literal and cell.row > 1:
        above = sheet.cell(row=cell.row - 1, column=cell.column).value
        if (
            isinstance(above, str)
            and above.strip()
            and not above.startswith("=")
        ):
            return above.strip()
    return ""


def _cached_formula_issues(
    formula_sheet: Any,
    value_book: Any,
    defined_name_references: Mapping[str, str],
) -> list[str]:
    """Detect visibly stale summary-card formulas after native recalculation."""

    issues: list[str] = []
    value_sheet = value_book[formula_sheet.title]
    merged_anchors = {
        (merged.min_row, merged.min_col)
        for merged in formula_sheet.merged_cells.ranges
    }
    for row, column in merged_anchors:
        cell = formula_sheet.cell(row=row, column=column)
        original_formula = cell.value
        if (
            not isinstance(original_formula, str)
            or not original_formula.startswith("=")
        ):
            continue
        formula = _expand_formula_defined_names(
            original_formula,
            defined_name_references,
        )
        cached = value_sheet[cell.coordinate].value
        cached_missing = cached in (None, "", 0, 0.0)
        direct = re.fullmatch(
            r"=\s*(?:(?:'((?:[^']|'')+)'|([^'!]+))!)?"
            r"\$?([A-Za-z]{1,3})\$?([1-9]\d*)\s*",
            formula,
        )
        referenced_value: Any = None
        if direct:
            sheet_name = (
                (direct.group(1) or "").replace("''", "'")
                or direct.group(2)
                or formula_sheet.title
            )
            if sheet_name in value_book.sheetnames:
                referenced_value = value_book[sheet_name][
                    f"{direct.group(3)}{direct.group(4)}"
                ].value
        expected_date: date | datetime | None = (
            referenced_value
            if isinstance(referenced_value, (date, datetime))
            else None
        )
        date_formula = re.fullmatch(
            r"=\s*DATE\(\s*(\d{4})\s*,\s*(\d{1,2})\s*,\s*(\d{1,2})\s*\)",
            formula,
            flags=re.IGNORECASE,
        )
        if date_formula:
            try:
                expected_date = date(
                    int(date_formula.group(1)),
                    int(date_formula.group(2)),
                    int(date_formula.group(3)),
                )
            except ValueError:
                expected_date = None
        if (
            isinstance(cached, (date, datetime))
            and expected_date is not None
            and (
                cached.year,
                cached.month,
                cached.day,
            )
            != (
                expected_date.year,
                expected_date.month,
                expected_date.day,
            )
        ):
            issues.append(
                f"{formula_sheet.title}!{cell.coordinate}: date formula "
                f"recalculates to {cached!r}, expected {expected_date!r}"
            )
            continue
        string_formula = (
            formula.lstrip().upper().startswith("=IF")
            and bool(re.search(r'"[^"]+"', formula))
        )
        if cached_missing and (
            referenced_value not in (None, "", 0, 0.0)
            or string_formula
        ):
            issues.append(
                f"{formula_sheet.title}!{cell.coordinate}: merged summary "
                f"formula recalculates to {cached!r}"
            )
    return issues


def _cross_engine_formula_issues(
    formula_sheet: Any,
    value_book: Any,
) -> list[str]:
    """Find formulas that depend on engine-specific implicit coercion."""

    issues: list[str] = []
    for row in formula_sheet.iter_rows():
        for cell in row:
            formula = cell.value
            if not isinstance(formula, str) or not formula.startswith("="):
                continue
            try:
                tokens = Tokenizer(formula).items
            except Exception:
                continue
            if not any(
                token.type == "OPERATOR-INFIX" and token.value == "-"
                for token in tokens
            ):
                continue
            references = [
                _formula_reference(formula_sheet.title, token.value)
                for token in tokens
                if token.type == "OPERAND" and token.subtype == "RANGE"
            ]
            referenced_values: list[Any] = []
            for reference in references:
                if reference is None:
                    continue
                sheet_name, coordinate = reference.rsplit("!", 1)
                if sheet_name in value_book.sheetnames:
                    referenced_values.append(
                        value_book[sheet_name][coordinate].value
                    )
            text_dates = [
                value
                for value in referenced_values
                if isinstance(value, str) and _calendar_date(value) is not None
            ]
            if len(text_dates) >= 2:
                issues.append(
                    f"{formula_sheet.title}!{cell.coordinate}: subtracts "
                    "text-form ISO dates; use DATEVALUE or exact equality"
                )
    return issues


_CONTROL_STATUS_LITERALS = frozenset(
    {
        "check",
        "explained",
        "fail",
        "match",
        "mismatch",
        "ok",
        "pass",
    }
)
_CONTROL_STATUS_HEADER_TOKENS = frozenset(
    {"check", "control", "outcome", "result", "status"}
)


def _meaningful_control_formula(formula: Any) -> bool:
    """Return whether a formula performs recognizable control logic."""

    if not isinstance(formula, str) or not formula.startswith("="):
        return False
    try:
        tokens = Tokenizer(formula).items
    except Exception:
        tokens = ()
    references = {
        token.value.replace("$", "").upper()
        for token in tokens
        if token.type == "OPERAND" and token.subtype == "RANGE"
    }
    if not references:
        return False
    has_comparison = any(
        token.type == "OPERATOR-INFIX"
        and token.value in {"<", "<=", "=", "<>", ">=", ">"}
        for token in tokens
    )
    has_reconciliation_operator = any(
        token.type == "OPERATOR-INFIX" and token.value == "-"
        for token in tokens
    )
    control_functions = {
        token.value.rstrip("(").upper()
        for token in tokens
        if token.type == "FUNC" and token.subtype == "OPEN"
    }
    has_conditional_control = bool(
        control_functions.intersection(
            {
                "AND",
                "COUNTIF",
                "IF",
                "IFS",
                "OR",
                "SUMIF",
                "SWITCH",
            }
        )
    )
    normalized = re.sub(r"\s+", "", formula).upper()
    has_status_literal = bool(
        re.search(
            r'"(?:CHECK|EXPLAINED|FAIL|MATCH|MISMATCH|OK|PASS)"',
            normalized,
        )
    )
    if has_comparison and (has_conditional_control or has_status_literal):
        return True
    if len(references) >= 2 and (
        has_comparison
        or has_reconciliation_operator
        or bool(
            control_functions.intersection(
                {"COUNTIF", "SUMIF"}
            )
        )
    ):
        return True
    return False


def _control_dependency_roles(
    cell: str,
    formula: str,
    formula_cells: dict[str, str],
    *,
    calculated_cells: set[str] | None = None,
    accepted_control_literals: set[str] | None = None,
    seen: set[str] | None = None,
) -> set[str]:
    """Classify a control formula's independent dependency paths.

    A valid reconciliation needs one path to the calculated Analysis result
    and a separate path to a native source/control total. Merely wrapping an
    Analysis cell in ``+0``, ``ROUND``, or a local alias is not independent
    control evidence.
    """

    if seen is None:
        seen = set()
    if cell in seen:
        return set()
    path = seen | {cell}
    current_sheet = cell.rsplit("!", 1)[0]
    try:
        tokens = Tokenizer(formula).items
    except Exception:
        return set()
    roles: set[str] = set()
    for token in tokens:
        if token.type != "OPERAND" or token.subtype != "RANGE":
            continue
        references = _formula_references(current_sheet, token.value)
        if not references:
            continue
        token_roles: set[str] = set()
        for reference in references:
            token_roles.update(
                _control_reference_roles(
                    reference,
                    formula_cells,
                    calculated_cells=calculated_cells,
                    accepted_control_literals=accepted_control_literals,
                    seen=path,
                )
            )
        # A source range or alias contaminated by an Analysis dependency (or
        # by a cycle) is not an independent control path. A separate clean
        # token in the control formula can still establish source evidence.
        if "analysis" in token_roles:
            roles.add("analysis")
        if "cycle" in token_roles:
            roles.add("cycle")
        if (
            "independent_source" in token_roles
            and "analysis" not in token_roles
            and "cycle" not in token_roles
        ):
            roles.add("independent_source")
    return roles


def _control_reference_roles(
    reference: str,
    formula_cells: dict[str, str],
    *,
    calculated_cells: set[str] | None,
    accepted_control_literals: set[str] | None,
    seen: set[str],
) -> set[str]:
    """Classify one dependency with bounded, cycle-safe graph traversal."""

    inherited_path = set(seen)
    if reference in inherited_path:
        return {"cycle"}

    maximum_nodes = len(formula_cells) + 1
    maximum_edges = 1_000_000
    pending = [reference]
    visited: set[str] = set()
    dependencies: dict[str, set[str]] = {}
    reaches_analysis = False
    reaches_source = False
    invalid = False
    edge_count = 0

    while pending:
        current = pending.pop()
        if current in inherited_path:
            invalid = True
            continue
        if current in visited:
            continue
        visited.add(current)
        dependencies.setdefault(current, set())
        if len(visited) > maximum_nodes:
            invalid = True
            break

        current_sheet = _normalize(current.rsplit("!", 1)[0])
        if current in (accepted_control_literals or set()):
            reaches_source = True
            continue
        if (
            calculated_cells is not None
            and current in calculated_cells
        ) or current_sheet == "analysis":
            reaches_analysis = True
            continue

        upstream = formula_cells.get(current)
        if upstream is None:
            if current_sheet not in {"control", "read me"}:
                reaches_source = True
            continue

        try:
            tokens = Tokenizer(upstream).items
        except Exception:
            invalid = True
            continue
        references: list[str] = []
        for token in tokens:
            if token.type != "OPERAND" or token.subtype != "RANGE":
                continue
            references.extend(
                _formula_references(
                    current.rsplit("!", 1)[0],
                    token.value,
                )
            )
        if (
            not references
            and current_sheet not in {"control", "read me"}
            and re.fullmatch(
                r"=(?:TRUE|FALSE)\(\)",
                re.sub(r"\s+", "", upstream).upper(),
            )
        ):
            # Boolean covenant truth may be encoded as a zero-input formula to
            # preserve its native type. It is a source terminal, unlike an
            # alias to Analysis or a self-reference discovered below.
            reaches_source = True
            continue

        for dependency in references:
            edge_count += 1
            if edge_count > maximum_edges:
                invalid = True
                pending.clear()
                break
            if dependency in inherited_path:
                invalid = True
                continue
            dependencies[current].add(dependency)
            dependencies.setdefault(dependency, set())
            if dependency not in visited:
                pending.append(dependency)

    indegree = {node: 0 for node in dependencies}
    for targets in dependencies.values():
        for target in targets:
            indegree[target] = indegree.get(target, 0) + 1
    ready = [node for node, degree in indegree.items() if degree == 0]
    processed = 0
    while ready:
        node = ready.pop()
        processed += 1
        for target in dependencies.get(node, set()):
            indegree[target] -= 1
            if indegree[target] == 0:
                ready.append(target)
    if processed != len(indegree):
        invalid = True

    roles: set[str] = set()
    if reaches_analysis:
        roles.add("analysis")
    if invalid:
        roles.add("cycle")
    if reaches_source and not reaches_analysis and not invalid:
        roles.add("independent_source")
    if not roles and _normalize(reference.rsplit("!", 1)[0]) not in {
        "control",
        "read me",
    }:
        roles.add("cycle")
    return roles


_PROMPT_AUTHORITATIVE_CONTROL_KEYS = frozenset(
    {
        "baseline_period_start",
        "baseline_period_end",
        "scenario_period",
        "status",
        "volume_change",
        "price_change",
        "material_cost_change",
    }
)


def _prompt_authoritative_control_literals(
    key: str,
    expected: Any,
    control_cell: str,
    literal_cells: Mapping[str, Any] | None,
) -> set[str]:
    """Find same-row Control literals that exactly restate prompt truth.

    These literals are legitimate independent expectations only for inputs
    supplied authoritatively by the task prompt. Monetary model outputs remain
    required to reconcile to a separate native source lineage.
    """

    if key not in _PROMPT_AUTHORITATIVE_CONTROL_KEYS or not literal_cells:
        return set()
    matched = re.fullmatch(
        r"(.+)!([A-Z]{1,3})([1-9]\d*)",
        control_cell,
    )
    if matched is None or _normalize(matched.group(1)) not in {
        "control",
        "controls",
        "check",
        "checks",
    }:
        return set()
    sheet, _column, row = matched.groups()
    return {
        literal_cell
        for literal_cell, actual in literal_cells.items()
        if literal_cell.rsplit("!", 1)[0] == sheet
        and re.fullmatch(
            rf"[A-Z]{{1,3}}{re.escape(row)}",
            literal_cell.rsplit("!", 1)[-1],
        )
        and _value_matches(actual, expected, key)
    }


def _central_metric_control_coverage(
    values: dict[str, Any],
    formula_cells: dict[str, str],
    formula_row_labels: dict[str, str],
    literal_cells: Mapping[str, Any] | None = None,
) -> tuple[bool, list[dict[str, Any]]]:
    """Require an independent native reconciliation for every central metric."""

    details: list[dict[str, Any]] = []
    for key in values:
        expected_label = _normalize(key.replace("_", " "))
        same_label_formula_cells = {
            cell
            for cell in formula_cells
            if _normalize(
                formula_row_labels.get(cell, "").replace("_", " ")
            )
            == expected_label
        }
        candidates: list[dict[str, Any]] = []
        for cell, formula in formula_cells.items():
            actual_label = _normalize(
                formula_row_labels.get(cell, "").replace("_", " ")
            )
            if (
                actual_label != expected_label
                or not _meaningful_control_formula(formula)
            ):
                continue
            accepted_control_literals = (
                _prompt_authoritative_control_literals(
                    key,
                    values[key],
                    cell,
                    literal_cells,
                )
            )
            role_trials = {
                calculated_cell: _control_dependency_roles(
                    cell,
                    formula,
                    formula_cells,
                    calculated_cells={calculated_cell},
                    accepted_control_literals=(
                        accepted_control_literals
                    ),
                )
                for calculated_cell in same_label_formula_cells - {cell}
            }
            roles = next(
                (
                    trial_roles
                    for trial_roles in role_trials.values()
                    if trial_roles
                    == {"analysis", "independent_source"}
                ),
                set().union(*role_trials.values())
                if role_trials
                else set(),
            )
            candidates.append(
                {
                    "cell": cell,
                    "meaningful": True,
                    "roles": sorted(roles),
                    "accepted_control_literals": sorted(
                        accepted_control_literals
                    ),
                    "role_trials": {
                        dependency: sorted(trial_roles)
                        for dependency, trial_roles in role_trials.items()
                    },
                }
            )
        controlled = any(
            candidate["meaningful"]
            and set(candidate["roles"])
            == {"analysis", "independent_source"}
            for candidate in candidates
        )
        details.append(
            {
                "metric_key": key,
                "controlled": controlled,
                "candidates": candidates,
            }
        )
    return all(detail["controlled"] for detail in details), details


def _hardcoded_control_status_cells(sheet: Any) -> list[str]:
    """Find literal outcomes only where the sheet presents a control result."""

    cells: list[str] = []
    for row in sheet.iter_rows():
        for cell in row:
            if (
                not isinstance(cell.value, str)
                or _normalize(cell.value) not in _CONTROL_STATUS_LITERALS
            ):
                continue
            adjacent_formula = any(
                _meaningful_control_formula(
                    sheet.cell(row=cell.row, column=column).value
                )
                for column in (cell.column - 1, cell.column + 1)
                if 1 <= column <= sheet.max_column
            )
            header_context = False
            for prior_row in range(cell.row - 1, 0, -1):
                candidate = sheet.cell(
                    row=prior_row,
                    column=cell.column,
                ).value
                if candidate in (None, ""):
                    continue
                if (
                    isinstance(candidate, str)
                    and (
                        candidate.startswith("=")
                        or _normalize(candidate) in _CONTROL_STATUS_LITERALS
                    )
                ):
                    continue
                candidate_tokens = set(_normalize(candidate).split())
                header_context = bool(
                    candidate_tokens.intersection(
                        _CONTROL_STATUS_HEADER_TOKENS
                    )
                )
                break
            if adjacent_formula or header_context:
                cells.append(f"{sheet.title}!{cell.coordinate}")
    return cells


_ACTION_REGISTER_EVIDENCE_LIMIT = 5


def _action_register_identifier(value: Any) -> str:
    """Normalize an Excel identifier without weakening text comparisons."""

    if isinstance(value, bool) or value is None:
        return ""
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float) and math.isfinite(value) and value.is_integer():
        return str(int(value))
    return str(value)


def _action_register_field_matches(
    field: str,
    actual: Any,
    expected: Any,
) -> bool:
    if field == "quantity":
        return (
            isinstance(actual, (int, float))
            and not isinstance(actual, bool)
            and math.isfinite(float(actual))
            and float(actual) == float(expected)
        )
    if field == "priority":
        return (
            isinstance(actual, (int, float))
            and not isinstance(actual, bool)
            and math.isfinite(float(actual))
            and float(actual).is_integer()
            and int(actual) == int(expected)
        )
    if field in {"release_date", "required_date"}:
        actual_date = _calendar_date(actual)
        expected_date = _calendar_date(expected)
        return (
            actual_date is not None
            and expected_date is not None
            and actual_date == expected_date
        )
    if field == "site_code":
        return (
            _action_register_identifier(actual)
            == _action_register_identifier(expected)
        )
    return isinstance(actual, str) and actual == str(expected)


def _action_register_sort_key(
    row: list[Any],
    header_indexes: dict[str, int],
) -> tuple[int, date, str] | None:
    priority = row[header_indexes["priority"]]
    required_date = _calendar_date(
        row[header_indexes["required_date"]]
    )
    planned_order_id = _action_register_identifier(
        row[header_indexes["planned_order_id"]]
    )
    if (
        isinstance(priority, bool)
        or not isinstance(priority, (int, float))
        or not math.isfinite(float(priority))
        or not float(priority).is_integer()
        or required_date is None
        or not planned_order_id
    ):
        return None
    return int(priority), required_date, planned_order_id


def _action_register_filter_refs(sheet: Any) -> list[str]:
    refs = [
        str(table.ref)
        for table in sheet.tables.values()
        if getattr(table, "ref", None)
    ]
    auto_filter_ref = getattr(sheet.auto_filter, "ref", None)
    if auto_filter_ref:
        refs.append(str(auto_filter_ref))
    return refs


def _action_register_filter_covers(
    reference: str,
    *,
    header_row: int,
    header_column: int,
    column_count: int,
    data_row_count: int,
) -> bool:
    try:
        min_column, min_row, max_column, max_row = range_boundaries(
            reference
        )
    except (TypeError, ValueError):
        return False
    return (
        min_row == header_row
        and min_column <= header_column
        and max_column >= header_column + column_count - 1
        and max_row >= header_row + data_row_count
    )


def _empty_action_register_evidence(
    error: str = "register not parsed",
) -> dict[str, Any]:
    return {
        "error": error,
        "header_candidate_count": 0,
        "header_candidate_samples": [],
        "selected_location": None,
        "filterable": False,
        "filter_refs": [],
        "row_count": 0,
        "expected_row_count": 0,
        "unique_id_count": 0,
        "expected_unique_id_count": 0,
        "duplicate_id_count": 0,
        "duplicate_id_samples": [],
        "missing_id_count": 0,
        "missing_id_samples": [],
        "extra_id_count": 0,
        "extra_id_samples": [],
        "field_mismatch_count": 0,
        "field_mismatch_samples": [],
        "first_order_mismatch_samples": [],
        "field_values_match": False,
        "sort_order_match": False,
    }


def _empty_detail_schedule_evidence(
    error: str = "detail schedule not parsed",
) -> dict[str, Any]:
    return {
        "sheet_present": False,
        "headers_match": False,
        "header_candidate_count": 0,
        "header_candidate_samples": [],
        "selected_location": None,
        "formula_backed_row_count": 0,
        "model_cell_count": 0,
        "model_cells": [],
        "model_formula_cell_count": 0,
        "model_formula_cells": [],
        "row_count": 0,
        "expected_row_count": 0,
        "keys_match": False,
        "duplicate_key_count": 0,
        "missing_keys": [],
        "extra_keys": [],
        "field_values_match": False,
        "field_mismatch_count": 0,
        "field_mismatch_samples": [],
        "error": error,
    }


def _detail_schedule_field_matches(
    column: str,
    actual: Any,
    expected: Any,
    *,
    row_context: dict[str, Any] | None = None,
) -> bool:
    """Compare an authoritative schedule field at its declared precision."""

    if expected in (None, ""):
        return actual in (None, "")
    if isinstance(expected, bool):
        return isinstance(actual, bool) and actual is expected
    normalized_column = _normalize(column)
    if "month" in set(normalized_column.split()):
        actual_period = _year_month(actual)
        expected_period = _year_month(expected)
        if actual_period is not None or expected_period is not None:
            return (
                actual_period is not None
                and expected_period is not None
                and actual_period == expected_period
            )
    if (
        isinstance(expected, (date, datetime))
        or any(
            token in normalized_column
            for token in ("date", "week ending", "month ending")
        )
    ):
        actual_date = _calendar_date(actual)
        expected_date = _calendar_date(expected)
        if actual_date is not None or expected_date is not None:
            return (
                actual_date is not None
                and expected_date is not None
                and actual_date == expected_date
            )
    if isinstance(expected, (int, float)) and not isinstance(expected, bool):
        if (
            not isinstance(actual, (int, float))
            or isinstance(actual, bool)
            or not math.isfinite(float(actual))
        ):
            return False
        expected_number = float(expected)
        actual_number = float(actual)
        if isinstance(expected, int):
            return actual_number.is_integer() and int(actual_number) == expected
        normalized_row = _normalize(
            " ".join(map(str, (row_context or {}).values()))
        )
        if (
            normalized_column in {"actual", "threshold", "headroom"}
            and "funded debt ltm adjusted ebitda" in normalized_row
        ):
            tolerance = 0.0000005
        elif "discount factor" in normalized_column:
            tolerance = 0.000000005
        elif "payback" in normalized_column:
            tolerance = 0.00005
        elif any(
            token in normalized_column
            for token in (
                "rate",
                "ratio",
                "margin",
                "percent",
                "score",
                "irr",
                "index",
                "yield",
                "utilization",
            )
        ):
            tolerance = 0.0000005
        elif any(
            token in normalized_column
            for token in ("quantity", "units", "hours")
        ):
            tolerance = 0.00005
        else:
            tolerance = 0.005
        epsilon = max(
            1.0,
            abs(actual_number),
            abs(expected_number),
        ) * 1e-12
        return abs(actual_number - expected_number) <= tolerance + epsilon
    normalized_actual = _normalize(actual)
    normalized_expected = _normalize(expected)
    if normalized_column == "maintenance classification":
        classification_aliases = {
            "planned preventive": {"planned", "planned preventive"},
            "other scheduled control work": {
                "other scheduled control work",
                "retained neither",
            },
        }
        if normalized_actual in classification_aliases.get(
            normalized_expected,
            {normalized_expected},
        ):
            return True
    return normalized_actual == normalized_expected


def _detail_schedule_key_part(column: str, value: Any) -> str:
    """Canonicalize only declared calendar dimensions for key matching."""

    normalized_column = _normalize(column)
    column_tokens = set(normalized_column.split())
    if "month" in column_tokens:
        period = _year_month(value)
        if period is not None:
            return f"{period[0]:04d}-{period[1]:02d}"
    if "date" in column_tokens or "week ending" in normalized_column:
        calendar_date = _calendar_date(value)
        if calendar_date is not None:
            return calendar_date.isoformat()
    return _normalize(value)


_DETAIL_SCHEDULE_HEADER_ALIASES: dict[str, frozenset[str]] = {
    "baseline month": frozenset(
        {"baseline 2025 month", "2025 baseline month"}
    ),
    "baseline material cost": frozenset({"baseline material"}),
    "baseline nonmaterial cost": frozenset({"baseline nonmaterial"}),
    "baseline operating expense": frozenset(
        {"operating expense", "baseline opex", "baseline op ex"}
    ),
    "scenario material cost": frozenset({"scenario material"}),
    "scenario nonmaterial cost": frozenset({"scenario nonmaterial"}),
    "scenario gross profit": frozenset({"gross profit"}),
    "scenario gross margin": frozenset({"gross margin"}),
    "scenario operating expense": frozenset(
        {"operating expense", "scenario opex", "scenario op ex"}
    ),
    "scenario operating income": frozenset({"operating income"}),
    "maintenance classification": frozenset({"planning classification"}),
    "maintenance cost": frozenset({"cost"}),
}


def _is_detail_total_boundary(value: Any) -> bool:
    """Recognize a professional trailing total without swallowing detail."""

    normalized = _normalize(value)
    if normalized in {"total", "grand total", "subtotal"}:
        return True
    tokens = set(normalized.split())
    return "annual" in tokens and bool(
        tokens.intersection({"sum", "total", "subtotal"})
    )


def _semantic_detail_header_positions(
    row: Iterable[Any],
    required_headers: Iterable[str],
) -> dict[str, list[int]]:
    """Match explicit business synonyms while keeping headers fail-closed."""

    positions: dict[str, list[int]] = {}
    normalized_row = [_normalize(value) for value in row]
    for required in required_headers:
        aliases = _DETAIL_SCHEDULE_HEADER_ALIASES.get(
            required,
            frozenset(),
        )
        for column_index, observed in enumerate(normalized_row):
            if observed == required or observed in aliases:
                positions.setdefault(required, []).append(column_index)
    return positions


def _shared_detail_columns_are_equivalent(
    column_indexes: Mapping[str, int],
    columns: list[str],
    expected_rows: list[list[Any]],
) -> bool:
    """Allow one semantic column to serve duplicate, identical model fields."""

    columns_by_index: dict[int, list[int]] = {}
    for column_index, column in enumerate(columns):
        columns_by_index.setdefault(
            column_indexes[column],
            [],
        ).append(column_index)
    return all(
        len(expected_indexes) == 1
        or all(
            all(
                row[index] == row[expected_indexes[0]]
                for index in expected_indexes[1:]
            )
            for row in expected_rows
        )
        for expected_indexes in columns_by_index.values()
    )


def _detail_schedule_evidence(
    path: Path,
    specification: dict[str, Any],
    *,
    workspace_root: Path,
) -> dict[str, Any]:
    columns = specification.get("columns")
    key_columns = specification.get("key_columns")
    expected_rows = specification.get("rows")
    if (
        not isinstance(columns, list)
        or not columns
        or not all(isinstance(column, str) for column in columns)
        or len({_normalize(column) for column in columns}) != len(columns)
        or not isinstance(key_columns, list)
        or not key_columns
        or not all(key in columns for key in key_columns)
        or not isinstance(expected_rows, list)
        or not all(
            isinstance(row, list) and len(row) == len(columns)
            for row in expected_rows
        )
    ):
        return _empty_detail_schedule_evidence(
            "invalid authoritative detail-schedule specification"
        )
    result = _empty_detail_schedule_evidence("detail schedule not found")
    result["expected_row_count"] = len(expected_rows)
    with validated_native_package_copy(
        path,
        workspace_root=workspace_root,
    ) as (snapshot, _inspection):
        recalculated = snapshot
        formula_workbook = load_workbook(
            snapshot, data_only=False, read_only=True
        )
        value_workbook = None
        try:
            recalculated = _recalculate_workbook(
                snapshot,
                snapshot.parent,
            )
            value_workbook = load_workbook(
                recalculated, data_only=True, read_only=True
            )
            result["sheet_present"] = True
            normalized_columns = [_normalize(column) for column in columns]
            required_headers = set(normalized_columns)
            candidates: list[dict[str, Any]] = []
            sheet_rows: dict[str, tuple[list[Any], list[Any]]] = {}
            for candidate_sheet_name in formula_workbook.sheetnames:
                if candidate_sheet_name not in value_workbook.sheetnames:
                    continue
                formula_rows = list(
                    formula_workbook[candidate_sheet_name].iter_rows(
                        values_only=True
                    )
                )
                value_rows = list(
                    value_workbook[candidate_sheet_name].iter_rows(
                        values_only=True
                    )
                )
                sheet_rows[candidate_sheet_name] = (
                    formula_rows,
                    value_rows,
                )
                for header_index, row in enumerate(formula_rows):
                    positions = _semantic_detail_header_positions(
                        row,
                        required_headers,
                    )
                    if not all(
                        len(positions.get(header, [])) == 1
                        for header in required_headers
                    ):
                        continue
                    column_indexes = {
                        column: positions[_normalize(column)][0]
                        for column in columns
                    }
                    if not _shared_detail_columns_are_equivalent(
                        column_indexes,
                        columns,
                        expected_rows,
                    ):
                        continue
                    candidates.append(
                        {
                            "sheet": candidate_sheet_name,
                            "header_index": header_index,
                            "column_indexes": column_indexes,
                        }
                    )
            result["header_candidate_count"] = len(candidates)
            result["header_candidate_samples"] = [
                (
                    f"{candidate['sheet']}!"
                    f"{candidate['header_index'] + 1}"
                )
                for candidate in candidates[:20]
            ]
            if len(candidates) != 1:
                result["error"] = (
                    "required semantic header table not found"
                    if not candidates
                    else "ambiguous duplicate detail-schedule tables"
                )
                return result
            candidate = candidates[0]
            sheet_name = candidate["sheet"]
            header_index = candidate["header_index"]
            column_indexes = candidate["column_indexes"]
            result["selected_location"] = (
                f"{sheet_name}!row {header_index + 1}"
            )
            result["headers_match"] = True
            formula_rows, value_rows = sheet_rows[sheet_name]
            key_indexes = [columns.index(key) for key in key_columns]

            def key_for(row: list[Any] | tuple[Any, ...]) -> tuple[str, ...]:
                return tuple(
                    _detail_schedule_key_part(key, row[index])
                    for key, index in zip(
                        key_columns,
                        key_indexes,
                        strict=True,
                    )
                )

            actual_rows: list[list[Any]] = []
            actual_row_indexes: list[int] = []
            started = False
            for index in range(
                header_index + 1,
                min(len(formula_rows), len(value_rows)),
            ):
                formula_row = formula_rows[index]
                occupied = any(
                    column_index < len(formula_row)
                    and formula_row[column_index] not in (None, "")
                    for column_index in column_indexes.values()
                )
                if not occupied:
                    if started:
                        break
                    continue
                started = True
                value_row = value_rows[index]
                key_values = [
                    (
                        value_row[column_indexes[key]]
                        if column_indexes[key] < len(value_row)
                        else None
                    )
                    for key in key_columns
                ]
                if (
                    _is_detail_total_boundary(key_values[0])
                    and all(
                        value in (None, "") for value in key_values[1:]
                    )
                ):
                    # A trailing native total is a professional table control,
                    # not an extra member of the authoritative detail
                    # population. Stop at that boundary; do not silently skip
                    # arbitrary non-authoritative detail rows.
                    break
                actual_row_indexes.append(index)
                actual_rows.append(
                    [
                        (
                            value_row[column_indexes[column]]
                            if column_indexes[column] < len(value_row)
                            else None
                        )
                        for column in columns
                    ]
                )
            model_formula_cells: list[str] = []
            model_cells = [
                (
                    f"{sheet_name}!"
                    f"{get_column_letter(column_indexes[column] + 1)}"
                    f"{row_index + 1}"
                )
                for row_index in actual_row_indexes
                for column in columns
            ]
            formula_backed_rows = 0
            for row_index in actual_row_indexes:
                row_formula_cells: list[str] = []
                formula_row = formula_rows[row_index]
                for column in columns:
                    column_index = column_indexes[column]
                    formula = (
                        formula_row[column_index]
                        if column_index < len(formula_row)
                        else None
                    )
                    if not _formula_is_calculation(formula):
                        continue
                    row_formula_cells.append(
                        f"{sheet_name}!"
                        f"{get_column_letter(column_index + 1)}"
                        f"{row_index + 1}"
                    )
                if row_formula_cells:
                    formula_backed_rows += 1
                    model_formula_cells.extend(row_formula_cells)
            result["formula_backed_row_count"] = formula_backed_rows
            result["model_cell_count"] = len(model_cells)
            result["model_cells"] = model_cells[:50_000]
            result["model_formula_cell_count"] = len(model_formula_cells)
            result["model_formula_cells"] = model_formula_cells[:5000]
            actual_keys = [key_for(row) for row in actual_rows]
            expected_keys = [key_for(row) for row in expected_rows]
            actual_counter = Counter(actual_keys)
            expected_counter = Counter(expected_keys)
            result["row_count"] = len(actual_rows)
            result["duplicate_key_count"] = sum(
                count - 1 for count in actual_counter.values() if count > 1
            )
            result["missing_keys"] = sorted(
                key
                for key, count in (expected_counter - actual_counter).items()
                for _ in range(count)
            )[:20]
            result["extra_keys"] = sorted(
                key
                for key, count in (actual_counter - expected_counter).items()
                for _ in range(count)
            )[:20]
            result["keys_match"] = actual_counter == expected_counter
            expected_by_key = {
                key_for(row): row for row in expected_rows
            }
            mismatches: list[dict[str, Any]] = []
            mismatch_count = 0
            for relative_row, (actual_key, actual_row) in enumerate(
                zip(actual_keys, actual_rows, strict=True),
                start=1,
            ):
                expected_row = expected_by_key.get(actual_key)
                if expected_row is None:
                    continue
                for column_index, column in enumerate(columns):
                    actual = actual_row[column_index]
                    expected = expected_row[column_index]
                    if _detail_schedule_field_matches(
                        column,
                        actual,
                        expected,
                        row_context=dict(zip(columns, expected_row, strict=True)),
                    ):
                        continue
                    mismatch_count += 1
                    if len(mismatches) < _ACTION_REGISTER_EVIDENCE_LIMIT:
                        mismatches.append(
                            {
                                "row": header_index + relative_row + 1,
                                "sheet": sheet_name,
                                "key": actual_key,
                                "field": column,
                                "actual": str(actual)[:120],
                                "expected": str(expected)[:120],
                            }
                        )
            result["field_mismatch_count"] = mismatch_count
            result["field_mismatch_samples"] = mismatches
            result["field_values_match"] = (
                result["keys_match"]
                and len(actual_rows) == len(expected_rows)
                and result["duplicate_key_count"] == 0
                and mismatch_count == 0
            )
            result["error"] = "detail schedule inspected"
            return result
        finally:
            formula_workbook.close()
            if value_workbook is not None:
                value_workbook.close()
            if recalculated != snapshot:
                shutil.rmtree(
                    recalculated.parents[1],
                    ignore_errors=True,
                )


def _detail_key_matches(
    paired_rows: list[list[Any]],
    specification: dict[str, Any],
) -> tuple[int, int, list[tuple[str, ...]]]:
    columns = specification.get("columns")
    key_columns = specification.get("key_columns")
    expected_rows = specification.get("rows")
    if (
        not isinstance(columns, list)
        or not isinstance(key_columns, list)
        or not isinstance(expected_rows, list)
        or not all(key in columns for key in key_columns)
    ):
        return 0, 0, []
    key_indexes = [columns.index(key) for key in key_columns]
    expected_keys = [
        tuple(_normalize(row[index]) for index in key_indexes)
        for row in expected_rows
        if isinstance(row, list) and len(row) == len(columns)
    ]
    normalized_rows = [
        f" {_normalize(' '.join(str(value) for value in row))} "
        for row in paired_rows
    ]
    missing = []
    matched = 0
    for key in expected_keys:
        if any(
            all(f" {part} " in row for part in key if part)
            for row in normalized_rows
        ):
            matched += 1
        else:
            missing.append(key)
    return matched, len(expected_keys), missing[:20]


def _detail_presentation_evidence(
    paired_rows: list[list[Any]],
    specification: dict[str, Any],
) -> dict[str, Any]:
    """Validate complete native table rows in a presentation."""

    columns = specification.get("columns")
    key_columns = specification.get("key_columns")
    expected_rows = specification.get("rows")
    result = {
        "headers_match": False,
        "row_count": 0,
        "expected_row_count": (
            len(expected_rows) if isinstance(expected_rows, list) else 0
        ),
        "keys_match": False,
        "missing_keys": [],
        "extra_keys": [],
        "field_values_match": False,
        "field_mismatch_count": 0,
        "field_mismatch_samples": [],
    }
    if (
        not isinstance(columns, list)
        or not columns
        or not all(isinstance(column, str) for column in columns)
        or not isinstance(key_columns, list)
        or not all(key in columns for key in key_columns)
        or not isinstance(expected_rows, list)
        or not all(
            isinstance(row, list) and len(row) == len(columns)
            for row in expected_rows
        )
    ):
        return result
    normalized_columns = [_normalize(column) for column in columns]
    candidate_rows: list[list[Any]] = []
    candidate_signatures: set[tuple[str, ...]] = set()
    for row in paired_rows:
        if len(row) != len(columns):
            continue
        normalized_row = [_normalize(value) for value in row]
        if normalized_row == normalized_columns:
            result["headers_match"] = True
            continue
        # Workbook evidence includes both formula-view and recalculated-value
        # views of the same native table. Compare authoritative detail truth
        # only to displayed values, and collapse identical duplicate views.
        if any(
            isinstance(value, str) and value.startswith("=")
            for value in row
        ):
            continue
        signature = tuple(repr(value) for value in row)
        if signature in candidate_signatures:
            continue
        candidate_signatures.add(signature)
        candidate_rows.append(list(row))

    key_indexes = [columns.index(key) for key in key_columns]

    def key_for(row: list[Any]) -> tuple[str, ...]:
        return tuple(_normalize(row[index]) for index in key_indexes)

    expected_by_key = {
        key_for(row): row for row in expected_rows
    }
    actual_by_key = {
        key_for(row): row for row in candidate_rows
        if key_for(row) in expected_by_key
    }
    expected_keys = Counter(key_for(row) for row in expected_rows)
    actual_keys = Counter(
        key_for(row)
        for row in candidate_rows
        if key_for(row) in expected_by_key
    )
    result["row_count"] = sum(actual_keys.values())
    result["missing_keys"] = sorted(
        key
        for key, count in (expected_keys - actual_keys).items()
        for _ in range(count)
    )[:20]
    result["extra_keys"] = sorted(
        key
        for key, count in (actual_keys - expected_keys).items()
        for _ in range(count)
    )[:20]
    result["keys_match"] = actual_keys == expected_keys

    mismatches: list[dict[str, Any]] = []
    mismatch_count = 0
    for key, expected_row in expected_by_key.items():
        actual_row = actual_by_key.get(key)
        if actual_row is None:
            continue
        for column_index, column in enumerate(columns):
            actual = actual_row[column_index]
            expected = expected_row[column_index]
            if isinstance(expected, (int, float)) and not isinstance(
                expected, bool
            ):
                matched = _displayed_claim_matches(
                    str(actual),
                    float(expected),
                )
            else:
                matched = (
                    _normalize(actual) == _normalize(expected)
                )
            if matched:
                continue
            mismatch_count += 1
            if len(mismatches) < _ACTION_REGISTER_EVIDENCE_LIMIT:
                mismatches.append(
                    {
                        "key": key,
                        "field": column,
                        "actual": str(actual)[:120],
                        "expected": str(expected)[:120],
                    }
                )
    result["field_mismatch_count"] = mismatch_count
    result["field_mismatch_samples"] = mismatches
    result["field_values_match"] = (
        result["headers_match"]
        and result["keys_match"]
        and mismatch_count == 0
    )
    return result


def _planning_action_register_evidence(
    path: Path,
    *,
    headers: list[str],
    expected_rows: list[dict[str, Any]],
    workspace_root: Path,
) -> dict[str, Any]:
    """Validate the complete task_028 native register with bounded evidence."""

    if (
        not headers
        or len(set(headers)) != len(headers)
        or "planned_order_id" not in headers
        or "priority" not in headers
        or "required_date" not in headers
    ):
        return _empty_action_register_evidence(
            "authoritative register headers are invalid"
        )
    if any(set(row) != set(headers) for row in expected_rows):
        return _empty_action_register_evidence(
            "authoritative register rows do not match the fixed schema"
        )

    expected_by_id = {
        _action_register_identifier(row["planned_order_id"]): row
        for row in expected_rows
    }
    if len(expected_by_id) != len(expected_rows):
        return _empty_action_register_evidence(
            "authoritative register contains duplicate planned-order IDs"
        )

    header_indexes = {
        header: index for index, header in enumerate(headers)
    }
    sorted_expected_rows = sorted(
        expected_rows,
        key=lambda row: (
            int(row["priority"]),
            _calendar_date(row["required_date"]),
            _action_register_identifier(row["planned_order_id"]),
        ),
    )
    expected_order = [
        _action_register_identifier(row["planned_order_id"])
        for row in sorted_expected_rows
    ]
    normalized_headers = [_normalize(header) for header in headers]
    selected_candidate: dict[str, Any] | None = None
    header_candidate_count = 0
    candidate_location_samples: list[str] = []

    with validated_native_package_copy(
        path,
        workspace_root=workspace_root,
    ) as (snapshot, _inspection):
        workbook = load_workbook(snapshot, data_only=False)
        try:
            for sheet in workbook.worksheets:
                for header_row, row in enumerate(
                    sheet.iter_rows(values_only=True),
                    start=1,
                ):
                    for offset in range(
                        max(0, len(row) - len(headers) + 1)
                    ):
                        candidate_headers = [
                            _normalize(value)
                            for value in row[offset : offset + len(headers)]
                        ]
                        if candidate_headers != normalized_headers:
                            continue
                        header_column = offset + 1
                        location = (
                            f"{sheet.title}!"
                            f"{get_column_letter(header_column)}{header_row}"
                        )
                        header_candidate_count += 1
                        if (
                            len(candidate_location_samples)
                            < _ACTION_REGISTER_EVIDENCE_LIMIT
                        ):
                            candidate_location_samples.append(location)
                        actual_rows: list[list[Any]] = []
                        row_count = 0
                        for data_row in range(
                            header_row + 1,
                            sheet.max_row + 1,
                        ):
                            values = [
                                sheet.cell(data_row, column).value
                                for column in range(
                                    header_column,
                                    header_column + len(headers),
                                )
                            ]
                            if all(value in (None, "") for value in values):
                                break
                            row_count += 1
                            if len(actual_rows) <= len(expected_rows):
                                actual_rows.append(values)

                        actual_ids = [
                            _action_register_identifier(
                                row_values[
                                    header_indexes["planned_order_id"]
                                ]
                            )
                            for row_values in actual_rows
                        ]
                        id_counts = Counter(actual_ids)
                        duplicate_ids = sorted(
                            planned_order_id
                            for planned_order_id, count in id_counts.items()
                            if count > 1
                        )
                        actual_id_set = set(actual_ids)
                        expected_id_set = set(expected_by_id)
                        missing_ids = sorted(
                            expected_id_set - actual_id_set
                        )
                        extra_ids = sorted(
                            actual_id_set - expected_id_set
                        )
                        mismatch_count = 0
                        matched_field_count = 0
                        mismatch_samples: list[dict[str, Any]] = []
                        for relative_row, row_values in enumerate(
                            actual_rows,
                            start=1,
                        ):
                            actual_id = actual_ids[relative_row - 1]
                            expected_row = expected_by_id.get(actual_id)
                            if expected_row is None:
                                continue
                            for field, column_index in header_indexes.items():
                                actual = row_values[column_index]
                                expected = expected_row[field]
                                if _action_register_field_matches(
                                    field,
                                    actual,
                                    expected,
                                ):
                                    matched_field_count += 1
                                    continue
                                mismatch_count += 1
                                if (
                                    len(mismatch_samples)
                                    < _ACTION_REGISTER_EVIDENCE_LIMIT
                                ):
                                    mismatch_samples.append(
                                        {
                                            "row": header_row + relative_row,
                                            "planned_order_id": actual_id,
                                            "field": field,
                                            "actual": str(actual)[:120],
                                            "expected": str(expected)[:120],
                                        }
                                    )

                        actual_sort_keys = [
                            _action_register_sort_key(
                                row_values,
                                header_indexes,
                            )
                            for row_values in actual_rows
                        ]
                        order_mismatches = [
                            {
                                "position": index + 1,
                                "actual": actual_id,
                                "expected": expected_order[index],
                            }
                            for index, actual_id in enumerate(
                                actual_ids[: len(expected_order)]
                            )
                            if actual_id != expected_order[index]
                        ][:_ACTION_REGISTER_EVIDENCE_LIMIT]
                        row_count_matches = row_count == len(expected_rows)
                        unique_ids_match = (
                            len(actual_id_set) == len(expected_id_set)
                            and not duplicate_ids
                            and actual_id_set == expected_id_set
                        )
                        field_values_match = (
                            row_count_matches
                            and unique_ids_match
                            and mismatch_count == 0
                        )
                        sort_order_match = (
                            row_count_matches
                            and len(actual_rows) == len(expected_rows)
                            and all(
                                sort_key is not None
                                for sort_key in actual_sort_keys
                            )
                            and actual_sort_keys
                            == sorted(actual_sort_keys)
                            and actual_ids == expected_order
                        )
                        filter_refs = _action_register_filter_refs(sheet)
                        filterable = any(
                            _action_register_filter_covers(
                                reference,
                                header_row=header_row,
                                header_column=header_column,
                                column_count=len(headers),
                                data_row_count=row_count,
                            )
                            for reference in filter_refs
                        )
                        candidate = {
                            "error": "",
                            "selected_location": location,
                            "filterable": filterable,
                            "filter_refs": filter_refs[
                                :_ACTION_REGISTER_EVIDENCE_LIMIT
                            ],
                            "row_count": row_count,
                            "expected_row_count": len(expected_rows),
                            "unique_id_count": len(actual_id_set),
                            "expected_unique_id_count": len(
                                expected_id_set
                            ),
                            "duplicate_id_count": len(duplicate_ids),
                            "duplicate_id_samples": duplicate_ids[
                                :_ACTION_REGISTER_EVIDENCE_LIMIT
                            ],
                            "missing_id_count": len(missing_ids),
                            "missing_id_samples": missing_ids[
                                :_ACTION_REGISTER_EVIDENCE_LIMIT
                            ],
                            "extra_id_count": len(extra_ids),
                            "extra_id_samples": extra_ids[
                                :_ACTION_REGISTER_EVIDENCE_LIMIT
                            ],
                            "field_mismatch_count": mismatch_count,
                            "field_mismatch_samples": mismatch_samples,
                            "first_order_mismatch_samples": (
                                order_mismatches
                            ),
                            "field_values_match": field_values_match,
                            "sort_order_match": sort_order_match,
                            "_matched_field_count": matched_field_count,
                        }
                        candidate_score = (
                            bool(candidate["field_values_match"]),
                            int(candidate["_matched_field_count"]),
                            int(candidate["unique_id_count"]),
                            bool(candidate["sort_order_match"]),
                            bool(candidate["filterable"]),
                        )
                        selected_score = (
                            (
                                bool(
                                    selected_candidate[
                                        "field_values_match"
                                    ]
                                ),
                                int(
                                    selected_candidate[
                                        "_matched_field_count"
                                    ]
                                ),
                                int(
                                    selected_candidate[
                                        "unique_id_count"
                                    ]
                                ),
                                bool(
                                    selected_candidate[
                                        "sort_order_match"
                                    ]
                                ),
                                bool(selected_candidate["filterable"]),
                            )
                            if selected_candidate is not None
                            else None
                        )
                        if (
                            selected_score is None
                            or candidate_score > selected_score
                        ):
                            selected_candidate = candidate
        finally:
            workbook.close()

    if selected_candidate is None:
        result = _empty_action_register_evidence(
            "no worksheet contains the semantic register header sequence"
        )
        result["expected_row_count"] = len(expected_rows)
        result["expected_unique_id_count"] = len(expected_by_id)
        return result
    selected = selected_candidate
    selected.pop("_matched_field_count", None)
    selected["header_candidate_count"] = header_candidate_count
    selected["header_candidate_samples"] = candidate_location_samples
    return selected


def _action_register_evidence_summary(evidence: dict[str, Any]) -> str:
    """Render stable, bounded criterion evidence without serializing rows."""

    summary = {
        key: evidence.get(key)
        for key in (
            "error",
            "header_candidate_count",
            "header_candidate_samples",
            "selected_location",
            "filterable",
            "filter_refs",
            "row_count",
            "expected_row_count",
            "unique_id_count",
            "expected_unique_id_count",
            "duplicate_id_count",
            "duplicate_id_samples",
            "missing_id_count",
            "missing_id_samples",
            "extra_id_count",
            "extra_id_samples",
            "field_mismatch_count",
            "field_mismatch_samples",
            "first_order_mismatch_samples",
        )
    }
    return json.dumps(summary, sort_keys=True, separators=(",", ":"))


def _workbook_evidence(
    path: Path,
    workspace_root: Path | None = None,
) -> dict[str, Any]:
    with validated_native_package_copy(
        path,
        workspace_root=workspace_root or path.parent,
    ) as (snapshot, _inspection):
        return _workbook_evidence_from_snapshot(snapshot)


def _workbook_evidence_from_snapshot(path: Path) -> dict[str, Any]:
    inspection_root = path.parent
    inspect_native_package(
        path,
        workspace_root=inspection_root,
    )
    formula_book = load_workbook(path, data_only=False)
    defined_name_references = _defined_name_references(formula_book)
    recalculated = path
    try:
        recalculated = _recalculate_workbook(path, inspection_root)
        inspect_native_package(
            recalculated,
            workspace_root=(
                inspection_root
                if recalculated == path
                else recalculated.parents[1]
            ),
        )
        value_book = load_workbook(recalculated, data_only=True)
    except Exception:
        formula_book.close()
        if recalculated != path:
            shutil.rmtree(recalculated.parents[1], ignore_errors=True)
        raise
    labels: list[str] = []
    values: list[Any] = []
    formulas: list[str] = []
    formula_cells: dict[str, str] = {}
    formula_formats: dict[str, str] = {}
    formula_row_labels: dict[str, str] = {}
    literal_cells = {
        f"{sheet.title}!{cell.coordinate}": cell.value
        for sheet in formula_book.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if cell.value is not None
        and not (
            isinstance(cell.value, str)
            and cell.value.startswith("=")
        )
    }
    value_cells = {
        f"{sheet.title}!{cell.coordinate}": cell.value
        for sheet in value_book.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if cell.value is not None
    }
    labeled_formula_results: list[dict[str, Any]] = []
    paired_rows: list[list[Any]] = []
    errors: list[str] = []
    visual_hygiene_issues: list[str] = []
    native_formula_issues: list[str] = []
    native_chart_issues: list[str] = []
    decision_blocks: list[str] = []
    hardcoded_control_status_cells: list[str] = []
    try:
        for sheet in formula_book.worksheets:
            page_setup_properties = sheet.sheet_properties.pageSetUpPr
            fit_to_page = (
                page_setup_properties.fitToPage
                if page_setup_properties is not None
                else None
            )
            if (
                fit_to_page is False
                and (
                    sheet.page_setup.fitToWidth is not None
                    or sheet.page_setup.fitToHeight is not None
                )
            ):
                visual_hygiene_issues.append(
                    f"{sheet.title}: fit-to-page dimensions are set while "
                    "fitToPage is disabled, so native print/PDF output can "
                    "spill across pages or crop content"
                )
            header_footer_fields = (
                ("odd header left", sheet.oddHeader.left.text),
                ("odd header center", sheet.oddHeader.center.text),
                ("odd header right", sheet.oddHeader.right.text),
                ("even header left", sheet.evenHeader.left.text),
                ("even header center", sheet.evenHeader.center.text),
                ("even header right", sheet.evenHeader.right.text),
                ("first header left", sheet.firstHeader.left.text),
                ("first header center", sheet.firstHeader.center.text),
                ("first header right", sheet.firstHeader.right.text),
                ("odd footer left", sheet.oddFooter.left.text),
                ("odd footer center", sheet.oddFooter.center.text),
                ("odd footer right", sheet.oddFooter.right.text),
                ("even footer left", sheet.evenFooter.left.text),
                ("even footer center", sheet.evenFooter.center.text),
                ("even footer right", sheet.evenFooter.right.text),
                ("first footer left", sheet.firstFooter.left.text),
                ("first footer center", sheet.firstFooter.center.text),
                ("first footer right", sheet.firstFooter.right.text),
            )
            for location, text in header_footer_fields:
                # Excel interprets ``&P`` as the current-page field. A literal
                # brand string accidentally written as ``&Pinehaven`` renders
                # as ``1inehaven``, ``2inehaven``, and so on.
                if text and "&pinehaven" in text.casefold():
                    visual_hygiene_issues.append(
                        f"{sheet.title}: malformed {location} page-code text"
                    )
            for chart_index, chart in enumerate(sheet._charts, start=1):
                for series_index, series in enumerate(chart.ser, start=1):
                    value_source = (
                        getattr(series, "val", None)
                        or getattr(series, "yVal", None)
                    )
                    category_source = getattr(series, "cat", None)
                    value_count = _chart_data_count(value_source)
                    category_count = _chart_data_count(category_source)
                    if (
                        value_count is not None
                        and category_count is not None
                        and value_count != category_count
                    ):
                        native_chart_issues.append(
                            f"{sheet.title}: chart {chart_index} series "
                            f"{series_index} has {value_count} values for "
                            f"{category_count} categories"
                        )
                    data_labels = getattr(series, "dLbls", None)
                    if data_labels is None:
                        continue
                    enabled = [
                        name
                        for name in (
                            "showLegendKey",
                            "showVal",
                            "showCatName",
                            "showSerName",
                            "showPercent",
                            "showBubbleSize",
                        )
                        if getattr(data_labels, name, None)
                    ]
                    if len(enabled) >= 3:
                        visual_hygiene_issues.append(
                            f"{sheet.title}: chart {chart_index} series "
                            f"{series_index} enables colliding labels "
                            f"{enabled!r}"
                        )
                numeric_values = [
                    value
                    for series in chart.ser
                    for value in _chart_numeric_values(series)
                ]
                if _mixed_sign_zero_axis_label_collision(
                    chart,
                    numeric_values,
                ):
                    visual_hygiene_issues.append(
                        f"{sheet.title}: chart {chart_index} uses mixed-sign "
                        "horizontal bars with category labels next to the "
                        "zero-crossing axis, which collide inside the plot"
                    )
                numeric_axis = next(
                    (
                        axis
                        for axis in (
                            getattr(chart, "x_axis", None),
                            getattr(chart, "y_axis", None),
                        )
                        if getattr(axis, "numFmt", None) is not None
                        and getattr(axis.numFmt, "formatCode", "General")
                        != "General"
                    ),
                    None,
                )
                axis_format = (
                    getattr(numeric_axis.numFmt, "formatCode", "")
                    if numeric_axis is not None
                    else ""
                )
                if (
                    numeric_values
                    and max(abs(value) for value in numeric_values) >= 1_000_000
                    and re.search(r"[0#]\.00", axis_format)
                    and ",," not in axis_format
                    and float(getattr(chart, "width", 0) or 0) <= 10.0
                ):
                    visual_hygiene_issues.append(
                        f"{sheet.title}: chart {chart_index} uses unscaled "
                        "million-dollar axis labels with fixed cents, which "
                        "collide at native chart width"
                    )
            native_formula_issues.extend(
                _cached_formula_issues(
                    sheet,
                    value_book,
                    defined_name_references,
                )
            )
            errors.extend(_cross_engine_formula_issues(sheet, value_book))
            value_sheet = value_book[sheet.title]
            remaining_visual_issue_slots = max(
                0,
                _WORKBOOK_MAX_CELL_VISUAL_ISSUES
                - sum(
                    "is likely clipped" in issue
                    for issue in visual_hygiene_issues
                ),
            )
            visual_hygiene_issues.extend(
                _workbook_cell_visual_issues(
                    sheet,
                    value_sheet,
                    limit=remaining_visual_issue_slots,
                )
            )
            last_section_context = ""
            last_column_headers: dict[int, str] = {}
            for row in sheet.iter_rows():
                row_values = [cell.value for cell in row if cell.value is not None]
                if row_values:
                    paired_rows.append(row_values)
                for cell in row:
                    value = cell.value
                    if isinstance(value, str):
                        labels.append(value)
                        if value.startswith("="):
                            formulas.append(value)
                            qualified_cell = (
                                f"{sheet.title}!{cell.coordinate}"
                            )
                            formula_cells[qualified_cell] = (
                                _expand_formula_defined_names(
                                    value,
                                    defined_name_references,
                                )
                            )
                            formula_formats[qualified_cell] = str(
                                cell.number_format or "General"
                            )
                            row_label = _native_formula_label(
                                sheet,
                                cell,
                                row,
                            )
                            if row_label:
                                formula_row_labels[qualified_cell] = row_label
                            label_value = row_label
                            if label_value:
                                column_header = last_column_headers.get(
                                    cell.column,
                                    "",
                                )
                                section_context = last_section_context
                                labeled_formula_results.append(
                                        {
                                            "cell": qualified_cell,
                                            "label": label_value,
                                            "column_header": column_header,
                                            "section_context": section_context,
                                            "formula": value,
                                            "value": (
                                                value_sheet[
                                                    cell.coordinate
                                                ].value
                                                if value_sheet[
                                                    cell.coordinate
                                                ].value
                                                is not None
                                                else _safe_normalizing_formula_value(
                                                    qualified_cell,
                                                    value,
                                                    literal_cells,
                                                )
                                            ),
                                        }
                                    )
                            try:
                                formula_tokens = Tokenizer(value).items
                            except Exception:
                                formula_tokens = ()
                            if any(
                                token.type == "OPERAND"
                                and token.subtype == "ERROR"
                                and token.value.upper() in FORMULA_ERRORS
                                for token in formula_tokens
                            ):
                                errors.append(value)
                    elif value is not None:
                        values.append(value)
                populated = [
                    cell.value
                    for cell in row
                    if cell.value not in (None, "")
                ]
                if (
                    len(populated) == 1
                    and isinstance(populated[0], str)
                    and not populated[0].startswith("=")
                ):
                    last_section_context = populated[0]
                if len(populated) >= 2:
                    for header_cell in row:
                        candidate = header_cell.value
                        if (
                            isinstance(candidate, str)
                            and not candidate.startswith("=")
                            and candidate.strip()
                        ):
                            last_column_headers[header_cell.column] = candidate
            hardcoded_control_status_cells.extend(
                _hardcoded_control_status_cells(sheet)
            )
        for sheet in value_book.worksheets:
            for row in sheet.iter_rows():
                row_values = [cell.value for cell in row if cell.value is not None]
                if row_values:
                    paired_rows.append(row_values)
                    decision_block = _labeled_decision_block(row_values)
                    if (
                        decision_block
                        and decision_block not in decision_blocks
                    ):
                        decision_blocks.append(decision_block)
                for cell in row:
                    if cell.value is not None:
                        values.append(cell.value)
                        if str(cell.value).upper() in FORMULA_ERRORS:
                            errors.append(str(cell.value))
        chart_count = sum(len(sheet._charts) for sheet in formula_book.worksheets)
        merged = sum(len(sheet.merged_cells.ranges) for sheet in formula_book.worksheets)
        return {
            "sheetnames": formula_book.sheetnames,
            "text": "\n".join(labels),
            "values": values,
            "formulas": formulas,
            "formula_cells": formula_cells,
            "formula_formats": formula_formats,
            "formula_row_labels": formula_row_labels,
            "literal_cells": literal_cells,
            "value_cells": value_cells,
            "labeled_formula_results": labeled_formula_results,
            "paired_rows": paired_rows,
            "errors": errors,
            "charts": chart_count,
            "merged_ranges": merged,
            "visual_hygiene_issues": visual_hygiene_issues,
            "native_formula_issues": native_formula_issues,
            "native_chart_issues": native_chart_issues,
            "decision_blocks": decision_blocks,
            "hardcoded_control_status_cells": (
                hardcoded_control_status_cells
            ),
        }
    finally:
        formula_book.close()
        value_book.close()
        if recalculated != path:
            shutil.rmtree(recalculated.parents[1], ignore_errors=True)


def _formula_reference(
    current_sheet: str,
    token: str,
) -> str | None:
    references = _formula_references(current_sheet, token)
    return references[0] if len(references) == 1 else None


def _formula_references(
    current_sheet: str,
    token: str,
    *,
    max_cells: int = 10_000,
) -> tuple[str, ...]:
    """Resolve one A1 token into bounded, fully qualified cell references."""

    value = token.replace("$", "")
    if "," in value:
        return ()
    if "!" in value:
        sheet, address = value.rsplit("!", 1)
        sheet = sheet.strip("'").replace("''", "'")
    else:
        sheet, address = current_sheet, value
    address = address.strip()
    if not re.fullmatch(
        r"[A-Za-z]{1,3}[1-9]\d*(?::[A-Za-z]{1,3}[1-9]\d*)?",
        address,
    ):
        return ()
    if ":" not in address:
        return (f"{sheet}!{address.upper()}",)
    minimum_column, minimum_row, maximum_column, maximum_row = (
        range_boundaries(address)
    )
    cell_count = (
        (maximum_column - minimum_column + 1)
        * (maximum_row - minimum_row + 1)
    )
    if cell_count > max_cells:
        return ()
    return tuple(
        f"{sheet}!{get_column_letter(column)}{row}"
        for row in range(minimum_row, maximum_row + 1)
        for column in range(minimum_column, maximum_column + 1)
    )


def _safe_normalizing_formula_value(
    cell: str,
    formula: str,
    literal_cells: dict[str, Any],
) -> Any:
    """Evaluate only bounded one-input normalization wrappers.

    Native recalculation remains authoritative. This fail-closed fallback lets
    reference fixtures and development hosts without the LibreOffice sandbox
    validate transparent ROUND/SUM, boolean, and text-normalization formulas.
    It intentionally does not interpret general spreadsheet expressions.
    """

    current_sheet = cell.rsplit("!", 1)[0]
    try:
        tokens = Tokenizer(formula).items
    except Exception:
        return None
    references = [
        _formula_reference(current_sheet, token.value)
        for token in tokens
        if token.type == "OPERAND" and token.subtype == "RANGE"
    ]
    if len(references) != 1 or references[0] not in literal_cells:
        return None
    source_value = literal_cells[references[0]]
    normalized = re.sub(r"\s+", "", formula).upper()
    escaped_reference = re.escape(
        next(
            token.value.replace("$", "").upper()
            for token in tokens
            if token.type == "OPERAND" and token.subtype == "RANGE"
        )
    )
    round_match = re.fullmatch(
        rf"=ROUND\(SUM\({escaped_reference}\),(\d+)\)",
        normalized,
    )
    if round_match:
        numeric = _number(source_value)
        return (
            round(numeric, int(round_match.group(1)))
            if numeric is not None
            else None
        )
    if re.fullmatch(
        rf"=IF\({escaped_reference},TRUE,FALSE\)",
        normalized,
    ):
        return bool(source_value)
    if re.fullmatch(rf'={escaped_reference}&""', normalized):
        return str(source_value)
    return None


def _formula_is_calculation(formula: Any) -> bool:
    """Return whether a formula calculates rather than merely aliases a cell."""

    if not isinstance(formula, str) or not formula.startswith("="):
        return False
    try:
        tokens = Tokenizer(formula).items
    except Exception:
        return False
    references = [
        token
        for token in tokens
        if token.type == "OPERAND" and token.subtype == "RANGE"
    ]
    if not references:
        return False
    return any(
        token.type in {"FUNC", "OPERATOR-INFIX", "OPERATOR-PREFIX"}
        for token in tokens
    )


def _formula_is_direct_reference(formula: Any) -> bool:
    """Return whether a formula is one transparent native cell reference."""

    if not isinstance(formula, str) or not formula.startswith("="):
        return False
    try:
        tokens = Tokenizer(formula).items
    except Exception:
        return False
    return (
        len(tokens) == 1
        and tokens[0].type == "OPERAND"
        and tokens[0].subtype == "RANGE"
        and bool(_formula_references("Analysis", tokens[0].value))
    )


def _direct_formula_reference(
    cell: str,
    formula: Any,
) -> str | None:
    """Resolve a transparent one-cell formula using its actual sheet."""

    if not isinstance(formula, str) or not formula.startswith("="):
        return None
    try:
        tokens = [
            token
            for token in Tokenizer(formula).items
            if token.type != "WHITE-SPACE"
        ]
    except Exception:
        return None
    if (
        len(tokens) != 1
        or tokens[0].type != "OPERAND"
        or tokens[0].subtype != "RANGE"
    ):
        return None
    references = _formula_references(
        cell.rsplit("!", 1)[0],
        tokens[0].value,
    )
    return references[0] if len(references) == 1 else None


def _formula_is_recursive_calculation(
    cell: str,
    formula: Any,
    formula_cells: Mapping[str, str],
) -> bool:
    """Follow transparent aliases until a native calculation is reached."""

    current_cell = cell
    current_formula = formula
    seen: set[str] = set()
    for _ in range(len(formula_cells) + 1):
        if current_cell in seen:
            return False
        seen.add(current_cell)
        if _formula_is_calculation(current_formula):
            return True
        reference = _direct_formula_reference(
            current_cell,
            current_formula,
        )
        if reference is None:
            return False
        upstream = formula_cells.get(reference)
        if upstream is None:
            return False
        current_cell = reference
        current_formula = upstream
    return False


def _source_lookup_key_literals(
    current_cell: str,
    formula: str,
) -> set[str]:
    """Allow a same-row literal lookup key only for an external source table.

    A control such as ``VLOOKUP(A7, Inputs!A:B, 2, 0)`` is source-backed: A7
    chooses the metric and the value comes from Inputs.  Treating the local key
    literal as an invalid analysis/control dependency rejects that legitimate
    lineage.  Keep the exception fail-closed to one-cell VLOOKUP keys and a
    table range on a non-Analysis/non-Control source sheet.
    """

    compact = re.sub(r"\s+", "", formula)
    match = re.fullmatch(
        r"=VLOOKUP\("
        r"(\$?[A-Za-z]{1,3}\$?[1-9]\d*),"
        r"((?:'(?:[^']|'')+'|[^,!]+)!"
        r"\$?[A-Za-z]{1,3}\$?[1-9]\d*:"
        r"\$?[A-Za-z]{1,3}\$?[1-9]\d*),"
        r"[1-9]\d*,(?:0|FALSE)\)",
        compact,
        flags=re.I,
    )
    if match is None:
        return set()
    current_sheet = current_cell.rsplit("!", 1)[0]
    table_references = _formula_references(current_sheet, match.group(2))
    if not table_references:
        return set()
    table_sheets = {
        _normalize(reference.rsplit("!", 1)[0])
        for reference in table_references
    }
    if (
        len(table_sheets) != 1
        or table_sheets
        & {
            "analysis",
            "control",
            "controls",
            "check",
            "checks",
            "read me",
        }
    ):
        return set()
    key = _formula_reference(current_sheet, match.group(1))
    return {key} if key is not None else set()


def _formula_reaches_source(
    cell: str,
    formula: str,
    formula_cells: dict[str, str],
    *,
    seen: set[str] | None = None,
) -> bool:
    """Trace a calculated claim to Inputs or another native source sheet."""

    if seen is None:
        seen = set()
    if cell in seen:
        return False
    reaches_source, invalid = _formula_lineage_state(
        cell,
        formula,
        formula_cells,
        seen=seen,
    )
    return reaches_source and not invalid


def _formula_reaches_targets(
    cell: str,
    formula: str,
    formula_cells: dict[str, str],
    targets: set[str],
    *,
    seen: set[str] | None = None,
) -> bool:
    """Trace a formula to any semantic model cell without sheet assumptions."""

    if seen is None:
        seen = set()
    if cell in seen:
        return False
    inherited_path = set(seen)
    pending: list[tuple[str, str]] = [(cell, formula)]
    visited: set[str] = set()
    edge_count = 0
    while pending:
        current_cell, current_formula = pending.pop()
        if current_cell in visited or current_cell in inherited_path:
            continue
        visited.add(current_cell)
        if len(visited) > len(formula_cells) + 1:
            return False
        current_sheet = current_cell.rsplit("!", 1)[0]
        try:
            tokens = Tokenizer(current_formula).items
        except Exception:
            return False
        for token in tokens:
            if token.type != "OPERAND" or token.subtype != "RANGE":
                continue
            for reference in _formula_references(
                current_sheet,
                token.value,
            ):
                edge_count += 1
                if edge_count > 1_000_000:
                    return False
                if reference in targets:
                    return True
                upstream = formula_cells.get(reference)
                if (
                    upstream is not None
                    and reference not in visited
                    and reference not in inherited_path
                ):
                    pending.append((reference, upstream))
    return False


def _formula_lineage_state(
    cell: str,
    formula: str,
    formula_cells: dict[str, str],
    *,
    seen: set[str],
) -> tuple[bool, bool]:
    """Return ``(reaches_source, invalid_path)`` for a bounded formula graph.

    Rolling schedules are directed acyclic graphs with many shared upstream
    cells. Recursive path-by-path traversal revisits those shared branches
    exponentially and can exhaust the Python stack before it reaches a
    source. Build the reachable graph once instead, then use a bounded
    topological pass to distinguish a valid shared dependency from a genuine
    cycle. Any cycle, malformed formula, internal literal, or traversal-budget
    overflow fails closed.
    """

    inherited_path = set(seen)
    if cell in inherited_path:
        return False, True

    # Formula cells are the complete finite graph supplied by workbook
    # evidence. Large native aggregations can legitimately expand the same
    # bounded source ranges across hundreds of formulas. Scale the traversal
    # budget with that finite graph, while retaining a hard ceiling for
    # hostile formulas containing many overlapping ranges.
    maximum_nodes = len(formula_cells) + 1
    maximum_edges = min(
        5_000_000,
        max(2_000_000, maximum_nodes * 256),
    )
    pending: list[tuple[str, str]] = [(cell, formula)]
    visited: set[str] = set()
    dependencies: dict[str, set[str]] = {}
    reaches_source = False
    invalid = False
    edge_count = 0

    while pending:
        current_cell, current_formula = pending.pop()
        if current_cell in visited:
            continue
        visited.add(current_cell)
        dependencies.setdefault(current_cell, set())
        if len(visited) > maximum_nodes:
            invalid = True
            break

        current_sheet = current_cell.rsplit("!", 1)[0]
        allowed_lookup_keys = _source_lookup_key_literals(
            current_cell,
            current_formula,
        )
        try:
            tokens = Tokenizer(current_formula).items
        except Exception:
            invalid = True
            continue
        for token in tokens:
            if token.type != "OPERAND" or token.subtype != "RANGE":
                continue
            references = _formula_references(current_sheet, token.value)
            if not references:
                continue
            for reference in references:
                edge_count += 1
                if edge_count > maximum_edges:
                    invalid = True
                    pending.clear()
                    break
                upstream_formula = formula_cells.get(reference)
                if upstream_formula is not None:
                    dependencies[current_cell].add(reference)
                    dependencies.setdefault(reference, set())
                    if reference in inherited_path:
                        invalid = True
                    elif reference not in visited:
                        pending.append((reference, upstream_formula))
                    continue

                referenced_sheet = _normalize(
                    reference.rsplit("!", 1)[0]
                )
                if reference in allowed_lookup_keys:
                    continue
                if referenced_sheet in {
                    "analysis",
                    "control",
                    "controls",
                    "check",
                    "checks",
                }:
                    invalid = True
                elif referenced_sheet != "read me":
                    reaches_source = True

    # Kahn's algorithm is iterative and detects every reachable directed
    # cycle, including a self-reference, without recursion-depth risk.
    indegree = {node: 0 for node in dependencies}
    for targets in dependencies.values():
        for target in targets:
            indegree[target] = indegree.get(target, 0) + 1
    ready = [node for node, degree in indegree.items() if degree == 0]
    processed = 0
    while ready:
        node = ready.pop()
        processed += 1
        for target in dependencies.get(node, set()):
            indegree[target] -= 1
            if indegree[target] == 0:
                ready.append(target)
    if processed != len(indegree):
        invalid = True

    return reaches_source, invalid


def _lineage_reference_state(
    reference: str,
    formula_cells: dict[str, str],
    *,
    seen: set[str],
) -> tuple[bool, bool]:
    """Trace a reference while treating cycles/internal literals as invalid."""

    if reference in seen:
        return False, True
    upstream_formula = formula_cells.get(reference)
    if upstream_formula is not None:
        return _formula_lineage_state(
            reference,
            upstream_formula,
            formula_cells,
            seen=seen,
        )
    referenced_sheet = _normalize(reference.rsplit("!", 1)[0])
    if referenced_sheet in {"analysis", "control"}:
        return False, True
    if referenced_sheet == "read me":
        return False, False
    return True, False


def _formula_has_unrounded_subtraction(formula: str) -> bool:
    """Detect subtraction left outside an exact ``ROUND(..., 2)`` envelope."""

    compact = re.sub(r"\s+", "", formula).upper()
    rounded_reference = (
        r"ROUND\("
        r"(?:(?:'[^']+'|[A-Z0-9_ ]+)!)?"
        r"\$?[A-Z]{1,3}\$?[1-9]\d*,2\)"
    )
    if re.fullmatch(
        rf"={rounded_reference}-{rounded_reference}",
        compact,
    ):
        return False

    upper = formula.upper()
    stripped: list[str] = []
    index = 0
    while index < len(formula):
        is_round = upper.startswith("ROUND(", index) and (
            index == 0
            or not (upper[index - 1].isalnum() or upper[index - 1] == "_")
        )
        if not is_round:
            stripped.append(formula[index])
            index += 1
            continue

        argument_start = index + len("ROUND(")
        depth = 1
        comma_positions: list[int] = []
        cursor = argument_start
        in_string = False
        while cursor < len(formula) and depth:
            character = formula[cursor]
            if character == '"':
                if (
                    in_string
                    and cursor + 1 < len(formula)
                    and formula[cursor + 1] == '"'
                ):
                    cursor += 2
                    continue
                in_string = not in_string
            elif not in_string:
                if character == "(":
                    depth += 1
                elif character == ")":
                    depth -= 1
                elif character == "," and depth == 1:
                    comma_positions.append(cursor)
            cursor += 1
        if (
            depth == 0
            and len(comma_positions) == 1
            and formula[comma_positions[0] + 1 : cursor - 1].strip() == "2"
        ):
            stripped.append("0")
            index = cursor
            continue
        stripped.append(formula[index])
        index += 1

    try:
        tokens = Tokenizer("".join(stripped)).items
    except Exception:
        return False
    return any(
        token.type == "OPERATOR-INFIX" and token.value == "-"
        for token in tokens
    )


def _formula_is_top_level_round_to_cents(formula: str) -> bool:
    """Return whether the complete formula result is ``ROUND(..., 2)``."""

    try:
        tokens = [
            token
            for token in Tokenizer(formula).items
            if token.type != "WHITE-SPACE"
        ]
    except Exception:
        return False
    if (
        len(tokens) < 4
        or tokens[0].type != "FUNC"
        or tokens[0].subtype != "OPEN"
        or tokens[0].value.upper() != "ROUND("
        or tokens[-1].type != "FUNC"
        or tokens[-1].subtype != "CLOSE"
    ):
        return False

    depth = 0
    argument_separators: list[int] = []
    for index, token in enumerate(tokens):
        if token.type in {"ARRAY", "FUNC", "PAREN"}:
            if token.subtype == "OPEN":
                depth += 1
                continue
            if token.subtype == "CLOSE":
                depth -= 1
                if depth < 0 or (depth == 0 and index != len(tokens) - 1):
                    return False
                continue
        if (
            depth == 1
            and token.type == "SEP"
            and token.subtype == "ARG"
        ):
            argument_separators.append(index)

    if depth != 0 or len(argument_separators) != 1:
        return False
    separator = argument_separators[0]
    decimal_tokens = tokens[separator + 1 : -1]
    return (
        separator > 1
        and len(decimal_tokens) == 1
        and decimal_tokens[0].type == "OPERAND"
        and decimal_tokens[0].subtype == "NUMBER"
        and decimal_tokens[0].value == "2"
    )


def _formula_is_cent_rounded_difference(formula: str) -> bool:
    """Return whether both sides of the top-level difference round to cents."""

    try:
        tokens = [
            token
            for token in Tokenizer(formula).items
            if token.type != "WHITE-SPACE"
        ]
    except Exception:
        return False
    depth = 0
    subtraction_indexes: list[int] = []
    for index, token in enumerate(tokens):
        if token.type in {"ARRAY", "FUNC", "PAREN"}:
            if token.subtype == "OPEN":
                depth += 1
                continue
            if token.subtype == "CLOSE":
                depth -= 1
                if depth < 0:
                    return False
                continue
        if (
            depth == 0
            and token.type == "OPERATOR-INFIX"
            and token.value == "-"
        ):
            subtraction_indexes.append(index)
    if depth != 0 or len(subtraction_indexes) != 1:
        return False
    subtraction = subtraction_indexes[0]
    if subtraction == 0 or subtraction == len(tokens) - 1:
        return False
    left = "=" + "".join(token.value for token in tokens[:subtraction])
    right = "=" + "".join(
        token.value for token in tokens[subtraction + 1 :]
    )
    return _formula_is_top_level_round_to_cents(
        left
    ) and _formula_is_top_level_round_to_cents(right)


def _unrounded_currency_control_issues(
    formula_cells: dict[str, str],
    formula_formats: dict[str, str] | None = None,
    literal_cells: dict[str, Any] | None = None,
    *,
    only_cells: set[str] | None = None,
) -> tuple[list[str], list[str]]:
    """Find zero/tolerance status tests fed by unrounded subtraction formulas."""

    checked: list[str] = []
    issues: list[str] = []

    def selected_if_branch(key: str, formula: str) -> str | None:
        """Select a literal-keyed IF branch without evaluating arithmetic."""

        if literal_cells is None:
            return None
        compact = re.sub(r"\s+", "", formula)
        if not compact.upper().startswith("=IF(") or not compact.endswith(")"):
            return None
        body = compact[4:-1]
        arguments: list[str] = []
        start = 0
        depth = 0
        in_string = False
        index = 0
        while index < len(body):
            character = body[index]
            if character == '"':
                if in_string and index + 1 < len(body) and body[index + 1] == '"':
                    index += 2
                    continue
                in_string = not in_string
            elif not in_string:
                if character == "(":
                    depth += 1
                elif character == ")":
                    depth -= 1
                    if depth < 0:
                        return None
                elif character == "," and depth == 0:
                    arguments.append(body[start:index])
                    start = index + 1
            index += 1
        arguments.append(body[start:])
        if in_string or depth != 0 or len(arguments) != 3:
            return None
        condition = re.fullmatch(
            r'(\$?[A-Za-z]{1,3}\$?[1-9]\d*)="([^"]*)"',
            arguments[0],
            flags=re.I,
        )
        if condition is None:
            return None
        sheet = key.rsplit("!", 1)[0]
        reference = _formula_reference(sheet, condition.group(1))
        if reference is None or reference not in literal_cells:
            return None
        matches = _normalize(literal_cells[reference]) == _normalize(
            condition.group(2)
        )
        return "=" + arguments[1 if matches else 2]

    def unsafe_upstream(
        key: str,
        seen: set[str],
    ) -> str | None:
        if key in seen:
            return None
        seen.add(key)
        formula = formula_cells.get(key)
        if not formula:
            return None
        effective_formula = selected_if_branch(key, formula) or formula
        try:
            tokens = Tokenizer(effective_formula).items
        except Exception:
            tokens = ()
        if (
            _formula_is_top_level_round_to_cents(effective_formula)
            or _formula_is_cent_rounded_difference(effective_formula)
        ):
            return None
        if _formula_has_unrounded_subtraction(effective_formula):
            return key
        sheet = key.rsplit("!", 1)[0]
        for token in tokens:
            if token.type != "OPERAND" or token.subtype != "RANGE":
                continue
            reference = _formula_reference(sheet, token.value)
            if reference is None:
                continue
            unsafe = unsafe_upstream(reference, seen)
            if unsafe is not None:
                return unsafe
        return None

    for key, formula in formula_cells.items():
        if only_cells is not None and key not in only_cells:
            continue
        normalized = re.sub(r"\s+", "", formula).upper()
        if "IF(" not in normalized or "ABS(" not in normalized:
            continue
        try:
            tokens = Tokenizer(formula).items
        except Exception:
            tokens = ()
        if not any(
            token.type == "OPERATOR-INFIX"
            and token.value in {"<", "<=", "=", "<>", ">=", ">"}
            for token in tokens
        ):
            continue
        checked.append(key)
        sheet = key.rsplit("!", 1)[0]
        for token in tokens:
            if token.type != "OPERAND" or token.subtype != "RANGE":
                continue
            reference = _formula_reference(sheet, token.value)
            if reference is None:
                continue
            unsafe = unsafe_upstream(reference, set())
            if unsafe is not None:
                if (
                    only_cells is not None
                    or
                    formula_formats is None
                    or "$" in formula_formats.get(unsafe, "")
                ):
                    issues.append(f"{key} depends on unrounded {unsafe}")
                break
        else:
            if _formula_has_unrounded_subtraction(formula):
                if (
                    only_cells is not None
                    or
                    formula_formats is None
                    or "$" in formula_formats.get(key, "")
                ):
                    issues.append(f"{key} contains an unrounded subtraction")
    return checked, issues


def _source_token(source: str) -> str:
    if source.startswith("ERP "):
        return "ERP"
    if "/" in source:
        return Path(source).stem
    return " ".join(source.split()[:4])


_ERP_SOURCE_ABBREVIATIONS = (
    (r"\bar\s*/\s*ap\b", "accounts receivable accounts payable"),
    (r"\bp\s*&\s*l\b", "profit loss"),
    (r"\bytd\b", "year to date"),
    (r"\bppv\b", "purchase price variance"),
    (r"\bbom\b", "bill of materials"),
    (r"\bwip\b", "work in process"),
    (r"\bpo\b", "purchase order"),
)
_ERP_GENERIC_SOURCE_TERMS = {
    "across",
    "all",
    "and",
    "authoritative",
    "erp",
    "module",
    "modules",
    "record",
    "records",
    "report",
    "reports",
    "the",
}


def _canonical_erp_source_terms(value: str) -> set[str]:
    canonical = value.casefold()
    for pattern, replacement in _ERP_SOURCE_ABBREVIATIONS:
        canonical = re.sub(pattern, replacement, canonical)
    return {
        term
        for term in _normalize(canonical).split()
        if term not in _ERP_GENERIC_SOURCE_TERMS
    }


def _source_present(text: str, source: str) -> bool:
    """Match a visible citation without requiring filename-literal wording."""

    normalized_text = _normalize(text)
    token = _source_token(source)
    normalized_token = _normalize(token)
    if normalized_token and normalized_token in normalized_text:
        if not source.startswith("ERP "):
            return True
        # Broad ERP contracts name several independent report components.
        # Every component must be visible; one coincidental word such as WIP
        # or cash cannot stand in for the full authoritative source family.
        source_terms = _canonical_erp_source_terms(source)
        visible_terms = _canonical_erp_source_terms(text)
        if source_terms and source_terms.issubset(visible_terms):
            return True

    path = Path(source)
    if "/" not in source:
        return False

    date_matches = False
    for match in re.finditer(r"\b(20\d{2})-(\d{2})-(\d{2})\b", path.stem):
        year, month, day = map(int, match.groups())
        if not 1 <= month <= 12:
            continue
        month_name = _MONTH_NAMES[month]
        month_short = month_name[:3]
        variants = (
            f"{year} {month:02d} {day:02d}",
            f"{year} {month} {day}",
            f"{day} {month_name} {year}",
            f"{day} {month_short} {year}",
            f"{month_name} {day} {year}",
            f"{month_short} {day} {year}",
        )
        if any(variant in normalized_text for variant in variants):
            date_matches = True
            break

    suffix_terms = {
        ".eml": {"email", "e mail", "message"},
        ".xlsx": {"workbook", "spreadsheet", "xlsx"},
        ".docx": {"document", "memo", "policy", "docx"},
        ".pptx": {"presentation", "deck", "pptx"},
    }.get(path.suffix.casefold(), set())
    kind_matches = any(term in normalized_text for term in suffix_terms)
    if path.suffix.casefold() == ".eml" and date_matches and kind_matches:
        return True

    ignored = {
        "the",
        "and",
        "for",
        "with",
        "from",
        "january",
        "february",
        "march",
        "april",
        "may",
        "june",
        "july",
        "august",
        "september",
        "october",
        "november",
        "december",
    }
    meaningful = {
        term
        for term in normalized_token.split()
        if len(term) >= 3 and not term.isdigit() and term not in ignored
    }
    hits = meaningful.intersection(normalized_text.split())
    required = min(3, max(2, len(meaningful)))
    return len(hits) >= required


def _text_contains(
    text: str,
    expected: Any,
    metric_key: str | None = None,
) -> bool:
    if isinstance(expected, bool):
        normalized = _normalize(text)
        wanted = "compliant" if expected else "not compliant"
        if _normalize(wanted) in normalized:
            return True
        return (
            "compliant" in _normalize(metric_key)
            and (
                ("pass" in normalized and expected)
                or ("fail" in normalized and not expected)
            )
        )
    if isinstance(expected, (int, float)) and not isinstance(expected, bool):
        text = text.translate(
            str.maketrans(
                {
                    "\N{MINUS SIGN}": "-",
                    "\N{FIGURE DASH}": "-",
                    "\N{EN DASH}": "-",
                    "\N{EM DASH}": "-",
                }
            )
        )
        for token in re.findall(
            (
                r"-?\$?\s*\(?-?\$?\d[\d,]*(?:\.\d+)?"
                r"\)?%?(?:[ \t]*[KMB](?![A-Za-z]))?\)?"
            ),
            text,
            flags=re.I,
        ):
            if _close(
                token,
                float(expected),
                metric_key,
                integer=isinstance(expected, int),
            ):
                return True
        return False
    normalized_text = f" {_normalize(text)} "
    expected_date = _calendar_date(expected)
    if expected_date is not None:
        month_name = _MONTH_NAMES[expected_date.month]
        month_abbreviation = month_name[:3]
        return any(
            f" {_normalize(candidate)} " in normalized_text
            for candidate in (
                expected_date.isoformat(),
                expected_date.strftime("%Y/%m/%d"),
                expected_date.strftime("%m/%d/%Y"),
                f"{month_name} {expected_date.day}, {expected_date.year}",
                f"{month_abbreviation} {expected_date.day}, {expected_date.year}",
                f"{expected_date.day} {month_name} {expected_date.year}",
                f"{expected_date.day} {month_abbreviation} {expected_date.year}",
            )
        )
    expected_period = _year_month(expected)
    if expected_period is not None:
        year, month = expected_period
        month_name = _MONTH_NAMES[month]
        month_abbreviation = month_name[:3]
        return any(
            f" {candidate} " in normalized_text
            for candidate in (
                f"{year} {month:02d}",
                f"{year} {month}",
                f"{month_name} {year}",
                f"{month_abbreviation} {year}",
                f"{year} {month_name}",
                f"{year} {month_abbreviation}",
            )
        )
    return f" {_normalize(expected)} " in normalized_text


def _paired_metric(
    rows: Iterable[Iterable[Any]],
    key: str,
    expected: Any,
    *,
    metric_key: str | None = None,
) -> bool:
    canonical_key = metric_key or key

    def metric_label(value: Any) -> str:
        # In finance workbooks, "/" is routinely used as the visible
        # equivalent of "per" (for example, "cost / employee").
        return _normalize(re.sub(r"\s*/\s*", " per ", str(value or "")))

    def value_matches(value: Any) -> bool:
        if _value_matches(value, expected, canonical_key):
            return True
        # A native row/cell may professionally retain presentation context
        # around the exact value ("$58,928.36 U", "7 years", or even
        # "YTD PPV $58,928.36 U"). It is still a visible label/value pairing;
        # requiring a numeric-only cell creates false negatives.
        if _text_contains(str(value), expected, canonical_key):
            return True
        if not isinstance(expected, str):
            return False
        # Accept a correctly paired business label that preserves an identifier
        # alongside the exact expected name, e.g. "CUST-0182 | Beacon ...".
        actual_text = _normalize(value)
        expected_text = _normalize(expected)
        return bool(
            expected_text
            and f" {expected_text} " in f" {actual_text} "
        )

    label = metric_label(key.replace("_", " "))
    for row in rows:
        row_values = list(row)
        direct_label = any(
            label == metric_label(value) or label in metric_label(value)
            for value in row_values
        )
        # Native tables and bounded cards often split a compound label across
        # adjacent cells/text frames ("TOP GROUP" + "$... revenue"). Match
        # all meaningful label tokens within that one native group, while
        # allowing the generic numeric suffix "value"/"amount" to be implicit.
        group_tokens = set(metric_label(" ".join(map(str, row_values))).split())
        required_tokens = {
            token for token in label.split() if token not in {"value", "amount"}
        }
        grouped_label = bool(required_tokens) and required_tokens <= group_tokens
        if not (direct_label or grouped_label):
            continue
        if any(value_matches(value) for value in row_values):
            return True
    return False


def _labeled_formula_result_consistency(
    results: Iterable[dict[str, Any]],
    label: str,
    expected: Any,
    metric_key: str,
) -> tuple[bool, list[dict[str, Any]]]:
    normalized_label = _normalize(label)
    non_result_column_tokens = {
        "check",
        "delta",
        "difference",
        "status",
        "tolerance",
        "variance",
    }
    matches = [
        dict(result)
        for result in results
        if _normalize(result.get("label")) == normalized_label
        and not (
            set(_normalize(result.get("column_header")).split())
            & non_result_column_tokens
        )
    ]
    return (
        bool(matches)
        and all(
            _number(result.get("value")) is not None
            and _value_matches(
                result.get("value"),
                expected,
                metric_key,
            )
            for result in matches
        ),
        matches,
    )


_CENTRAL_PASSTHROUGH_METRIC_KEYS = frozenset(
    {
        "baseline_period_start",
        "baseline_period_end",
        "scenario_period",
        "status",
        "volume_change",
        "price_change",
        "material_cost_change",
    }
)


def _repeated_formula_column_header_claim(
    result: Mapping[str, Any],
    formula_cells: Mapping[str, str] | None,
) -> bool:
    """Identify a table field header borrowed as a first-row metric label.

    ``_native_formula_label`` deliberately lets a formula immediately below a
    header inherit that header for compact one-value cards.  In a multi-row
    formula table, however, the same layout means the header names the whole
    field rather than asserting that the first scenario row is the central
    metric.  Ignore only that repeated-column case.  A workbook that omits a
    separate exact-labeled central claim still fails through the missing-claim
    check in ``_formula_metric_claim_consistency``.
    """

    if formula_cells is None:
        return False
    label = _normalize(result.get("label"))
    column_header = _normalize(result.get("column_header"))
    if not label or label != column_header:
        return False
    cell = str(result.get("cell") or "")
    matched = re.fullmatch(
        r"([^!]+)!([A-Z]{1,3})([1-9]\d*)",
        cell,
    )
    if matched is None:
        return False
    sheet, column, row = matched.groups()
    return f"{sheet}!{column}{int(row) + 1}" in formula_cells


def _formula_metric_claim_consistency(
    results: Iterable[dict[str, Any]],
    values: dict[str, Any],
    formula_cells: dict[str, str] | None = None,
    *,
    independently_controlled_keys: Iterable[str] = (),
) -> tuple[bool, list[dict[str, Any]]]:
    """Validate a source-backed calculated claim for every central metric.

    A correct static decision sentence must not mask a contradictory native
    Analysis/Control result. Control-only outputs such as differences and
    statuses are excluded because their expected value is deliberately zero
    or PASS rather than the central metric itself.
    """

    non_result_column_tokens = {
        "check",
        "delta",
        "difference",
        "reconciliation",
        "status",
        "tolerance",
        "variance",
    }
    checks: list[dict[str, Any]] = []
    controlled_keys = {str(key) for key in independently_controlled_keys}
    expected_by_label = {
        _normalize(key.replace("_", " ")): (key, expected)
        for key, expected in values.items()
    }
    covered_keys: set[str] = set()
    for result in results:
        normalized_label = _normalize(
            str(result.get("label") or "").replace("_", " ")
        )
        matched_expected = expected_by_label.get(normalized_label)
        if matched_expected is None:
            continue
        cell = str(result.get("cell") or "")
        formula = str(
            (formula_cells or {}).get(cell)
            or result.get("formula")
            or ""
        )
        control_sheet = _normalize(cell.rsplit("!", 1)[0]) in {
            "control",
            "controls",
            "check",
            "checks",
        }
        column_header = _normalize(result.get("column_header"))
        if (
            control_sheet
            and set(column_header.split()).intersection(
                {"result", "outcome"}
            )
            and _meaningful_control_formula(formula)
        ):
            # A Control "Result" field is the PASS/FAIL-style outcome of a
            # reconciliation, not another assertion of the row's metric.
            continue
        if (
            set(column_header.split())
            & non_result_column_tokens
            and (
                column_header != normalized_label
                or (
                    control_sheet
                    and _meaningful_control_formula(formula)
                )
            )
        ):
            continue
        if _repeated_formula_column_header_claim(result, formula_cells):
            continue
        key, expected = matched_expected
        context = _normalize(result.get("section_context"))
        if (
            "detail derived" in column_header
            and "reconciliation" in context
            and "rounding" in context
        ):
            # A source-detail reconciliation is deliberately not the
            # authoritative central claim. Individually rounded source lines
            # may bridge by a few cents while the central metric ties exactly.
            # Keep the exception narrowly bounded so a materially wrong
            # formula cannot evade consistency grading merely by using these
            # labels.
            actual_number = _number(result.get("value"))
            expected_number = _number(expected)
            if (
                actual_number is not None
                and expected_number is not None
                and abs(actual_number - expected_number) <= 0.05 + 1e-9
            ):
                continue
        if (
            column_header == "actual"
            and _normalize(key).split()
            and _normalize(key).split()[0] in {"maximum", "minimum"}
        ):
            # Covenant test summaries often label the test by its limit
            # ("Maximum leverage") while the Actual column intentionally
            # displays leverage, not the maximum threshold.
            continue
        actual = result.get("value")
        matched = _value_matches(actual, expected, key)
        calculated = True
        source_backed = True
        if formula_cells is not None:
            calculated = (
                _formula_is_recursive_calculation(
                    cell,
                    formula,
                    formula_cells,
                )
                or (
                    (
                        key in _CENTRAL_PASSTHROUGH_METRIC_KEYS
                        or key in controlled_keys
                    )
                    and _formula_is_direct_reference(formula)
                )
            )
            source_backed = bool(
                cell
                and _formula_reaches_source(
                    cell,
                    formula,
                    formula_cells,
                )
            ) or key in controlled_keys
            # Formula consistency evaluates calculated claims, not the
            # independent literal source totals they reconcile against.
            # Ignore literal rows here; a missing calculated claim is added
            # below as an explicit failure, while native-claim checks cover
            # contradictory visible literals elsewhere in the workbook.
            if not calculated:
                continue
            matched = matched and calculated and source_backed
        covered_keys.add(key)
        key_tokens = set(_normalize(key).split())
        sign_convention_context = any(
            phrase in context
            for phrase in (
                "bridge",
                "cash flow",
                "cashflow",
                "present value",
                "outflow",
            )
        )
        sign_convention_metric = bool(
            key_tokens.intersection(
                {
                    "investment",
                    "cost",
                    "maintenance",
                    "capex",
                    "expenditure",
                    "outflow",
                }
            )
        )
        actual_number = _number(actual)
        if (
            not matched
            and sign_convention_context
            and sign_convention_metric
            and actual_number is not None
            and isinstance(expected, (int, float))
            and not isinstance(expected, bool)
            and actual_number * float(expected) < 0
        ):
            matched = _close(
                abs(actual_number),
                abs(float(expected)),
                key,
            )
        checks.append(
            {
                "cell": result.get("cell"),
                "label": result.get("label"),
                "column_header": result.get("column_header"),
                "section_context": result.get("section_context"),
                "actual": actual,
                "expected": expected,
                "metric_key": key,
                "calculated": calculated,
                "source_backed": source_backed,
                "matched": matched,
            }
        )
    for key, expected in values.items():
        if key in covered_keys:
            continue
        checks.append(
            {
                "cell": None,
                "label": key.replace("_", " "),
                "column_header": None,
                "section_context": None,
                "actual": None,
                "expected": expected,
                "metric_key": key,
                "calculated": False,
                "source_backed": False,
                "matched": False,
                "issue": (
                    "no exact-labeled calculated claim on Analysis or Control"
                ),
            }
        )
    return all(check["matched"] for check in checks), checks


def _metric_label_options(
    key: str,
    label_aliases: Mapping[str, Iterable[str]] | None = None,
) -> tuple[str, ...]:
    raw_aliases = (label_aliases or {}).get(key, ())
    if isinstance(raw_aliases, (str, bytes)):
        raise ValueError(
            f"label aliases for {key!r} must be an iterable of strings"
        )
    try:
        aliases = tuple(raw_aliases)
    except TypeError as exc:
        raise ValueError(
            f"label aliases for {key!r} must be an iterable of strings"
        ) from exc
    if any(
        not isinstance(alias, str) or not alias.strip()
        for alias in aliases
    ):
        raise ValueError(
            f"label aliases for {key!r} must be non-empty strings"
        )
    options = (key, *aliases)
    normalized = [
        _normalize(option.replace("_", " "))
        for option in options
        if isinstance(option, str)
    ]
    if any(not label for label in normalized) or len(set(normalized)) != len(
        normalized
    ):
        raise ValueError(f"invalid or duplicate label aliases for {key!r}")
    return options


def _materialize_metric_label_aliases(
    values: Mapping[str, Any],
    label_aliases: Mapping[str, Iterable[str]] | None = None,
) -> dict[str, tuple[str, ...]]:
    """Return one validated, reusable alias registry.

    Materializing at the grading boundary prevents generators from producing
    different results across the presence and contradiction checks.
    """

    if label_aliases is None:
        return {}
    if not isinstance(label_aliases, Mapping):
        raise ValueError("label aliases must be a mapping keyed by metric")
    unexpected = sorted(set(label_aliases) - set(values), key=str)
    if unexpected:
        raise ValueError(
            f"label aliases reference unknown metric keys: {unexpected!r}"
        )
    materialized: dict[str, tuple[str, ...]] = {}
    for key in label_aliases:
        if not isinstance(key, str) or not key.strip():
            raise ValueError("label alias metric keys must be non-empty strings")
        materialized[key] = _metric_label_options(
            key,
            {key: label_aliases[key]},
        )[1:]
    return materialized


def _validated_metric_labels(
    values: Mapping[str, Any],
    label_aliases: Mapping[str, Iterable[str]] | None = None,
) -> dict[str, tuple[str, Any]]:
    aliases = _materialize_metric_label_aliases(values, label_aliases)
    expected_by_label: dict[str, tuple[str, Any]] = {}
    for key, expected in values.items():
        for label in _metric_label_options(key, aliases):
            normalized = _normalize(label.replace("_", " "))
            previous = expected_by_label.get(normalized)
            if previous is not None and previous[0] != key:
                raise ValueError(
                    f"metric label alias collision for {normalized!r}: "
                    f"{previous[0]!r} vs {key!r}"
                )
            expected_by_label[normalized] = (key, expected)
    return expected_by_label


def _native_metric_claim_consistency(
    rows: Iterable[Iterable[Any]],
    values: dict[str, Any],
    *,
    label_aliases: Mapping[str, Iterable[str]] | None = None,
) -> tuple[bool, list[dict[str, Any]]]:
    """Validate quantified native rows whose label is a declared metric label.

    This targets decision tables and bounded KPI cards. Definitions and prose
    that merely mention a metric are ignored, as are formula/control rows.
    """

    aliases = _materialize_metric_label_aliases(values, label_aliases)
    expected_by_label = _validated_metric_labels(values, aliases)
    native_header_labels = {
        "action",
        "calculation",
        "comment",
        "control",
        "decision",
        "description",
        "due date",
        "evidence",
        "logic",
        "metric",
        "owner",
        "result",
        "source",
        "status",
        "value",
    }
    checks: list[dict[str, Any]] = []
    for row in rows:
        row_values = list(row)
        if any(
            isinstance(value, str) and value.startswith("=")
            for value in row_values
        ):
            continue
        normalized_cells = [
            _normalize(str(value).replace("_", " ")) for value in row_values
        ]
        exact_metric_label_count = sum(
            normalized_cell in expected_by_label
            for normalized_cell in normalized_cells
        )
        nonempty_cells = [
            normalized_cell
            for normalized_cell in normalized_cells
            if normalized_cell
        ]
        if (
            exact_metric_label_count
            and len(nonempty_cells) >= 2
            and all(
                normalized_cell in native_header_labels
                for normalized_cell in nonempty_cells
            )
        ):
            continue
        for label_index, normalized_cell in enumerate(normalized_cells):
            matched_expected = expected_by_label.get(normalized_cell)
            if matched_expected is None:
                continue
            key, expected = matched_expected
            candidates: list[Any] = []
            for value_index, value in enumerate(row_values):
                if value_index == label_index:
                    continue
                if isinstance(expected, bool):
                    if isinstance(value, bool) or _normalize(value) in {
                        "true",
                        "false",
                        "yes",
                        "no",
                        "compliant",
                        "not compliant",
                    }:
                        candidates.append(value)
                elif isinstance(expected, (int, float)) and not isinstance(
                    expected,
                    bool,
                ):
                    parsed = _number(value)
                    if parsed is None:
                        continue
                    text_value = str(value).strip()
                    # The row may inherit "USD millions" from a separate
                    # table header (for example, "$92.550"). Without that
                    # header in this bounded row the scale is ambiguous, so
                    # skip it rather than misclassifying a scaled display as
                    # a contradictory $92.55 claim.
                    ambiguous_large_money_scale = (
                        abs(float(expected)) >= 1_000_000
                        and abs(parsed) < 1_000
                        and re.search(
                            r"[KMB]\s*$",
                            text_value,
                            flags=re.I,
                        )
                        is None
                    )
                    if (
                        not ambiguous_large_money_scale
                        or "$" in text_value
                    ):
                        candidates.append(value)
                elif _calendar_date(expected) is not None:
                    if _calendar_date(value) is not None:
                        candidates.append(value)
                elif isinstance(expected, str):
                    if (
                        exact_metric_label_count == 1
                        and normalized_cells[value_index]
                        not in expected_by_label
                        and isinstance(value, str)
                        and value.strip()
                    ):
                        candidates.append(value)
            if not candidates:
                # A native two-or-more-cell group with one exact metric label
                # is an explicit claim even when its paired value is
                # non-numeric (for example, "materially wrong"). Record that
                # contradiction instead of allowing a different correct claim
                # elsewhere to mask it. Label-only text boxes and table label
                # columns containing several metric names remain definitions,
                # not claims, and are ignored.
                comparable_expected = (
                    isinstance(expected, bool)
                    or (
                        isinstance(expected, (int, float))
                        and not isinstance(expected, bool)
                    )
                    or _calendar_date(expected) is not None
                    or isinstance(expected, str)
                )
                if (
                    comparable_expected
                    and len(row_values) > 1
                    and exact_metric_label_count == 1
                ):
                    checks.append(
                        {
                            "label": row_values[label_index],
                            "row": row_values,
                            "candidates": [],
                            "expected": expected,
                            "matched": False,
                            "issue": "exact-labeled native group has no comparable value",
                        }
                    )
                continue
            matched = any(
                _value_matches(candidate, expected, key)
                or (
                    isinstance(expected, str)
                    and isinstance(candidate, str)
                    and _text_contains(candidate, expected, key)
                )
                or (
                    isinstance(expected, (int, float))
                    and not isinstance(expected, bool)
                    and isinstance(candidate, str)
                    and _displayed_claim_matches(
                        candidate,
                        float(expected),
                    )
                )
                or (
                    isinstance(expected, (int, float))
                    and not isinstance(expected, bool)
                    and isinstance(candidate, str)
                    and _scaled_currency_claim_matches(
                        candidate,
                        float(expected),
                    )
                )
                for candidate in candidates
            )
            row_context = _normalize(
                " ".join(
                    str(value)
                    for value_index, value in enumerate(row_values)
                    if value_index != label_index
                )
            )
            key_tokens = set(_normalize(key).split())
            sign_convention_context = _has_decision_phrase(
                row_context,
                "subtract",
                "deduct",
                "deduction",
                "cash outflow",
                "use of cash",
                "cash flow bridge",
            )
            sign_convention_metric = (
                _normalize(key) == "tax depreciation"
                or bool(
                    key_tokens.intersection(
                        {
                            "investment",
                            "cost",
                            "capex",
                            "expenditure",
                            "outflow",
                        }
                    )
                )
            )
            if (
                not matched
                and sign_convention_context
                and sign_convention_metric
                and isinstance(expected, (int, float))
                and not isinstance(expected, bool)
            ):
                matched = any(
                    (actual := _number(candidate)) is not None
                    and actual * float(expected) < 0
                    and (
                        _close(abs(actual), abs(float(expected)), key)
                        or (
                            isinstance(candidate, str)
                            and _displayed_claim_matches(
                                candidate,
                                float(expected),
                                absolute=True,
                            )
                        )
                    )
                    for candidate in candidates
                )
            checks.append(
                {
                    "label": row_values[label_index],
                    "row": row_values,
                    "candidates": candidates,
                    "expected": expected,
                    "matched": matched,
                }
            )
    return all(check["matched"] for check in checks), checks


def _bounded_metric_value(
    rows: Iterable[Iterable[Any]],
    key: str,
    expected: Any,
    *,
    metric_key: str | None = None,
) -> bool:
    """Find a value in a native block that also identifies its metric.

    The value criterion is intentionally a little less strict than
    ``_paired_metric``: a bounded card may abbreviate one component of a
    compound metric label. It may not, however, borrow a label from an
    unrelated part of the artifact. That distinction prevents a raw detail
    row containing (for example) the integer 156 from satisfying
    ``vendor_count`` merely because the flattened workbook contains that
    label somewhere else.
    """

    canonical_key = metric_key or key
    label_tokens = {
        token
        for token in _normalize(
            re.sub(r"\s*/\s*", " per ", key.replace("_", " "))
        ).split()
        if token not in {"value", "amount"}
    }
    if not label_tokens:
        return False
    # A two-token metric cannot safely drop half of its identity:
    # ``quality`` alone must not let a raw ``Quality exposure`` row satisfy
    # ``quality normalization``. Longer compound labels may omit one token in
    # a compact card while still remaining meaningfully scoped.
    required_hits = (
        len(label_tokens)
        if len(label_tokens) <= 2
        else len(label_tokens) - 1
    )
    for row in rows:
        row_values = list(row)
        if not any(
            _value_matches(value, expected, canonical_key)
            or _text_contains(str(value), expected, canonical_key)
            for value in row_values
        ):
            continue
        group_tokens = set(
            _normalize(
                re.sub(
                    r"\s*/\s*",
                    " per ",
                    " ".join(map(str, row_values)),
                )
            ).split()
        )
        if len(label_tokens.intersection(group_tokens)) >= required_hits:
            return True
    return False


def _native_row_sequence_matches(
    rows: Iterable[Iterable[Any]],
    label: Any,
    expected: Iterable[tuple[str, Any]],
) -> bool:
    """Match a native table row by label and ordered field values."""

    normalized_label = _normalize(label)
    expected_values = list(expected)
    for candidate in rows:
        row = list(candidate)
        if (
            not row
            or _normalize(row[0]) != normalized_label
            or len(row) < len(expected_values) + 1
        ):
            continue
        if all(
            _value_matches(row[index], value, key)
            for index, (key, value) in enumerate(
                expected_values,
                start=1,
            )
        ):
            return True
    return False


def _semantic_native_table_truth(
    tables: Iterable[Iterable[Iterable[Any]]],
    *,
    column_aliases: dict[str, Iterable[str]],
    key_fields: tuple[str, ...],
    expected_rows: list[dict[str, Any]],
) -> dict[str, Any]:
    """Validate one unambiguous native table independent of column order."""

    result = {
        "candidate_count": 0,
        "row_count": 0,
        "expected_row_count": len(expected_rows),
        "keys_match": False,
        "field_values_match": False,
        "mismatches": [],
    }
    if (
        not column_aliases
        or not key_fields
        or not all(field in column_aliases for field in key_fields)
        or any(set(row) != set(column_aliases) for row in expected_rows)
    ):
        return result
    normalized_aliases = {
        field: {_normalize(alias) for alias in aliases}
        for field, aliases in column_aliases.items()
    }
    candidates: list[tuple[list[list[Any]], dict[str, int]]] = []
    for raw_table in tables:
        table = [list(row) for row in raw_table]
        if not table:
            continue
        header = table[0]
        positions: dict[str, int] = {}
        valid = True
        for field, aliases in normalized_aliases.items():
            matches = [
                index
                for index, value in enumerate(header)
                if _normalize(value) in aliases
            ]
            if len(matches) != 1:
                valid = False
                break
            positions[field] = matches[0]
        if valid:
            candidates.append((table, positions))
    result["candidate_count"] = len(candidates)
    if len(candidates) != 1:
        return result

    table, positions = candidates[0]
    actual_rows = [
        {
            field: (
                row[index]
                if index < len(row)
                else None
            )
            for field, index in positions.items()
        }
        for row in table[1:]
        if any(value not in (None, "") for value in row)
    ]
    result["row_count"] = len(actual_rows)

    def key_for(row: dict[str, Any]) -> tuple[str, ...]:
        keys: list[str] = []
        for field in key_fields:
            value = row[field]
            normalized_field = _normalize(field)
            if "year" in normalized_field:
                year_match = re.fullmatch(
                    r"(?:year\s*)?(\d+)",
                    _normalize(value),
                )
                if year_match:
                    keys.append(year_match.group(1))
                    continue
            numeric = _number(value)
            if numeric is not None:
                keys.append(f"{numeric:.12g}")
                continue
            keys.append(_normalize(value))
        return tuple(keys)

    expected_by_key = {
        key_for(row): row for row in expected_rows
    }
    actual_by_key = {
        key_for(row): row for row in actual_rows
    }
    expected_counter = Counter(key_for(row) for row in expected_rows)
    actual_counter = Counter(key_for(row) for row in actual_rows)
    result["keys_match"] = actual_counter == expected_counter
    mismatches: list[dict[str, Any]] = []
    for key, expected in expected_by_key.items():
        actual = actual_by_key.get(key)
        if actual is None:
            continue
        for field, expected_value in expected.items():
            if field in key_fields:
                # Key equality is already enforced through its semantic
                # canonicalizer (for example, `1` and `Year 1`).
                continue
            if _value_matches(actual[field], expected_value, field):
                continue
            if len(mismatches) < _ACTION_REGISTER_EVIDENCE_LIMIT:
                mismatches.append(
                    {
                        "key": key,
                        "field": field,
                        "actual": actual[field],
                        "expected": expected_value,
                    }
                )
    result["mismatches"] = mismatches
    result["field_values_match"] = (
        result["keys_match"]
        and len(actual_rows) == len(expected_rows)
        and not mismatches
    )
    return result


_DECISION_NUMBER_PATTERN = re.compile(
    (
        r"-?\$?\s*\(?-?\$?\d[\d,]*(?:\.\d+)?"
        r"\)?%?(?:[ \t]*(?:[KMB](?![A-Za-z])|thousand|million|billion))?\)?"
    ),
    flags=re.I,
)


def _decision_number_tokens(text: str) -> list[tuple[str, str]]:
    tokens: list[tuple[str, str]] = []
    for match in _DECISION_NUMBER_PATTERN.finditer(str(text or "")):
        token = match.group(0).strip()
        if not token:
            continue
        kind = (
            "percent"
            if "%" in token
            else "currency"
            if "$" in token
            else "plain"
        )
        tokens.append((token, kind))
    return tokens


def _decision_number_match(token: str, expected: float, kind: str) -> bool:
    del kind
    return _displayed_claim_matches(token, expected)


_DECISION_PLAIN_METRIC_PATTERNS: dict[str, tuple[str, ...]] = {
    "inventory_turns": (
        (
            r"(?P<number>-?\d[\d,]*(?:\.\d+)?)\s*x?\s+"
            r"(?:inventory\s+)?turns?\b"
        ),
        (
            r"\binventory\s+turns?\s*(?:are|is|of|at|=|:)?\s*"
            r"(?P<number>-?\d[\d,]*(?:\.\d+)?)"
        ),
    ),
    "days_inventory": (
        (
            r"(?P<number>-?\d[\d,]*(?:\.\d+)?)\s+days?\s+"
            r"(?:of\s+inventory|on\s+hand)\b"
        ),
        (
            r"\binventory\s+turns?\b[^\n.;]{0,60}?"
            r"(?P<number>-?\d[\d,]*(?:\.\d+)?)\s+days?\b"
        ),
    ),
    "scenario_overloaded_centers": (
        (
            r"(?P<number>\d[\d,]*)\s*(?:of\s+\d[\d,]*\s+)?"
            r"(?:work\s+)?centers?\s+(?:remain|are|still)\s+overloaded\b"
        ),
    ),
    "work_centers": (
        (
            r"\d[\d,]*\s+of\s+(?P<number>\d[\d,]*)\s+"
            r"(?:work\s+)?centers?\b"
        ),
    ),
    "planned_orders": (
        r"(?P<number>\d[\d,]*)\s+planned\s+orders?\b",
    ),
    "fully_depreciated_assets": (
        (
            r"(?P<number>\d[\d,]*)\s+assets?\s+(?:are\s+)?"
            r"fully\s+depreciated\b"
        ),
        (
            r"\bfully\s+depreciated\s+assets?\s*(?:total|are|=|:)?\s*"
            r"(?P<number>\d[\d,]*)"
        ),
    ),
    "top_asset_downtime": (
        r"(?P<number>-?\d[\d,]*(?:\.\d+)?)\s+downtime\s+hours?\b",
        r"(?P<number>-?\d[\d,]*(?:\.\d+)?)\s+hours?\s+of\s+downtime\b",
    ),
    "reconciliation_failures": (
        (
            r"\breconciliation\s+failures?\s*(?:total|are|=|:)?\s*"
            r"(?P<number>\d[\d,]*)"
        ),
        r"\bfailures?\s+total\s+(?P<number>\d[\d,]*)",
    ),
}


def _decision_plain_metric_claims_match(
    text: str,
    values: dict[str, Any],
) -> bool:
    """Validate explicitly stated counts, days, turns, and hours exactly."""

    for key, patterns in _DECISION_PLAIN_METRIC_PATTERNS.items():
        expected = values.get(key)
        if (
            not isinstance(expected, (int, float))
            or isinstance(expected, bool)
        ):
            continue
        for pattern in patterns:
            for match in re.finditer(pattern, text, flags=re.I):
                token = match.group("number")
                actual = _number(token)
                if actual is None:
                    return False
                if isinstance(expected, int):
                    if not actual.is_integer() or int(actual) != expected:
                        return False
                elif not _displayed_claim_matches(token, float(expected)):
                    return False

    action_counts = [
        values.get(key)
        for key in (
            "expedite_count",
            "defer_count",
            "increase_count",
            "decrease_count",
        )
    ]
    if all(
        isinstance(value, int) and not isinstance(value, bool)
        for value in action_counts
    ):
        expected_actions = sum(action_counts)
        for match in re.finditer(
            (
                r"(?P<number>\d[\d,]*)\s+explicit\s+"
                r"MRP\s+action\s+messages?\b"
            ),
            text,
            flags=re.I,
        ):
            actual = _number(match.group("number"))
            if (
                actual is None
                or not actual.is_integer()
                or int(actual) != expected_actions
            ):
                return False
    return True


def _decision_entity_supported(
    text: str,
    key: str,
    expected: str,
    rows: Iterable[Iterable[Any]],
) -> bool:
    if _text_contains(text, expected, key):
        return True

    expected_tokens = set(_normalize(expected).split())
    ignored_alias_tokens = {
        "action",
        "address",
        "approve",
        "capacity",
        "center",
        "centers",
        "commitments",
        "contain",
        "containment",
        "control",
        "cost",
        "decision",
        "delivery",
        "department",
        "escalate",
        "gate",
        "highest",
        "immediate",
        "lowest",
        "manage",
        "overloaded",
        "plan",
        "prioritize",
        "protect",
        "recover",
        "recovery",
        "relieve",
        "score",
        "supplier",
        "utilization",
        "vendor",
    }
    alias_tokens = {
        token
        for token in _normalize(text).split()
        if len(token) >= 3
        and not token.isdigit()
        and token not in expected_tokens
        and token not in ignored_alias_tokens
    }
    if not alias_tokens:
        return False
    for row in rows:
        row_values = list(row)
        if not any(_value_matches(value, expected, key) for value in row_values):
            continue
        row_tokens = set(_normalize(" ".join(map(str, row_values))).split())
        if alias_tokens.intersection(row_tokens):
            return True
    return False


def _has_decision_phrase(text: str, *phrases: str) -> bool:
    return any(_normalize(phrase) in text for phrase in phrases)


def _supplier_mitigation_selected(text: str) -> bool:
    """Require one affirmative, resolved mitigation in a bounded clause."""

    action_pattern = (
        r"\b(?:approve|approves|approved|approving|authorize|authorizes|"
        r"authorized|authorizing|adopt|adopts|adopted|adopting|begin|"
        r"begins|beginning|build|builds|building|built|commit|commits|"
        r"committed|committing|direct|directs|directed|directing|"
        r"establish|establishes|established|establishing|execute|executes|"
        r"executed|executing|implement|implements|implemented|implementing|"
        r"initiate|initiates|initiated|initiating|launch|launches|launched|"
        r"launching|proceed|proceeds|proceeded|proceeding|qualify|qualifies|"
        r"qualified|qualifying|recommend|recommends|recommended|"
        r"recommending|require|requires|required|requiring|select|selects|"
        r"selected|selecting|start|starts|started|starting)\b"
    )
    mitigation_patterns = (
        r"\b(?:supplier|delivery)?\s*recovery\s+"
        r"(?:plan|action|program)\b",
        r"\b(?:supplier\s+)?corrective\s+action(?:\s+plan|\s+program)?\b",
        r"\bscar\b",
        r"\b(?:dual|second|alternate|alternative|backup)\s+"
        r"(?:source|supplier)\b",
        r"\b(?:safety|buffer)\s+stock\b",
        r"\bcontainment\b",
    )
    negative_pattern = (
        r"\b(?:do\s+not|don\s+t|not|never|reject\w*|declin\w*|"
        r"defer\w*|paus\w*)\b"
    )
    modal_pattern = (
        r"\b(?:may|might|could|should|would|can|consider\w*|"
        r"evaluat\w*)\b"
    )

    raw_clauses = re.split(r"(?:[.;]|\r?\n)+", str(text or ""))
    for raw_clause in raw_clauses:
        clause = _normalize(raw_clause)
        if not clause:
            continue
        mitigation_matches = [
            match
            for pattern in mitigation_patterns
            for match in re.finditer(pattern, clause, flags=re.I)
        ]
        if not mitigation_matches:
            continue
        action_matches = list(
            re.finditer(action_pattern, clause, flags=re.I)
        )
        if not action_matches or re.search(
            negative_pattern,
            clause,
            flags=re.I,
        ):
            continue

        # "Could launch" and "consider launching" are options, not selected
        # actions. A modal elsewhere in the clause does not negate a separate
        # imperative such as "Launch ... with escalation options."
        unresolved_modal = False
        for modal in re.finditer(modal_pattern, clause, flags=re.I):
            if any(
                0 <= action.start() - modal.end() <= 24
                for action in action_matches
            ):
                unresolved_modal = True
                break
        if unresolved_modal:
            continue

        # Reject unresolved mitigation menus while retaining valid wording
        # such as "launch recovery with containment or escalation options."
        if len(mitigation_matches) >= 2:
            first_start = min(match.start() for match in mitigation_matches)
            last_end = max(match.end() for match in mitigation_matches)
            between = clause[first_start:last_end]
            if (
                re.search(r"\bor\b", between)
                or "/" in raw_clause
                or "," in raw_clause
            ):
                continue

        # The action and mitigation must be close enough to govern the same
        # bounded clause, whichever order a professional sentence uses.
        for action in action_matches:
            for mitigation in mitigation_matches:
                if action.end() <= mitigation.start():
                    bridge = clause[action.end():mitigation.start()]
                elif mitigation.end() <= action.start():
                    bridge = clause[mitigation.end():action.start()]
                else:
                    bridge = ""
                bridge_tokens = bridge.split()
                if len(bridge_tokens) > 12:
                    continue
                if re.search(
                    r"\b(?:but|whereas|while|discuss\w*|evaluat\w*|"
                    r"consider\w*|review\w*|outlin\w*|present\w*|"
                    r"propos\w*)\b",
                    bridge,
                ):
                    continue
                return True
    return False


def _decision_semantic_equivalence(
    text: str,
    decision: str,
) -> bool:
    """Recognize narrow, auditable business-decision paraphrases."""

    if all(
        token in decision
        for token in ("subledger", "reconcil", "pre close")
    ):
        return (
            "control" in text
            and "reconcil" in text
            and _has_decision_phrase(
                text,
                "before lock",
                "before locking",
                "do not lock",
                "hold period lock",
                "keep open",
                "pre close",
            )
        )

    if all(
        phrase in decision
        for phrase in ("inventory turns", "days on hand", "working capital")
    ):
        return (
            "inventory turns" in text
            and "days" in text
            and _has_decision_phrase(
                text,
                "aging disposition",
                "aged inventory",
                "inventory aging",
                "inventory recoverability",
                "inventory disposition",
                "slow moving inventory",
                "release working capital",
            )
        )

    if all(phrase in decision for phrase in ("backlog", "overdue", "capacity")):
        return (
            "backlog" in text
            and _has_decision_phrase(
                text,
                "convert",
                "conversion",
                "re promise",
                "reprioritize",
                "prioritize",
                "capacity",
                "overdue",
            )
            and _has_decision_phrase(
                text,
                "action",
                "mitigate",
                "reduce",
                "recover",
                "prioritize",
                "re promise",
            )
        )

    if (
        all(phrase in decision for phrase in ("supplier", "delivery"))
        and _supplier_mitigation_selected(decision)
    ):
        return (
            _has_decision_phrase(text, "supplier", "vendor")
            and _has_decision_phrase(
                text,
                "delivery",
                "on time",
                "lead time",
                "quality",
            )
            and _supplier_mitigation_selected(text)
        )

    if all(phrase in decision for phrase in ("copq", "scrap", "quality")):
        return (
            _has_decision_phrase(text, "copq", "cost of poor quality")
            and "scrap" in text
            and _has_decision_phrase(
                text,
                "quality",
                "open quality",
                "quality order",
                "rm 0102",
            )
            and _has_decision_phrase(
                text,
                "prioritize",
                "resolve",
                "reduce",
                "recover",
                "contain",
            )
        )

    if all(
        phrase in decision
        for phrase in (
            "gate unmitigated commitments",
            "top constraint",
            "utilization",
        )
    ):
        return (
            _has_decision_phrase(
                text,
                "do not accept the unconstrained",
                "do not accept unconstrained",
            )
            and _has_decision_phrase(
                text,
                "constraint",
                "bottleneck",
                "overloaded",
            )
        )

    if all(
        token in decision
        for token in ("mrp", "action", "planned orders")
    ):
        return (
            all(token in text for token in ("explicit", "mrp", "all"))
            and _has_decision_phrase(
                text,
                "execute",
                "authorize",
                "disposition",
            )
        )

    if "prioritize disposition" in decision:
        return (
            "disposition" in text
            and (
                "first" in text.split()
                or _has_decision_phrase(
                    text,
                    "priority disposition",
                    "prioritize disposition",
                )
            )
        )

    if "consensus demand baseline" in decision:
        return (
            "2026 06 sop" in text
            and "baseline" in text
            and _has_decision_phrase(text, "approve", "authorize", "lock")
            and _has_decision_phrase(
                text,
                "jul dec 2026",
                "dec 2026",
                "december 2026",
                "2026 12",
            )
        )

    if all(token in decision for token in ("drv", "aft", "margin")):
        return (
            all(token in text for token in ("drv", "aft"))
            and _has_decision_phrase(text, "protect")
            and _has_decision_phrase(
                text,
                "correct",
                "prioritize",
                "recovery",
                "target",
            )
        )

    if (
        "price action" in decision
        and "material cost shock" in decision
        and "does not offset" in decision
    ):
        return (
            "price" in text
            and _has_decision_phrase(
                text,
                "insufficient",
                "not sufficient",
                "does not offset",
                "do not approve",
            )
            and _has_decision_phrase(
                text,
                "material",
                "gross profit",
                "margin",
                "price floor",
                "price realization",
            )
        )

    if "minimum cash buffer" in decision and "revolver" in decision:
        return (
            "incremental" in text
            and _has_decision_phrase(text, "revolver", "borrowing", "draw")
            and _has_decision_phrase(
                text,
                "no incremental",
                "zero incremental",
                "maximum incremental revolver need is 0",
            )
            and _has_decision_phrase(text, "minimum cash", "cash buffer")
        )

    if "leverage covenant" in decision and "not compliant" in decision:
        noncompliance = _has_decision_phrase(
            text,
            "not compliant",
            "noncompliant",
            "overall compliant false",
            "leverage fails",
            "leverage breach",
            "do not certify compliance",
            "withhold clean compliance certification",
        )
        lender_action = _has_decision_phrase(
            text,
            "engage lenders",
            "lender action",
            "lender counsel",
            "lender waiver",
            "lender planning",
            "waiver planning",
            "waiver direction",
            "waiver cure",
        )
        return noncompliance and lender_action

    if (
        ("tax provision" in decision or "tax benefit" in decision)
        and "tax benefit" in text
    ):
        return (
            _has_decision_phrase(text, "approve", "record", "recognize")
            and _has_decision_phrase(text, "pre close", "preclose")
        )

    if all(
        phrase in decision
        for phrase in ("price increase", "gross profit", "gross margin")
    ):
        return (
            "price" in text
            and _has_decision_phrase(text, "floor", "minimum", "at least")
            and _has_decision_phrase(text, "target", "gross margin")
            and _has_decision_phrase(
                text,
                "insufficient",
                "do not approve",
                "not sufficient",
            )
        )

    return False


def _decision_concept_match(
    text: str,
    decision: str,
    *,
    values: dict[str, Any] | None = None,
    paired_rows: Iterable[Iterable[Any]] = (),
    require_explicit_numbers: bool = False,
) -> bool:
    normalized_text = _normalize(text)
    normalized_decision = _normalize(decision)
    if not normalized_text or not normalized_decision:
        return False
    supplier_delivery_decision = all(
        phrase in normalized_decision for phrase in ("supplier", "delivery")
    ) and _supplier_mitigation_selected(normalized_decision)
    supplier_action_equivalence = False
    if supplier_delivery_decision:
        supplier_scope = _has_decision_phrase(
            normalized_text,
            "supplier",
            "vendor",
        )
        performance_scope = _has_decision_phrase(
            normalized_text,
            "delivery",
            "on time",
            "lead time",
            "quality",
        )
        if not (
            supplier_scope
            and performance_scope
            and _supplier_mitigation_selected(normalized_text)
        ):
            return False
        supplier_action_equivalence = True
    negative_approval_phrases = (
        "do not approve",
        "not approve",
        "reject",
        "decline",
        "defer approval",
        "gate approval",
        "pause approval",
    )
    price_floor_equivalence = (
        all(
            phrase in normalized_decision
            for phrase in ("price increase", "gross profit", "gross margin")
        )
        and "price" in normalized_text
        and _has_decision_phrase(
            normalized_text,
            "floor",
            "minimum",
            "at least",
        )
        and _has_decision_phrase(
            normalized_text,
            "insufficient",
            "do not approve",
            "not sufficient",
        )
    )
    if "do not approve" in normalized_decision and not any(
        phrase in normalized_text
        for phrase in negative_approval_phrases
    ):
        return False
    if (
        "approve" in normalized_decision
        and "do not approve" not in normalized_decision
        and (
            any(
                phrase in normalized_text
                for phrase in negative_approval_phrases
            )
            or not any(
                phrase in normalized_text
                for phrase in ("approve", "authorize", "proceed", "advance")
            )
        )
        and not price_floor_equivalence
    ):
        return False
    if "not compliant" in normalized_decision and not any(
        phrase in normalized_text
        for phrase in (
            "not compliant",
            "noncompliant",
            "breach",
            "covenant fail",
            "leverage fail",
            "do not certify compliance",
            "withhold clean compliance certification",
        )
    ):
        return False

    values = values or {}
    rows = [list(row) for row in paired_rows]
    if not _decision_plain_metric_claims_match(text, values):
        return False
    entity_key_tokens = {
        "asset",
        "constraint",
        "customer",
        "department",
        "family",
        "group",
        "item",
        "largest",
        "month",
        "order",
        "site",
        "supplier",
        "top",
        "vendor",
        "worst",
    }
    for key, expected in values.items():
        if not isinstance(expected, str):
            continue
        key_tokens = set(_normalize(key).split())
        if not key_tokens.intersection(entity_key_tokens):
            continue
        if not _text_contains(decision, expected, key):
            continue
        if (
            supplier_action_equivalence
            and key_tokens.intersection({"supplier", "vendor", "worst"})
        ):
            # A resolved supplier delivery/quality mitigation is an accepted
            # action alternative even when it identifies the affected scope
            # rather than repeating the gold supplier display name.
            continue
        if not _decision_entity_supported(text, key, expected, rows):
            return False

    # If the submitted decision states currency or percentages, validate the
    # displayed claims against the corresponding numbers in the gold decision.
    # The submitted display precision controls the rounding window.
    expected_numbers: dict[str, list[float]] = {"currency": [], "percent": []}
    gold_number_tokens = _decision_number_tokens(decision)
    for token, kind in gold_number_tokens:
        if kind not in expected_numbers:
            continue
        expected = _number(token)
        if expected is not None:
            expected_numbers[kind].append(expected)
    submitted_numbers = _decision_number_tokens(text)
    for kind, expected_values in expected_numbers.items():
        candidates = [
            token for token, candidate_kind in submitted_numbers
            if candidate_kind == kind
        ]
        tax_benefit_absolute_match = (
            kind == "currency"
            and "tax provision" in normalized_decision
            and "tax benefit" in normalized_text
        )
        zero_revolver_semantic_match = (
            kind == "currency"
            and "revolver" in normalized_decision
            and _has_decision_phrase(
                normalized_text,
                "no incremental revolver",
                "no incremental draw",
                "no incremental borrowing",
            )
        )
        if expected_values and require_explicit_numbers and not candidates:
            return False
        if not expected_values or not candidates:
            continue

        def expected_is_matched(expected: float) -> bool:
            if zero_revolver_semantic_match and abs(expected) <= 0.005:
                return True
            return any(
                _displayed_claim_matches(
                    token,
                    expected,
                    absolute=tax_benefit_absolute_match,
                )
                for token in candidates
            )

        matches = [
            expected_is_matched(expected) for expected in expected_values
        ]
        # When every gold claim is repeated, every one must be correct. A
        # partial executive statement may state one correct claim and rely on
        # its paired native value chain for the rest (for example task 086).
        if (
            len(candidates) >= len(expected_values)
            and not all(matches)
        ) or (
            len(candidates) < len(expected_values)
            and not any(matches)
        ):
            return False

    if (
        "approve" in normalized_decision
        and "do not approve" not in normalized_decision
        and set(normalized_decision.split()) <= {"approve", "the", "project"}
    ):
        return True

    standalone_valuation = all(
        phrase in normalized_decision
        for phrase in (
            "standalone valuation",
            "enterprise value",
            "equity value",
            "net debt",
        )
    )
    if standalone_valuation:
        equity_value = values.get("equity_value")
        stated_equity = (
            isinstance(equity_value, (int, float))
            and not isinstance(equity_value, bool)
            and any(
                kind == "currency"
                and _decision_number_match(
                    token,
                    float(equity_value),
                    kind,
                )
                for token, kind in submitted_numbers
            )
        )
        chain_supported = all(
            key in values and _paired_metric(rows, key, values[key])
            for key in ("enterprise_value", "net_debt", "equity_value")
        )
        if (
            "standalone" in normalized_text
            and "equity value" in normalized_text
            and stated_equity
            and chain_supported
        ):
            return True

    executive_program = all(
        phrase in normalized_decision
        for phrase in (
            "margin",
            "backlog to cash",
            "working capital",
            "discretionary capital",
        )
    )
    if executive_program:
        working_capital_action = (
            "working capital" in normalized_text
            or (
                "inventory" in normalized_text
                and "wip" in normalized_text
                and _has_decision_phrase(
                    normalized_text,
                    "inventory wip conversion",
                    "inventory aging reduction",
                )
            )
        )
        return (
            "margin" in normalized_text
            and all(
                phrase in normalized_text
                for phrase in ("backlog", "cash")
            )
            and working_capital_action
            and "discretionary" in normalized_text
            and _has_decision_phrase(
                normalized_text,
                "june close",
                "june pre close",
                "pre close",
            )
        )

    stop = {
        "the", "and", "for", "with", "that", "this", "from", "until",
        "use", "only", "into", "named", "result", "stated",
    }
    tokens = [
        token
        for token in _normalize(decision).split()
        if len(token) >= 4 and token not in stop
    ]
    if not tokens:
        return False
    haystack = set(normalized_text.split())
    lexical_match = sum(token in haystack for token in tokens) >= max(
        2, math.ceil(len(tokens) * 0.45)
    )
    if lexical_match:
        return True

    if _decision_semantic_equivalence(
        normalized_text,
        normalized_decision,
    ):
        return True

    concept_groups = (
        {
            "address",
            "contain",
            "containment",
            "control",
            "correct",
            "escalate",
            "gate",
            "launch",
            "manage",
            "prioritize",
            "recover",
            "recovery",
            "relieve",
        },
        {
            "material",
            "output",
            "production",
            "scrap",
            "usage",
            "variance",
            "wip",
        },
        {
            "bottleneck",
            "capacity",
            "center",
            "centers",
            "constraint",
            "overloaded",
            "utilization",
        },
    )
    decision_tokens = set(normalized_decision.split())
    shared_groups = sum(
        bool(group.intersection(decision_tokens))
        and bool(group.intersection(haystack))
        for group in concept_groups
    )
    return shared_groups >= 2


def _decision_action_row_match(
    text: str,
    decision: str,
    *,
    values: dict[str, Any] | None = None,
) -> bool:
    """Match one action row without borrowing entity or amount evidence."""

    return _decision_concept_match(
        text,
        decision,
        values=values,
        paired_rows=((text,),),
        require_explicit_numbers=True,
    )


def _project_site_scope_match(
    text: str,
    site_name: str,
    site_code: str,
) -> bool:
    normalized = f" {_normalize(text)} "
    normalized_name = _normalize(site_name)
    normalized_code = _normalize(site_code)
    return (
        bool(normalized_name and normalized_code)
        and f" {normalized_name} " in normalized
        and re.search(
            rf"\bsite(?:\s+code)?\s+{re.escape(normalized_code)}\b",
            normalized,
        )
        is not None
    )


def _scoped_site_approval_match(
    text: str,
    site_name: str,
    site_code: str,
) -> bool:
    """Require an affirmative approval and the project site in one native block."""

    positive_terms = ("approve", "approval", "authorize", "authorization", "proceed")
    negative_terms = (
        "do not approve",
        "not approve",
        "reject",
        "decline",
        "defer approval",
        "pause approval",
    )
    for block in str(text or "").splitlines():
        normalized = _normalize(block)
        if not normalized or any(term in normalized for term in negative_terms):
            continue
        if (
            any(term in normalized for term in positive_terms)
            and _project_site_scope_match(normalized, site_name, site_code)
        ):
            return True
    return False


def _artifact_answer_criteria(
    task_id: str, answer: Any, expected_path: str
) -> tuple[list[Criterion], dict[str, Any]]:
    mapping, exact_json, evidence = _strict_json(answer)
    required_keys = {
        "deliverable",
        "status",
        "central_decision",
        "sources_checked",
    }
    exact_completion_schema = set(mapping) == required_keys
    sources_checked = mapping.get("sources_checked")
    completion_types_valid = (
        all(
            isinstance(mapping.get(key), str)
            and bool(mapping[key].strip())
            for key in ("deliverable", "status", "central_decision")
        )
        and isinstance(sources_checked, list)
        and bool(sources_checked)
        and all(
            isinstance(source, str) and bool(source.strip())
            for source in sources_checked
        )
    )
    criteria = [
        _criterion(
            task_id,
            "completion_json",
            "The completion response contains a valid JSON object",
            bool(mapping),
            evidence,
            category="structure",
            weight=1,
        ),
        _criterion(
            task_id,
            "completion_json_only",
            (
                "The completion response is JSON without wrapper prose and "
                "uses exactly the four required keys and value types"
            ),
            exact_json and exact_completion_schema and completion_types_valid,
            (
                f"{evidence}; actual_keys={sorted(mapping)!r}; "
                f"required_keys={sorted(required_keys)!r}; "
                f"types_valid={completion_types_valid}"
            ),
            category="structure",
            weight=1,
        ),
        _criterion(
            task_id,
            "completion_path",
            "The completion response identifies the exact required deliverable",
            isinstance(mapping.get("deliverable"), str)
            and _normalize(mapping["deliverable"]) == _normalize(expected_path),
            f"actual={mapping.get('deliverable')!r}; expected={expected_path!r}",
            category="structure",
            weight=3,
        ),
        _criterion(
            task_id,
            "completion_status",
            "The completion response marks the deliverable complete",
            isinstance(mapping.get("status"), str)
            and _normalize(mapping["status"]) == "complete",
            f"actual={mapping.get('status')!r}",
            category="structure",
            weight=1,
        ),
    ]
    return criteria, mapping


def _completion_sources_match(
    completion: Mapping[str, Any],
    authoritative_sources: Iterable[str],
) -> bool:
    checked = completion.get("sources_checked")
    if not isinstance(checked, list):
        return False
    normalized_checked = {
        _normalize(value)
        for value in checked
        if isinstance(value, str) and _normalize(value)
    }
    if len(normalized_checked) < 2:
        return False
    matched_sources: set[int] = set()
    source_list = list(authoritative_sources)
    for value in checked:
        if not isinstance(value, str) or not value.strip():
            continue
        for index, source in enumerate(source_list):
            if _source_present(value, source):
                matched_sources.add(index)
                break
    return len(matched_sources) >= 2


def _completion_central_decision_match(
    completion: Mapping[str, Any],
    decision: str,
    *,
    values: dict[str, Any] | None = None,
    paired_rows: Iterable[Iterable[Any]] = (),
) -> bool:
    actual = completion.get("central_decision")
    return (
        isinstance(actual, str)
        and bool(actual.strip())
        and _decision_concept_match(
            actual,
            decision,
            values=values,
            paired_rows=paired_rows,
        )
    )


def _metric_artifact_criteria(
    task_id: str,
    values: dict[str, Any],
    *,
    text: str,
    paired_rows: list[list[Any]],
    formulas: list[str] | None = None,
    formula_metric_checks: list[dict[str, Any]] | None = None,
    label_aliases: Mapping[str, Iterable[str]] | None = None,
) -> list[Criterion]:
    criteria: list[Criterion] = []
    aliases = _materialize_metric_label_aliases(values, label_aliases)
    _validated_metric_labels(values, aliases)
    for key, expected in values.items():
        slug = _normalize(key).replace(" ", "_")
        labels = _metric_label_options(key, aliases)
        value_present = any(
            _bounded_metric_value(
                paired_rows,
                label,
                expected,
                metric_key=key,
            )
            for label in labels
        )
        paired = any(
            _paired_metric(
                paired_rows,
                label,
                expected,
                metric_key=key,
            )
            for label in labels
        )
        criteria.extend(
            [
                _criterion(
                    task_id,
                    f"{slug}__value",
                    f"The deliverable contains the correct `{key}` value",
                    value_present,
                    (
                        "expected value found"
                        if value_present
                        else f"expected={expected!r}"
                    ),
                    weight=10,
                ),
                _criterion(
                    task_id,
                    f"{slug}__paired",
                    (
                        f"The `{key}` label and correct value are paired in a "
                        "native row, table, or bounded text block"
                    ),
                    paired,
                    "label/value pairing checked",
                    category="auditability",
                    weight=5,
                    failure_cap=0.69,
                ),
            ]
        )
    if formulas is not None:
        lineage_keys = {
            str(check.get("metric_key"))
            for check in (formula_metric_checks or [])
            if check.get("calculated") and check.get("source_backed")
        }
        expected_keys = set(values)
        criteria.append(
            _criterion(
                task_id,
                "metric_formula_lineage",
                (
                    "Every central metric has an exact-labeled calculated "
                    "claim with visible native source lineage"
                ),
                bool(expected_keys) and lineage_keys == expected_keys,
                (
                    f"formula_count={len(formulas)}; "
                    f"lineage_keys={sorted(lineage_keys)!r}; "
                    f"expected_keys={sorted(expected_keys)!r}"
                ),
                category="auditability",
                weight=5,
                failure_cap=0.69,
            )
        )
    return criteria


_NON_CURRENCY_ORDINARY_METRICS = {
    "average age years",
    "days inventory",
    "inventory turns",
    "preventive downtime",
    "priority vendor average lead time days",
    "top asset downtime",
    "unplanned downtime",
}


def _currency_metric_keys(values: dict[str, Any]) -> set[str]:
    """Identify central metrics whose controls must normalize to cents."""

    currency_keys: set[str] = set()
    for key, value in values.items():
        # All integer central values in the task catalog are counts or whole
        # periods. Monetary central values are represented as floats, even
        # when their displayed amount happens to be whole dollars.
        if not isinstance(value, float):
            continue
        normalized = _normalize(key)
        if (
            _is_rate_metric(normalized)
            or _is_quantity_metric(normalized)
            or normalized in _NON_CURRENCY_ORDINARY_METRICS
        ):
            continue
        currency_keys.add(key)
    return currency_keys


def _task_model_quality_criteria(
    task_id: str,
    gold: dict[str, Any],
    evidence: dict[str, Any],
) -> list[Criterion]:
    supported_tasks = {
        "task_068",
        "task_076",
        "task_077",
        "task_086",
        "task_087",
        "task_095",
    }
    if task_id not in supported_tasks:
        return []
    calculated_keys = {
        "task_068": {
            "make_total_cost",
            "buy_total_cost",
            "annual_savings",
            "payback_years",
            "payback_status",
            "three_year_undiscounted_savings",
            "three_year_npv",
        },
        "task_076": {
            "customer_collections",
            "other_inflows",
            "supplier_payments",
            "payroll_and_benefits",
            "interest",
            "capital_expenditures",
            "other_outflows",
            "minimum_cash",
            "minimum_cash_week",
            "ending_cash",
            "maximum_revolver",
        },
        "task_077": {
            "leverage",
            "leverage_headroom",
            "cash",
            "undrawn_revolver",
            "liquidity",
            "compliant",
        },
        "task_086": {
            "base_cash_flow",
            "enterprise_value",
            "equity_value",
        },
        "task_087": {
            "normalized_ebitda",
            "quality_of_earnings_ratio",
        },
        "task_095": {
            "baseline_revenue",
            "baseline_gross_margin",
            "scenario_revenue",
            "scenario_gross_profit",
            "scenario_gross_margin",
            "scenario_operating_income",
            "incremental_working_capital",
        },
    }[task_id] & set(gold["values"])
    formula_cells = evidence.get("formula_cells", {})
    row_labels = evidence.get("formula_row_labels", {})
    formula_results_by_cell = {
        str(item.get("cell")): item.get("value")
        for item in evidence.get("labeled_formula_results", [])
        if item.get("cell")
    }
    calculated_claim_cells: dict[str, set[str]] = {}
    for key in sorted(calculated_keys):
        label = _normalize(key.replace("_", " "))
        calculated_claim_cells[key] = {
            cell
            for cell, formula in formula_cells.items()
            if _normalize(
                row_labels.get(cell, "").replace("_", " ")
            )
            == label
            and _formula_is_recursive_calculation(
                cell,
                formula,
                formula_cells,
            )
            and _value_matches(
                formula_results_by_cell.get(cell),
                gold["values"][key],
                key,
            )
            and _formula_reaches_source(
                cell,
                formula,
                formula_cells,
            )
        }
    claim_lineage = {
        key: bool(cells)
        for key, cells in calculated_claim_cells.items()
    }
    controls_complete, control_details = _central_metric_control_coverage(
        gold["values"],
        formula_cells,
        row_labels,
        evidence.get("literal_cells", {}),
    )
    detail_evidence = evidence.get("detail_schedule_evidence", {})
    expected_rows = int(detail_evidence.get("expected_row_count", 0))
    model_formula_cells = set(
        detail_evidence.get("model_formula_cells", [])
    )
    model_cells = set(
        detail_evidence.get("model_cells", model_formula_cells)
    )
    schedule_required_keys = set(calculated_keys)
    if task_id == "task_086":
        schedule_required_keys.discard("base_cash_flow")
    elif task_id == "task_077":
        schedule_required_keys.difference_update(
            {"cash", "undrawn_revolver", "liquidity"}
        )
    elif task_id == "task_095":
        schedule_required_keys.difference_update(
            {"baseline_revenue", "baseline_gross_margin"}
        )
    calculated_schedule_links = {
        key: any(
            _formula_reaches_targets(
                cell,
                formula_cells[cell],
                formula_cells,
                model_cells,
            )
            for cell in calculated_claim_cells.get(key, set())
        )
        for key in sorted(schedule_required_keys)
    }
    model_rows_formula_backed = (
        expected_rows > 0
        and detail_evidence.get("header_candidate_count") == 1
        and detail_evidence.get("formula_backed_row_count")
        == expected_rows
        and bool(model_formula_cells)
    )
    return [
        _criterion(
            task_id,
            "task_model_source_separation",
            (
                "The model separates calculated claims from independent "
                "source totals and reconciles every central metric"
            ),
            controls_complete,
            (
                "incomplete_controls="
                f"{[item for item in control_details if not item['controlled']][:10]!r}"
            ),
            category="auditability",
            weight=10,
            failure_cap=0.49,
        ),
        _criterion(
            task_id,
            "task_model_schedule_lineage",
            (
                "Calculated central outputs have recursive source lineage "
                "through a semantically discovered, formula-driven schedule"
            ),
            all(claim_lineage.values())
            and model_rows_formula_backed
            and all(calculated_schedule_links.values()),
            (
                f"claim_lineage={claim_lineage!r}; "
                f"calculated_schedule_links={calculated_schedule_links!r}; "
                f"model_rows_formula_backed={model_rows_formula_backed}; "
                f"model_formula_cell_count={len(model_formula_cells)}"
            ),
            category="core_finance",
            weight=10,
            failure_cap=0.49,
        ),
    ]


def _dcf_sensitivity_table_evidence(
    evidence: dict[str, Any],
    gold: dict[str, Any],
) -> dict[str, Any]:
    """Find a DCF sensitivity table by semantic headers, not coordinates."""

    literal_cells = evidence.get("literal_cells", {})
    value_cells = evidence.get("value_cells", {})
    formula_cells = evidence.get("formula_cells", {})
    supporting = gold.get("supporting_values", {})
    sensitivity = supporting.get("terminal_growth_sensitivity", [])
    result = {
        "table_found": False,
        "matched_cases": 0,
        "formula_backed_cases": 0,
        "base_case_matched": False,
    }
    if not isinstance(sensitivity, list) or len(sensitivity) != 5:
        return result

    def parts(qualified: str) -> tuple[str, int, int] | None:
        if "!" not in qualified:
            return None
        sheet, coordinate = qualified.rsplit("!", 1)
        try:
            column, row, _, _ = range_boundaries(coordinate)
        except ValueError:
            return None
        return sheet, row, column

    headers: dict[tuple[str, int], dict[str, int]] = {}
    for qualified, value in literal_cells.items():
        parsed = parts(qualified)
        if parsed is None or not isinstance(value, str):
            continue
        sheet, row, column = parsed
        tokens = set(_normalize(value).split())
        semantic = (
            "terminal_growth"
            if {"terminal", "growth"} <= tokens
            else "enterprise_value"
            if {"enterprise", "value"} <= tokens
            else "equity_value"
            if {"equity", "value"} <= tokens
            else None
        )
        if semantic:
            headers.setdefault((sheet, row), {})[semantic] = column

    for (sheet, header_row), columns in headers.items():
        if set(columns) != {
            "terminal_growth",
            "enterprise_value",
            "equity_value",
        }:
            continue
        result["table_found"] = True
        matched_expected: set[int] = set()
        formula_backed: set[int] = set()
        for row in range(header_row + 1, header_row + 31):
            growth_cell = (
                f"{sheet}!"
                f"{get_column_letter(columns['terminal_growth'])}{row}"
            )
            enterprise_cell = (
                f"{sheet}!"
                f"{get_column_letter(columns['enterprise_value'])}{row}"
            )
            equity_cell = (
                f"{sheet}!"
                f"{get_column_letter(columns['equity_value'])}{row}"
            )
            growth = value_cells.get(
                growth_cell,
                literal_cells.get(growth_cell),
            )
            enterprise = value_cells.get(enterprise_cell)
            equity = value_cells.get(equity_cell)
            for index, expected in enumerate(sensitivity):
                if index in matched_expected:
                    continue
                if not (
                    _close(
                        growth,
                        expected["terminal_growth"],
                        "terminal_growth",
                    )
                    and _close(
                        enterprise,
                        expected["enterprise_value"],
                        "enterprise_value",
                    )
                    and _close(
                        equity,
                        expected["equity_value"],
                        "equity_value",
                    )
                ):
                    continue
                matched_expected.add(index)
                if _close(
                    expected["terminal_growth"],
                    gold["values"]["terminal_growth"],
                    "terminal_growth",
                ):
                    result["base_case_matched"] = (
                        _close(
                            enterprise,
                            gold["values"]["enterprise_value"],
                            "enterprise_value",
                        )
                        and _close(
                            equity,
                            gold["values"]["equity_value"],
                            "equity_value",
                        )
                    )
                enterprise_formula = formula_cells.get(enterprise_cell)
                equity_formula = formula_cells.get(equity_cell)
                if (
                    _formula_is_calculation(enterprise_formula)
                    and _formula_is_calculation(equity_formula)
                    and _formula_reaches_source(
                        enterprise_cell,
                        enterprise_formula,
                        formula_cells,
                    )
                    and _formula_reaches_source(
                        equity_cell,
                        equity_formula,
                        formula_cells,
                    )
                ):
                    formula_backed.add(index)
                break
        result["matched_cases"] = max(
            result["matched_cases"],
            len(matched_expected),
        )
        result["formula_backed_cases"] = max(
            result["formula_backed_cases"],
            len(formula_backed),
        )
    return result


def _task_086_model_criteria(
    task_id: str,
    gold: dict[str, Any],
    evidence: dict[str, Any],
) -> list[Criterion]:
    """Enforce the complete, formula-driven DCF schedule and sensitivity."""

    if task_id != "task_086":
        return []
    supporting = gold.get("supporting_values")
    detail = gold.get("detail_schedule")
    if not isinstance(supporting, dict) or not isinstance(detail, dict):
        return []

    detail_evidence = evidence.get("detail_schedule_evidence")
    if not isinstance(detail_evidence, dict):
        detail_evidence = _detail_presentation_evidence(
            evidence.get("paired_rows", []),
            detail,
        )
    sensitivity_evidence = _dcf_sensitivity_table_evidence(
        evidence,
        gold,
    )
    return [
        _criterion(
            task_id,
            "dcf_full_schedule_truth",
            (
                "All five forecast rows, including full-precision terminal "
                "value and present value, agree with authoritative DCF truth"
            ),
            bool(detail_evidence["field_values_match"]),
            (
                f"row_count={detail_evidence['row_count']}; "
                f"field_mismatch_count="
                f"{detail_evidence['field_mismatch_count']}; "
                f"samples={detail_evidence['field_mismatch_samples']!r}"
            ),
            category="core_finance",
            weight=10,
            failure_cap=0.49,
        ),
        _criterion(
            task_id,
            "dcf_terminal_growth_sensitivity_truth",
            (
                "The native five-case terminal-growth sensitivity ties exact "
                "enterprise and equity values and includes the 2.50% base case"
            ),
            sensitivity_evidence["matched_cases"] == 5
            and sensitivity_evidence["base_case_matched"],
            f"sensitivity_evidence={sensitivity_evidence!r}",
            category="core_finance",
            weight=10,
            failure_cap=0.49,
        ),
        _criterion(
            task_id,
            "dcf_terminal_growth_sensitivity_formulas",
            (
                "Each sensitivity enterprise/equity result is a native "
                "calculation with visible source lineage"
            ),
            sensitivity_evidence["formula_backed_cases"] == 5,
            f"sensitivity_evidence={sensitivity_evidence!r}",
            category="auditability",
            weight=10,
            failure_cap=0.49,
        ),
    ]


def _task_068_085_model_criteria(
    task_id: str,
    gold: dict[str, Any],
    evidence: dict[str, Any],
) -> list[Criterion]:
    """Enforce repaired treasury mechanics without fixed sheet coordinates."""

    if task_id not in {"task_068", "task_076", "task_077"}:
        return []

    formula_cells = evidence.get("formula_cells", {})
    row_labels = evidence.get("formula_row_labels", {})
    literal_cells = evidence.get("literal_cells", {})
    value_cells = evidence.get("value_cells", {})
    paired_rows = evidence.get("paired_rows", [])
    analysis_cells = {
        _normalize(row_labels.get(cell, "").replace("_", " ")): cell
        for cell in formula_cells
        if cell.startswith("Analysis!")
    }

    def analysis_cell(key: str) -> str:
        return analysis_cells.get(
            _normalize(key.replace("_", " ")),
            "",
        )

    def profile(key: str, *, allowed_zero: bool = False) -> dict[str, Any]:
        cell = analysis_cell(key)
        return _capital_dependency_profile(
            [cell] if cell else [],
            formula_cells,
            literal_zero_cells={cell} if cell and allowed_zero else set(),
        )

    def parse_cell(cell: str) -> tuple[str, int, int] | None:
        if "!" not in cell:
            return None
        sheet, address = cell.rsplit("!", 1)
        try:
            column, row, maximum_column, maximum_row = range_boundaries(
                address
            )
        except (TypeError, ValueError):
            return None
        if column != maximum_column or row != maximum_row:
            return None
        return sheet, column, row

    def find_table(
        required_headers: set[str],
    ) -> tuple[str, int, dict[str, int]] | None:
        candidates: dict[tuple[str, int], dict[str, int]] = {}
        for cell, value in literal_cells.items():
            parsed = parse_cell(cell)
            if parsed is None:
                continue
            sheet, column, row = parsed
            header = _normalize(str(value).replace("_", " "))
            if header not in required_headers:
                continue
            candidates.setdefault((sheet, row), {})[header] = column
        matches = [
            (sheet, row, columns)
            for (sheet, row), columns in candidates.items()
            if required_headers <= set(columns)
        ]
        if not matches:
            return None
        return max(matches, key=lambda item: len(item[2]))

    def qualified(sheet: str, column: int, row: int) -> str:
        return f"{sheet}!{get_column_letter(column)}{row}"

    def reaches_cell(model_profile: dict[str, Any], cell: str) -> bool:
        return cell in model_profile.get("references", set())

    def reaches_column(
        model_profile: dict[str, Any],
        sheet: str,
        column: int,
    ) -> bool:
        for reference in model_profile.get("references", set()):
            parsed = parse_cell(reference)
            if parsed is None:
                continue
            reference_sheet, reference_column, _row = parsed
            if (
                _normalize(reference_sheet) == _normalize(sheet)
                and reference_column == column
            ):
                return True
        return False

    def recalculated_results(
        keys: tuple[str, ...],
    ) -> tuple[bool, dict[str, Any]]:
        checks: dict[str, Any] = {}
        populated = 0
        matched = True
        for key in keys:
            cell = analysis_cell(key)
            actual = value_cells.get(cell) if cell else None
            available = actual is not None
            populated += int(available)
            key_matched = (
                _value_matches(actual, gold["values"][key], key)
                if available
                else None
            )
            if available:
                matched = matched and bool(key_matched)
            checks[key] = {
                "cell": cell,
                "available": available,
                "matched": key_matched,
            }
        # Development hosts may not expose recalculated formula caches. Once
        # any target result is present, require the complete set and exact
        # authoritative values; release grading recalculates before this check.
        return (
            (populated == 0 or (populated == len(keys) and matched)),
            checks,
        )

    def profile_text(model_profile: dict[str, Any]) -> str:
        return _normalize(" ".join(model_profile.get("formula_texts", [])))

    if task_id == "task_068":
        required_keys = (
            "payback_years",
            "payback_status",
            "three_year_undiscounted_savings",
        )
        if not set(required_keys) <= set(gold["values"]):
            return []
        schedule_headers = {
            "year",
            "annual savings",
            "transition cost",
            "net cash flow",
        }
        table = find_table(schedule_headers)
        schedule_zero_cells = (
            {
                cell
                for cell, formula in formula_cells.items()
                if cell.startswith(f"{table[0]}!")
                and re.sub(r"\s+", "", str(formula)) in {
                    "=0",
                    "=+0",
                    "=-0",
                }
            }
            if table is not None
            else set()
        )
        payback_profile = _capital_dependency_profile(
            [analysis_cell("payback_years")],
            formula_cells,
            literal_zero_cells=schedule_zero_cells,
        )
        status_profile = _capital_dependency_profile(
            [analysis_cell("payback_status")],
            formula_cells,
            literal_zero_cells=schedule_zero_cells,
        )
        undiscounted_profile = _capital_dependency_profile(
            [analysis_cell("three_year_undiscounted_savings")],
            formula_cells,
            literal_zero_cells=schedule_zero_cells,
        )
        result_match, result_checks = recalculated_results(required_keys)
        schedule_found = table is not None
        payback_dependencies = status_dependencies = False
        undiscounted_dependencies = False
        if table is not None:
            sheet, header_row, columns = table
            savings_column = columns["annual savings"]
            transition_column = columns["transition cost"]
            cash_flow_column = columns["net cash flow"]
            payback_dependencies = (
                reaches_column(payback_profile, sheet, savings_column)
                and (
                    reaches_column(
                        payback_profile,
                        sheet,
                        transition_column,
                    )
                    or reaches_column(
                        payback_profile,
                        sheet,
                        cash_flow_column,
                    )
                )
            )
            status_dependencies = (
                reaches_column(status_profile, sheet, savings_column)
                and (
                    reaches_column(
                        status_profile,
                        sheet,
                        transition_column,
                    )
                    or reaches_column(
                        status_profile,
                        sheet,
                        cash_flow_column,
                    )
                )
            )
            horizon_cells = {
                qualified(sheet, cash_flow_column, row)
                for row in range(header_row + 1, header_row + 5)
            }
            cash_flow_horizon = horizon_cells <= set(
                undiscounted_profile.get("references", set())
            )
            savings_transition_horizon = (
                reaches_column(
                    undiscounted_profile,
                    sheet,
                    savings_column,
                )
                and reaches_column(
                    undiscounted_profile,
                    sheet,
                    transition_column,
                )
                and _capital_profile_uses_expected_literal(
                    undiscounted_profile,
                    3,
                )
                and "*" in undiscounted_profile.get("operators", set())
                and "-" in undiscounted_profile.get("operators", set())
            )
            undiscounted_dependencies = (
                cash_flow_horizon or savings_transition_horizon
            )
        conditional_functions = {
            "CHOOSE",
            "IF",
            "IFS",
            "INDEX",
            "LOOKUP",
            "SWITCH",
            "XLOOKUP",
        }
        payback_text = profile_text(payback_profile)
        status_text = profile_text(status_profile)
        recovery_logic = (
            schedule_found
            and payback_profile["clean"]
            and status_profile["clean"]
            and bool(
                conditional_functions & payback_profile["functions"]
            )
            and bool(
                conditional_functions & status_profile["functions"]
            )
            and payback_dependencies
            and status_dependencies
            and "no payback" in payback_text
            and "no recovery" in status_text
            and "recovered" in status_text
            and "not recovered" in status_text
            and all(
                (
                    cell := analysis_cell(key)
                )
                and _formula_reaches_source(
                    cell,
                    formula_cells.get(cell, ""),
                    formula_cells,
                )
                for key in ("payback_years", "payback_status")
            )
            and result_match
        )
        undiscounted_logic = (
            schedule_found
            and undiscounted_profile["clean"]
            and undiscounted_dependencies
            and (
                cell := analysis_cell(
                    "three_year_undiscounted_savings"
                )
            )
            and _formula_reaches_source(
                cell,
                formula_cells.get(cell, ""),
                formula_cells,
            )
            and result_match
        )
        return [
            _criterion(
                task_id,
                "make_buy_payback_recovery_logic",
                (
                    "Payback and recovery status are formula-backed, and "
                    "non-positive annual savings return No payback/no recovery"
                ),
                recovery_logic,
                (
                    f"schedule_found={schedule_found}; "
                    f"payback_clean={payback_profile['clean']}; "
                    f"status_clean={status_profile['clean']}; "
                    f"payback_dependencies={payback_dependencies}; "
                    f"status_dependencies={status_dependencies}; "
                    f"result_checks={result_checks!r}"
                ),
                category="core_finance",
                weight=10,
                failure_cap=0.49,
            ),
            _criterion(
                task_id,
                "make_buy_undiscounted_horizon",
                (
                    "Three-year undiscounted savings formula includes year 0 "
                    "through year 3 net cash flow"
                ),
                undiscounted_logic,
                (
                    f"schedule_found={schedule_found}; "
                    f"clean={undiscounted_profile['clean']}; "
                    f"dependencies={undiscounted_dependencies}; "
                    f"result_checks={result_checks!r}"
                ),
                category="core_finance",
                weight=10,
                failure_cap=0.49,
            ),
        ]

    if task_id == "task_076":
        required_keys = {"other_inflows", "other_outflows"}
        if not required_keys <= set(gold["values"]):
            return []
        detail = gold.get("detail_schedule", {})
        row_count = len(detail.get("rows", []))
        schedule_headers = {
            _normalize(column.replace("_", " "))
            for column in detail.get("columns", [])
        }
        table = find_table(schedule_headers)
        pre_financing_met = repayment_met = closing_met = False
        cumulative_met = summary_lanes_met = False
        result_match, result_checks = recalculated_results(
            (
                "other_inflows",
                "other_outflows",
                "minimum_cash",
                "ending_cash",
                "maximum_revolver",
            )
        )
        if table is not None and row_count:
            sheet, header_row, columns = table
            first_row = header_row + 1
            last_row = header_row + row_count
            initial_repayment_cell = qualified(
                sheet,
                columns["revolver repayment"],
                first_row,
            )
            allowed_zero_cells = {initial_repayment_cell}
            lane_headers = (
                "opening cash",
                "customer collections",
                "other inflows",
                "supplier payments",
                "payroll and benefits",
                "interest",
                "capital expenditures",
                "other outflows",
            )
            pre_financing_met = all(
                (
                    model_profile := _capital_dependency_profile(
                        [
                            qualified(
                                sheet,
                                columns["pre financing cash"],
                                row,
                            )
                        ],
                        formula_cells,
                        literal_zero_cells=allowed_zero_cells,
                    )
                )["clean"]
                and all(
                    reaches_cell(
                        model_profile,
                        qualified(sheet, columns[header], row),
                    )
                    for header in lane_headers
                )
                for row in range(first_row, last_row + 1)
            )
            repayment_met = all(
                (
                    model_profile := _capital_dependency_profile(
                        [
                            qualified(
                                sheet,
                                columns["revolver repayment"],
                                row,
                            )
                        ],
                        formula_cells,
                        literal_zero_cells=allowed_zero_cells,
                    )
                )["clean"]
                and (
                    row == first_row
                    or (
                        {"MIN", "MAX"} <= model_profile["functions"]
                        and reaches_cell(
                            model_profile,
                            qualified(
                                sheet,
                                columns["cumulative revolver"],
                                row - 1,
                            ),
                        )
                        and reaches_cell(
                            model_profile,
                            qualified(
                                sheet,
                                columns["pre financing cash"],
                                row,
                            ),
                        )
                        and reaches_cell(
                            model_profile,
                            qualified(
                                sheet,
                                columns["minimum cash buffer"],
                                row,
                            ),
                        )
                    )
                )
                for row in range(first_row, last_row + 1)
            )
            closing_met = all(
                (
                    model_profile := _capital_dependency_profile(
                        [
                            qualified(
                                sheet,
                                columns["closing cash"],
                                row,
                            )
                        ],
                        formula_cells,
                        literal_zero_cells=allowed_zero_cells,
                    )
                )["clean"]
                and all(
                    reaches_cell(
                        model_profile,
                        qualified(sheet, columns[header], row),
                    )
                    for header in (
                        "pre financing cash",
                        "revolver draw",
                        "revolver repayment",
                    )
                )
                for row in range(first_row, last_row + 1)
            )
            cumulative_met = all(
                (
                    model_profile := _capital_dependency_profile(
                        [
                            qualified(
                                sheet,
                                columns["cumulative revolver"],
                                row,
                            )
                        ],
                        formula_cells,
                        literal_zero_cells=allowed_zero_cells,
                    )
                )["clean"]
                and all(
                    reaches_cell(
                        model_profile,
                        qualified(sheet, columns[header], row),
                    )
                    for header in (
                        "revolver draw",
                        "revolver repayment",
                    )
                )
                and (
                    row == first_row
                    or reaches_cell(
                        model_profile,
                        qualified(
                            sheet,
                            columns["cumulative revolver"],
                            row - 1,
                        ),
                    )
                )
                for row in range(first_row, last_row + 1)
            )
            summary_lanes_met = all(
                (
                    model_profile := profile(key)
                )["clean"]
                and reaches_column(
                    model_profile,
                    sheet,
                    columns[_normalize(key.replace("_", " "))],
                )
                for key in required_keys
            )
        supporting = gold.get("supporting_values")
        debt_profile_met = (
            isinstance(supporting, dict)
            and all(
                any(
                    _paired_metric(paired_rows, label, supporting[key])
                    for label in labels
                )
                for key, labels in (
                    ("revolver_outstanding", ("revolver outstanding", "revolver")),
                    (
                        "equipment_loan_outstanding",
                        ("equipment loan outstanding", "equipment loan"),
                    ),
                    ("term_loan_outstanding", ("term loan outstanding", "term loan")),
                )
            )
            and all(
                _source_present(evidence.get("text", ""), source)
                for source in (
                    "Treasury/2025 ABL Facility Summary.pdf",
                    "Treasury/Equipment Facility Summary.pdf",
                    "Treasury/2022 Sponsor Term Loan Summary.pdf",
                )
            )
        )
        return [
            _criterion(
                task_id,
                "cash_forecast_complete_rollforward",
                (
                    "Every weekly row shows all inflow/outflow lanes and "
                    "formula-rolls pre-financing, draw, repayment, closing "
                    "cash, and cumulative revolver"
                ),
                pre_financing_met
                and repayment_met
                and closing_met
                and cumulative_met
                and summary_lanes_met
                and result_match,
                (
                    f"schedule_found={table is not None}; "
                    f"pre_financing={pre_financing_met}; "
                    f"repayment={repayment_met}; closing={closing_met}; "
                    f"cumulative={cumulative_met}; "
                    f"summary_lanes={summary_lanes_met}; "
                    f"result_checks={result_checks!r}"
                ),
                category="core_finance",
                weight=10,
                failure_cap=0.49,
            ),
            _criterion(
                task_id,
                "cash_forecast_active_debt_profile",
                (
                    "The package shows all three active debt instruments with "
                    "outstanding principal and executed-source citations"
                ),
                debt_profile_met,
                "complete active debt profile checked",
                category="provenance",
                weight=10,
                failure_cap=0.49,
            ),
        ]

    required_keys = {"cash", "undrawn_revolver", "liquidity"}
    if not required_keys <= set(gold["values"]):
        return []
    component_profiles = {
        key: profile(key)
        for key in required_keys
    }
    cash_references = set(component_profiles["cash"]["references"])
    undrawn_references = set(
        component_profiles["undrawn_revolver"]["references"]
    )
    liquidity_references = set(
        component_profiles["liquidity"]["references"]
    )
    cash_terminals = {
        reference
        for reference in cash_references
        if reference not in formula_cells
    }
    undrawn_terminals = {
        reference
        for reference in undrawn_references
        if reference not in formula_cells
    }
    liquidity_terminals = {
        reference
        for reference in liquidity_references
        if reference not in formula_cells
    }
    result_match, result_checks = recalculated_results(
        ("cash", "undrawn_revolver", "liquidity")
    )
    bridge_met = (
        all(model_profile["clean"] for model_profile in component_profiles.values())
        and bool(cash_terminals)
        and bool(undrawn_terminals)
        and cash_terminals <= liquidity_terminals
        and undrawn_terminals <= liquidity_terminals
        and bool(
            {"SUM"} & component_profiles["liquidity"]["functions"]
            or "+"
            in component_profiles["liquidity"].get("operators", set())
        )
        and all(
            (
                cell := analysis_cell(key)
            )
            and _formula_reaches_source(
                cell,
                formula_cells.get(cell, ""),
                formula_cells,
            )
            for key in required_keys
        )
        and result_match
    )
    return [
        _criterion(
            task_id,
            "liquidity_component_bridge",
            (
                "Cash and undrawn revolver are shown separately and "
                "formula-summed to total liquidity"
            ),
            bridge_met,
            (
                f"cash_terminals={sorted(cash_terminals)!r}; "
                f"undrawn_terminals={sorted(undrawn_terminals)!r}; "
                f"liquidity_terminals={sorted(liquidity_terminals)!r}; "
                f"result_checks={result_checks!r}"
            ),
            category="core_finance",
            weight=10,
            failure_cap=0.49,
        )
    ]


def _capital_formula_tokens(formula: Any) -> tuple[Any, ...]:
    """Tokenize a capital-model formula, failing closed on malformed input."""

    if not isinstance(formula, str) or not formula.startswith("="):
        return ()
    try:
        return tuple(Tokenizer(formula).items)
    except Exception:
        return ()


def _capital_formula_edges(
    cell: str,
    formula: Any,
) -> tuple[set[str], set[str]]:
    """Return bounded cell references and referenced sheets for a formula."""

    current_sheet = cell.rsplit("!", 1)[0]
    references: set[str] = set()
    sheets: set[str] = set()
    for token in _capital_formula_tokens(formula):
        if token.type != "OPERAND" or token.subtype != "RANGE":
            continue
        raw = str(token.value).replace("$", "")
        if "!" in raw:
            sheet, address = raw.rsplit("!", 1)
            sheet = sheet.strip("'").replace("''", "'")
        else:
            sheet, address = current_sheet, raw
        sheets.add(sheet)
        if re.fullmatch(r"[A-Za-z]{1,3}[1-9]\d*", address):
            references.add(f"{sheet}!{address.upper()}")
            continue
        try:
            min_col, min_row, max_col, max_row = range_boundaries(address)
        except (TypeError, ValueError):
            continue
        if not all(
            isinstance(value, int)
            for value in (min_col, min_row, max_col, max_row)
        ):
            continue
        area = (max_col - min_col + 1) * (max_row - min_row + 1)
        if area > 2_048:
            continue
        for row in range(min_row, max_row + 1):
            for column in range(min_col, max_col + 1):
                references.add(
                    f"{sheet}!{get_column_letter(column)}{row}"
                )
    return references, sheets


def _capital_formula_is_tainted(
    formula: Any,
    *,
    allow_literal_zero: bool = False,
) -> bool:
    """Reject constant/dead branches and literal overrides posing as lineage."""

    tokens = _capital_formula_tokens(formula)
    if not tokens:
        return True
    compact = re.sub(r"\s+", "", str(formula)).upper()
    has_reference = any(
        token.type == "OPERAND" and token.subtype == "RANGE"
        for token in tokens
    )
    has_function = any(
        token.type == "FUNC" and token.subtype == "OPEN"
        for token in tokens
    )
    if (
        not has_reference
        and not has_function
        and not (allow_literal_zero and compact in {"=0", "=+0", "=-0"})
    ):
        return True
    constant_condition_patterns = (
        r"(?:^|[^A-Z])IF\((?:TRUE|FALSE|0|1)(?:,|;)",
        (
            r"(?:^|[^A-Z])IF\([+-]?\d+(?:\.\d+)?"
            r"(?:=|<>|<=|>=|<|>)[+-]?\d+(?:\.\d+)?(?:,|;)"
        ),
        r"(?:^|[^A-Z])IF\(NOT\((?:TRUE|FALSE)\)(?:,|;)",
        r"(?:^|[^A-Z])IFS\((?:TRUE|FALSE|0|1)(?:,|;)",
    )
    if any(re.search(pattern, compact) for pattern in constant_condition_patterns):
        return True
    if re.search(r"\*0(?:\.0+)?(?=[+\-),;]|$)", compact):
        return True
    if re.search(r"(?<=[=(,+;\-])0(?:\.0+)?\*", compact):
        return True
    return False


def _capital_dependency_profile(
    start_cells: Iterable[str],
    formula_cells: dict[str, str],
    *,
    literal_zero_cells: set[str] | None = None,
) -> dict[str, Any]:
    """Trace a bounded formula graph and summarize its model semantics."""

    allowed_zero = literal_zero_cells or set()
    stack = list(dict.fromkeys(start_cells))
    visited: set[str] = set()
    references: set[str] = set()
    sheets: set[str] = set()
    functions: set[str] = set()
    operators: set[str] = set()
    numeric_literals: list[float] = []
    formula_texts: list[str] = []
    clean = bool(stack)
    while stack and len(visited) < 4_096:
        cell = stack.pop()
        if cell in visited:
            continue
        visited.add(cell)
        formula = formula_cells.get(cell)
        if formula is None:
            continue
        formula_texts.append(str(formula))
        clean = clean and not _capital_formula_is_tainted(
            formula,
            allow_literal_zero=cell in allowed_zero,
        )
        tokens = _capital_formula_tokens(formula)
        functions.update(
            token.value.rstrip("(").upper()
            for token in tokens
            if token.type == "FUNC" and token.subtype == "OPEN"
        )
        operators.update(
            token.value
            for token in tokens
            if token.type.startswith("OPERATOR")
        )
        for token in tokens:
            if token.type != "OPERAND" or token.subtype != "NUMBER":
                continue
            try:
                numeric_literals.append(float(token.value))
            except (TypeError, ValueError):
                pass
        cell_references, cell_sheets = _capital_formula_edges(cell, formula)
        references.update(cell_references)
        sheets.update(cell_sheets)
        stack.extend(
            reference
            for reference in cell_references
            if reference in formula_cells and reference not in visited
        )
    return {
        "clean": clean and not stack,
        "visited": visited,
        "references": references,
        "sheets": sheets,
        "functions": functions,
        "operators": operators,
        "numeric_literals": numeric_literals,
        "formula_texts": formula_texts,
    }


def _capital_profile_reaches(
    profile: dict[str, Any],
    sheet: str,
) -> bool:
    expected = _normalize(sheet)
    return any(
        _normalize(candidate) == expected
        for candidate in profile["sheets"]
    )


def _capital_profile_uses_expected_literal(
    profile: dict[str, Any],
    expected: Any,
) -> bool:
    if not isinstance(expected, (int, float)) or isinstance(expected, bool):
        return False
    return any(
        math.isclose(
            abs(literal),
            abs(float(expected)),
            rel_tol=1e-12,
            abs_tol=1e-9,
        )
        for literal in profile["numeric_literals"]
    )


def _capital_labeled_formula_cells(
    sheet: str | None,
    label: str,
    formula_cells: dict[str, str],
    row_labels: dict[str, str],
) -> list[str]:
    wanted = _normalize(label.replace("_", " "))
    return [
        cell
        for cell in formula_cells
        if (sheet is None or cell.rsplit("!", 1)[0] == sheet)
        and _normalize(
            row_labels.get(cell, "").replace("_", " ")
        )
        == wanted
    ]


def _capital_calculated_results_match(
    evidence: dict[str, Any],
    values: dict[str, Any],
    candidate_cells: dict[str, Iterable[str]],
) -> tuple[bool, list[dict[str, Any]]]:
    """Validate native calculated values when the host exposes cached results."""

    non_result_tokens = {
        "check",
        "delta",
        "difference",
        "reconciliation",
        "status",
        "tolerance",
        "variance",
    }
    checks: list[dict[str, Any]] = []
    all_match = True
    for key, cells in candidate_cells.items():
        wanted = _normalize(key.replace("_", " "))
        allowed_cells = set(cells)
        populated = [
            result
            for result in evidence.get("labeled_formula_results", [])
            if str(result.get("cell", "")) in allowed_cells
            and _normalize(
                str(result.get("label", "")).replace("_", " ")
            )
            == wanted
            and not (
                set(_normalize(result.get("column_header")).split())
                & non_result_tokens
            )
            and result.get("value") is not None
        ]
        matched = all(
            _value_matches(result.get("value"), values[key], key)
            for result in populated
        )
        all_match = all_match and matched
        checks.append(
            {
                "metric_key": key,
                "result_count": len(populated),
                "matched": matched,
            }
        )
    return all_match, checks


def _capital_sheets_with_labels(
    labels: Iterable[str],
    formula_cells: dict[str, str],
    row_labels: dict[str, str],
) -> list[str]:
    """Discover formula-table sheets from their semantic row populations."""

    wanted = {_normalize(label) for label in labels}
    found: dict[str, set[str]] = {}
    for cell in formula_cells:
        label = _normalize(row_labels.get(cell, ""))
        if label not in wanted:
            continue
        sheet = cell.rsplit("!", 1)[0]
        found.setdefault(sheet, set()).add(label)
    return sorted(
        sheet
        for sheet, present in found.items()
        if wanted <= present
    )


def _capital_profile_reaches_external_sheet(
    profile: dict[str, Any],
    excluded_sheets: Iterable[str],
) -> bool:
    excluded = {_normalize(sheet) for sheet in excluded_sheets if sheet}
    return any(
        _normalize(sheet) not in excluded
        for sheet in profile["sheets"]
    )


def _capital_independent_control_coverage(
    values: dict[str, Any],
    formula_cells: dict[str, str],
    row_labels: dict[str, str],
) -> tuple[bool, list[dict[str, Any]]]:
    """Find two-sided result-versus-source controls through full lineage."""

    return _central_metric_control_coverage(
        values,
        formula_cells,
        row_labels,
    )


def _capital_literal_rows_by_sheet(
    evidence: dict[str, Any],
) -> dict[str, list[list[Any]]]:
    """Reconstruct native literal rows for sheet-scoped semantic discovery."""

    cells_by_row: dict[tuple[str, int], dict[int, Any]] = {}
    for qualified, value in evidence.get("literal_cells", {}).items():
        if "!" not in qualified:
            continue
        sheet, coordinate = qualified.rsplit("!", 1)
        try:
            column, row, _, _ = range_boundaries(coordinate)
        except (TypeError, ValueError):
            continue
        if not isinstance(column, int) or not isinstance(row, int):
            continue
        cells_by_row.setdefault((sheet, row), {})[column] = value
    rows_by_sheet: dict[str, list[list[Any]]] = {}
    for (sheet, _row), values in sorted(cells_by_row.items()):
        if not values:
            continue
        rows_by_sheet.setdefault(sheet, []).append(
            [
                values.get(column)
                for column in range(min(values), max(values) + 1)
            ]
        )
    return rows_by_sheet


def _capital_sensitivity_case(label: Any) -> str | None:
    raw = str(label or "").casefold()
    normalized = _normalize(raw)
    if "base" in normalized or "baseline" in normalized:
        return "base"
    benefit = any(
        term in normalized
        for term in (
            "annual benefit",
            "annual savings",
            "net benefit",
            "saving",
        )
    )
    rate = "discount rate" in normalized or "wacc" in normalized
    down = (
        "-" in raw
        or any(
            term in normalized
            for term in ("down", "downside", "decrease", "lower", "reduction")
        )
        or (benefit and "90" in normalized)
    )
    up = (
        "+" in raw
        or any(
            term in normalized
            for term in ("up", "upside", "increase", "higher")
        )
        or (benefit and "110" in normalized)
    )
    if benefit and down:
        return "benefit_down"
    if benefit and up:
        return "benefit_up"
    if rate and down:
        return "rate_down"
    if rate and up:
        return "rate_up"
    return None


def _capital_sensitivity_shift_met(
    case: str,
    profile: dict[str, Any],
) -> bool:
    if case == "base":
        return True
    # The public task requires decision-useful sensitivities but does not
    # prescribe a hidden +/-10% or +/-1-point magnitude. Require a genuine
    # directional adjustment in the source-backed formula graph.
    return bool(profile.get("clean") and profile.get("references"))


def _capital_governance_evidence(rows: Iterable[Iterable[Any]]) -> bool:
    """Require governance concepts in separate, structured native rows."""

    normalized_rows = [
        [_normalize(value) for value in row if value not in (None, "")]
        for row in rows
    ]
    normalized_rows = [row for row in normalized_rows if len(row) >= 2]

    def joined(row: list[str]) -> str:
        return " ".join(row)

    row_texts = [joined(row) for row in normalized_rows]
    all_text = " ".join(row_texts)
    header_met = (
        (
            any(
                "alternative" in text and "disposition" in text
                for text in row_texts
            )
            and any(
                ("risk" in text or "measure" in text)
                and "evidence" in text
                and ("owner" in text or "timing" in text)
                for text in row_texts
            )
        )
        or any(
            "governance" in text
            and ("evidence" in text or "requirement" in text)
            and ("owner" in text or "executive" in text)
            and ("timing" in text or "milestone" in text)
            for text in row_texts
        )
        or any(
            "alternative" in text
            and ("risk" in text or "measure" in text)
            and "disposition" in text
            and "evidence" in text
            and ("owner" in text or "executive" in text)
            and ("timing" in text or "milestone" in text)
            for text in row_texts
        )
    )
    owner_terms = (
        "vp operations",
        "vice president operations",
        "operations vice president",
    )
    date_terms = (
        "2026 07 31",
        "july 31 2026",
        "31 july 2026",
    )
    alternatives_met = (
        any(
            any(
                term in text
                for term in (
                    "status quo",
                    "baseline",
                    "current state",
                    "no investment",
                )
            )
            for text in row_texts
        )
        and any(
            "automation" in text or "automated" in text
            for text in row_texts
        )
        and "dayton" in all_text
        and "200" in all_text
        and any(term in all_text for term in owner_terms)
        and any(term in all_text for term in date_terms)
    )
    risk_rows = row_texts
    commissioning_risk = any(
        any(
            term in text
            for term in ("commission", "installation", "startup", "ramp")
        )
        and any(
            term in text
            for term in ("disruption", "capacity", "delivery", "production")
        )
        for text in risk_rows
    )
    benefit_risk = any(
        any(term in text for term in ("benefit", "saving", "adoption"))
        and any(
            term in text
            for term in (
                "realization",
                "realized",
                "below",
                "performance",
                "shortfall",
            )
        )
        for text in risk_rows
    )
    audit_rows = [
        text
        for text in row_texts
        if any(term in text for term in ("audit", "review", "measure"))
    ]
    post_audit_met = any(
        any(term in text for term in ("saving", "benefit"))
        and (
            "maintenance" in text
            or "operating cost" in text
            or "ongoing cost" in text
            or "capacity" in text
        )
        and any(
            term in text
            for term in ("after", "month", "in service", "post implementation")
        )
        for text in audit_rows
    )
    return (
        header_met
        and alternatives_met
        and commissioning_risk
        and benefit_risk
        and post_audit_met
    )


def _capital_model_formula_row_labels(
    task_id: str,
    evidence: Mapping[str, Any],
    formula_cells: Mapping[str, str],
    existing_labels: Mapping[str, str],
) -> dict[str, str]:
    """Propagate explicit table row keys for the capital-model checks.

    The general workbook extractor intentionally labels ordinary formulas only
    from adjacent text.  In the task 066/067 model tables, however, one row key
    governs several formulas to its right: Project ID, Combination ID, Case,
    or Scenario.  Limit the wider lookup to those task-declared key headers so
    unrelated distant text cannot become formula-lineage evidence.
    """

    allowed_headers = {
        "task_066": {"project id", "combination id"},
        "task_067": {"case", "scenario"},
    }.get(task_id, set())
    labels = {
        str(cell): str(label)
        for cell, label in existing_labels.items()
    }
    if not allowed_headers:
        return labels

    literal_cells = evidence.get("literal_cells", {})
    if not isinstance(literal_cells, Mapping):
        return labels
    key_headers: list[tuple[str, int, int]] = []
    for qualified, value in literal_cells.items():
        if (
            not isinstance(qualified, str)
            or "!" not in qualified
            or _normalize(value) not in allowed_headers
        ):
            continue
        sheet, coordinate = qualified.rsplit("!", 1)
        try:
            minimum_column, minimum_row, maximum_column, maximum_row = (
                range_boundaries(coordinate)
            )
        except (TypeError, ValueError):
            continue
        if (
            minimum_column == maximum_column
            and minimum_row == maximum_row
        ):
            key_headers.append((sheet, minimum_row, minimum_column))

    for qualified in formula_cells:
        if not isinstance(qualified, str) or "!" not in qualified:
            continue
        sheet, coordinate = qualified.rsplit("!", 1)
        try:
            column, row, maximum_column, maximum_row = range_boundaries(
                coordinate
            )
        except (TypeError, ValueError):
            continue
        if column != maximum_column or row != maximum_row:
            continue
        candidates: list[tuple[int, str]] = []
        for header_sheet, header_row, key_column in key_headers:
            if (
                header_sheet != sheet
                or row <= header_row
                or column <= key_column
            ):
                continue
            row_key = literal_cells.get(
                f"{sheet}!{get_column_letter(key_column)}{row}"
            )
            if (
                isinstance(row_key, str)
                and row_key.strip()
                and not row_key.startswith("=")
            ):
                candidates.append((header_row, row_key.strip()))
        if candidates:
            labels[qualified] = max(
                candidates,
                key=lambda candidate: candidate[0],
            )[1]
    return labels


def _task_066_067_model_criteria(
    task_id: str,
    gold: dict[str, Any],
    evidence: dict[str, Any],
) -> list[Criterion]:
    """Enforce dependency-backed capital mechanics for tasks 066 and 067."""

    if task_id not in {"task_066", "task_067"}:
        return []

    formula_cells = {
        str(cell): str(formula)
        for cell, formula in evidence.get("formula_cells", {}).items()
    }
    row_labels = _capital_model_formula_row_labels(
        task_id,
        evidence,
        formula_cells,
        evidence.get("formula_row_labels", {}),
    )
    controls_complete, control_coverage = (
        _capital_independent_control_coverage(
            gold["values"],
            formula_cells,
            row_labels,
        )
    )
    common = [
        _criterion(
            task_id,
            "capital_model_source_separation",
            (
                "Every central capital metric has a live result-versus-"
                "benchmark reconciliation with an independent literal source"
            ),
            controls_complete,
            f"control_coverage={control_coverage!r}",
            category="auditability",
            weight=10,
            failure_cap=0.49,
        )
    ]

    def coordinate_row(cell: str) -> int | None:
        match = re.search(r"![A-Za-z]{1,3}([1-9]\d*)$", cell)
        return int(match.group(1)) if match else None

    def sheet_formula_count(sheet: str) -> int:
        return sum(
            cell.rsplit("!", 1)[0] == sheet
            for cell in formula_cells
        )

    def candidate_sheet(labels: Iterable[str]) -> str:
        candidates = _capital_sheets_with_labels(
            labels,
            formula_cells,
            row_labels,
        )
        return max(
            candidates,
            key=sheet_formula_count,
            default="",
        )

    if task_id == "task_066":
        detail = gold.get("detail_schedule", {})
        detail_rows = detail.get("rows", [])
        project_ids = [str(row[0]) for row in detail_rows]
        project_count = len(project_ids)
        combination_count = 1 << project_count
        combination_ids = [
            f"COMBO-{mask:03d}"
            for mask in range(combination_count)
        ]
        project_sheet = candidate_sheet(project_ids)
        combination_sheet = candidate_sheet(combination_ids)
        paired_rows = evidence.get("paired_rows", [])
        combination_checks: dict[str, bool] = {}
        combination_profiles: list[dict[str, Any]] = []
        for mask, combination_id in enumerate(combination_ids):
            cells = _capital_labeled_formula_cells(
                combination_sheet or None,
                combination_id,
                formula_cells,
                row_labels,
            )
            profiles = [
                _capital_dependency_profile([cell], formula_cells)
                for cell in cells
            ]
            combination_profiles.extend(profiles)
            cell_rows = {
                row
                for cell in cells
                if (row := coordinate_row(cell)) is not None
            }
            expected_flags = tuple(
                1 if mask & (1 << offset) else 0
                for offset in range(project_count)
            )
            observed_flags = {
                tuple(
                    _number(value)
                    for value in row[1 : project_count + 1]
                )
                for row in paired_rows
                if row
                and _normalize(row[0]) == _normalize(combination_id)
                and len(row) >= project_count + 1
            }
            flags_match = observed_flags == {expected_flags}
            direct_references: set[str] = set()
            for cell in cells:
                references, _ = _capital_formula_edges(
                    cell,
                    formula_cells[cell],
                )
                direct_references.update(references)
            literal_same_row_references = {
                reference
                for reference in direct_references
                if reference.rsplit("!", 1)[0] == combination_sheet
                and reference not in formula_cells
                and coordinate_row(reference) in cell_rows
            }
            portfolio_calculations = [
                profile
                for profile in profiles
                if _capital_profile_reaches(profile, project_sheet)
                and (
                    "*" in profile["operators"]
                    or bool(
                        profile["functions"]
                        & {"SUMIF", "SUMIFS", "SUMPRODUCT"}
                    )
                )
            ]
            budget_check = any(
                _capital_profile_reaches_external_sheet(
                    profile,
                    {combination_sheet, project_sheet},
                )
                and bool(
                    profile["operators"]
                    & {"<", "<=", "=", "<>", ">=", ">"}
                )
                for profile in profiles
            )
            combination_checks[combination_id] = (
                bool(project_sheet)
                and bool(combination_sheet)
                and len(cells) >= 4
                and len(cell_rows) == 1
                and flags_match
                and all(profile["clean"] for profile in profiles)
                and len(literal_same_row_references) >= project_count
                and len(portfolio_calculations) >= 2
                and budget_check
            )

        lookup_functions = {
            "FILTER",
            "INDEX",
            "LOOKUP",
            "MATCH",
            "TAKE",
            "XLOOKUP",
            "XMATCH",
        }
        aggregate_functions = {
            "COUNTIF",
            "COUNTIFS",
            "FILTER",
            "INDEX",
            "MATCH",
            "SUMIF",
            "SUMIFS",
            "SUMPRODUCT",
            "XLOOKUP",
            "XMATCH",
        }
        winner_functions = {
            "LARGE",
            "MAX",
            "MAXIFS",
            "RANK",
            "RANK.EQ",
            "SORT",
            "SORTBY",
        }
        decision_cells: list[str] = []
        decision_profiles: list[dict[str, Any]] = []
        budget_cells: list[str] = []
        budget_profiles: list[dict[str, Any]] = []
        decision_row_checks: dict[str, bool] = {}
        budget_row_checks: dict[str, bool] = {}
        for project_id in project_ids:
            project_cells = _capital_labeled_formula_cells(
                project_sheet or None,
                project_id,
                formula_cells,
                row_labels,
            )
            valid_decisions: list[tuple[str, dict[str, Any]]] = []
            valid_budgets: list[tuple[str, dict[str, Any]]] = []
            for cell in project_cells:
                formula = formula_cells[cell]
                profile = _capital_dependency_profile(
                    [cell],
                    formula_cells,
                )
                tokens = _capital_formula_tokens(formula)
                direct_functions = {
                    token.value.rstrip("(").upper()
                    for token in tokens
                    if token.type == "FUNC" and token.subtype == "OPEN"
                }
                direct_operators = {
                    token.value
                    for token in tokens
                    if token.type.startswith("OPERATOR")
                }
                profile_text = _normalize(
                    " ".join(profile["formula_texts"])
                )
                if (
                    profile["clean"]
                    and _capital_profile_reaches(
                        profile,
                        combination_sheet,
                    )
                    and bool(profile["functions"] & lookup_functions)
                    and "select" in profile_text
                    and "defer" in profile_text
                    and not (
                        direct_functions
                        & {
                            "COUNTIF",
                            "COUNTIFS",
                            "SUMIF",
                            "SUMIFS",
                            "SUMPRODUCT",
                        }
                    )
                    and "-" not in direct_operators
                ):
                    valid_decisions.append((cell, profile))
                if (
                    profile["clean"]
                    and "-" in profile["operators"]
                    and _capital_profile_reaches_external_sheet(
                        profile,
                        {project_sheet, combination_sheet},
                    )
                    and (
                        _capital_profile_reaches(profile, project_sheet)
                        or _capital_profile_reaches(
                            profile,
                            combination_sheet,
                        )
                    )
                    and bool(
                        profile["functions"]
                        & {
                            "COUNTIF",
                            "COUNTIFS",
                            "SUMIF",
                            "SUMIFS",
                            "SUMPRODUCT",
                        }
                    )
                ):
                    valid_budgets.append((cell, profile))
            decision_row_checks[project_id] = len(valid_decisions) == 1
            budget_row_checks[project_id] = len(valid_budgets) == 1
            if len(valid_decisions) == 1:
                decision_cells.append(valid_decisions[0][0])
                decision_profiles.append(valid_decisions[0][1])
            if len(valid_budgets) == 1:
                budget_cells.append(valid_budgets[0][0])
                budget_profiles.append(valid_budgets[0][1])

        winner_logic = any(
            bool(profile["functions"] & winner_functions)
            for profile in combination_profiles + decision_profiles
        )
        enumeration_met = (
            project_count > 0
            and all(combination_checks.values())
            and winner_logic
        )
        decision_lineage = (
            len(decision_cells) == project_count
            and all(decision_row_checks.values())
        )
        budget_lineage = (
            len(budget_cells) == project_count
            and all(budget_row_checks.values())
        )
        optimized_decisions_met = (
            decision_lineage and winner_logic and budget_lineage
        )

        selection_cells: dict[str, list[str]] = {}
        selection_lineage: dict[str, bool] = {}
        for key in (
            "selected_project_count",
            "selected_investment",
            "selected_npv",
            "remaining_budget",
        ):
            valid: list[str] = []
            for cell in _capital_labeled_formula_cells(
                None,
                key,
                formula_cells,
                row_labels,
            ):
                profile = _capital_dependency_profile(
                    [cell],
                    formula_cells,
                )
                reaches_model = (
                    _capital_profile_reaches(profile, project_sheet)
                    or _capital_profile_reaches(
                        profile,
                        combination_sheet,
                    )
                )
                if key == "remaining_budget":
                    semantic = (
                        "-" in profile["operators"]
                        and _capital_profile_reaches_external_sheet(
                            profile,
                            {project_sheet, combination_sheet},
                        )
                        and bool(
                            profile["functions"] & aggregate_functions
                        )
                        and not _capital_profile_uses_expected_literal(
                            profile,
                            gold["values"][key],
                        )
                    )
                else:
                    semantic = (
                        bool(profile["functions"] & aggregate_functions)
                        and (
                            key == "selected_project_count"
                            or not _capital_profile_uses_expected_literal(
                                profile,
                                gold["values"][key],
                            )
                        )
                    )
                if profile["clean"] and reaches_model and semantic:
                    valid.append(cell)
            selection_cells[key] = valid
            selection_lineage[key] = bool(valid)
        result_match, result_checks = _capital_calculated_results_match(
            evidence,
            gold["values"],
            selection_cells,
        )
        selection_met = all(selection_lineage.values()) and result_match
        return common + [
            _criterion(
                task_id,
                "portfolio_combination_model",
                (
                    "Every row-bound combination ID carries the correct mask "
                    "and clean formulas for investment, NPV, feasibility, "
                    "and the winning objective"
                ),
                enumeration_met,
                (
                    f"project_sheet={project_sheet!r}; "
                    f"combination_sheet={combination_sheet!r}; "
                    f"combination_checks={combination_checks!r}; "
                    f"winner_logic={winner_logic}"
                ),
                category="core_finance",
                weight=10,
                failure_cap=0.49,
            ),
            _criterion(
                task_id,
                "portfolio_optimized_decisions",
                (
                    "Project decisions follow the calculated winning "
                    "combination, and remaining budget subtracts the "
                    "formula-selected investment from the authorized budget"
                ),
                optimized_decisions_met,
                (
                    f"decision_rows={decision_row_checks!r}; "
                    f"budget_rows={budget_row_checks!r}; "
                    f"winner_logic={winner_logic}"
                ),
                category="decision",
                weight=10,
                failure_cap=0.49,
            ),
            _criterion(
                task_id,
                "portfolio_selection_lineage",
                (
                    "Selected count, investment, NPV, and remaining budget "
                    "have clean optimized dependencies and agree with "
                    "available native calculated results"
                ),
                selection_met,
                (
                    f"selection_lineage={selection_lineage!r}; "
                    f"result_checks={result_checks!r}"
                ),
                category="core_finance",
                weight=10,
                failure_cap=0.49,
            ),
        ]

    detail_rows = gold.get("detail_schedule", {}).get("rows", [])
    expected_case_counts = Counter(
        _normalize(row[0])
        for row in detail_rows
    )
    schedule_candidates = _capital_sheets_with_labels(
        expected_case_counts,
        formula_cells,
        row_labels,
    )
    schedule_sheet = max(
        schedule_candidates,
        key=sheet_formula_count,
        default="",
    )
    schedule_groups: dict[tuple[str, int], list[str]] = {}
    for cell in formula_cells:
        if cell.rsplit("!", 1)[0] != schedule_sheet:
            continue
        case = _normalize(row_labels.get(cell, ""))
        row_number = coordinate_row(cell)
        if case in expected_case_counts and row_number is not None:
            schedule_groups.setdefault((case, row_number), []).append(cell)
    literal_zero_cells = {
        cell
        for (case, _row), cells in schedule_groups.items()
        if all(
            _number(row[2]) == 0
            for row in detail_rows
            if _normalize(row[0]) == case
        )
        for cell in cells
        if re.sub(r"\s+", "", formula_cells[cell]).upper()
        in {"=0", "=+0", "=-0"}
    }
    zero_cases = {
        case
        for case in expected_case_counts
        if all(
            _number(row[2]) == 0
            for row in detail_rows
            if _normalize(row[0]) == case
        )
    }
    schedule_checks: dict[str, bool] = {}
    cash_flow_kinds: Counter[tuple[str, str]] = Counter()
    for (case, row_number), cells in schedule_groups.items():
        row_met = False
        row_cash_kind = ""
        for present_value_cell in cells:
            present_value_profile = _capital_dependency_profile(
                [present_value_cell],
                formula_cells,
                literal_zero_cells=literal_zero_cells,
            )
            if not (
                present_value_profile["clean"]
                and (
                    "*" in present_value_profile["operators"]
                    or "PRODUCT" in present_value_profile["functions"]
                )
            ):
                continue
            local_references = {
                reference
                for reference in present_value_profile["references"]
                if reference in cells
            }
            for discount_cell in local_references:
                discount_profile = _capital_dependency_profile(
                    [discount_cell],
                    formula_cells,
                    literal_zero_cells=literal_zero_cells,
                )
                direct_discount_references, _ = _capital_formula_edges(
                    discount_cell,
                    formula_cells[discount_cell],
                )
                year_references = {
                    reference
                    for reference in direct_discount_references
                    if reference.rsplit("!", 1)[0] == schedule_sheet
                    and coordinate_row(reference) == row_number
                    and reference not in formula_cells
                }
                discount_met = (
                    discount_profile["clean"]
                    and _capital_profile_reaches_external_sheet(
                        discount_profile,
                        {schedule_sheet},
                    )
                    and bool(year_references)
                    and (
                        {"/", "^"} <= discount_profile["operators"]
                        or "POWER" in discount_profile["functions"]
                    )
                )
                if not discount_met:
                    continue
                for cash_cell in local_references - {discount_cell}:
                    cash_profile = _capital_dependency_profile(
                        [cash_cell],
                        formula_cells,
                        literal_zero_cells=literal_zero_cells,
                    )
                    external_literals = {
                        reference
                        for reference in cash_profile["references"]
                        if reference.rsplit("!", 1)[0] != schedule_sheet
                        and reference not in formula_cells
                    }
                    compact_cash = re.sub(
                        r"\s+",
                        "",
                        formula_cells[cash_cell],
                    ).upper()
                    if (
                        case in zero_cases
                        and compact_cash in {"=0", "=+0", "=-0"}
                    ):
                        cash_kind = "zero"
                    elif (
                        cash_profile["clean"]
                        and "-" in cash_profile["operators"]
                        and len(external_literals) == 1
                    ):
                        cash_kind = "initial"
                    elif (
                        cash_profile["clean"]
                        and "-" in cash_profile["operators"]
                        and len(external_literals) >= 2
                    ):
                        cash_kind = "benefit"
                    else:
                        continue
                    if case in zero_cases and cash_kind != "zero":
                        continue
                    if case not in zero_cases and cash_kind == "zero":
                        continue
                    row_met = True
                    row_cash_kind = cash_kind
                    break
                if row_met:
                    break
            if row_met:
                break
        schedule_checks[f"{case}@{row_number}"] = row_met
        if row_met:
            cash_flow_kinds[(case, row_cash_kind)] += 1

    actual_case_counts = Counter(
        case
        for case, _row in schedule_groups
    )
    cash_flow_population_met = all(
        (
            cash_flow_kinds[(case, "zero")] == count
            if case in zero_cases
            else (
                cash_flow_kinds[(case, "initial")] == 1
                and cash_flow_kinds[(case, "benefit")] == count - 1
            )
        )
        for case, count in expected_case_counts.items()
    )

    calculated_cells: dict[str, list[str]] = {}
    calculated_lineage: dict[str, bool] = {}
    metric_functions = {
        "npv": {"NPV", "SUM", "SUMPRODUCT", "XNPV"},
        "irr": {"IRR", "XIRR"},
        "payback_years": {"LOOKUP", "MATCH", "XLOOKUP", "XMATCH"},
        "profitability_index": set(),
    }
    for key in ("npv", "irr", "payback_years", "profitability_index"):
        valid: list[str] = []
        for cell in _capital_labeled_formula_cells(
            None,
            key,
            formula_cells,
            row_labels,
        ):
            profile = _capital_dependency_profile(
                [cell],
                formula_cells,
                literal_zero_cells=literal_zero_cells,
            )
            if key == "payback_years":
                semantic = (
                    "/" in profile["operators"]
                    or bool(profile["functions"] & metric_functions[key])
                )
            elif key == "profitability_index":
                semantic = "/" in profile["operators"]
            else:
                semantic = bool(
                    profile["functions"] & metric_functions[key]
                )
            if (
                profile["clean"]
                and _capital_profile_reaches(profile, schedule_sheet)
                and semantic
                and not _capital_profile_uses_expected_literal(
                    profile,
                    gold["values"][key],
                )
            ):
                valid.append(cell)
        calculated_cells[key] = valid
        calculated_lineage[key] = bool(valid)
    calculated_results_match, calculated_result_checks = (
        _capital_calculated_results_match(
            evidence,
            gold["values"],
            calculated_cells,
        )
    )
    schedule_lineage_met = (
        bool(schedule_sheet)
        and actual_case_counts == expected_case_counts
        and all(schedule_checks.values())
        and cash_flow_population_met
        and all(calculated_lineage.values())
        and calculated_results_match
    )

    # Task 067 discloses a status-quo/investment comparison and a year 0-7
    # cash-flow/PV schedule, but it does not require an extra 16-row table with
    # a repeated Case column. Accept the unique authoritative detail schedule
    # already validated above, then require all four valuation outputs to
    # traverse that schedule's native formula cells.
    detail_evidence = evidence.get("detail_schedule_evidence", {})
    if (
        isinstance(detail_evidence, Mapping)
        and detail_evidence.get("headers_match")
        and detail_evidence.get("keys_match")
        and detail_evidence.get("field_values_match")
        and detail_evidence.get("row_count")
        == int(gold["values"]["life_years"]) + 1
        and detail_evidence.get("formula_backed_row_count")
        == int(gold["values"]["life_years"]) + 1
    ):
        selected_location = str(
            detail_evidence.get("selected_location") or ""
        )
        semantic_schedule_sheet = selected_location.split("!", 1)[0]
        schedule_targets = {
            str(cell)
            for cell in detail_evidence.get("model_formula_cells", [])
            if str(cell) in formula_cells
        }
        semantic_calculated_cells: dict[str, list[str]] = {}
        semantic_lineage: dict[str, bool] = {}
        for key in ("npv", "irr", "payback_years", "profitability_index"):
            candidates: list[str] = []
            for cell in _capital_labeled_formula_cells(
                None,
                key,
                formula_cells,
                row_labels,
            ):
                profile = _capital_dependency_profile([cell], formula_cells)
                if (
                    profile["clean"]
                    and _formula_reaches_targets(
                        cell,
                        formula_cells[cell],
                        formula_cells,
                        schedule_targets,
                    )
                    and not _capital_profile_uses_expected_literal(
                        profile,
                        gold["values"][key],
                    )
                ):
                    candidates.append(cell)
            semantic_calculated_cells[key] = candidates
            semantic_lineage[key] = bool(candidates)
        semantic_results_match, semantic_result_checks = (
            _capital_calculated_results_match(
                evidence,
                gold["values"],
                semantic_calculated_cells,
            )
        )
        schedule_lineage_met = (
            bool(semantic_schedule_sheet)
            and len(schedule_targets)
            >= 2 * (int(gold["values"]["life_years"]) + 1)
            and all(semantic_lineage.values())
            and semantic_results_match
        )
        if schedule_lineage_met:
            schedule_sheet = semantic_schedule_sheet
            calculated_lineage = semantic_lineage
            calculated_result_checks = semantic_result_checks

    sensitivity_by_sheet: dict[str, dict[str, list[str]]] = {}
    for cell in formula_cells:
        case = _capital_sensitivity_case(row_labels.get(cell, ""))
        if case is None:
            continue
        sheet = cell.rsplit("!", 1)[0]
        sensitivity_by_sheet.setdefault(sheet, {}).setdefault(
            case,
            [],
        ).append(cell)
    required_sensitivity_cases = {
        "base",
        "benefit_down",
        "benefit_up",
        "rate_down",
        "rate_up",
    }
    sensitivity_sheet = max(
        (
            sheet
            for sheet, cases in sensitivity_by_sheet.items()
            if required_sensitivity_cases <= set(cases)
        ),
        key=sheet_formula_count,
        default="",
    )
    sensitivity_cells = sensitivity_by_sheet.get(
        sensitivity_sheet,
        {case: [] for case in required_sensitivity_cases},
    )
    sensitivity_checks: dict[str, bool] = {}
    for case in sorted(required_sensitivity_cases):
        cells = sensitivity_cells.get(case, [])
        row_numbers = {
            row
            for cell in cells
            if (row := coordinate_row(cell)) is not None
        }
        profile = _capital_dependency_profile(
            cells,
            formula_cells,
            literal_zero_cells=literal_zero_cells,
        )
        valuation_met = False
        decision_met = False
        for cell in cells:
            cell_profile = _capital_dependency_profile(
                [cell],
                formula_cells,
                literal_zero_cells=literal_zero_cells,
            )
            local_formula_references = {
                reference
                for reference in cell_profile["references"]
                if reference.rsplit("!", 1)[0] == sensitivity_sheet
                and reference in formula_cells
                and coordinate_row(reference) in row_numbers
            }
            valuation_met = valuation_met or (
                (
                    len(local_formula_references) >= 2
                    or bool(
                        cell_profile["functions"]
                        & {"NPV", "PV", "SUMPRODUCT", "XNPV"}
                    )
                )
                and (
                    bool(cell_profile["operators"] & {"/", "^"})
                    or bool(
                        cell_profile["functions"]
                        & {"NPV", "PV", "SUMPRODUCT", "XNPV"}
                    )
                )
                and _capital_profile_reaches_external_sheet(
                    cell_profile,
                    {sensitivity_sheet},
                )
            )
            normalized_formula = _normalize(formula_cells[cell])
            decision_met = decision_met or (
                "IF" in cell_profile["functions"]
                and bool(
                    cell_profile["operators"]
                    & {"<", "<=", "=", "<>", ">=", ">"}
                )
                and any(
                    term in normalized_formula
                    for term in (
                        "approve",
                        "negative npv",
                        "positive npv",
                        "reject",
                        "proceed",
                        "do not proceed",
                    )
                )
            )
        sensitivity_checks[case] = (
            len(cells) >= 4
            and len(row_numbers) == 1
            and profile["clean"]
            and valuation_met
            and decision_met
            and _capital_sensitivity_shift_met(case, profile)
        )
    sensitivity_met = (
        bool(sensitivity_sheet)
        and all(sensitivity_checks.values())
    )

    literal_rows_by_sheet = _capital_literal_rows_by_sheet(evidence)
    governance_sheets = [
        sheet
        for sheet, rows in literal_rows_by_sheet.items()
        if _capital_governance_evidence(rows)
    ]
    implementation_met = bool(governance_sheets)
    return common + [
        _criterion(
            task_id,
            "business_case_schedule_lineage",
            (
                "Each discovered schedule row has clean input-to-cash-flow-"
                "to-present-value dependencies, and NPV, IRR, payback, and "
                "PI derive from that schedule and calculated results"
            ),
            schedule_lineage_met,
            (
                f"schedule_sheet={schedule_sheet!r}; "
                f"schedule_checks={schedule_checks!r}; "
                f"cash_flow_population={cash_flow_kinds!r}; "
                f"calculated_lineage={calculated_lineage!r}; "
                f"result_checks={calculated_result_checks!r}"
            ),
            category="core_finance",
            weight=10,
            failure_cap=0.49,
        ),
        _criterion(
            task_id,
            "business_case_sensitivity",
            (
                "Row-bound base, benefit-up/down, and rate-up/down cases "
                "recalculate NPV and a decision from the source-backed model"
            ),
            sensitivity_met,
            (
                f"sensitivity_sheet={sensitivity_sheet!r}; "
                f"sensitivity_checks={sensitivity_checks!r}"
            ),
            category="core_finance",
            weight=10,
            failure_cap=0.49,
        ),
        _criterion(
            task_id,
            "business_case_implementation_governance",
            (
                "Structured implementation rows compare alternatives, state "
                "commissioning and benefit risks, assign Operations "
                "leadership and the July 31 gate, and define a post-audit"
            ),
            implementation_met,
            f"governance_sheets={governance_sheets!r}",
            category="decision",
            weight=10,
            failure_cap=0.49,
        ),
    ]


def _xlsx_criteria(
    task_id: str,
    answer: Any,
    workspace_root: Path,
    gold: dict[str, Any],
) -> list[Criterion]:
    target = gold["artifact"]["path"]
    path = workspace_root / target
    criteria, completion = _artifact_answer_criteria(task_id, answer, target)
    readable = False
    evidence: dict[str, Any] = {
        "sheetnames": [],
        "text": "",
        "values": [],
        "formulas": [],
        "formula_cells": {},
        "formula_formats": {},
        "formula_row_labels": {},
        "literal_cells": {},
        "value_cells": {},
        "labeled_formula_results": [],
        "paired_rows": [],
        "errors": ["missing"],
        "charts": 0,
        "merged_ranges": 0,
        "visual_hygiene_issues": ["missing"],
        "native_formula_issues": ["missing"],
        "native_chart_issues": ["missing"],
        "decision_blocks": [],
        "hardcoded_control_status_cells": ["missing"],
    }
    error = "required workbook is missing"
    artifact_present = os.path.lexists(path)
    register_evidence = _empty_action_register_evidence(
        "task does not require an action register"
    )
    detail_specification = gold.get("detail_schedule")
    detail_evidence = _empty_detail_schedule_evidence(
        "task does not require a detail schedule"
    )
    if artifact_present:
        try:
            evidence = _workbook_evidence(path, workspace_root)
            readable = True
            error = "workbook parsed"
        except NativePackageSafetyError as exc:
            error = f"native package safety failure: {exc}"
        except Exception as exc:
            error = f"{type(exc).__name__}: {exc}"
    if task_id == "task_028":
        supporting_values = gold.get("supporting_values")
        headers = (
            supporting_values.get("action_register_columns")
            if isinstance(supporting_values, dict)
            else None
        )
        expected_rows = (
            supporting_values.get("action_register_rows")
            if isinstance(supporting_values, dict)
            else None
        )
        if (
            readable
            and isinstance(headers, list)
            and all(isinstance(header, str) for header in headers)
            and isinstance(expected_rows, list)
            and all(isinstance(row, dict) for row in expected_rows)
        ):
            try:
                register_evidence = _planning_action_register_evidence(
                    path,
                    headers=headers,
                    expected_rows=expected_rows,
                    workspace_root=workspace_root,
                )
            except NativePackageSafetyError as exc:
                register_evidence = _empty_action_register_evidence(
                    f"native package safety failure: {exc}"
                )
            except Exception as exc:
                register_evidence = _empty_action_register_evidence(
                    f"{type(exc).__name__}: {exc}"
                )
        else:
            register_evidence = _empty_action_register_evidence(
                "workbook or authoritative register truth is unavailable"
            )
            if isinstance(expected_rows, list):
                register_evidence["expected_row_count"] = len(
                    expected_rows
                )
    if (
        artifact_present
        and readable
        and isinstance(detail_specification, dict)
    ):
        try:
            detail_evidence = _detail_schedule_evidence(
                path,
                detail_specification,
                workspace_root=workspace_root,
            )
        except NativePackageSafetyError as exc:
            detail_evidence = _empty_detail_schedule_evidence(
                f"native package safety failure: {exc}"
            )
        except Exception as exc:
            detail_evidence = _empty_detail_schedule_evidence(
                f"{type(exc).__name__}: {exc}"
            )
        detail_evidence["expected_row_count"] = len(
            detail_specification.get("rows", [])
        )
    evidence["detail_schedule_evidence"] = detail_evidence
    criteria.extend(
        [
            _criterion(
                task_id,
                "artifact_exists",
                "The exact required workbook exists",
                artifact_present,
                target,
                category="integrity",
                weight=10,
            ),
            _criterion(
                task_id,
                "artifact_readable",
                "The workbook parses as a genuine XLSX",
                readable,
                error,
                category="integrity",
                weight=10,
            ),
            _criterion(
                task_id,
                "required_sheet_set",
                "The workbook contains Read Me, Inputs, Analysis, and Control",
                all(
                    name in evidence["sheetnames"]
                    for name in ("Read Me", "Inputs", "Analysis", "Control")
                ),
                f"sheets={evidence['sheetnames']!r}",
                category="structure",
                weight=5,
                failure_cap=0.69,
            ),
        ]
    )
    for name in ("Read Me", "Inputs", "Analysis", "Control"):
        criteria.append(
            _criterion(
                task_id,
                f"sheet_{_normalize(name).replace(' ', '_')}",
                f"The `{name}` worksheet is present",
                name in evidence["sheetnames"],
                f"sheets={evidence['sheetnames']!r}",
                category="structure",
                weight=1,
            )
        )
    if task_id == "task_028":
        register_summary = _action_register_evidence_summary(
            register_evidence
        )
        criteria.extend(
            [
                _criterion(
                    task_id,
                    "action_register_headers",
                    (
                        "A native worksheet contains the exact required "
                        "action-register header sequence"
                    ),
                    register_evidence["header_candidate_count"] >= 1,
                    register_summary,
                    category="structure",
                    weight=5,
                    failure_cap=0.49,
                ),
                _criterion(
                    task_id,
                    "action_register_filterable",
                    (
                        "The complete action register has a native table or "
                        "AutoFilter covering every required column and row"
                    ),
                    bool(register_evidence["filterable"]),
                    register_summary,
                    category="structure",
                    weight=3,
                    failure_cap=0.69,
                ),
                _criterion(
                    task_id,
                    "action_register_row_count",
                    (
                        "The action register contains exactly the complete "
                        "4,000-row planned-order action population"
                    ),
                    (
                        register_evidence["row_count"]
                        == register_evidence["expected_row_count"]
                        == 4_000
                    ),
                    register_summary,
                    category="core_finance",
                    weight=10,
                    failure_cap=0.49,
                ),
                _criterion(
                    task_id,
                    "action_register_unique_planned_orders",
                    (
                        "All 4,000 expected planned-order IDs appear exactly "
                        "once, with no missing, duplicate, or fabricated IDs"
                    ),
                    (
                        register_evidence["unique_id_count"]
                        == register_evidence["expected_unique_id_count"]
                        == 4_000
                        and register_evidence["duplicate_id_count"] == 0
                        and register_evidence["missing_id_count"] == 0
                        and register_evidence["extra_id_count"] == 0
                    ),
                    register_summary,
                    category="core_finance",
                    weight=10,
                    failure_cap=0.49,
                ),
                _criterion(
                    task_id,
                    "action_register_field_values",
                    (
                        "Every required field in every action-register row "
                        "matches authoritative planned-order truth"
                    ),
                    bool(register_evidence["field_values_match"]),
                    register_summary,
                    category="core_finance",
                    weight=10,
                    failure_cap=0.49,
                ),
                _criterion(
                    task_id,
                    "action_register_sort_order",
                    (
                        "The complete register is sorted by priority, "
                        "required date, and planned-order ID"
                    ),
                    bool(register_evidence["sort_order_match"]),
                    register_summary,
                    category="auditability",
                    weight=5,
                    failure_cap=0.49,
                ),
            ]
        )
    if isinstance(detail_specification, dict):
        detail_summary = (
            f"selected_location={detail_evidence['selected_location']!r}; "
            f"header_candidate_count="
            f"{detail_evidence['header_candidate_count']}; "
            f"header_candidate_samples="
            f"{detail_evidence['header_candidate_samples']!r}; "
            f"headers_match={detail_evidence['headers_match']}; "
            f"row_count={detail_evidence['row_count']}; "
            f"expected_row_count={detail_evidence['expected_row_count']}; "
            f"duplicate_key_count={detail_evidence['duplicate_key_count']}; "
            f"missing_keys={detail_evidence['missing_keys']!r}; "
            f"extra_keys={detail_evidence['extra_keys']!r}; "
            f"field_values_match={detail_evidence['field_values_match']}; "
            f"field_mismatch_count="
            f"{detail_evidence['field_mismatch_count']}; "
            f"field_mismatch_samples="
            f"{detail_evidence['field_mismatch_samples']!r}; "
            f"error={detail_evidence['error']!r}"
        )
        criteria.extend(
            [
                _criterion(
                    task_id,
                    "detail_schedule_sheet",
                    (
                        "A native worksheet contains one discoverable "
                        "detail-schedule table"
                    ),
                    detail_evidence["header_candidate_count"] == 1,
                    detail_summary,
                    category="structure",
                    weight=5,
                    failure_cap=0.69,
                ),
                _criterion(
                    task_id,
                    "detail_schedule_headers",
                    (
                        "The unique detail schedule uses every required "
                        "semantic header, in any column order"
                    ),
                    bool(detail_evidence["headers_match"]),
                    detail_summary,
                    category="structure",
                    weight=5,
                    failure_cap=0.49,
                ),
                _criterion(
                    task_id,
                    "detail_schedule_population",
                    (
                        "The detail schedule contains the complete "
                        "authoritative population with no duplicate, missing, "
                        "or fabricated dimension keys"
                    ),
                    (
                        detail_evidence["row_count"]
                        == detail_evidence["expected_row_count"]
                        and bool(detail_evidence["keys_match"])
                        and detail_evidence["duplicate_key_count"] == 0
                    ),
                    detail_summary,
                    category="core_finance",
                    weight=10,
                    failure_cap=0.49,
                ),
                _criterion(
                    task_id,
                    "detail_schedule_field_values",
                    (
                        "Every field in every authoritative detail-schedule "
                        "row agrees with recalculated truth"
                    ),
                    bool(detail_evidence["field_values_match"]),
                    detail_summary,
                    category="core_finance",
                    weight=10,
                    failure_cap=0.49,
                ),
            ]
        )
    _controls_complete, central_control_details = (
        _central_metric_control_coverage(
            gold["values"],
            evidence["formula_cells"],
            evidence["formula_row_labels"],
            evidence["literal_cells"],
        )
    )
    independently_controlled_keys = {
        str(detail["metric_key"])
        for detail in central_control_details
        if detail.get("controlled")
    }
    formula_metrics_consistent, formula_metric_checks = (
        _formula_metric_claim_consistency(
            evidence["labeled_formula_results"],
            gold["values"],
            evidence["formula_cells"],
            independently_controlled_keys=independently_controlled_keys,
        )
    )
    criteria.extend(
        _metric_artifact_criteria(
            task_id,
            gold["values"],
            text=evidence["text"],
            paired_rows=evidence["paired_rows"],
            formulas=evidence["formulas"],
            formula_metric_checks=formula_metric_checks,
        )
    )
    criteria.append(
        _criterion(
            task_id,
            "formula_metric_consistency",
            (
                "Every central metric has an exact-labeled source-backed "
                "native calculation that agrees with truth"
            ),
            formula_metrics_consistent,
            f"checks={formula_metric_checks[:20]!r}",
            category="core_finance",
            weight=10,
            failure_cap=0.49,
        )
    )
    if task_id == "task_007" and isinstance(
        gold.get("supporting_values"),
        dict,
    ):
        claim_checks: list[dict[str, Any]] = []
        consistent = True
        for label, key in (
            ("negative residual order count", "negative_wip_residual_count"),
            ("negative residual value", "negative_wip_residual_value"),
        ):
            claim_consistent, checks = _optional_labeled_claim_consistency(
                evidence["paired_rows"],
                label=label,
                expected=gold["supporting_values"][key],
                metric_key=key,
                label_aliases=(
                    "negative residual total",
                    "negative WIP exceptions",
                    "negative WIP residuals",
                ),
            )
            consistent = consistent and claim_consistent
            claim_checks.extend(checks)
        criteria.append(
            _criterion(
                task_id,
                "negative_wip_residual_claims",
                (
                    "Any quantified negative WIP residual count or value "
                    "reconciles to the atomic open-order WIP population"
                ),
                consistent,
                f"claim_checks={claim_checks!r}",
                category="core_finance",
                weight=10,
                failure_cap=0.49,
            )
        )
    if gold["bundle"] == "make_buy":
        consistent, repeated = _labeled_formula_result_consistency(
            evidence["labeled_formula_results"],
            "Annual savings",
            gold["values"]["annual_savings"],
            "annual_savings",
        )
        criteria.append(
            _criterion(
                task_id,
                "annual_savings_formula_consistency",
                (
                    "Every formula-driven Annual savings result reconciles "
                    "to make total cost less buy total cost"
                ),
                consistent,
                f"labeled_formula_results={repeated!r}",
                category="auditability",
                weight=10,
                failure_cap=0.49,
            )
        )
    formula_count = len(evidence["formulas"])
    controls_complete, control_coverage = (
        _central_metric_control_coverage(
            gold["values"],
            evidence["formula_cells"],
            evidence["formula_row_labels"],
            evidence["literal_cells"],
        )
    )
    meaningful_control_formula_cells = sorted(
        {
            str(candidate["cell"])
            for detail in control_coverage
            for candidate in detail["candidates"]
            if set(candidate["roles"])
            == {"analysis", "independent_source"}
        }
    )
    currency_keys = _currency_metric_keys(gold["values"])
    expected_currency_labels = {
        _normalize(key.replace("_", " ")): key for key in currency_keys
    }
    currency_status_cells = {
        key
        for key, formula in evidence["formula_cells"].items()
        if _normalize(
            evidence["formula_row_labels"].get(key, "").replace("_", " ")
        )
        in expected_currency_labels
        and "IF(" in re.sub(r"\s+", "", formula).upper()
        and "ABS(" in re.sub(r"\s+", "", formula).upper()
    }
    checked_controls, precision_issues = _unrounded_currency_control_issues(
        evidence["formula_cells"],
        evidence["formula_formats"],
        evidence["literal_cells"],
        only_cells=currency_status_cells,
    )
    controlled_currency_keys = {
        expected_currency_labels[
            _normalize(
                evidence["formula_row_labels"].get(cell, "").replace("_", " ")
            )
        ]
        for cell in checked_controls
        if _normalize(
            evidence["formula_row_labels"].get(cell, "").replace("_", " ")
        )
        in expected_currency_labels
    }
    currency_controls_complete = (
        not currency_keys
        or controlled_currency_keys == currency_keys
    )
    # Keep the public rubric denominator stable even when an artifact is
    # missing or contains no monetary zero/tolerance comparisons. The rollout
    # auditor reconstructs criterion metadata without a submission; emitting
    # this row only when formulas happen to be present makes valid workbook
    # traces look like they used an undeclared rubric.
    criteria.append(
        _criterion(
            task_id,
            "currency_precision_controls",
            (
                "Every monetary central metric has a native zero/tolerance "
                "control that rounds differences to cents before PASS/FAIL"
            ),
            currency_controls_complete and not precision_issues,
            (
                f"currency_keys={sorted(currency_keys)!r}; "
                f"controlled_currency_keys="
                f"{sorted(controlled_currency_keys)!r}; "
                f"checked={checked_controls!r}; "
                f"issues={precision_issues[:10]!r}"
            ),
            category="auditability",
            weight=10,
            failure_cap=0.49,
        )
    )
    criteria.extend(
        [
            _criterion(
                task_id,
                "formula_volume",
                "The workbook contains substantive formula-driven analysis",
                formula_count >= max(8, len(gold["value_keys"])),
                f"formula_count={formula_count}",
                category="auditability",
                weight=5,
                failure_cap=0.69,
            ),
            _criterion(
                task_id,
                "formula_errors",
                "The workbook contains no visible formula errors",
                not evidence["errors"],
                f"errors={evidence['errors'][:10]!r}",
                category="auditability",
                weight=10,
                failure_cap=0.69,
            ),
            _criterion(
                task_id,
                "native_chart",
                (
                    "The workbook contains a native chart whose series values "
                    "align with its category ranges"
                ),
                evidence["charts"] >= 1
                and not evidence["native_chart_issues"],
                (
                    f"chart_count={evidence['charts']}; "
                    f"issues={evidence['native_chart_issues'][:10]!r}"
                ),
                category="structure",
                weight=3,
            ),
            _criterion(
                task_id,
                "native_formula_results",
                (
                    "Native recalculation produces usable values in merged "
                    "summary-card formulas"
                ),
                not evidence["native_formula_issues"],
                f"issues={evidence['native_formula_issues'][:10]!r}",
                category="auditability",
                weight=10,
                failure_cap=0.69,
            ),
            _criterion(
                task_id,
                "visual_hygiene",
                (
                    "The workbook avoids obvious header/footer control-code "
                    "corruption, chart-label collisions, and clipped text"
                ),
                not evidence["visual_hygiene_issues"],
                f"issues={evidence['visual_hygiene_issues'][:10]!r}",
                category="structure",
                weight=3,
            ),
            _criterion(
                task_id,
                "control_formulas",
                (
                    "Every central metric has an independently sourced native "
                    "reconciliation, and control outcomes are not "
                    "hard-coded"
                ),
                controls_complete
                and bool(meaningful_control_formula_cells)
                and not evidence["hardcoded_control_status_cells"],
                (
                    "meaningful_control_formula_cells="
                    f"{meaningful_control_formula_cells[:10]!r}; "
                    "incomplete_metric_controls="
                    f"{[row for row in control_coverage if not row['controlled']][:10]!r}; "
                    "hardcoded_control_status_cells="
                    f"{evidence['hardcoded_control_status_cells'][:10]!r}"
                ),
                category="auditability",
                weight=5,
                failure_cap=0.69,
            ),
            _criterion(
                task_id,
                "source_lineage",
                "The workbook visibly documents source, date/version, and cutoff lineage",
                all(
                    token in _normalize(evidence["text"])
                    for token in ("source", "version", "cutoff")
                ),
                "lineage tokens checked",
                category="provenance",
                weight=5,
            ),
            _criterion(
                task_id,
                "completion_sources",
                (
                    "The completion response names at least two distinct "
                    "authoritative checked sources"
                ),
                _completion_sources_match(completion, gold["sources"]),
                f"sources_checked={completion.get('sources_checked')!r}",
                category="provenance",
                weight=3,
            ),
        ]
    )
    for index, source in enumerate(gold["sources"], start=1):
        token = _source_token(source)
        present = _source_present(evidence["text"], source)
        criteria.append(
            _criterion(
                task_id,
                f"source_{index:02d}",
                f"The workbook cites authoritative source `{token}`",
                present,
                f"token={token!r}",
                category="provenance",
                weight=3,
            )
        )
    decision = _decision_concept_match(
        "\n".join(evidence["decision_blocks"]).strip(),
        gold["central_decision"],
        values=gold["values"],
        paired_rows=evidence["paired_rows"],
    )
    completion_decision = _completion_central_decision_match(
        completion,
        gold["central_decision"],
        values=gold["values"],
        paired_rows=evidence["paired_rows"],
    )
    criteria.append(
        _criterion(
            task_id,
            "central_decision",
            (
                "A native workbook decision block and the completion response "
                "reach the environment-aligned central decision"
            ),
            decision and completion_decision,
            (
                f"expected concept={gold['central_decision']!r}; "
                f"native_matched={decision}; "
                f"completion_matched={completion_decision}"
            ),
            category="decision",
            weight=10,
            semantic=True,
            failure_cap=0.49,
        )
    )
    criteria.extend(
        _task_model_quality_criteria(task_id, gold, evidence)
    )
    criteria.extend(
        _task_086_model_criteria(task_id, gold, evidence)
    )
    criteria.extend(
        _task_068_085_model_criteria(task_id, gold, evidence)
    )
    criteria.extend(
        _task_066_067_model_criteria(task_id, gold, evidence)
    )
    return criteria


def _docx_page_count(
    path: Path,
    workspace_root: Path | None = None,
) -> int | None:
    with validated_native_package_copy(
        path,
        workspace_root=workspace_root or path.parent,
    ) as (snapshot, _inspection):
        return _docx_page_count_from_snapshot(snapshot)


def _flat_scalar_metric_mapping(value: Any) -> bool:
    """Return whether a mapping is safe for generic label/value grading."""

    return (
        isinstance(value, dict)
        and bool(value)
        and all(
            item is not None
            and isinstance(
                item,
                (bool, int, float, str, date, datetime),
            )
            for item in value.values()
        )
    )


def _docx_page_count_from_snapshot(path: Path) -> int | None:
    inspect_native_package(
        path,
        workspace_root=path.parent,
    )
    temporary = Path(tempfile.mkdtemp(prefix="pinehaven-grade-docx-"))
    try:
        input_dir = temporary / "input"
        output_dir = temporary / "output"
        profile_dir = temporary / "profile"
        for directory in (
            input_dir,
            output_dir,
            profile_dir,
            temporary / "cache",
            temporary / "config",
            temporary / "home",
        ):
            directory.mkdir(mode=0o700)
        copied = input_dir / path.name
        shutil.copy2(path, copied)
        inspect_native_package(copied, workspace_root=temporary)
        completed = run_sandboxed_libreoffice(
            sandbox_root=temporary,
            input_path=copied,
            output_dir=output_dir,
            profile_dir=profile_dir,
            convert_to="pdf",
        )
        pdf = output_dir / f"{path.stem}.pdf"
        if completed is None or completed.returncode != 0:
            return None
        try:
            pdf_stat = pdf.lstat()
        except OSError:
            return None
        if (
            stat.S_ISREG(pdf_stat.st_mode)
            and pdf_stat.st_size <= 64 * 1024 * 1024
        ):
            return len(PdfReader(pdf).pages)
        return None
    finally:
        shutil.rmtree(temporary, ignore_errors=True)


def _docx_visual_issues(
    path: Path,
    workspace_root: Path | None = None,
) -> list[str]:
    """Return conservative OOXML defects that break native Word rendering."""

    try:
        with validated_native_package_copy(
            path,
            workspace_root=workspace_root or path.parent,
        ) as (snapshot, _inspection):
            return _docx_visual_issues_from_snapshot(snapshot)
    except (
        OSError,
        zipfile.BadZipFile,
        ElementTree.ParseError,
        NativePackageSafetyError,
    ) as exc:
        return [f"OOXML visual validation failed: {type(exc).__name__}"]


def _docx_visual_issues_from_snapshot(path: Path) -> list[str]:
    issues: list[str] = []
    word_namespace = (
        "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    )
    document_relationship_namespace = (
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    )
    package_relationship_namespace = (
        "http://schemas.openxmlformats.org/package/2006/relationships"
    )
    run_tag = f"{{{word_namespace}}}r"
    run_properties_tag = f"{{{word_namespace}}}rPr"
    vanished_tags = {
        f"{{{word_namespace}}}vanish",
        f"{{{word_namespace}}}webHidden",
    }
    simple_field_tag = f"{{{word_namespace}}}fldSimple"
    text_tag = f"{{{word_namespace}}}t"
    instruction_tag = f"{{{word_namespace}}}instrText"
    field_character_tag = f"{{{word_namespace}}}fldChar"
    instruction_attribute = f"{{{word_namespace}}}instr"
    field_type_attribute = f"{{{word_namespace}}}fldCharType"
    value_attribute = f"{{{word_namespace}}}val"
    section_properties_tag = f"{{{word_namespace}}}sectPr"
    footer_reference_tag = f"{{{word_namespace}}}footerReference"
    footer_type_attribute = f"{{{word_namespace}}}type"
    relationship_id_attribute = (
        f"{{{document_relationship_namespace}}}id"
    )
    title_page_tag = f"{{{word_namespace}}}titlePg"
    even_odd_tag = f"{{{word_namespace}}}evenAndOddHeaders"
    relationship_tag = (
        f"{{{package_relationship_namespace}}}Relationship"
    )
    relationship_id = "Id"
    relationship_type = "Type"
    relationship_target = "Target"
    relationship_target_mode = "TargetMode"

    def enabled_property(element: ElementTree.Element) -> bool:
        return _normalize(element.get(value_attribute, "true")) not in {
            "0",
            "false",
            "no",
            "off",
        }

    def first_ancestor(
        element: ElementTree.Element,
        parents: Mapping[ElementTree.Element, ElementTree.Element],
        tag: str,
    ) -> ElementTree.Element | None:
        current = parents.get(element)
        while current is not None:
            if current.tag == tag:
                return current
            current = parents.get(current)
        return None

    def run_is_hidden(run: ElementTree.Element | None) -> bool:
        if run is None:
            return False
        properties = run.find(run_properties_tag)
        if properties is None:
            return False
        return any(
            child.tag in vanished_tags and enabled_property(child)
            for child in properties
        )

    def page_instruction(instruction: str) -> bool:
        return (
            re.match(r"^\s*PAGE(?:\s|$)", instruction, flags=re.I)
            is not None
        )

    def positive_page_result(
        elements: Iterable[ElementTree.Element],
        parents: Mapping[ElementTree.Element, ElementTree.Element],
    ) -> bool:
        text_elements = [
            element for element in elements if element.tag == text_tag
        ]
        visible_elements = [
            element
            for element in text_elements
            if not run_is_hidden(first_ancestor(element, parents, run_tag))
        ]
        # Native PAGE fields do not need a cached display result. An absent or
        # explicitly empty visible cache remains structurally valid; a hidden
        # cache does not manufacture visible evidence.
        if not text_elements:
            return True
        if not visible_elements:
            return False
        visible = "".join(
            element.text or ""
            for element in visible_elements
        ).strip()
        return visible == "" or re.fullmatch(r"[1-9]\d*", visible) is not None

    def footer_page_field_status(
        footer_root: ElementTree.Element,
    ) -> tuple[bool, bool]:
        parents = {
            child: parent
            for parent in footer_root.iter()
            for child in parent
        }
        simple_fields = list(footer_root.iter(simple_field_tag))
        instruction_elements = list(footer_root.iter(instruction_tag))
        visible_text = " ".join(
            element.text or ""
            for element in footer_root.iter(text_tag)
            if not run_is_hidden(first_ancestor(element, parents, run_tag))
        )
        has_page_claim = (
            re.search(r"\bpage\b", visible_text, flags=re.I) is not None
            or any(
                re.search(
                    r"\bPAGE\b",
                    field.get(instruction_attribute, ""),
                    flags=re.I,
                )
                is not None
                for field in simple_fields
            )
            or re.search(
                r"\bPAGE\b",
                "".join(
                    element.text or ""
                    for element in instruction_elements
                ),
                flags=re.I,
            )
            is not None
        )

        for field in simple_fields:
            if first_ancestor(field, parents, run_tag) is not None:
                continue
            if (
                page_instruction(field.get(instruction_attribute, ""))
                and positive_page_result(field.iter(text_tag), parents)
            ):
                return has_page_claim, True

        stack: list[dict[str, Any]] = []
        for element in footer_root.iter():
            if element.tag == field_character_tag:
                field_type = _normalize(
                    element.get(field_type_attribute, "")
                )
                if field_type == "begin":
                    stack.append(
                        {
                            "instruction": [],
                            "separated": False,
                            "results": [],
                            "malformed": False,
                        }
                    )
                elif field_type == "separate":
                    if stack and not stack[-1]["separated"]:
                        stack[-1]["separated"] = True
                    elif stack:
                        stack[-1]["malformed"] = True
                elif field_type == "end" and stack:
                    frame = stack.pop()
                    if (
                        not frame["malformed"]
                        and frame["separated"]
                        and page_instruction(
                            "".join(frame["instruction"])
                        )
                        and positive_page_result(
                            frame["results"],
                            parents,
                        )
                    ):
                        return has_page_claim, True
            elif stack and element.tag == instruction_tag:
                if stack[-1]["separated"]:
                    stack[-1]["malformed"] = True
                else:
                    stack[-1]["instruction"].append(element.text or "")
            elif (
                stack
                and stack[-1]["separated"]
                and element.tag == text_tag
            ):
                stack[-1]["results"].append(element)
        return has_page_claim, False

    def resolve_footer_target(target: str) -> str | None:
        if not target:
            return None
        if target.startswith("/"):
            resolved = posixpath.normpath(target.lstrip("/"))
        else:
            resolved = posixpath.normpath(posixpath.join("word", target))
        if (
            resolved.startswith("../")
            or resolved == ".."
            or not re.fullmatch(r"word/footer[^/]*\.xml", resolved)
        ):
            return None
        return resolved

    try:
        inspect_native_package(
            path,
            workspace_root=path.parent,
        )
        with zipfile.ZipFile(path) as package:
            names = set(package.namelist())
            required_parts = {
                "word/document.xml",
                "word/_rels/document.xml.rels",
            }
            missing_parts = sorted(required_parts - names)
            if missing_parts:
                issues.append(
                    f"document is missing footer relationship parts: "
                    f"{missing_parts!r}"
                )
                return issues

            document_root = ElementTree.fromstring(
                package.read("word/document.xml")
            )
            relationships_root = ElementTree.fromstring(
                package.read("word/_rels/document.xml.rels")
            )
            footer_relationships: dict[str, str] = {}
            for relationship in relationships_root.iter(relationship_tag):
                if not relationship.get(relationship_type, "").endswith(
                    "/footer"
                ):
                    continue
                relation_id = relationship.get(relationship_id, "")
                target = resolve_footer_target(
                    relationship.get(relationship_target, "")
                )
                external = (
                    _normalize(
                        relationship.get(relationship_target_mode, "")
                    )
                    == "external"
                )
                if relation_id and target and not external:
                    footer_relationships[relation_id] = target

            even_and_odd = False
            if "word/settings.xml" in names:
                settings_root = ElementTree.fromstring(
                    package.read("word/settings.xml")
                )
                setting = settings_root.find(f".//{even_odd_tag}")
                even_and_odd = (
                    setting is not None and enabled_property(setting)
                )

            sections = list(document_root.iter(section_properties_tag))
            if not sections:
                issues.append("document has no section properties")
                return issues

            footer_validity: dict[str, tuple[bool, bool]] = {}
            effective_references: dict[str, str] = {}
            for section_number, section in enumerate(sections, start=1):
                for reference in section.findall(footer_reference_tag):
                    reference_type = _normalize(
                        reference.get(footer_type_attribute, "default")
                    )
                    relation_id = reference.get(
                        relationship_id_attribute,
                        "",
                    )
                    target = footer_relationships.get(relation_id)
                    if reference_type not in {"default", "first", "even"}:
                        issues.append(
                            f"section {section_number} has unknown footer "
                            f"type {reference_type!r}"
                        )
                    elif target is None:
                        issues.append(
                            f"section {section_number} {reference_type} "
                            "footer reference is missing or invalid"
                        )
                        effective_references.pop(reference_type, None)
                    else:
                        effective_references[reference_type] = target

                required_types = {"default"}
                title_page = section.find(title_page_tag)
                if title_page is not None and enabled_property(title_page):
                    required_types.add("first")
                if even_and_odd:
                    required_types.add("even")

                for reference_type in sorted(required_types):
                    target = effective_references.get(reference_type)
                    if target is None:
                        issues.append(
                            f"section {section_number} has no effective "
                            f"{reference_type} footer reference"
                        )
                        continue
                    if target not in names:
                        issues.append(
                            f"section {section_number} {reference_type} "
                            f"footer part {target!r} is missing"
                        )
                        continue
                    if target not in footer_validity:
                        footer_root = ElementTree.fromstring(
                            package.read(target)
                        )
                        footer_validity[target] = (
                            footer_page_field_status(footer_root)
                        )
                    has_page_claim, valid_page_field = (
                        footer_validity[target]
                    )
                    if has_page_claim and not valid_page_field:
                        issues.append(
                            f"section {section_number} {reference_type} "
                            f"footer {target} has no valid native PAGE "
                            "field"
                        )
    except (
        OSError,
        zipfile.BadZipFile,
        ElementTree.ParseError,
        NativePackageSafetyError,
    ) as exc:
        issues.append(f"OOXML visual validation failed: {type(exc).__name__}")
    return issues


def _docx_criteria(
    task_id: str,
    answer: Any,
    workspace_root: Path,
    gold: dict[str, Any],
) -> list[Criterion]:
    target = gold["artifact"]["path"]
    path = workspace_root / target
    criteria, completion = _artifact_answer_criteria(task_id, answer, target)
    readable = False
    text = ""
    rows: list[list[Any]] = []
    decision_blocks: list[str] = []
    decision_action_rows: list[str] = []
    structure: dict[str, Any] = {
        "title_candidates": [],
        "title_present": False,
        "headings": [],
        "date_present": False,
        "audience_present": False,
        "calculation_rows": {},
        "calculation_coverage": set(),
        "calculation_lineage_rows": {},
        "calculation_lineage_coverage": set(),
        "calculation_lineage_issues": [],
        "expected_metric_labels": {
            _normalize(key.replace("_", " ")) for key in gold["values"]
        },
        "owner_timing_blocks": [],
    }
    table_count = 0
    page_count: int | None = None
    visual_issues: list[str] = ["required document is missing"]
    error = "required document is missing"
    artifact_present = os.path.lexists(path)
    if artifact_present:
        try:
            with validated_native_package_copy(
                path,
                workspace_root=workspace_root,
            ) as (snapshot, _inspection):
                document = Document(snapshot)
                text, rows = _flatten_docx(document)
                decision_blocks = _docx_decision_blocks(document)
                decision_action_rows = _docx_decision_action_rows(document)
                structure = _docx_native_structure(
                    document,
                    title=gold["title"],
                    metric_keys=gold["values"],
                    sources=gold["sources"],
                )
                table_count = len(_native_docx_snapshot(document).tables)
                readable = True
                page_count = _docx_page_count(
                    snapshot,
                    snapshot.parent,
                )
                visual_issues = _docx_visual_issues(
                    snapshot,
                    snapshot.parent,
                )
                error = "document parsed"
        except NativePackageSafetyError as exc:
            error = f"native package safety failure: {exc}"
        except Exception as exc:
            error = f"{type(exc).__name__}: {exc}"
    criteria.extend(
        [
            _criterion(
                task_id,
                "artifact_exists",
                "The exact required document exists",
                artifact_present,
                target,
                category="integrity",
                weight=10,
            ),
            _criterion(
                task_id,
                "artifact_readable",
                "The document parses as a genuine DOCX",
                readable,
                error,
                category="integrity",
                weight=10,
            ),
            _criterion(
                task_id,
                "page_limit",
                "The decision memorandum is no more than three rendered pages",
                page_count is not None and page_count <= 3,
                f"page_count={page_count!r}",
                category="structure",
                weight=3,
            ),
            _criterion(
                task_id,
                "native_table",
                "The memorandum contains a native decision table",
                table_count >= 1,
                f"table_count={table_count}",
                category="structure",
                weight=3,
            ),
            _criterion(
                task_id,
                "visual_hygiene",
                "The document uses valid native Word fields and layout markup",
                not visual_issues,
                f"issues={visual_issues[:20]!r}",
                category="structure",
                weight=3,
                failure_cap=0.69,
            ),
            _criterion(
                task_id,
                "title",
                "The memorandum has the required native Word title",
                structure["title_present"],
                f"title_candidates={structure['title_candidates']!r}",
                category="structure",
                weight=3,
            ),
            _criterion(
                task_id,
                "date",
                "The memorandum visibly states the June 30, 2026 date",
                structure["date_present"],
                "native date text checked",
                category="structure",
                weight=3,
            ),
            _criterion(
                task_id,
                "audience",
                "The memorandum visibly identifies a nonempty audience",
                structure["audience_present"],
                "native Audience label checked",
                category="structure",
                weight=3,
            ),
            _criterion(
                task_id,
                "calculation_chain",
                (
                    "A native table traces every central metric through "
                    "source/input, calculation logic, and result"
                ),
                structure["calculation_lineage_coverage"]
                == structure["expected_metric_labels"],
                (
                    "lineage_covered="
                    f"{sorted(structure['calculation_lineage_coverage'])!r}; "
                    "expected="
                    f"{sorted(structure['expected_metric_labels'])!r}; "
                    "issues="
                    f"{structure['calculation_lineage_issues']!r}"
                ),
                category="auditability",
                weight=5,
                failure_cap=0.69,
            ),
        ]
    )
    required_sections = (
        "Executive conclusion",
        "Evidence",
        "Economics",
        "Risks and controls",
        "Sources",
    )
    native_headings = {
        _normalize(heading) for heading in structure["headings"]
    }
    for section in required_sections:
        present = _normalize(section) in native_headings
        criteria.append(
            _criterion(
                task_id,
                f"section_{_normalize(section).replace(' ', '_')}",
                (
                    f"The memorandum uses a native Word heading for "
                    f"`{section}`"
                ),
                present,
                f"native_headings={structure['headings']!r}",
                category="structure",
                weight=3,
            )
        )
    criteria.extend(
        _metric_artifact_criteria(
            task_id,
            gold["values"],
            text=text,
            paired_rows=rows,
        )
    )
    task_parameters = gold.get("parameters", {})
    required_owner = str(
        task_parameters.get("decision_owner", "")
    ).strip()
    required_timing = next(
        (
            str(task_parameters[key]).strip()
            for key in (
                "decision_date",
                "action_due_date",
                "implementation_gate_date",
            )
            if str(task_parameters.get(key, "")).strip()
        ),
        "",
    )
    owner_timing_matches = bool(structure["owner_timing_blocks"]) and any(
        (
            not required_owner
            or _normalize(required_owner) in _normalize(block)
        )
        and (
            not required_timing
            or _text_contains(block, required_timing)
        )
        for block in structure["owner_timing_blocks"]
    )
    native_metrics_consistent, native_metric_checks = (
        _native_metric_claim_consistency(rows, gold["values"])
    )
    criteria.append(
        _criterion(
            task_id,
            "native_metric_consistency",
            (
                "Every quantified native table/card row exactly labeled as "
                "a central metric agrees with authoritative truth"
            ),
            native_metrics_consistent,
            f"checks={native_metric_checks[:20]!r}",
            category="core_finance",
            weight=10,
            failure_cap=0.49,
        )
    )
    if _flat_scalar_metric_mapping(gold.get("supporting_values")):
        supporting_values = gold["supporting_values"]
        supporting_label_aliases = _materialize_metric_label_aliases(
            supporting_values,
            gold.get("supporting_label_aliases", {}),
        )
        criteria.extend(
            _metric_artifact_criteria(
                task_id,
                supporting_values,
                text=text,
                paired_rows=rows,
                label_aliases=supporting_label_aliases,
            )
        )
        supporting_consistent, supporting_checks = (
            _native_metric_claim_consistency(
                rows,
                supporting_values,
                label_aliases=supporting_label_aliases,
            )
        )
        criteria.append(
            _criterion(
                task_id,
                "supporting_metric_consistency",
                (
                    "Every required decision-support metric is paired in a "
                    "native row and agrees with authoritative truth"
                ),
                supporting_consistent
                and all(
                    any(
                        _paired_metric(
                            rows,
                            label,
                            expected,
                            metric_key=key,
                        )
                        for label in _metric_label_options(
                            key,
                            supporting_label_aliases,
                        )
                    )
                    for key, expected in supporting_values.items()
                ),
                f"checks={supporting_checks[:20]!r}",
                category="core_finance",
                weight=10,
                failure_cap=0.49,
            )
        )
    criteria.extend(
        [
            _criterion(
                task_id,
                "preclose_label",
                "The memorandum labels June 2026 as pre-close",
                "pre close" in _normalize(text),
                "pre-close language checked",
                category="provenance",
                weight=5,
            ),
            _criterion(
                task_id,
                "owner_and_timing",
                (
                    "The memorandum associates a specific accountable owner "
                    "and concrete timing in one native block, using any "
                    "task-supplied owner and decision date exactly"
                ),
                owner_timing_matches,
                (
                    "owner_timing_blocks="
                    f"{structure['owner_timing_blocks'][:10]!r}; "
                    f"required_owner={required_owner!r}; "
                    f"required_timing={required_timing!r}"
                ),
                category="decision",
                weight=5,
            ),
            _criterion(
                task_id,
                "fact_assumption_separation",
                "The memorandum distinguishes facts from assumptions",
                "assumption" in _normalize(text) and "fact" in _normalize(text),
                "fact/assumption labels checked",
                category="auditability",
                weight=5,
                failure_cap=0.69,
            ),
            _criterion(
                task_id,
                "completion_sources",
                (
                    "The completion response names at least two distinct "
                    "authoritative checked sources"
                ),
                _completion_sources_match(completion, gold["sources"]),
                f"sources_checked={completion.get('sources_checked')!r}",
                category="provenance",
                weight=3,
            ),
        ]
    )
    for index, source in enumerate(gold["sources"], start=1):
        token = _source_token(source)
        present = _source_present(text, source)
        criteria.append(
            _criterion(
                task_id,
                f"source_{index:02d}",
                f"The memorandum cites authoritative source `{token}`",
                present,
                f"token={token!r}",
                category="provenance",
                weight=3,
            )
        )
    if task_id == "task_039":
        native_decision_text = "\n".join(decision_blocks).strip()
        gp_rate = gold["values"]["required_price_change_for_gp"]
        margin_rate = gold["values"][
            "required_price_change_for_margin"
        ]
        gp_paired = _decision_rate_semantic_pair(
            native_decision_text,
            expected=gp_rate,
            metric_key="required_price_change_for_gp",
            own_concept=(
                r"(?:baseline\s+)?gross\s+profit"
                r"(?:\s+dollars?)?"
            ),
            competing_concept=(
                r"(?:baseline\s+)?gross\s+margin"
                r"(?:\s+(?:rate|percent(?:age)?))?"
            ),
        )
        margin_paired = _decision_rate_semantic_pair(
            native_decision_text,
            expected=margin_rate,
            metric_key="required_price_change_for_margin",
            own_concept=(
                r"(?:baseline\s+)?gross\s+margin"
                r"(?:\s+(?:rate|percent(?:age)?))?"
            ),
            competing_concept=(
                r"(?:baseline\s+)?gross\s+profit"
                r"(?:\s+dollars?)?"
            ),
        )
        recommends_price_action = bool(
            re.search(
                (
                    r"\b(?:raise|increase|set|approve|implement)\w*"
                    r".{0,80}\b(?:drv\s+)?pric\w*"
                    r"|\b(?:drv\s+)?pric\w*.{0,80}"
                    r"\b(?:raise|increase|set|approve|implement)\w*"
                ),
                native_decision_text,
                flags=re.I,
            )
        )
        criteria.append(
            _criterion(
                task_id,
                "break_even_price_recommendation",
                (
                    "The native decision recommends a DRV price action and "
                    "pairs the exact preserve-gross-profit and "
                    "preserve-gross-margin increases with their distinct "
                    "objectives"
                ),
                gp_paired and margin_paired and recommends_price_action,
                (
                    f"gp_rate={gp_rate!r}; gp_paired={gp_paired}; "
                    f"margin_rate={margin_rate!r}; "
                    f"margin_paired={margin_paired}; "
                    f"price_action={recommends_price_action}"
                ),
                category="decision",
                weight=10,
                failure_cap=0.49,
            )
        )
    if task_id == "task_069":
        site_name = str(gold["parameters"]["site_name"])
        site_code = str(gold["parameters"]["site_code"])
        scoped_blocks = "\n".join(
            " ".join(str(value) for value in row)
            for row in rows
        )
        criteria.extend(
            [
                _criterion(
                    task_id,
                    "project_site_scope",
                    (
                        "The memorandum identifies the authorized project "
                        f"scope as {site_name} (site {site_code})"
                    ),
                    _project_site_scope_match(text, site_name, site_code),
                    f"required_scope={site_name!r}/site {site_code}",
                    category="decision",
                    weight=10,
                    failure_cap=0.49,
                ),
                _criterion(
                    task_id,
                    "scoped_central_decision",
                    (
                        "A native memorandum block affirmatively approves the "
                        f"{site_name} site {site_code} project"
                    ),
                    _scoped_site_approval_match(
                        scoped_blocks,
                        site_name,
                        site_code,
                    ),
                    (
                        "affirmative approval and authorized site checked "
                        "within the same paragraph or native table row"
                    ),
                    category="decision",
                    weight=10,
                    failure_cap=0.49,
                ),
            ]
        )
    decision_text = "\n".join(decision_blocks).strip()
    decision_matched = _decision_concept_match(
        decision_text,
        gold["central_decision"],
        values=gold["values"],
        paired_rows=rows,
    ) or any(
        _decision_action_row_match(
            row,
            gold["central_decision"],
            values=gold["values"],
        )
        for row in decision_action_rows
    )
    completion_decision_matched = _completion_central_decision_match(
        completion,
        gold["central_decision"],
        values=gold["values"],
        paired_rows=rows,
    )
    criteria.append(
        _criterion(
            task_id,
            "central_decision",
            (
                "A native memorandum decision block and the completion "
                "response reach the environment-aligned central decision"
            ),
            decision_matched and completion_decision_matched,
            (
                f"expected concept={gold['central_decision']!r}; "
                f"native_matched={decision_matched}; "
                f"completion_matched={completion_decision_matched}"
            ),
            category="decision",
            weight=10,
            semantic=True,
            failure_cap=0.49,
        )
    )
    return criteria


def _pptx_criteria(
    task_id: str,
    answer: Any,
    workspace_root: Path,
    gold: dict[str, Any],
) -> list[Criterion]:
    target = gold["artifact"]["path"]
    path = workspace_root / target
    detail_specification = gold.get("detail_schedule")
    criteria, completion = _artifact_answer_criteria(task_id, answer, target)
    readable = False
    text = ""
    rows: list[list[Any]] = []
    decision_blocks: list[str] = []
    chart_evidence: list[dict[str, Any]] = []
    native_tables: list[list[list[str]]] = []
    chart_count = table_count = slide_count = 0
    title_candidates: list[list[str]] = []
    error = "required presentation is missing"
    placeholder_count = 0
    visual_issues: list[str] = ["missing"]
    ooxml_issues: list[str] = ["missing"]
    artifact_present = os.path.lexists(path)
    if artifact_present:
        try:
            with validated_native_package_copy(
                path,
                workspace_root=workspace_root,
            ) as (snapshot, _inspection):
                presentation = Presentation(snapshot)
                slide_count = len(presentation.slides)
                (
                    text,
                    chart_count,
                    table_count,
                    title_candidates,
                ) = _flatten_pptx(presentation)
                chart_evidence = _pptx_chart_evidence(presentation)
                native_tables = [
                    [
                        [cell.text for cell in row.cells]
                        for row in shape.table.rows
                    ]
                    for slide in presentation.slides
                    for shape in slide.shapes
                    if getattr(shape, "has_table", False)
                ]
                rows = _pptx_paired_rows(
                    presentation,
                    (
                        detail_specification
                        if isinstance(detail_specification, dict)
                        else None
                    ),
                )
                decision_blocks = _pptx_decision_blocks(presentation)
                for slide in presentation.slides:
                    for shape in slide.shapes:
                        if getattr(shape, "is_placeholder", False):
                            placeholder_count += 1
                visual_issues = _pptx_visual_issues(presentation)
                ooxml_issues = pptx_axis_id_issues(
                    snapshot,
                    workspace_root=snapshot.parent,
                )
                readable = True
                error = "presentation parsed"
        except NativePackageSafetyError as exc:
            error = f"native package safety failure: {exc}"
        except Exception as exc:
            error = f"{type(exc).__name__}: {exc}"
    criteria.extend(
        [
            _criterion(
                task_id,
                "artifact_exists",
                "The exact required presentation exists",
                artifact_present,
                target,
                category="integrity",
                weight=10,
            ),
            _criterion(
                task_id,
                "artifact_readable",
                (
                    "The presentation parses as genuine PPTX and uses valid "
                    "unsigned chart-axis identifiers"
                ),
                readable and not ooxml_issues,
                f"{error}; ooxml_issues={ooxml_issues[:10]!r}",
                category="integrity",
                weight=10,
            ),
            _criterion(
                task_id,
                "slide_count",
                "The presentation contains exactly five slides",
                slide_count == 5,
                f"slide_count={slide_count}",
                category="structure",
                weight=5,
            ),
        ]
    )
    expected_titles = ["Decision", "Evidence", "Economics", "Risks", "Actions"]
    for index, title in enumerate(expected_titles):
        candidates = (
            title_candidates[index] if index < len(title_candidates) else []
        )
        criteria.append(
            _criterion(
                task_id,
                f"slide_{index + 1:02d}_title",
                f"Slide {index + 1} is titled `{title}`",
                any(
                    _normalize(candidate) == _normalize(title)
                    for candidate in candidates
                ),
                f"candidates={candidates!r}",
                category="structure",
                weight=3,
            )
        )
    criteria.extend(
        _metric_artifact_criteria(
            task_id,
            gold["values"],
            text=text,
            paired_rows=rows,
        )
    )
    if task_id == "task_079" and isinstance(
        gold.get("supporting_values"),
        dict,
    ):
        supporting_values = gold["supporting_values"]
        criteria.extend(
            _metric_artifact_criteria(
                task_id,
                supporting_values,
                text=text,
                paired_rows=rows,
            )
        )
        (
            supporting_metrics_consistent,
            supporting_metric_checks,
        ) = _native_metric_claim_consistency(rows, supporting_values)
        criteria.append(
            _criterion(
                task_id,
                "lender_update_supporting_truth",
                (
                    "The lender update natively pairs the complete active-debt "
                    "profile and 13-week cash outlook with authoritative values"
                ),
                supporting_metrics_consistent
                and all(
                    _paired_metric(rows, key, expected)
                    for key, expected in supporting_values.items()
                ),
                f"checks={supporting_metric_checks[:20]!r}",
                category="core_finance",
                weight=10,
                failure_cap=0.49,
            )
        )
    native_metrics_consistent, native_metric_checks = (
        _native_metric_claim_consistency(rows, gold["values"])
    )
    criteria.append(
        _criterion(
            task_id,
            "native_metric_consistency",
            (
                "Every quantified native table/card row exactly labeled as "
                "a central metric agrees with authoritative truth"
            ),
            native_metrics_consistent,
            f"checks={native_metric_checks[:20]!r}",
            category="core_finance",
            weight=10,
            failure_cap=0.49,
        )
    )
    if isinstance(detail_specification, dict):
        matched_keys, expected_keys, missing_keys = _detail_key_matches(
            rows,
            detail_specification,
        )
        detail_presentation = _detail_presentation_evidence(
            rows,
            detail_specification,
        )
        criteria.append(
            _criterion(
                task_id,
                "detail_dimension_population",
                (
                    "The deck's native tables contain every required "
                    "dimension row from the authoritative detail schedule"
                ),
                expected_keys > 0 and matched_keys == expected_keys,
                (
                    f"matched_keys={matched_keys}; "
                    f"expected_keys={expected_keys}; "
                    f"missing_keys={missing_keys!r}"
                ),
                category="core_finance",
                weight=10,
                failure_cap=0.49,
            )
        )
        criteria.append(
            _criterion(
                task_id,
                "detail_native_field_values",
                (
                    "The deck's complete native detail table contains the "
                    "authoritative value for every field in every dimension "
                    "row at displayed precision"
                ),
                bool(detail_presentation["field_values_match"]),
                (
                    f"headers_match="
                    f"{detail_presentation['headers_match']}; "
                    f"row_count={detail_presentation['row_count']}; "
                    f"expected_row_count="
                    f"{detail_presentation['expected_row_count']}; "
                    f"keys_match={detail_presentation['keys_match']}; "
                    f"field_mismatch_count="
                    f"{detail_presentation['field_mismatch_count']}; "
                    f"field_mismatch_samples="
                    f"{detail_presentation['field_mismatch_samples']!r}"
                ),
                category="core_finance",
                weight=10,
                failure_cap=0.49,
            )
        )
    if task_id == "task_029" and isinstance(
        gold.get("supporting_values"),
        dict,
    ):
        supporting_values = gold["supporting_values"]
        criteria.extend(
            _metric_artifact_criteria(
                task_id,
                supporting_values,
                text=text,
                paired_rows=rows,
            )
        )
        (
            supporting_metrics_consistent,
            supporting_metric_checks,
        ) = _native_metric_claim_consistency(rows, supporting_values)
        criteria.append(
            _criterion(
                task_id,
                "supporting_metric_consistency",
                (
                    "Every required S&OP support metric is present in a "
                    "native table/card and agrees with authoritative truth"
                ),
                supporting_metrics_consistent
                and all(
                    _paired_metric(rows, key, expected)
                    for key, expected in supporting_values.items()
                ),
                f"checks={supporting_metric_checks[:30]!r}",
                category="core_finance",
                weight=10,
                failure_cap=0.49,
            )
        )
        consistent, claim_checks = _sop_past_due_claim_consistency(
            rows,
            supporting_values,
        )
        criteria.append(
            _criterion(
                task_id,
                "past_due_backlog_claims",
                (
                    "Required past-due backlog claims use the strict "
                    "pre-cutoff population and reconcile to authoritative "
                    "units, value, and percentage"
                ),
                consistent
                and all(
                    _paired_metric(rows, key, supporting_values[key])
                    for key in (
                        "past_due_backlog_units",
                        "past_due_backlog_value",
                        "past_due_backlog_percent",
                    )
                ),
                f"claim_checks={claim_checks!r}",
                category="core_finance",
                weight=10,
                failure_cap=0.49,
            )
        )
    criteria.extend(
        [
            _criterion(
                task_id,
                "native_chart",
                "The deck contains at least one native editable chart",
                chart_count >= 1,
                f"chart_count={chart_count}",
                category="structure",
                weight=3,
            ),
            _criterion(
                task_id,
                "native_table",
                "The deck contains at least one native editable table",
                table_count >= 1,
                f"table_count={table_count}",
                category="structure",
                weight=3,
            ),
            _criterion(
                task_id,
                "preclose_label",
                (
                    "The deck labels June 2026 as pre-close when its "
                    "calculation basis requires that status"
                ),
                gold.get("bundle") not in {"close_control", "pnl", "executive"}
                or "pre close" in _normalize(text),
                (
                    "pre-close language checked"
                    if gold.get("bundle")
                    in {"close_control", "pnl", "executive"}
                    else "not required by this presentation's calculation basis"
                ),
                category="provenance",
                weight=5,
            ),
            _criterion(
                task_id,
                "owner_and_timing",
                "The Actions slide names owners and timing",
                "owner" in _normalize(text)
                and any(token in _normalize(text) for token in ("timing", "due", "date", "by ")),
                "owner/timing language checked",
                category="decision",
                weight=5,
            ),
            _criterion(
                task_id,
                "no_template_placeholders",
                "The deck contains no visible template placeholder text",
                not re.search(
                    r"\b(?:todo|tbd|lorem ipsum|insert (?:chart|text)|placeholder)\b",
                    text,
                    flags=re.I,
                ),
                "placeholder text checked",
                category="structure",
                weight=3,
            ),
            _criterion(
                task_id,
                "visual_hygiene",
                (
                    "The deck keeps native shapes on-canvas and preserves "
                    "consistent text styling across manual line breaks"
                ),
                not visual_issues,
                f"issues={visual_issues[:20]!r}",
                category="structure",
                weight=3,
                failure_cap=0.69,
            ),
            _criterion(
                task_id,
                "completion_sources",
                (
                    "The completion response names at least two distinct "
                    "authoritative checked sources"
                ),
                _completion_sources_match(completion, gold["sources"]),
                f"sources_checked={completion.get('sources_checked')!r}",
                category="provenance",
                weight=3,
            ),
        ]
    )
    for index, source in enumerate(gold["sources"], start=1):
        token = _source_token(source)
        criteria.append(
            _criterion(
                task_id,
                f"source_{index:02d}",
                f"The deck cites authoritative source `{token}`",
                _source_present(text, source),
                f"token={token!r}",
                category="provenance",
                weight=3,
            )
        )
    native_decision_matched = _decision_concept_match(
        "\n".join(decision_blocks).strip(),
        gold["central_decision"],
        values=gold["values"],
        paired_rows=rows,
    )
    completion_decision_matched = _completion_central_decision_match(
        completion,
        gold["central_decision"],
        values=gold["values"],
        paired_rows=rows,
    )
    criteria.append(
        _criterion(
            task_id,
            "central_decision",
            (
                "A native Decision or Actions slide and the completion "
                "response reach the environment-aligned central decision"
            ),
            native_decision_matched and completion_decision_matched,
            (
                f"expected concept={gold['central_decision']!r}; "
                f"native_matched={native_decision_matched}; "
                f"completion_matched={completion_decision_matched}"
            ),
            category="decision",
            weight=10,
            semantic=True,
            failure_cap=0.49,
        )
    )
    normalized_text = _normalize(text)
    if task_id == "task_079":
        required_terms = (
            "leverage",
            "tangible net worth",
            "liquidity",
            "lender",
            "treasurer",
            "controller",
            "cash outlook",
            "revolver",
            "equipment loan",
            "term loan",
        )
        criteria.append(
            _criterion(
                task_id,
                "lender_update_decision_content",
                (
                    "The lender update covers both covenant tests, liquidity, "
                    "lender escalation, and accountable finance owners"
                ),
                all(term in normalized_text for term in required_terms),
                f"required_terms={required_terms!r}",
                category="decision",
                weight=10,
                failure_cap=0.49,
            )
        )
    elif task_id == "task_090":
        required_terms = (
            "normalization",
            "maintenance",
            "terminal",
            "enterprise value",
            "net debt",
            "equity value",
        )
        criteria.append(
            _criterion(
                task_id,
                "valuation_story_content",
                (
                    "The valuation deck connects normalization and "
                    "maintenance capital through terminal value to the "
                    "enterprise-to-equity bridge"
                ),
                all(term in normalized_text for term in required_terms),
                f"required_terms={required_terms!r}",
                category="decision",
                weight=10,
                failure_cap=0.49,
            )
        )
        supporting = gold.get("supporting_values")
        if isinstance(supporting, dict):
            bridge_keys = (
                "reported_ebitda",
                "production_variance_normalization",
                "quality_normalization",
                "maintenance_normalization",
                "normalized_ebitda",
                "maintenance_capital_expenditures",
            )
            bridge_met = all(
                _paired_metric(rows, key, supporting[key])
                for key in bridge_keys
            )
            forecast_evidence = _semantic_native_table_truth(
                native_tables,
                column_aliases={
                    "year_component": (
                        "Year / component",
                        "Forecast year / component",
                    ),
                    "cash_flow_value": (
                        "Cash flow / value",
                        "Cash flow or terminal value",
                    ),
                    "present_value": ("Present value", "PV"),
                },
                key_fields=("year_component",),
                expected_rows=[
                    {
                        "year_component": item["forecast_year"],
                        "cash_flow_value": item["cash_flow"],
                        "present_value": item["present_value"],
                    }
                    for item in supporting["explicit_forecast"]
                ]
                + [
                    {
                        "year_component": "Terminal value",
                        "cash_flow_value": supporting["terminal_value"],
                        "present_value": supporting[
                            "terminal_present_value"
                        ],
                    }
                ],
            )
            sensitivity = supporting["terminal_growth_sensitivity"]
            sensitivity_chart_met = any(
                "terminal growth" in _normalize(chart["title"])
                and len(chart["categories"]) == len(sensitivity)
                and len(chart["series"]) == 1
                and len(chart["series"][0]["values"]) == len(sensitivity)
                and all(
                    any(
                        _value_matches(
                            category,
                            item["terminal_growth"],
                            "terminal_growth",
                        )
                        and _close(
                            value,
                            item["enterprise_value"],
                            "enterprise_value",
                        )
                        for category, value in zip(
                            chart["categories"],
                            chart["series"][0]["values"],
                            strict=True,
                        )
                    )
                    for item in sensitivity
                )
                for chart in chart_evidence
            )
            criteria.extend(
                [
                    _criterion(
                        task_id,
                        "valuation_qoe_maintenance_bridge",
                        (
                            "The native evidence table ties every QoE "
                            "normalization and maintenance-capital input"
                        ),
                        bridge_met,
                        f"bridge_met={bridge_met}",
                        category="core_finance",
                        weight=10,
                        failure_cap=0.49,
                    ),
                    _criterion(
                        task_id,
                        "valuation_forecast_terminal_truth",
                        (
                            "The native economics table presents all five "
                            "forecast cash flows/PVs and terminal value/PV"
                        ),
                        forecast_evidence["field_values_match"],
                        f"forecast_evidence={forecast_evidence!r}",
                        category="core_finance",
                        weight=10,
                        failure_cap=0.49,
                    ),
                    _criterion(
                        task_id,
                        "valuation_sensitivity_chart_truth",
                        (
                            "The native sensitivity chart contains all five "
                            "exact terminal-growth enterprise values"
                        ),
                        sensitivity_chart_met,
                        f"chart_evidence={chart_evidence!r}",
                        category="core_finance",
                        weight=10,
                        failure_cap=0.49,
                    ),
                ]
            )
    elif task_id == "task_098":
        supporting = gold.get("supporting_values")
        downside = (
            supporting.get("fy26_h2_downside")
            if isinstance(supporting, dict)
            else None
        )
        plan = (
            supporting.get("fy27_upside_plan")
            if isinstance(supporting, dict)
            else None
        )
        periods_present = all(
            term in normalized_text
            for term in (
                "fy26 ytd",
                "fy26 h2 downside",
                "fy27 upside plan",
            )
        )
        scenario_evidence = (
            _semantic_native_table_truth(
                native_tables,
                column_aliases={
                    "scenario": ("Scenario", "Period / scenario"),
                    "scenario_revenue": ("Revenue", "Scenario revenue"),
                    "scenario_gross_profit": (
                        "Gross profit",
                        "Scenario gross profit",
                    ),
                    "scenario_gross_margin": (
                        "Gross margin",
                        "Scenario gross margin",
                    ),
                    "scenario_operating_income": (
                        "Operating income",
                        "Scenario operating income",
                    ),
                    "incremental_working_capital": (
                        "Incremental WC",
                        "Incremental working capital",
                    ),
                    "starting_liquidity": (
                        "Starting liquidity",
                        "June 30 total liquidity",
                    ),
                    "working_capital_liquidity_impact": (
                        "WC liquidity impact",
                        "Working capital liquidity impact",
                    ),
                    "pro_forma_liquidity": (
                        "Pro forma liquidity",
                    ),
                },
                key_fields=("scenario",),
                expected_rows=[
                    {
                        "scenario": label,
                        **{
                            key: scenario[key]
                            for key in (
                                "scenario_revenue",
                                "scenario_gross_profit",
                                "scenario_gross_margin",
                                "scenario_operating_income",
                                "incremental_working_capital",
                                "starting_liquidity",
                                "working_capital_liquidity_impact",
                                "pro_forma_liquidity",
                            )
                        },
                    }
                    for label, scenario in (
                        ("FY26 H2 downside", downside),
                        ("FY27 upside plan", plan),
                    )
                ],
            )
            if isinstance(downside, dict) and isinstance(plan, dict)
            else {}
        )
        economics_paired = bool(
            scenario_evidence.get("field_values_match")
        )
        liquidity_math_met = (
            isinstance(downside, dict)
            and isinstance(plan, dict)
            and all(
                _close(
                    scenario["pro_forma_liquidity"],
                    scenario["starting_liquidity"]
                    - scenario["incremental_working_capital"],
                    "pro_forma_liquidity",
                )
                and _close(
                    scenario["working_capital_liquidity_impact"],
                    -scenario["incremental_working_capital"],
                    "working_capital_liquidity_impact",
                )
                for scenario in (downside, plan)
            )
            and all(
                token in normalized_text
                for token in (
                    "pro forma liquidity",
                    "total liquidity",
                    "minus incremental working capital",
                    "positive working capital is a use",
                    "negative working capital is a source",
                )
            )
        )
        criteria.extend(
            [
                _criterion(
                    task_id,
                    "board_scenario_period_architecture",
                    (
                        "The board deck separates FY26 YTD current state, "
                        "FY26 H2 downside, and the FY27 upside plan"
                    ),
                    periods_present,
                    (
                        "required periods: FY26 YTD, FY26 H2 downside, "
                        "FY27 upside plan"
                    ),
                    category="decision",
                    weight=10,
                    failure_cap=0.49,
                ),
                _criterion(
                    task_id,
                    "board_forward_scenario_economics",
                    (
                        "Both forward scenarios pair exact revenue, gross "
                        "profit/margin, operating income, working capital, "
                        "and pro forma liquidity in one native row"
                    ),
                    economics_paired,
                    f"scenario_evidence={scenario_evidence!r}",
                    category="core_finance",
                    weight=10,
                    failure_cap=0.49,
                ),
                _criterion(
                    task_id,
                    "board_scenario_liquidity_bridge",
                    (
                        "The deck defines and applies pro forma liquidity as "
                        "starting liquidity less incremental working capital"
                    ),
                    liquidity_math_met,
                    f"liquidity_math_met={liquidity_math_met}",
                    category="core_finance",
                    weight=10,
                    failure_cap=0.49,
                ),
            ]
        )
    return criteria


def _audit_rows(connection: sqlite3.Connection) -> list[dict[str, Any]]:
    connection.row_factory = sqlite3.Row
    rows = [
        dict(row)
        for row in connection.execute(
            "SELECT * FROM audit_events ORDER BY id"
        ).fetchall()
    ]
    for row in rows:
        try:
            row["details"] = json.loads(row.pop("details_json"))
        except (json.JSONDecodeError, TypeError):
            row["details"] = {}
    return rows


def _stable_erp_evidence(value: Any) -> Any:
    """Remove only wall-clock fields from deterministic criterion evidence."""

    if isinstance(value, dict):
        return {
            key: _stable_erp_evidence(item)
            for key, item in value.items()
            if key not in {"created_at", "event_at"}
        }
    if isinstance(value, list):
        return [_stable_erp_evidence(item) for item in value]
    if isinstance(value, tuple):
        return tuple(_stable_erp_evidence(item) for item in value)
    return value


def _sequence_matches(actual: Any, expected: list[str]) -> bool:
    if isinstance(actual, str):
        actual_values = [
            part.strip()
            for part in re.split(r"[,;>]", actual)
            if part.strip()
        ]
    elif isinstance(actual, list):
        actual_values = [str(value) for value in actual]
    else:
        return False
    if len(actual_values) != len(expected):
        return False

    def canonical(value: str) -> str:
        normalized = _normalize(
            value.replace("→", " to ").replace("->", " to ")
        )
        if (
            "status change" in normalized
            or "status transition" in normalized
            or (
                " to " in f" {normalized} "
                and any(
                    status in normalized
                    for status in ("scheduled", "released", "started")
                )
            )
        ):
            for status in ("started", "released", "scheduled"):
                if status in normalized:
                    return f"status change {status}"
            return "status change"
        return normalized

    for actual_value, expected_value in zip(actual_values, expected, strict=True):
        actual_action = canonical(actual_value)
        expected_action = canonical(expected_value)
        if actual_action == expected_action:
            continue
        # When the answer reports the generic audited action, the exact target
        # status remains independently enforced by the returned `status`, the
        # ERP object, and the ordered database audit events.
        if (
            actual_action == "status change"
            and expected_action.startswith("status change ")
        ):
            continue
        return False
    return True


def _erp_state(
    gold: dict[str, Any], database_path: Path
) -> tuple[dict[str, Any], list[Criterion], list[dict[str, Any]]]:
    task_id = gold["task_id"]
    bundle = gold["bundle"]
    params = gold["parameters"]
    expected: dict[str, Any] = {}
    state_criteria: list[Criterion] = []
    with sqlite3.connect(f"file:{database_path}?mode=ro", uri=True) as connection:
        connection.row_factory = sqlite3.Row
        audit = _audit_rows(connection)
        if bundle in {"erp_draft_journal", "erp_posted_journal"}:
            header_rows = connection.execute(
                "SELECT * FROM journal_headers WHERE reference=? ORDER BY id",
                (params["reference"],),
            ).fetchall()
            header = dict(header_rows[0]) if len(header_rows) == 1 else {}
            lines = (
                [
                    dict(row)
                    for row in connection.execute(
                        "SELECT * FROM journal_lines WHERE journal_id=? ORDER BY id",
                        (header["id"],),
                    )
                ]
                if header
                else []
            )
            journal_audit = [
                row
                for row in audit
                if row["entity"] == "journal"
                and row["identifier"] == params["reference"]
            ]
            actions = [row["action"] for row in journal_audit]
            expected = {
                "reference": params["reference"],
                "status": "Posted" if bundle == "erp_posted_journal" else "Draft",
                "debits": _number(sum(float(row["debit"]) for row in lines)) or 0.0,
                "credits": _number(sum(float(row["credit"]) for row in lines)) or 0.0,
                (
                    "audit_actions"
                    if bundle == "erp_posted_journal"
                    else "audit_action"
                ): actions if bundle == "erp_posted_journal" else (actions[0] if actions else ""),
            }
            expected_lines = params["lines"]

            def exact_line(
                row: dict[str, Any],
                expected_line: dict[str, Any],
            ) -> bool:
                return (
                    row["account_code"] == expected_line["account_code"]
                    and row["department_code"]
                    == expected_line.get("department_code")
                    and row["site_code"] == expected_line.get("site_code")
                    and row["item_id"] == expected_line.get("item_id")
                    and row["production_order_id"]
                    == expected_line.get("production_order_id")
                    and row["customer_id"] == expected_line.get("customer_id")
                    and row["vendor_id"] == expected_line.get("vendor_id")
                    and _close(
                        row["debit"],
                        expected_line.get("debit", 0.0),
                        "debit_amount",
                    )
                    and _close(
                        row["credit"],
                        expected_line.get("credit", 0.0),
                        "credit_amount",
                    )
                    and row["line_memo"]
                    == expected_line.get("line_memo", params["memo"])
                )

            line_match = len(lines) == len(expected_lines) and all(
                any(
                    exact_line(row, expected_line)
                    for row in lines
                )
                for expected_line in expected_lines
            )

            expected_audit_actions = (
                ["create_draft", "post"]
                if bundle == "erp_posted_journal"
                else ["create_draft"]
            )
            audit_details_exact = bool(header) and (
                len(journal_audit) == len(expected_audit_actions)
                and actions == expected_audit_actions
            )
            audit_times: list[datetime] = []
            if audit_details_exact:
                expected_details = [
                    {
                        "journal_id": header["id"],
                        "posting_date": params["posting_date"],
                        "lines": len(expected_lines),
                    }
                ]
                if bundle == "erp_posted_journal":
                    expected_details.append(
                        {
                            "journal_id": header["id"],
                            "period": params["posting_date"][:7],
                        }
                    )
                for row, details in zip(
                    journal_audit,
                    expected_details,
                    strict=True,
                ):
                    if (
                        row["actor"] != "rl-agent"
                        or row.get("details") != details
                    ):
                        audit_details_exact = False
                        break
                    try:
                        event_at = datetime.fromisoformat(row["event_at"])
                    except (TypeError, ValueError):
                        audit_details_exact = False
                        break
                    if event_at.tzinfo is None:
                        audit_details_exact = False
                        break
                    audit_times.append(event_at)

            created_at_exact = False
            if header and audit_details_exact and audit_times:
                try:
                    created_at = datetime.fromisoformat(header["created_at"])
                except (TypeError, ValueError):
                    created_at = None
                if created_at is not None and created_at.tzinfo is not None:
                    elapsed = (
                        audit_times[0] - created_at
                    ).total_seconds()
                    created_at_exact = 0 <= elapsed <= 60
                    if len(audit_times) > 1:
                        created_at_exact = (
                            created_at_exact
                            and audit_times == sorted(audit_times)
                        )
            state_criteria.extend(
                [
                    _criterion(
                        task_id,
                        "erp_header",
                        "Exactly one journal with the approved reference exists",
                        len(header_rows) == 1,
                        (
                            f"matching_header_count={len(header_rows)}; "
                            f"header={_stable_erp_evidence(header)!r}"
                        ),
                        category="integrity",
                        weight=10,
                    ),
                    _criterion(
                        task_id,
                        "erp_journal_header_values",
                        "Journal date, period, memo, source, and status are exact",
                        bool(header)
                        and header["posting_date"] == params["posting_date"]
                        and header["period"] == params["posting_date"][:7]
                        and header["memo"] == params["memo"]
                        and header["source"] == "Agent Journal"
                        and header["status"] == expected["status"]
                        and created_at_exact,
                        f"header={_stable_erp_evidence(header)!r}",
                        weight=10,
                    ),
                    _criterion(
                        task_id,
                        "erp_journal_lines",
                        "Journal lines exactly match approved accounts, sites, and amounts",
                        line_match,
                        f"lines={lines!r}",
                        weight=10,
                    ),
                    _criterion(
                        task_id,
                        "erp_journal_balanced",
                        "Journal debits and credits balance exactly",
                        _close(
                            expected["debits"],
                            expected["credits"],
                            "journal_balance_amount",
                        )
                        and expected["debits"] > 0,
                        f"debits={expected['debits']}; credits={expected['credits']}",
                        category="auditability",
                        weight=10,
                    ),
                    _criterion(
                        task_id,
                        "erp_journal_audit",
                        "Journal audit sequence is exact",
                        audit_details_exact,
                        f"audit={_stable_erp_evidence(journal_audit)!r}",
                        category="integrity",
                        weight=10,
                    ),
                ]
            )
        elif bundle == "erp_purchase_order":
            created = [
                row
                for row in audit
                if row["entity"] == "purchase_order" and row["action"] == "create"
            ]
            identifier = created[0]["identifier"] if len(created) == 1 else ""
            header_row = connection.execute(
                "SELECT * FROM purchase_orders WHERE po_number=?", (identifier,)
            ).fetchone()
            header = dict(header_row) if header_row else {}
            lines = (
                [
                    dict(row)
                    for row in connection.execute(
                        "SELECT * FROM purchase_order_lines WHERE po_id=? ORDER BY line_number",
                        (header["id"],),
                    )
                ]
                if header
                else []
            )
            total = sum(row["order_quantity"] * row["unit_price"] for row in lines)
            expected = {
                "po_number": identifier,
                "status": "Open",
                "line_count": len(lines),
                "order_value": round(total, 2),
                "audit_action": "create",
            }
            header_match = bool(header) and all(
                header[key] == params[key]
                for key in (
                    "vendor_id",
                    "site_code",
                    "order_date",
                    "expected_date",
                    "buyer_employee_id",
                )
            )
            line_match = len(lines) == len(params["lines"]) and all(
                any(
                    row["item_id"] == item["item_id"]
                    and _close(
                        row["order_quantity"],
                        item["quantity"],
                        "order_quantity",
                    )
                    and _close(
                        row["unit_price"],
                        item["unit_price"],
                        "unit_price",
                    )
                    for row in lines
                )
                for item in params["lines"]
            )
            buyer_row = connection.execute(
                """
                SELECT * FROM employees WHERE id=?
                """,
                (params["buyer_employee_id"],),
            ).fetchone()
            buyer = dict(buyer_row) if buyer_row else {}
            buyer_authorized = (
                bool(buyer)
                and bool(buyer["active"])
                and buyer["site_code"] == params["site_code"]
                and buyer["department_code"] == "SCM"
                and any(
                    token in str(buyer["title"]).casefold()
                    for token in ("buyer", "planner")
                )
            )
            authorized_item_ids = {
                str(row["id"])
                for row in connection.execute(
                    """
                    SELECT id FROM items WHERE primary_vendor_id=?
                    """,
                    (params["vendor_id"],),
                )
            }
            vendor_items_authorized = all(
                str(item["item_id"]) in authorized_item_ids
                for item in params["lines"]
            )
            state_criteria.extend(
                [
                    _criterion(task_id, "erp_po_single", "Exactly one PO was created", len(created) == 1, f"audit={_stable_erp_evidence(created)!r}", category="integrity", weight=10),
                    _criterion(task_id, "erp_po_header", "PO header exactly matches the authorization", header_match, f"header={header!r}", weight=10),
                    _criterion(task_id, "erp_po_lines", "PO lines exactly match item, quantity, and price authorization", line_match, f"lines={lines!r}", weight=10),
                    _criterion(task_id, "erp_po_value", "PO extension equals the approved total", _close(total, sum(x["quantity"] * x["unit_price"] for x in params["lines"]), "order_value"), f"total={total}", weight=10),
                    _criterion(
                        task_id,
                        "erp_po_buyer_authorization",
                        (
                            "The buyer is active, belongs to the PO site and "
                            "Supply Chain, and has a planner/buyer role"
                        ),
                        buyer_authorized,
                        f"buyer={_stable_erp_evidence(buyer)!r}",
                        category="integrity",
                        weight=10,
                        failure_cap=0.49,
                    ),
                    _criterion(
                        task_id,
                        "erp_po_vendor_item_authorization",
                        (
                            "Every PO item is authorized to the selected "
                            "primary vendor"
                        ),
                        vendor_items_authorized,
                        (
                            f"authorized_item_ids="
                            f"{sorted(authorized_item_ids)!r}; "
                            f"requested={[item['item_id'] for item in params['lines']]!r}"
                        ),
                        category="integrity",
                        weight=10,
                        failure_cap=0.49,
                    ),
                    _criterion(task_id, "erp_po_zero_receipts", "New PO has no receipt or invoice activity", all(row["received_quantity"] == 0 and row["invoiced_quantity"] == 0 for row in lines), f"lines={lines!r}", category="integrity", weight=5),
                ]
            )
        elif bundle == "erp_production_order":
            created = [
                row
                for row in audit
                if row["entity"] == "production_order" and row["action"] == "create"
            ]
            identifier = created[0]["identifier"] if len(created) == 1 else ""
            order_row = connection.execute(
                "SELECT * FROM production_orders WHERE order_number=?",
                (identifier,),
            ).fetchone()
            order = dict(order_row) if order_row else {}
            actions = [
                row["action"]
                + (
                    f":{row['details'].get('to')}"
                    if row["action"] == "status_change"
                    else ""
                )
                for row in audit
                if row["entity"] == "production_order"
                and row["identifier"] == identifier
            ]
            final_status = params.get("final_status", "Scheduled")
            expected = {
                "order_number": identifier,
                "status": final_status,
                "item_id": params["item_id"],
                "site_code": params["site_code"],
                "order_quantity": params["order_quantity"],
                "standard_unit_cost": order.get("standard_unit_cost", 0.0),
                "audit_actions": actions,
            }
            effective_bom = connection.execute(
                """
                SELECT id FROM bom_headers
                WHERE item_id=? AND site_code=? AND status='Active'
                  AND effective_from<=?
                  AND (effective_to IS NULL OR effective_to>=?)
                ORDER BY effective_from DESC LIMIT 1
                """,
                (
                    params["item_id"],
                    params["site_code"],
                    params["scheduled_start"],
                    params["scheduled_start"],
                ),
            ).fetchone()
            effective_routing = connection.execute(
                """
                SELECT id FROM routing_headers
                WHERE item_id=? AND site_code=? AND status='Active'
                  AND effective_from<=?
                  AND (effective_to IS NULL OR effective_to>=?)
                ORDER BY effective_from DESC LIMIT 1
                """,
                (
                    params["item_id"],
                    params["site_code"],
                    params["scheduled_start"],
                    params["scheduled_start"],
                ),
            ).fetchone()
            site_cost = connection.execute(
                """
                SELECT standard_material_cost + standard_labor_cost
                     + standard_variable_overhead + standard_fixed_overhead
                     + standard_outside_processing
                FROM item_sites
                WHERE item_id=? AND site_code=?
                """,
                (params["item_id"], params["site_code"]),
            ).fetchone()
            source_reference = str(params["source_reference"])
            expected_source_type = (
                "Planned Order"
                if source_reference.startswith("PLAN-")
                else "Manual"
            )
            planned_order_row = (
                connection.execute(
                    "SELECT * FROM planned_orders WHERE id=?",
                    (source_reference,),
                ).fetchone()
                if expected_source_type == "Planned Order"
                else None
            )
            planned_order = (
                dict(planned_order_row) if planned_order_row else {}
            )
            plan_link_match = (
                expected_source_type != "Planned Order"
                or (
                    bool(planned_order)
                    and planned_order["item_id"] == params["item_id"]
                    and planned_order["site_code"] == params["site_code"]
                    and _close(
                        planned_order["quantity"],
                        params["order_quantity"],
                        "order_quantity",
                    )
                    and planned_order["release_date"]
                    == params["scheduled_start"]
                    and planned_order["required_date"]
                    == params["scheduled_finish"]
                )
            )
            exact = (
                bool(order)
                and effective_bom is not None
                and effective_routing is not None
                and site_cost is not None
                and all(
                    (
                        _close(
                            order[key],
                            params[key],
                            "order_quantity",
                        )
                        if key == "order_quantity"
                        else order[key] == params[key]
                    )
                    for key in (
                        "item_id",
                        "site_code",
                        "order_quantity",
                        "scheduled_start",
                        "scheduled_finish",
                        "source_reference",
                    )
                )
                and order["bom_id"] == effective_bom["id"]
                and order["routing_id"] == effective_routing["id"]
                and order["source_type"] == expected_source_type
                and order["created_date"] == "2026-06-30"
                and _close(
                    order["standard_unit_cost"],
                    site_cost[0],
                    "standard_unit_cost",
                )
                and _close(
                    order["completed_quantity"],
                    0.0,
                    "completed_quantity",
                )
                and _close(
                    order["scrapped_quantity"],
                    0.0,
                    "scrapped_quantity",
                )
                and all(
                    _close(order[key], 0.0, key)
                    for key in (
                        "actual_material_cost",
                        "actual_labor_cost",
                        "actual_variable_overhead",
                        "actual_fixed_overhead",
                        "actual_outside_processing",
                    )
                )
                and order["status"] == final_status
                and order.get("actual_finish") is None
                and (
                    (
                        order.get("actual_start")
                        == max("2026-06-30", params["scheduled_start"])
                    )
                    if final_status == "Started"
                    else order.get("actual_start") is None
                )
            )
            materials = (
                [
                    dict(row)
                    for row in connection.execute(
                        """
                        SELECT * FROM production_order_materials
                        WHERE production_order_id=? ORDER BY line_number
                        """,
                        (order["id"],),
                    )
                ]
                if order
                else []
            )
            expected_materials = (
                [
                    dict(row)
                    for row in connection.execute(
                        """
                        SELECT b.line_number,b.component_item_id,
                               b.quantity_per,b.scrap_factor,
                               i.standard_material_cost
                        FROM bom_lines b
                        JOIN item_sites i
                          ON i.item_id=b.component_item_id
                         AND i.site_code=?
                        WHERE b.bom_id=? ORDER BY b.line_number
                        """,
                        (order["site_code"], order["bom_id"]),
                    )
                ]
                if order
                else []
            )
            material_match = (
                len(materials) == len(expected_materials)
                and bool(materials)
                and all(
                    actual["line_number"] == expected_row["line_number"]
                    and actual["component_item_id"]
                    == expected_row["component_item_id"]
                    and _close(
                        actual["planned_quantity"],
                        float(order["order_quantity"])
                        * float(expected_row["quantity_per"])
                        * (1 + float(expected_row["scrap_factor"])),
                        "planned_quantity",
                    )
                    and _close(
                        actual["standard_unit_cost"],
                        expected_row["standard_material_cost"],
                        "standard_unit_cost",
                    )
                    and _close(
                        actual["issued_quantity"],
                        0.0,
                        "issued_quantity",
                    )
                    and _close(
                        actual["actual_unit_cost"],
                        expected_row["standard_material_cost"],
                        "actual_unit_cost",
                    )
                    and actual["issue_date"] is None
                    and actual["substitution"] == 0
                    for actual, expected_row in zip(
                        materials,
                        expected_materials,
                    )
                )
            )
            operations = (
                [
                    dict(row)
                    for row in connection.execute(
                        """
                        SELECT * FROM production_order_operations
                        WHERE production_order_id=? ORDER BY operation_number
                        """,
                        (order["id"],),
                    )
                ]
                if order
                else []
            )
            expected_operations = (
                [
                    dict(row)
                    for row in connection.execute(
                        """
                        SELECT o.operation_number,o.work_center_code,
                               o.setup_hours,o.run_hours_per_unit,
                               w.labor_rate,w.machine_rate
                        FROM routing_operations o
                        JOIN work_centers w ON w.code=o.work_center_code
                        WHERE o.routing_id=? ORDER BY o.operation_number
                        """,
                        (order["routing_id"],),
                    )
                ]
                if order
                else []
            )
            operation_match = (
                len(operations) == len(expected_operations)
                and bool(operations)
                and all(
                    actual["operation_number"]
                    == expected_row["operation_number"]
                    and actual["work_center_code"]
                    == expected_row["work_center_code"]
                    and _close(
                        actual["planned_setup_hours"],
                        expected_row["setup_hours"],
                        "planned_setup_hours",
                    )
                    and _close(
                        actual["planned_run_hours"],
                        float(expected_row["run_hours_per_unit"])
                        * float(order["order_quantity"]),
                        "planned_run_hours",
                    )
                    and _close(
                        actual["labor_rate_standard"],
                        expected_row["labor_rate"],
                        "labor_rate_standard",
                    )
                    and _close(
                        actual["machine_rate_standard"],
                        expected_row["machine_rate"],
                        "machine_rate_standard",
                    )
                    and _close(
                        actual["actual_setup_hours"],
                        0.0,
                        "actual_setup_hours",
                    )
                    and _close(
                        actual["actual_run_hours"],
                        0.0,
                        "actual_run_hours",
                    )
                    and _close(
                        actual["labor_rate_actual"],
                        expected_row["labor_rate"],
                        "labor_rate_actual",
                    )
                    and actual["completed_date"] is None
                    and actual["status"] == "Scheduled"
                    for actual, expected_row in zip(
                        operations,
                        expected_operations,
                    )
                )
            )
            variance_row = (
                connection.execute(
                    """
                    SELECT * FROM production_variances
                    WHERE production_order_id=?
                    """,
                    (order["id"],),
                ).fetchone()
                if order
                else None
            )
            variance = dict(variance_row) if variance_row else {}
            variance_match = (
                bool(variance)
                and all(
                    _close(variance[key], 0.0, key)
                    for key in (
                        "material_price_variance",
                        "material_usage_variance",
                        "labor_rate_variance",
                        "labor_efficiency_variance",
                        "variable_overhead_variance",
                        "fixed_overhead_volume_variance",
                        "scrap_variance",
                        "substitution_variance",
                        "total_variance",
                    )
                )
                and variance["settled_date"] is None
            )
            order_activity_zero = bool(order) and all(
                _close(order[key], 0.0, key)
                for key in (
                    "completed_quantity",
                    "scrapped_quantity",
                    "actual_material_cost",
                    "actual_labor_cost",
                    "actual_variable_overhead",
                    "actual_fixed_overhead",
                    "actual_outside_processing",
                )
            ) and order["actual_finish"] is None
            material_activity_zero = bool(materials) and all(
                _close(row["issued_quantity"], 0.0, "issued_quantity")
                and row["issue_date"] is None
                and row["substitution"] == 0
                for row in materials
            )
            operation_activity_zero = bool(operations) and all(
                _close(row["actual_setup_hours"], 0.0, "actual_setup_hours")
                and _close(row["actual_run_hours"], 0.0, "actual_run_hours")
                and row["completed_date"] is None
                and row["status"] == "Scheduled"
                for row in operations
            )
            expected_actions = ["create"]
            if final_status in {"Released", "Started"}:
                expected_actions.append("status_change:Released")
            if final_status == "Started":
                expected_actions.append("status_change:Started")
            state_criteria.extend(
                [
                    _criterion(task_id, "erp_mo_single", "Exactly one production order was created", len(created) == 1, f"audit={_stable_erp_evidence(created)!r}", category="integrity", weight=10),
                    _criterion(task_id, "erp_mo_header", "Production order header and final status exactly match authorization", exact, f"order={order!r}", weight=10),
                    _criterion(
                        task_id,
                        "erp_mo_source_link",
                        (
                            "A PLAN-* production order preserves the exact "
                            "approved planned-order item, site, quantity, "
                            "release date, and required date"
                        ),
                        plan_link_match,
                        (
                            f"source_type={expected_source_type!r}; "
                            f"planned_order={planned_order!r}"
                        ),
                        category="auditability",
                        weight=10,
                    ),
                    _criterion(task_id, "erp_mo_bom_copy", "The complete effective BOM was copied exactly into planned material lines", material_match, f"materials={materials!r}; expected={expected_materials!r}", category="auditability", weight=5),
                    _criterion(task_id, "erp_mo_routing_copy", "The complete effective routing was copied exactly into planned operations", operation_match, f"operations={operations!r}; expected={expected_operations!r}", category="auditability", weight=5),
                    _criterion(task_id, "erp_mo_zero_activity", "The newly scheduled order has no fabricated issue, completion, actual-cost, or variance activity", order_activity_zero and material_activity_zero and operation_activity_zero and variance_match, f"order={order!r}; variance={variance!r}", category="integrity", weight=10),
                    _criterion(task_id, "erp_mo_audit", "The production lifecycle audit sequence is exact", actions == expected_actions, f"actions={actions!r}; expected={expected_actions!r}", category="integrity", weight=10),
                ]
            )
        elif bundle == "erp_quality_hold":
            holds = [
                row
                for row in audit
                if row["entity"] == "inventory" and row["action"] == "place_hold"
            ]
            identifier = holds[0]["identifier"] if len(holds) == 1 else ""
            order_row = connection.execute(
                "SELECT * FROM quality_orders WHERE id=?", (identifier,)
            ).fetchone()
            order = dict(order_row) if order_row else {}
            actions = [
                row["action"]
                for row in audit
                if row["entity"] == "inventory"
                and row["identifier"] == identifier
            ]
            release_events = [
                row
                for row in audit
                if row["entity"] == "inventory"
                and row["identifier"] == identifier
                and row["action"] == "release_hold"
            ]
            final_status = "Closed" if params.get("release") else "Open"
            balance: dict[str, Any] = {}
            if order and "/" in str(order.get("reference_id") or ""):
                warehouse_code, location_code = order["reference_id"].split("/", 1)
                balance_row = connection.execute(
                    """
                    SELECT * FROM inventory_balances
                    WHERE item_id=? AND site_code=? AND warehouse_code=?
                      AND location_code=? AND lot_number=?
                    """,
                    (
                        order["item_id"],
                        order["site_code"],
                        warehouse_code,
                        location_code,
                        order.get("lot_number") or "",
                    ),
                ).fetchone()
                balance = dict(balance_row) if balance_row else {}
            standard_cost = float(balance.get("standard_unit_cost") or 0.0)
            expected_exposure = round(
                float(params["quantity"]) * standard_cost,
                2,
            )
            expected = {
                "quality_order_id": identifier,
                "status": final_status,
                "placed_quantity": params["quantity"],
                "estimated_financial_exposure": expected_exposure,
                "audit_actions": actions,
            }
            exact = (
                bool(order)
                and order["item_id"] == params["item_id"]
                and order["site_code"] == params["site_code"]
                and _close(
                    order["quantity_inspected"],
                    params["quantity"],
                    "held_quantity",
                )
                and _close(
                    order["quantity_failed"],
                    0.0,
                    "failed_quantity",
                )
                and order["disposition"]
                == (
                    params.get("disposition")
                    if params.get("release")
                    else params["reason"]
                )
                and order["status"] == final_status
            )
            expected_actions = (
                ["place_hold", "release_hold"]
                if params.get("release")
                else ["place_hold"]
            )
            place_details = holds[0].get("details", {}) if len(holds) == 1 else {}
            before_hold = _number(place_details.get("quality_hold_before"))
            after_hold = _number(place_details.get("quality_hold_after"))
            audit_quantity_exact = (
                before_hold is not None
                and after_hold is not None
                and _close(
                    after_hold - before_hold,
                    params["quantity"],
                    "held_quantity",
                )
            )
            current_hold = _number(balance.get("quality_hold_quantity"))
            inventory_exact = False
            if audit_quantity_exact and current_hold is not None:
                if params.get("release") and len(release_events) == 1:
                    release_details = release_events[0].get("details", {})
                    release_before = _number(
                        release_details.get("quality_hold_before")
                    )
                    release_after = _number(
                        release_details.get("quality_hold_after")
                    )
                    inventory_exact = (
                        release_before is not None
                        and release_after is not None
                        and _close(
                            release_before,
                            after_hold,
                            "held_quantity",
                        )
                        and _close(
                            release_after,
                            before_hold,
                            "held_quantity",
                        )
                        and _close(
                            current_hold,
                            before_hold,
                            "held_quantity",
                        )
                    )
                elif not params.get("release"):
                    inventory_exact = _close(
                        current_hold,
                        after_hold,
                        "held_quantity",
                    )
            state_criteria.extend(
                [
                    _criterion(task_id, "erp_hold_single", "Exactly one inventory hold was created", len(holds) == 1, f"holds={_stable_erp_evidence(holds)!r}", category="integrity", weight=10),
                    _criterion(task_id, "erp_hold_values", "Hold item, site, quantity, reason/disposition, and status are exact", exact, f"quality_order={order!r}", weight=10),
                    _criterion(task_id, "erp_hold_exposure", "Financial exposure equals held quantity times the selected inventory standard cost", bool(order) and _close(order["estimated_financial_exposure"], expected_exposure, "estimated_financial_exposure"), f"actual={order.get('estimated_financial_exposure')!r}; expected={expected_exposure!r}", weight=10),
                    _criterion(task_id, "erp_hold_audit", "Hold/release audit sequence is exact", actions == expected_actions, f"actions={actions!r}", category="integrity", weight=10),
                    _criterion(task_id, "erp_hold_final_inventory", "Released holds exactly restore prior blocked quantity; unreleased holds remain exactly blocked", inventory_exact, f"balance={balance!r}; place={place_details!r}; release={_stable_erp_evidence(release_events)!r}", category="integrity", weight=5),
                ]
            )
        else:
            raise ValueError(f"Unsupported ERP bundle {bundle}")
    return expected, state_criteria, audit


def _erp_criteria(
    task_id: str,
    answer: Any,
    database_path: Path,
    gold: dict[str, Any],
) -> list[Criterion]:
    mapping, exact_json, evidence = _strict_json(answer)
    expected, state_criteria, audit = _erp_state(gold, database_path)
    keys = gold["value_keys"]
    criteria = [
        _criterion(task_id, "json_parse", "The ERP completion response contains valid JSON", bool(mapping), evidence, category="structure", weight=3),
        _criterion(task_id, "json_only", "The ERP completion response is JSON without prose", exact_json, evidence, category="structure", weight=1),
        _criterion(task_id, "exact_key_order", "The ERP completion response uses exactly the required keys in order", list(mapping) == keys, f"actual={list(mapping)!r}; expected={keys!r}", category="decision", weight=10, failure_cap=0.49),
        _criterion(task_id, "authorized_audit_count", "The task produced only the required audited workflow events", len(audit) in ({2} if gold["bundle"] in {"erp_posted_journal"} else {3} if gold["task_id"] == "task_099" else {2} if gold["task_id"] in {"task_080", "task_100"} else {1}), f"audit_count={len(audit)}", category="integrity", weight=10),
    ]
    for key in keys:
        actual = mapping.get(key)
        target = expected.get(key)
        if isinstance(target, list):
            matched = _sequence_matches(actual, target)
            valid = isinstance(actual, (list, str))
        else:
            matched = _value_matches(actual, target, key)
            valid = (
                _number(actual) is not None
                if isinstance(target, (int, float)) and not isinstance(target, bool)
                else isinstance(actual, str) and bool(actual.strip())
            )
        slug = _normalize(key).replace(" ", "_")
        criteria.extend(
            [
                _criterion(task_id, f"{slug}__present_valid", f"`{key}` is present with the required usable type", key in mapping and valid, f"actual={actual!r}", category="auditability", weight=5),
                _criterion(task_id, f"{slug}__correct", f"`{key}` agrees with the audited ERP state", matched, f"actual={actual!r}; expected={target!r}", weight=10),
            ]
        )
    return criteria + state_criteria


def _task_specific_document_post_criteria(
    task_id: str,
    workspace_root: Path,
    gold: dict[str, Any],
) -> list[Criterion]:
    required_terms_by_task = {
        "task_078": (
            "leverage",
            "tangible net worth",
            "liquidity",
            "lender",
            "treasurer",
            "controller",
            "cash outlook",
            "fact",
            "assumption",
            "owner",
            "timing",
        ),
        "task_088": (
            "tax benefit",
            "current tax",
            "deferred tax",
            "tax director",
            "controller",
            "fact",
            "assumption",
            "owner",
            "timing",
        ),
        "task_089": (
            "reported ebitda",
            "normalized",
            "maintenance capital",
            "terminal",
            "net debt",
            "sensitivity",
            "fact",
            "assumption",
            "owner",
            "timing",
        ),
        "task_096": (
            "pre close",
            "margin",
            "backlog to cash",
            "working capital",
            "treasurer",
            "vp operations",
            "vp sales",
            "complete",
            "open",
            "hold",
            "approve",
            "fact",
            "assumption",
            "owner",
            "timing",
        ),
        "task_097": (
            "fy27",
            "2025",
            "monthly",
            "operating income",
            "working capital",
            "vp sales",
            "vp supply chain",
            "fact",
            "assumption",
            "owner",
            "timing",
        ),
    }
    required_terms = required_terms_by_task.get(task_id)
    if required_terms is None:
        return []
    path = workspace_root / gold["artifact"]["path"]
    text = ""
    native_rows: list[list[str]] = []
    native_tables: list[list[list[str]]] = []
    paragraph_texts: list[str] = []
    error = "required memorandum is missing"
    if os.path.lexists(path):
        try:
            with validated_native_package_copy(
                path,
                workspace_root=workspace_root,
            ) as (snapshot, _inspection):
                document = Document(snapshot)
                text, native_rows = _flatten_docx(document)
                native_snapshot = _native_docx_snapshot(document)
                native_tables = [
                    [list(row) for row in table.rows]
                    for table in native_snapshot.tables
                ]
                paragraph_texts = [
                    paragraph.text.strip()
                    for paragraph in native_snapshot.paragraphs
                    if paragraph.table_depth == 0
                    and paragraph.text.strip()
                ]
                error = "memorandum parsed"
        except Exception as exc:
            error = f"{type(exc).__name__}: {exc}"
    normalized_text = _normalize(text)
    criteria = [
        _criterion(
            task_id,
            "task_specific_decision_content",
            (
                "The memorandum contains the task-specific economics, "
                "period architecture, risks, owners, and dated actions"
            ),
            all(term in normalized_text for term in required_terms),
            f"{error}; required_terms={required_terms!r}",
            category="decision",
            weight=10,
            failure_cap=0.49,
        )
    ]
    if task_id == "task_078":
        normalized_paragraphs = [
            _normalize(paragraph) for paragraph in paragraph_texts
        ]
        explicit_owner_timing = any(
            "owner" in paragraph and "timing" in paragraph
            for paragraph in normalized_paragraphs
        )
        explicit_fact_assumption = any(
            paragraph.startswith("fact")
            and "assumption" in paragraph
            for paragraph in normalized_paragraphs
        )
        criteria.append(
            _criterion(
                task_id,
                "liquidity_memo_labels_retained",
                (
                    "The enhanced memorandum retains explicit Owner/Timing "
                    "and Fact/Assumption labels in native paragraphs"
                ),
                explicit_owner_timing and explicit_fact_assumption,
                (
                    f"owner_timing={explicit_owner_timing}; "
                    f"fact_assumption={explicit_fact_assumption}"
                ),
                category="decision",
                weight=10,
                failure_cap=0.49,
            )
        )
    if task_id in {"task_088", "task_089", "task_096", "task_097"}:
        normalized_paragraphs = [
            _normalize(paragraph) for paragraph in paragraph_texts
        ]
        explicit_owner_timing = any(
            "owner" in paragraph and "timing" in paragraph
            for paragraph in normalized_paragraphs
        )
        explicit_fact_assumption = any(
            paragraph.startswith("fact")
            and "assumption" in paragraph
            for paragraph in normalized_paragraphs
        )
        criteria.append(
            _criterion(
                task_id,
                "memo_labels_retained",
                (
                    "The enhanced memorandum retains explicit Owner/Timing "
                    "and Fact/Assumption labels in native paragraphs"
                ),
                explicit_owner_timing and explicit_fact_assumption,
                (
                    f"owner_timing={explicit_owner_timing}; "
                    f"fact_assumption={explicit_fact_assumption}"
                ),
                category="decision",
                weight=10,
                failure_cap=0.49,
            )
        )

    supporting = gold.get("supporting_values")
    if not isinstance(supporting, dict):
        return criteria

    if task_id == "task_089":
        bridge_evidence = _semantic_native_table_truth(
            native_tables,
            column_aliases={
                "component": ("Component", "Bridge component"),
                "amount": ("Amount", "Value"),
                "classification": (
                    "Classification",
                    "Fact / assumption",
                ),
            },
            key_fields=("component",),
            expected_rows=[
                {
                    "component": label,
                    "amount": supporting[key],
                    "classification": classification,
                }
                for label, key, classification in (
                    ("Reported EBITDA", "reported_ebitda", "Fact"),
                    (
                        "Production variance normalization",
                        "production_variance_normalization",
                        "Fact",
                    ),
                    (
                        "Quality normalization",
                        "quality_normalization",
                        "Fact",
                    ),
                    (
                        "Corrective maintenance normalization",
                        "maintenance_normalization",
                        "Fact",
                    ),
                    (
                        "Normalized EBITDA",
                        "normalized_ebitda",
                        "Calculated",
                    ),
                    (
                        "Maintenance capital expenditures",
                        "maintenance_capital_expenditures",
                        "Fact",
                    ),
                )
            ],
        )
        forecast_evidence = _semantic_native_table_truth(
            native_tables,
            column_aliases={
                "year_component": (
                    "Year / component",
                    "Forecast year / component",
                ),
                "cash_flow_value": (
                    "Cash flow / value",
                    "Cash flow or terminal value",
                ),
                "present_value": ("Present value", "PV"),
            },
            key_fields=("year_component",),
            expected_rows=[
                {
                    "year_component": item["forecast_year"],
                    "cash_flow_value": item["cash_flow"],
                    "present_value": item["present_value"],
                }
                for item in supporting["explicit_forecast"]
            ]
            + [
                {
                    "year_component": "Terminal",
                    "cash_flow_value": supporting["terminal_value"],
                    "present_value": supporting[
                        "terminal_present_value"
                    ],
                }
            ],
        )
        sensitivity_evidence = _semantic_native_table_truth(
            native_tables,
            column_aliases={
                "terminal_growth": (
                    "Terminal growth",
                    "Terminal growth rate",
                ),
                "enterprise_value": ("Enterprise value", "EV"),
                "equity_value": ("Equity value",),
            },
            key_fields=("terminal_growth",),
            expected_rows=[
                {
                    "terminal_growth": item["terminal_growth"],
                    "enterprise_value": item["enterprise_value"],
                    "equity_value": item["equity_value"],
                }
                for item in supporting["terminal_growth_sensitivity"]
            ],
        )
        normalized_ebitda = _number(
            supporting.get("normalized_ebitda")
        )
        maintenance_capex = _number(
            supporting.get("maintenance_capital_expenditures")
        )
        base_cash_flow = _number(
            gold.get("values", {}).get("base_cash_flow")
        )
        cash_tax_rate = _number(supporting.get("cash_tax_rate"))
        implied_cash_tax_rate = None
        if (
            normalized_ebitda is not None
            and maintenance_capex is not None
            and base_cash_flow is not None
            and abs(normalized_ebitda) > 1e-12
        ):
            implied_cash_tax_rate = 1.0 - (
                base_cash_flow + maintenance_capex
            ) / normalized_ebitda
        parameters = gold.get("parameters", {})
        cash_tax_component_rows = {
            "normalized_ebitda": _paired_metric(
                native_rows,
                "normalized_ebitda",
                supporting.get("normalized_ebitda"),
            ),
            "maintenance_capital_expenditures": _paired_metric(
                native_rows,
                "maintenance_capital_expenditures",
                supporting.get("maintenance_capital_expenditures"),
            ),
            "base_cash_flow": _paired_metric(
                native_rows,
                "base_cash_flow",
                gold.get("values", {}).get("base_cash_flow"),
            ),
        }
        cash_tax_assumption_met = (
            cash_tax_rate is not None
            and implied_cash_tax_rate is not None
            and isinstance(parameters, dict)
            and _value_matches(
                parameters.get("cash_tax_rate"),
                cash_tax_rate,
                "cash_tax_rate",
            )
            and _close(
                implied_cash_tax_rate,
                cash_tax_rate,
                "cash_tax_rate",
            )
            and all(cash_tax_component_rows.values())
            and "cash tax" in normalized_text
            and "one less the cash tax rate" in normalized_text
        )
        criteria.extend(
            [
                _criterion(
                    task_id,
                    "valuation_cash_tax_assumption",
                    (
                        "The native valuation bridge makes the exact cash-tax "
                        "assumption auditable from normalized EBITDA, "
                        "maintenance capital, and base cash flow"
                    ),
                    cash_tax_assumption_met,
                    (
                        f"cash_tax_rate={cash_tax_rate!r}; "
                        "task_parameter="
                        f"{parameters.get('cash_tax_rate') if isinstance(parameters, dict) else None!r}; "
                        f"implied_cash_tax_rate={implied_cash_tax_rate!r}; "
                        f"component_rows={cash_tax_component_rows!r}"
                    ),
                    category="core_finance",
                    weight=10,
                    failure_cap=0.49,
                ),
                _criterion(
                    task_id,
                    "valuation_qoe_maintenance_bridge",
                    (
                        "The native table ties the complete QoE normalization "
                        "and maintenance-capital bridge"
                    ),
                    bridge_evidence["field_values_match"],
                    f"bridge_evidence={bridge_evidence!r}",
                    category="core_finance",
                    weight=10,
                    failure_cap=0.49,
                ),
                _criterion(
                    task_id,
                    "valuation_forecast_terminal_truth",
                    (
                        "The native table presents all five forecast cash "
                        "flows/PVs and exact terminal value/PV"
                    ),
                    forecast_evidence["field_values_match"],
                    f"forecast_evidence={forecast_evidence!r}",
                    category="core_finance",
                    weight=10,
                    failure_cap=0.49,
                ),
                _criterion(
                    task_id,
                    "valuation_sensitivity_truth",
                    (
                        "The native five-case terminal-growth sensitivity "
                        "ties exact enterprise and equity values"
                    ),
                    sensitivity_evidence["field_values_match"],
                    f"sensitivity_evidence={sensitivity_evidence!r}",
                    category="core_finance",
                    weight=10,
                    failure_cap=0.49,
                ),
            ]
        )
    elif task_id == "task_096":
        controls_evidence = _semantic_native_table_truth(
            native_tables,
            column_aliases={
                "control": ("Control", "Close control"),
                "status": ("Status", "Control status"),
                "evidence": ("Evidence", "Result"),
            },
            key_fields=("control",),
            expected_rows=[
                {
                    "control": item["control"],
                    "status": item["status"],
                    "evidence": item["evidence"],
                }
                for item in supporting["control_register"]
            ],
        )
        residual = next(
            item
            for item in supporting["control_register"]
            if item["control"] == "Negative WIP residual review"
        )
        exception_count_met = bool(
            re.search(
                rf"\b{int(residual['exception_count'])}\s+orders?\b",
                text,
                flags=re.I,
            )
        )
        decisions_evidence = _semantic_native_table_truth(
            native_tables,
            column_aliases={
                "decision": ("Decision", "Approve / hold"),
                "gate": ("Gate", "Decision gate"),
                "owner": ("Owner", "Accountable owner"),
                "due": ("Due", "Due date", "Timing"),
            },
            key_fields=("decision", "gate"),
            expected_rows=[
                {
                    "decision": item["decision"],
                    "gate": item["gate"],
                    "owner": item["owner"],
                    "due": item["due"],
                }
                for item in supporting["decision_register"]
            ],
        )
        exposure_register = supporting.get("exposure_register")
        exposure_rows = (
            exposure_register
            if isinstance(exposure_register, list)
            and all(isinstance(item, dict) for item in exposure_register)
            else []
        )
        exposure_by_label = {
            _normalize(item.get("exposure")): item
            for item in exposure_rows
        }
        authoritative_exposures = {
            "margin": gold["values"]["ytd_operating_income"],
            "backlog to cash": gold["values"]["backlog"],
            "working capital": (
                gold["values"]["inventory"] + gold["values"]["wip"]
            ),
            "liquidity": gold["values"]["total_liquidity"],
        }
        exposure_native_support = {
            "margin": _paired_metric(
                native_rows,
                "ytd_operating_income",
                authoritative_exposures["margin"],
            ),
            "backlog to cash": _paired_metric(
                native_rows,
                "backlog",
                authoritative_exposures["backlog to cash"],
            ),
            "working capital": (
                _paired_metric(
                    native_rows,
                    "inventory",
                    gold["values"]["inventory"],
                )
                and _paired_metric(
                    native_rows,
                    "wip",
                    gold["values"]["wip"],
                )
                and _text_contains(
                    text,
                    authoritative_exposures["working capital"],
                    "working_capital",
                )
            ),
            "liquidity": _paired_metric(
                native_rows,
                "total_liquidity",
                authoritative_exposures["liquidity"],
            ),
        }
        exposure_checks = {
            label: (
                label in exposure_by_label
                and isinstance(
                    exposure_by_label[label].get("basis"),
                    str,
                )
                and bool(exposure_by_label[label]["basis"].strip())
                and _value_matches(
                    exposure_by_label[label].get("amount"),
                    expected,
                    f"{label}_exposure",
                )
                and exposure_native_support[label]
                and label in normalized_text
            )
            for label, expected in authoritative_exposures.items()
        }
        exposure_register_met = (
            len(exposure_rows) == len(authoritative_exposures)
            and len(exposure_by_label) == len(authoritative_exposures)
            and all(exposure_checks.values())
        )
        criteria.extend(
            [
                _criterion(
                    task_id,
                    "close_control_register_truth",
                    (
                        "All four close controls have exact Complete/Open "
                        "status and the quantified WIP exception"
                    ),
                    controls_evidence["field_values_match"]
                    and exception_count_met,
                    (
                        f"controls_evidence={controls_evidence!r}; "
                        f"exception_count_met={exception_count_met}"
                    ),
                    category="core_finance",
                    weight=10,
                    failure_cap=0.49,
                ),
                _criterion(
                    task_id,
                    "close_exposure_register_truth",
                    (
                        "The memorandum retains the exact margin, "
                        "backlog-to-cash, working-capital, and liquidity "
                        "exposures with native support"
                    ),
                    exposure_register_met,
                    (
                        f"exposure_checks={exposure_checks!r}; "
                        f"native_support={exposure_native_support!r}; "
                        f"labels={sorted(exposure_by_label)!r}"
                    ),
                    category="core_finance",
                    weight=10,
                    failure_cap=0.49,
                ),
                _criterion(
                    task_id,
                    "close_approve_hold_register",
                    (
                        "The native decision register contains every exact "
                        "APPROVE/HOLD gate, owner, and due date"
                    ),
                    decisions_evidence["field_values_match"],
                    f"decisions_evidence={decisions_evidence!r}",
                    category="decision",
                    weight=10,
                    failure_cap=0.49,
                ),
            ]
        )
    elif task_id == "task_097":
        milestones = supporting.get("monthly_milestones", [])
        milestones_evidence = _semantic_native_table_truth(
            native_tables,
            column_aliases={
                "month": ("Month", "FY27 month"),
                "revenue": ("Revenue",),
                "gross_profit": ("Gross profit", "GP"),
                "operating_income": ("Operating income", "OI"),
                "incremental_working_capital": (
                    "Incremental WC",
                    "Incremental working capital",
                ),
                "owner": ("Owner", "Accountable executive"),
                "board_gate": ("Board gate", "Gate / review"),
            },
            key_fields=("month",),
            expected_rows=[
                {
                    "month": item["month"],
                    "revenue": item["revenue"],
                    "gross_profit": item["gross_profit"],
                    "operating_income": item["operating_income"],
                    "incremental_working_capital": item[
                        "incremental_working_capital"
                    ],
                    "owner": item["owner"],
                    "board_gate": item["board_gate"],
                }
                for item in milestones
            ]
            if isinstance(milestones, list)
            else [],
        )
        milestones_met = (
            isinstance(milestones, list)
            and len(milestones) == 12
            and milestones_evidence["field_values_match"]
        )
        annual_tie_outs_met = all(
            round(
                sum(float(item[support_key]) for item in milestones),
                2,
            )
            == round(float(gold["values"][annual_key]), 2)
            for support_key, annual_key in (
                ("revenue", "scenario_revenue"),
                ("gross_profit", "scenario_gross_profit"),
                ("operating_income", "scenario_operating_income"),
                (
                    "incremental_working_capital",
                    "incremental_working_capital",
                ),
            )
        )
        value_creation_priorities = supporting.get(
            "value_creation_priorities"
        )
        priority_rows = (
            value_creation_priorities
            if isinstance(value_creation_priorities, list)
            and all(
                isinstance(item, dict)
                for item in value_creation_priorities
            )
            else []
        )
        priority_by_label = {
            _normalize(item.get("priority")): item
            for item in priority_rows
        }
        priority_definitions = {
            "price and volume realization": {
                "amount": (
                    gold["values"]["scenario_revenue"]
                    - gold["values"]["baseline_revenue"]
                ),
                "concepts": ("price volume",),
                "native": (
                    _paired_metric(
                        native_rows,
                        "baseline_revenue",
                        gold["values"]["baseline_revenue"],
                    )
                    and _paired_metric(
                        native_rows,
                        "scenario_revenue",
                        gold["values"]["scenario_revenue"],
                    )
                ),
            },
            "operating profit delivery": {
                "amount": gold["values"]["scenario_operating_income"],
                "concepts": ("operating income",),
                "native": _paired_metric(
                    native_rows,
                    "scenario_operating_income",
                    gold["values"]["scenario_operating_income"],
                ),
            },
            "growth funding": {
                "amount": gold["values"]["incremental_working_capital"],
                "concepts": ("funding", "working capital"),
                "native": _paired_metric(
                    native_rows,
                    "incremental_working_capital",
                    gold["values"]["incremental_working_capital"],
                ),
            },
        }

        def priority_owner_visible(owner: Any) -> bool:
            owner_parts = [
                _normalize(part)
                for part in re.split(r"\s+and\s+", str(owner or ""))
                if _normalize(part)
            ]
            return bool(owner_parts) and all(
                part in normalized_text for part in owner_parts
            )

        priority_checks: dict[str, bool] = {}
        for label, definition in priority_definitions.items():
            row = priority_by_label.get(label, {})
            priority_checks[label] = (
                bool(row)
                and _value_matches(
                    row.get("quantified_driver"),
                    definition["amount"],
                    "quantified_driver",
                )
                and bool(definition["native"])
                and all(
                    concept in normalized_text
                    for concept in definition["concepts"]
                )
                and priority_owner_visible(row.get("owner"))
                and _text_contains(text, row.get("gate"))
            )
        priorities_met = (
            len(priority_rows) == len(priority_definitions)
            and len(priority_by_label) == len(priority_definitions)
            and all(priority_checks.values())
        )

        capacity_risk = supporting.get("capacity_risk")
        volume_change = gold["values"]["volume_change"]
        capacity_risk_met = (
            isinstance(capacity_risk, str)
            and "capacity" in _normalize(capacity_risk)
            and "volume" in _normalize(capacity_risk)
            and _text_contains(
                capacity_risk,
                volume_change,
                "volume_change",
            )
            and "capacity" in normalized_text
            and "volume" in normalized_text
            and _paired_metric(
                native_rows,
                "volume_change",
                volume_change,
            )
        )
        funding_risk = supporting.get("funding_risk")
        funding_risk_met = (
            _value_matches(
                funding_risk,
                gold["values"]["incremental_working_capital"],
                "funding_risk",
            )
            and "funding" in normalized_text
            and "working capital" in normalized_text
            and _paired_metric(
                native_rows,
                "incremental_working_capital",
                funding_risk,
            )
        )
        criteria.extend(
            [
                _criterion(
                    task_id,
                    "fy27_monthly_milestone_truth",
                    (
                        "The memorandum includes all twelve exact monthly "
                        "revenue, GP, OI, working-capital, owner, and gate rows"
                    ),
                    milestones_met,
                    (
                        f"milestones_met={milestones_met}; "
                        f"table_evidence={milestones_evidence!r}"
                    ),
                    category="core_finance",
                    weight=10,
                    failure_cap=0.49,
                ),
                _criterion(
                    task_id,
                    "fy27_monthly_annual_tie_out",
                    (
                        "The twelve displayed monthly money rows tie exactly "
                        "to the approved annual plan at cents"
                    ),
                    milestones_met and annual_tie_outs_met,
                    (
                        f"milestones_met={milestones_met}; "
                        f"annual_tie_outs_met={annual_tie_outs_met}"
                    ),
                    category="auditability",
                    weight=10,
                    failure_cap=0.49,
                ),
                _criterion(
                    task_id,
                    "fy27_value_creation_priorities_truth",
                    (
                        "The memorandum supports all three exact quantified "
                        "FY27 value-creation priorities and retains their "
                        "owners and gates"
                    ),
                    priorities_met,
                    (
                        f"priority_checks={priority_checks!r}; "
                        f"labels={sorted(priority_by_label)!r}"
                    ),
                    category="decision",
                    weight=10,
                    failure_cap=0.49,
                ),
                _criterion(
                    task_id,
                    "fy27_capacity_risk_truth",
                    (
                        "The capacity risk explicitly agrees with the "
                        "approved FY27 volume increase"
                    ),
                    capacity_risk_met,
                    (
                        f"capacity_risk={capacity_risk!r}; "
                        f"volume_change={volume_change!r}"
                    ),
                    category="decision",
                    weight=10,
                    failure_cap=0.49,
                ),
                _criterion(
                    task_id,
                    "fy27_funding_risk_truth",
                    (
                        "The funding risk agrees exactly with incremental "
                        "working-capital need and is natively supported"
                    ),
                    funding_risk_met,
                    (
                        f"funding_risk={funding_risk!r}; "
                        "incremental_working_capital="
                        f"{gold['values']['incremental_working_capital']!r}"
                    ),
                    category="decision",
                    weight=10,
                    failure_cap=0.49,
                ),
            ]
        )
    return criteria


def grade_task(
    task_id: str,
    answer: Any,
    workspace_root: str | Path,
    database_path: str | Path | None = None,
) -> dict[str, Any]:
    gold = load_gold(task_id)
    mode = gold["output_mode"]
    workspace = Path(workspace_root)
    if mode == "console":
        criteria = _console_criteria(task_id, answer, gold)
    elif mode == "spreadsheet":
        criteria = _xlsx_criteria(task_id, answer, workspace, gold)
    elif mode == "document":
        criteria = _docx_criteria(task_id, answer, workspace, gold)
        criteria.extend(
            _task_specific_document_post_criteria(
                task_id,
                workspace,
                gold,
            )
        )
    elif mode == "presentation":
        criteria = _pptx_criteria(task_id, answer, workspace, gold)
    elif mode == "erp":
        if database_path is None:
            raise ValueError("ERP grading requires the mutable runtime database")
        criteria = _erp_criteria(task_id, answer, Path(database_path), gold)
    else:
        raise ValueError(f"Unknown output mode {mode}")
    rows = [criterion.as_dict() for criterion in criteria]
    return {
        "task_id": task_id,
        "criteria": rows,
        "criteria_met": sum(row["value"] for row in rows),
        "criteria_total": len(rows),
        "strict_pass": bool(rows) and all(row["value"] for row in rows),
        "reward": (
            sum(row["value"] for row in rows) / len(rows) if rows else 0.0
        ),
    }
