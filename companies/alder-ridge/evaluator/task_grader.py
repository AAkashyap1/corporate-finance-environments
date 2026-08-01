from __future__ import annotations

import json
import math
import os
import re
import shutil
import subprocess
import tempfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Callable

from openpyxl import load_workbook
from docx import Document
from pypdf import PdfReader
from pptx import Presentation
from evaluator.semantic import (
    contains_concept,
    ordered_semantic_list_matches,
    semantic_equal,
    semantic_value_matches,
)
from evaluator.hybrid_semantic import semantic_requirement
from evaluator.rubric import apply_reward_policy, attach_default_policy
from alder_ridge_company.paths import resolve_project_root, resolve_seed_root


ROOT = resolve_project_root(__file__)
GOLD_PATH = Path(__file__).resolve().parent / "reference" / "tasks_001_025.json"
SEED_WORKSPACE = resolve_seed_root(__file__) / "workspace"


@dataclass(frozen=True)
class Criterion:
    id: str
    description: str
    met: bool
    evidence: str
    category: str | None = None
    weight: int | None = None
    semantic: bool | None = None
    failure_cap: float | None = None


def load_reference(task_id: str | None = None) -> dict[str, Any]:
    payload = json.loads(GOLD_PATH.read_text(encoding="utf-8"))
    return payload[task_id] if task_id else payload


def _normalize(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(value or "").lower()).strip()


def _number(value: Any) -> float | None:
    if isinstance(value, bool):
        return float(value)
    if isinstance(value, (int, float)) and math.isfinite(float(value)):
        return float(value)
    if not isinstance(value, str):
        return None
    text = value.strip()
    percent = text.endswith("%")
    negative = text.startswith("(") and text.endswith(")")
    text = text.strip("() ")
    suffix_match = re.search(r"([kmb])\s*$", text, flags=re.I)
    multiplier = {"k": 1_000.0, "m": 1_000_000.0, "b": 1_000_000_000.0}.get(
        suffix_match.group(1).lower() if suffix_match else "",
        1.0,
    )
    if suffix_match:
        text = text[: suffix_match.start()]
    text = text.replace("$", "").replace(",", "").replace("×", "").replace("x", "")
    text = text.strip("% ")
    try:
        result = float(text)
    except ValueError:
        return None
    if negative:
        result = -result
    result *= multiplier
    return result / 100 if percent else result


def _answer_mapping(answer: Any) -> dict[str, Any]:
    if isinstance(answer, dict):
        return answer
    text = str(answer or "").strip()
    candidates = [text]
    candidates.extend(re.findall(r"```(?:json)?\s*(\{.*?\})\s*```", text, flags=re.I | re.S))
    brace = re.search(r"\{.*\}", text, flags=re.S)
    if brace:
        candidates.append(brace.group(0))
    for candidate in reversed(candidates):
        try:
            parsed = json.loads(candidate)
            if isinstance(parsed, dict):
                return parsed
        except (json.JSONDecodeError, TypeError):
            pass

    mapping: dict[str, Any] = {}
    for line in text.splitlines():
        match = re.match(r"\s*[-*]?\s*([A-Za-z][A-Za-z0-9 _/-]{1,60})\s*[:=]\s*(.+?)\s*$", line)
        if match:
            mapping[re.sub(r"\s+", "_", match.group(1).strip().lower())] = match.group(2).strip()
    return mapping


def _get(mapping: dict[str, Any], key: str) -> Any:
    if key in mapping:
        return mapping[key]
    wanted = _normalize(key)
    for candidate, value in mapping.items():
        if _normalize(candidate) == wanted:
            return value
    return None


def _close(actual: Any, expected: float, *, abs_tol: float = .02, rel_tol: float = 1e-6) -> bool:
    value = _number(actual)
    if value is None:
        return False
    return abs(value - expected) <= max(abs_tol, abs(expected) * rel_tol)


def _numeric_criteria(
    mapping: dict[str, Any],
    criterion_id: str,
    description: str,
    expected: dict[str, float],
    *,
    abs_tol: float = .02,
    rel_tol: float = 0.0,
) -> list[Criterion]:
    criteria: list[Criterion] = []
    for key, target in expected.items():
        actual = _get(mapping, key)
        met = _close(actual, target, abs_tol=abs_tol, rel_tol=rel_tol)
        criteria.append(
            Criterion(
                f"{criterion_id}__{_normalize(key).replace(' ', '_')}",
                f"{description}: `{key}` is correct",
                met,
                "reported value matched" if met else f"{key}={actual!r}; expected {target}",
                category="core_finance",
                weight=10,
                semantic=False,
            )
        )
    return criteria


def _field_numeric_criteria(
    mapping: dict[str, Any],
    criterion_id: str,
    description: str,
    expected: dict[str, tuple[float, float]],
) -> list[Criterion]:
    criteria: list[Criterion] = []
    for key, (target, abs_tol) in expected.items():
        actual = _get(mapping, key)
        met = _close(actual, target, abs_tol=abs_tol + 1e-9, rel_tol=0.0)
        criteria.append(
            Criterion(
                f"{criterion_id}__{_normalize(key).replace(' ', '_')}",
                f"{description}: `{key}` is correct",
                met,
                "reported value matched at required precision"
                if met else f"{key}={actual!r}; expected {target} ± {abs_tol}",
                category="core_finance",
                weight=10,
                semantic=False,
            )
        )
    return criteria


def _percentage_criterion(
    mapping: dict[str, Any],
    criterion_id: str,
    description: str,
    key: str,
    expected_ratio: float,
    *,
    abs_tol: float = .00005,
) -> Criterion:
    actual = _get(mapping, key)
    matched = _close(actual, expected_ratio, abs_tol=abs_tol, rel_tol=0.0) or _close(
        actual, expected_ratio * 100, abs_tol=abs_tol * 100, rel_tol=0.0
    )
    return Criterion(
        criterion_id,
        description,
        matched,
        "percentage matched as a decimal ratio or percentage points"
        if matched
        else f"{key}={actual!r}; expected {expected_ratio} or {expected_ratio * 100}",
        category="core_finance",
        weight=10,
        semantic=False,
    )


def _string_criteria(
    mapping: dict[str, Any], criterion_id: str, description: str, expected: dict[str, str]
) -> list[Criterion]:
    criteria = []
    for key, target in expected.items():
        actual = _get(mapping, key)
        met = semantic_value_matches(actual, target)
        criteria.append(
            Criterion(
                f"{criterion_id}__{_normalize(key).replace(' ', '_')}",
                f"{description}: `{key}` is correct",
                met,
                "professional conclusion matched" if met else f"{key}={actual!r}; expected {target!r}",
                category="decision",
                weight=10,
                semantic=True,
                failure_cap=0.49,
            )
        )
    return criteria


def _list_criteria(
    mapping: dict[str, Any], criterion_id: str, description: str, key: str, expected: list[str]
) -> list[Criterion]:
    actual = _get(mapping, key)
    def flatten(value: Any) -> list[str]:
        if isinstance(value, dict):
            return [item for key, child in value.items() for item in (str(key), *flatten(child))]
        if isinstance(value, (list, tuple, set)):
            return [item for child in value for item in flatten(child)]
        return [str(value)] if value is not None else []

    flattened = flatten(actual)
    expected_ids = [token.upper() for token in expected if re.fullmatch(r"CX-\d{2}-\d{3}", token, flags=re.I)]
    if expected_ids:
        actual_values = []
        for value in flattened:
            actual_values.extend(token.upper() for token in re.findall(r"CX-\d{2}-\d{3}", value, flags=re.I))
        actual_values = list(dict.fromkeys(actual_values))
        met = sorted(actual_values) == sorted(expected_ids)
    else:
        actual_values = [part.strip() for value in flattened for part in re.split(r"[,;]", value) if part.strip()]
        met = ordered_semantic_list_matches(actual, expected)
    criteria = [
        Criterion(
            f"{criterion_id}__item_{index:02d}_{_normalize(item).replace(' ', '_')}",
            f"{description}: required item {item!r} is correctly included",
            any(semantic_value_matches(actual_item, item) for actual_item in actual_values),
            f"actual={actual_values!r}; required={item!r}",
            category="decision",
            weight=10,
            semantic=True,
            failure_cap=0.49,
        )
        for index, item in enumerate(expected, start=1)
    ]
    criteria.append(
        Criterion(
            f"{criterion_id}__complete_set",
            f"{description}: the set has no missing, extra, or duplicated items",
            met,
            f"actual={actual_values!r}; expected={expected!r}",
            category="decision",
            weight=10,
            semantic=True,
            failure_cap=0.49,
        )
    )
    return criteria


def _result(criteria: list[Criterion]) -> dict[str, Any]:
    met_count = sum(criterion.met for criterion in criteria)
    reward = met_count / len(criteria) if criteria else 0.0
    return {
        "reward": round(reward, 6),
        "strict_pass": bool(criteria) and met_count == len(criteria),
        "criteria_met": met_count,
        "criteria_total": len(criteria),
        "criteria": [
            {
                "id": criterion.id,
                "description": criterion.description,
                "value": int(criterion.met),
                "evidence": criterion.evidence,
                **(
                    {
                        "category": criterion.category,
                        "weight": criterion.weight,
                        "semantic": criterion.semantic,
                        **(
                            {"failure_cap": criterion.failure_cap}
                            if criterion.failure_cap is not None else {}
                        ),
                    }
                    if criterion.category is not None else {}
                ),
            }
            for criterion in criteria
        ],
    }


# Keep the rubric denominator stable even when an agent deletes, corrupts, or
# never creates the required deliverable.  A missing file is zero credit on
# every must-have criterion, not a different one-criterion version of the task.
FILE_CRITERION_SPECS: dict[str, tuple[tuple[str, str], ...]] = {
    "task_004": (
        ("sourced_inputs", "Required source inputs and protected evidence are correct"),
        ("formulas", "Project calculations are driven by workbook references"),
        ("project_results", "All four project results are correct"),
        ("totals", "The total/control row is formula-driven"),
    ),
    "task_008": (
        ("preservation", "Protected base-case and scenario content is preserved"),
        ("assumption_amounts", "Scenario amounts are correct"),
        ("formula_rollforward", "The downside schedule is formula-driven"),
        ("cash_results", "All weekly cash and financing results are correct"),
        ("chart", "The requested comparison chart is present"),
    ),
    "task_012": (
        ("preservation", "Protected instructions and source labels are preserved"),
        ("source_pull", "Source values are correct"),
        ("bridge", "The variance bridge is correct and formula-driven"),
        ("chart", "The requested comparison chart is present"),
    ),
    "task_015": (
        ("preservation", "The original deck is preserved and one slide is appended"),
        ("structure", "The appended slide has the requested structure"),
        ("metrics", "The covenant table contains the correct metrics"),
        ("status", "The proposed/unposted and compliance status is correct"),
    ),
    "task_023": (
        ("structure", "The source sheets are preserved and Borrowing Base is added"),
        ("covenant_formulas", "Covenant outputs are driven by workbook references"),
        ("covenant_results", "Covenant outputs are correct"),
        ("borrowing_base", "The borrowing-base summary is correct"),
        ("support_and_lineage", "Calculations and source lineage are auditable"),
    ),
    "task_024": (
        ("structure", "The requested workbook and table structure is present"),
        ("request_population", "All request rows and classifications are correct"),
        ("request_formulas", "Request calculations are formula-driven"),
        ("summary", "All summary outputs are correct"),
        ("summary_lineage", "Summary outputs are formula-driven"),
        ("sources_controls", "Controlling sources and controls are documented"),
    ),
    "task_025": (
        ("identity", "The title and internal-working status are correct"),
        ("table", "The decision table contains all required workstreams"),
        ("wip", "The WIP decision is correct"),
        ("liquidity", "The downside-liquidity decision is correct"),
        ("borrowing_base", "The borrowing-base decision is correct"),
        ("covenants", "The covenant decision is correct"),
        ("capex", "The capex decision and classifications are correct"),
        ("sources", "The controlling sources are identified"),
        ("length", "The note is no more than two pages"),
    ),
}


def _legacy_file_atomic_specs(task_id: str) -> list[tuple[str, str]]:
    """Stable atomic denominator for missing or unreadable legacy artifacts."""

    gold = load_reference(task_id)
    specs: list[tuple[str, str]] = []
    if task_id == "task_004":
        specs.extend((f"preservation__{_normalize(sheet).replace(' ', '_')}", f"Preserve {sheet}") for sheet in ("READ ME first", "Risk Review", "Evidence Map"))
        for expected in gold["source_inputs"]:
            specs.extend((f"source__{expected['project_id'].casefold()}__{label}", f"{expected['project_id']} source {label}") for label in ("current_contract", "cost_to_date", "pm_etc", "documented_etc_overlay", "billings"))
        for expected in gold["projects"]:
            specs.extend((f"formula__{expected['project_id'].casefold()}__{label}", f"{expected['project_id']} formula {label}") for label in ("close_etc", "eac", "percent_complete", "earned", "asset", "liability", "margin", "margin_percent"))
        for expected in gold["projects"]:
            specs.extend((f"result__{expected['project_id'].casefold()}__{label}", f"{expected['project_id']} result {label}") for label in ("close_etc", "eac", "earned", "asset", "liability", "margin", "percent_complete", "margin_percent"))
        specs.extend((f"total_formula__{coordinate.casefold()}", f"Total formula {coordinate}") for coordinate in [f"{column}11" for column in "CDEFGHJKLMN"])
    elif task_id == "task_008":
        specs.extend((f"preservation__{_normalize(sheet).replace(' ', '_')}", f"Preserve {sheet}") for sheet in ("Base Case Import", "Scenario Assumptions"))
        specs.extend((f"assumption__{label}", f"Scenario assumption {label}") for label in ("eastbank_collection", "cascade_collection", "redmond_collection", "accelerated_po", "insurance_deposit"))
        for expected in gold["weeks"]:
            specs.extend((f"result__{expected['week_ending']}__{key}", f"{expected['week_ending']} {key}") for key in ("cash_before_financing", "revolver_draw", "ending_cash", "ending_revolver"))
        specs.extend((f"formula__{column.casefold()}{row}", f"Formula {column}{row}") for row in range(6, 19) for column in "BCDEFGHIJK")
        specs.append(("chart", "Required liquidity chart"))
    elif task_id == "task_012":
        specs.extend((f"preservation__{_normalize(sheet).replace(' ', '_')}", f"Preserve {sheet}") for sheet in ("read me - DC marks", "Source Pull"))
        specs.extend((f"source_pull__{label}", f"Source Pull {label}") for label in ("posted_revenue", "proposed_wip_adjustment", "pro_forma_revenue", "plan_revenue", "gross_profit", "plan_gross_profit", "operating_income", "plan_operating_income"))
        specs.extend((f"bridge_value__{label}", f"Bridge value {label}") for label in ("pro_forma_revenue", "plan_revenue", "revenue_variance", "gross_profit", "plan_gross_profit", "gross_profit_variance", "operating_income", "plan_operating_income", "operating_income_variance", "revenue_variance_percent", "gross_margin_rate_variance", "operating_income_variance_percent"))
        specs.extend((f"bridge_formula__{column.casefold()}{row}", f"Bridge formula {column}{row}") for row in range(6, 9) for column in "BCDEFG")
        specs.append(("chart", "Required comparison chart"))
    elif task_id == "task_015":
        specs.append(("preservation__slide_count", "Exactly one slide appended"))
        specs.extend((f"preservation__slide_{index:02d}", f"Preserve slide {index}") for index in range(1, 6))
        specs.extend((("structure__title", "Required slide title"), ("structure__table", "Required PowerPoint table")))
        for metric in ("leverage", "fccr", "tangible_net_worth"):
            specs.extend((f"metric__{metric}__{field}", f"{metric} {field}") for field in ("posted", "pro_forma", "threshold", "status"))
        specs.extend((("status__proposed", "Proposed WIP status"), ("status__unposted", "Unposted WIP status"), ("status__compliant", "Compliance status")))
    elif task_id == "task_023":
        specs.extend((f"structure__sheet_{_normalize(name).replace(' ', '_')}", f"Required sheet {name}") for name in ("Executed Terms", "Q2 Headroom Working", "Borrowing Base"))
        specs.extend((("structure__no_extra_sheets", "No extra sheets"), ("preservation__executed_terms", "Preserve executed terms"), ("preservation__q2_headroom", "Preserve headroom inputs")))
        specs.extend((f"covenant_formula__{column.casefold()}{row}", f"Covenant formula {column}{row}") for row in range(6, 9) for column in "CDEFG")
        specs.extend((f"covenant_value__{column.casefold()}{row}", f"Covenant value {column}{row}") for row in range(6, 9) for column in "CDEFG")
        labels = ("Invoice Count", "Gross Open AR", "Pre-Concentration Eligible AR", "Customer Cap", "Post-Concentration Eligible AR", "Advance Rate", "Borrowing Base")
        for label in labels:
            slug = _normalize(label).replace(" ", "_")
            specs.extend(((f"borrowing_base_label__{slug}", f"Borrowing-base label {label}"), (f"borrowing_base_value__{slug}", f"Borrowing-base value {label}")))
        specs.extend((f"borrowing_base_formula__b{row}", f"Borrowing-base formula B{row}") for row in range(5, 9))
        specs.append(("support__heading", "Covenant-support heading"))
        specs.extend((f"support__{_normalize(label).replace(' ', '_')}", f"Covenant support {label}") for label in ("Funded Debt", "Posted EBITDA", "Pro Forma EBITDA", "Posted FCCR", "Pro Forma FCCR", "Posted Tangible Net Worth", "Pro Forma Tangible Net Worth"))
    elif task_id == "task_024":
        specs.extend((("structure__funding_screen", "Funding Screen sheet"), ("structure__sources_controls", "Sources & Controls sheet"), ("structure__no_extra_sheets", "No extra sheets"), ("structure__row_count", "Five request rows")))
        headers = ("Request ID", "Asset", "Documented Decision", "Request Amount", "Approved Amount", "Funding", "Annual Savings", "Simple Payback", "Equipment-Line Proceeds", "Cash Required", "Finance Classification")
        specs.extend((f"structure__header_{_normalize(header).replace(' ', '_')}", f"Header {header}") for header in headers)
        for request in gold["requests"]:
            rid = request["request_id"].casefold()
            specs.append((f"request__{rid}__present", f"Request {rid} present"))
            specs.extend((f"request__{rid}__{field}", f"Request {rid} {field}") for field in ("request_amount", "approved_amount", "annual_savings", "simple_payback", "equipment_line_proceeds", "cash_required", "asset", "decision", "funding", "classification"))
            specs.extend((f"request_formula__{rid}__{field}", f"Request {rid} formula {field}") for field in ("simple_payback", "equipment_line_proceeds", "cash_required"))
        summary_labels = ("Approved Spend", "Approved Annual Savings", "Equipment-Line Proceeds", "Cash Required", "Facility Remaining", "AOP Remaining After Approved Requests", "Downside Maximum Revolver", "Revolver Capacity After Downside", "Pro Forma Leverage")
        for label in summary_labels:
            slug = _normalize(label).replace(" ", "_")
            specs.extend(((f"summary_label__{slug}", f"Summary label {label}"), (f"summary_value__{slug}", f"Summary value {label}"), (f"summary_formula__{slug}", f"Summary formula {label}")))
        for token in ("capex asks", "fy26 op plan", "fleet", "equipment line proposal", "downside assumptions", "usbank amdt2"):
            specs.append((f"source__{_normalize(token).replace(' ', '_')}", f"Source {token}"))
        specs.extend((("source__accounting_lineage", "Accounting lineage"), ("controls__formula_count", "Formula-driven controls")))
    elif task_id == "task_025":
        specs.extend((("identity__title", "Required title"), ("identity__status", "Internal-working status"), ("table__real_word_table", "Real Word table")))
        specs.extend((f"table__row_{_normalize(row).replace(' ', '_')}", f"Decision row {row}") for row in ("proposed june wip", "downside liquidity", "borrowing base", "covenant", "midyear capex"))
        specs.extend((f"wip__{field}", f"WIP {field}") for field in ("amount", "proposed", "unposted"))
        specs.extend((f"liquidity__{field}", f"Liquidity {field}") for field in ("minimum_cash", "first_draw_week", "maximum_revolver"))
        specs.extend((("borrowing_base__amount", "Borrowing-base amount"), ("covenants__posted_leverage", "Posted leverage"), ("covenants__pro_forma_leverage", "Pro-forma leverage"), ("covenants__compliance", "Covenant compliance"), ("capex__cash_required", "Capex cash required"), ("capex__classification_labels", "Capex classification labels")))
        capex = gold["capex"]
        for category, ids in (("release_now", capex["release_now"]), ("conditional", capex["conditional"]), ("hold_or_defer", capex["hold_or_defer"])):
            specs.extend((f"capex__{request_id.casefold()}__{category}", f"{request_id} {category}") for request_id in ids)
        specs.extend((f"source__{group}", f"Source {group}") for group in ("accounting", "wip", "liquidity", "borrowing_base", "covenants", "capex"))
        specs.append(("length", "Two-page limit"))
    else:
        raise KeyError(task_id)
    ids = [criterion_id for criterion_id, _ in specs]
    if len(ids) != len(set(ids)):
        raise ValueError(f"duplicate legacy atomic criterion ids for {task_id}")
    return specs


def _file_failure(task_id: str, evidence: str) -> dict[str, Any]:
    return _result([
        Criterion(criterion_id, description, False, evidence)
        for criterion_id, description in _legacy_file_atomic_specs(task_id)
    ])


def _legacy_file_semantic_ids(task_id: str) -> set[str]:
    ids = {criterion_id for criterion_id, _ in _legacy_file_atomic_specs(task_id)}
    if task_id == "task_015":
        return {
            "structure__title", "status__compliant", "status__proposed", "status__unposted",
            "metric__leverage__status", "metric__fccr__status",
            "metric__tangible_net_worth__status",
        }
    if task_id == "task_023":
        return {criterion_id for criterion_id in ids if criterion_id.startswith("borrowing_base_label__")}
    if task_id == "task_024":
        return {
            criterion_id
            for criterion_id in ids
            if (
                re.match(r"request__.+__(asset|decision|funding|classification)$", criterion_id)
                or criterion_id.startswith("summary_label__")
                or criterion_id.startswith("source__")
            )
        }
    if task_id == "task_025":
        fixed = {
            "identity__title", "identity__status", "wip__proposed", "wip__unposted",
            "liquidity__first_draw_week", "covenants__compliance",
            "capex__classification_labels",
        }
        return fixed | {
            criterion_id
            for criterion_id in ids
            if criterion_id.startswith("source__")
            or re.match(r"capex__cx-\d{2}-\d{3}__", criterion_id)
        }
    return set()


def _canonicalize_legacy_file_policy(task_id: str, result: dict[str, Any]) -> dict[str, Any]:
    """Make missing, corrupt, and complete artifact rubrics policy-identical.

    Artifact evidence descriptions necessarily differ between a missing file
    and a parsed file, so importance and semantic routing must never be inferred
    from whichever description happens to be present on that execution path.
    """

    canonical = attach_default_policy(_file_failure(task_id, "canonical policy"))
    semantic_ids = _legacy_file_semantic_ids(task_id)
    policy_by_id: dict[str, dict[str, Any]] = {}
    for row in canonical["criteria"]:
        criterion_id = str(row["id"])
        row["semantic"] = criterion_id in semantic_ids
        # Source naming is genuine semantic provenance, not formula lineage or
        # a central decision, even when its id includes "accounting_lineage".
        if criterion_id in semantic_ids and criterion_id.startswith("source__"):
            row.update({"category": "provenance", "weight": 3})
            row.pop("failure_cap", None)
        elif task_id == "task_025" and criterion_id == "identity__status":
            row.update({"category": "controls", "weight": 3})
            row.pop("failure_cap", None)
        elif task_id == "task_025" and (
            criterion_id in {
                "wip__proposed", "wip__unposted", "covenants__compliance",
                "capex__classification_labels",
            }
            or re.match(r"capex__cx-\d{2}-\d{3}__", criterion_id)
        ):
            row.update({"category": "decision", "weight": 10, "failure_cap": 0.49})
        elif row["semantic"] and row.get("category") == "decision":
            row.update({"weight": 10, "failure_cap": 0.49})
        policy_by_id[criterion_id] = {
            key: row[key]
            for key in ("category", "weight", "semantic", "failure_cap")
            if key in row
        }

    actual_ids = {str(row["id"]) for row in result.get("criteria", [])}
    if actual_ids != set(policy_by_id):
        raise ValueError(
            f"legacy artifact rubric ids drifted for {task_id}: "
            f"missing={sorted(set(policy_by_id) - actual_ids)}, "
            f"extra={sorted(actual_ids - set(policy_by_id))}"
        )
    review = result.get("semantic_review")
    if isinstance(review, dict):
        reviewed = {
            str(row.get("criterion_id"))
            for row in review.get("criteria", [])
            if isinstance(row, dict)
        }
        if reviewed != semantic_ids:
            raise ValueError(
                f"legacy semantic route drifted for {task_id}: "
                f"missing={sorted(semantic_ids - reviewed)}, extra={sorted(reviewed - semantic_ids)}"
            )
    for row in result["criteria"]:
        row.pop("failure_cap", None)
        row.update(policy_by_id[str(row["id"])])
    return result


def _grade_numeric(task_id: str, answer: Any) -> dict[str, Any]:
    gold = load_reference(task_id)
    mapping = _answer_mapping(answer)
    criteria: list[Criterion]
    if task_id == "task_001":
        project = gold["project"]
        criteria = [
            *_numeric_criteria(mapping, "wip_chain", "ARM-2409 WIP calculation chain is correct", {
                "current_contract": project["current_contract"],
                "posted_cost": project["cost_to_date"],
                "billings": project["billings"],
                "pm_etc": gold["pm_etc"],
                "required_etc_adjustment": gold["required_etc_adjustment"],
                "close_etc": project["estimated_cost_to_complete"],
                "eac": project["estimated_cost_at_completion"],
                "earned_revenue": project["earned_revenue"],
                "contract_asset": project["underbilling"],
            }, abs_tol=.02),
            _percentage_criterion(mapping, "wip_chain__percent_complete", "ARM-2409 percent complete is correct", "percent_complete", project["percent_complete"]),
        ]
    elif task_id == "task_002":
        project = gold["project"]
        criteria = [
            *_numeric_criteria(mapping, "wip_chain", "ARM-2417 WIP and margin calculation chain is correct", {
                "current_contract": project["current_contract"],
                "posted_cost": project["cost_to_date"],
                "billings": project["billings"],
                "pm_etc": gold["pm_etc"],
                "required_etc_adjustment": gold["required_etc_adjustment"],
                "close_etc": project["estimated_cost_to_complete"],
                "eac": project["estimated_cost_at_completion"],
                "earned_revenue": project["earned_revenue"],
                "estimated_margin": project["estimated_total_margin"],
            }, abs_tol=.02),
            _percentage_criterion(mapping, "wip_chain__percent_complete", "ARM-2417 percent complete is correct", "percent_complete", project["percent_complete"]),
            _percentage_criterion(
                mapping,
                "margin_percent",
                "ARM-2417 estimated margin percentage is correct after the pending-change-order policy overlay",
                "margin_percent",
                project["estimated_margin_percent"],
            ),
        ]
    elif task_id == "task_003":
        criteria = [
            *_numeric_criteria(mapping, "gross_wip", "Gross contract asset and liability are correct", {
                "contract_asset": gold["contract_asset"],
                "contract_liability": gold["contract_liability"],
            }),
            *_numeric_criteria(mapping, "net_adjustment", "The net WIP revenue adjustment is correct", {
                "net_wip_revenue_adjustment": gold["net_wip_revenue_adjustment"],
            }),
            *_numeric_criteria(mapping, "journal_entry", "The proposed three-line journal is balanced and correctly classified", gold["journal"]),
        ]
    elif task_id == "task_005":
        criteria = [
            *_numeric_criteria(mapping, "liquidity", "Base 13-week liquidity result is correct", {
                "opening_cash": gold["opening_cash"],
                "minimum_cash": gold["minimum_cash"],
                "liquidity_floor": gold["liquidity_floor"],
                "maximum_revolver": gold["maximum_revolver"],
                "ending_cash_2026_09_25": gold["ending_cash_2026_09_25"],
            }),
            *_string_criteria(mapping, "liquidity", "Base 13-week minimum-cash timing is correct", {
                "minimum_cash_week": gold["minimum_cash_week"],
            }),
        ]
    elif task_id == "task_006":
        criteria = [
            *_numeric_criteria(mapping, "scenario_inputs", "Named collections, accelerated PO amount and insurance deposit are correct", {
                "named_collections_total": gold["named_collection_total"],
                "accelerated_po_amount": gold["accelerated_po_amount"],
                "insurance_deposit": gold["assumptions"]["insurance_deposit"],
            }),
            *_numeric_criteria(mapping, "minimum_cash", "Minimum cash before financing is correct", {"minimum_cash_before_financing": gold["minimum_cash_before_financing"]}),
            *_string_criteria(mapping, "timing", "The minimum and first-draw weeks are correct", {
                "minimum_cash_week": gold["minimum_cash_week"],
                "first_draw_week": gold["first_draw_week"],
            }),
            *_numeric_criteria(mapping, "revolver", "Maximum and ending revolver balances are correct", {
                "maximum_revolver": gold["maximum_revolver"],
                "ending_revolver": gold["ending_revolver"],
            }),
        ]
    elif task_id == "task_007":
        criteria = [
            *_numeric_criteria(mapping, "concentration", "Named-customer AR concentration components are correct", {
                "gross_open_ar": gold["gross_open_ar"],
                "eastbank_data_ar": gold["customer_amounts"]["Eastbank Data"],
                "cascade_health_ar": gold["customer_amounts"]["Cascade Health"],
                "redmond_unified_ar": gold["customer_amounts"]["Redmond Unified"],
                "named_total": gold["named_total"],
            }),
            _percentage_criterion(
                mapping,
                "concentration",
                "The three-customer concentration percentage is correct",
                "named_percent_of_ar",
                gold["named_percent_of_ar"],
            ),
            *_string_criteria(mapping, "concentration", "Largest named-customer balance is correctly identified", {
                "largest_customer": gold["largest_customer"],
            }),
        ]
    elif task_id == "task_009":
        criteria = [
            *_numeric_criteria(mapping, "revenue_bridge", "Posted revenue, WIP adjustment and pro-forma revenue are correct", {
                "posted_revenue": gold["posted_revenue"],
                "proposed_wip_adjustment": gold["proposed_wip_adjustment"],
                "pro_forma_revenue": gold["pro_forma_revenue"],
            }),
            *_field_numeric_criteria(mapping, "revenue_variance", "June plan revenue and variance are correct", {
                "plan_revenue": (gold["plan_revenue"], .02),
                "revenue_variance": (gold["revenue_variance"], .02),
                "revenue_variance_percent": (gold["revenue_variance_percent"], .00005),
            }),
            *_field_numeric_criteria(mapping, "gross_profit", "Gross profit and gross-margin variance are correct", {
                "gross_profit": (gold["gross_profit"], .02),
                "plan_gross_profit": (gold["plan_gross_profit"], .02),
                "gross_margin": (gold["gross_margin"], .00005),
                "plan_gross_margin": (gold["plan_gross_margin"], .00005),
                "gross_margin_variance_bps": (gold["gross_margin_variance_bps"], .02),
            }),
            *_numeric_criteria(mapping, "operating_income", "Operating income and plan variance are correct", {
                "operating_income": gold["operating_income"],
                "plan_operating_income": gold["plan_operating_income"],
                "operating_income_variance": gold["operating_income_variance"],
            }),
        ]
    elif task_id == "task_010":
        criteria = [
            *_list_criteria(mapping, "top_three_projects", "The three highest-rework projects are correctly ranked", "top_three_projects", [row["project_id"] for row in gold["top_three_projects"]]),
            *_numeric_criteria(mapping, "top_three", "Top-three rework hours, rate, cost and concentration are correct", {
                "first_project_hours": gold["top_three_projects"][0]["rework_hours"],
                "second_project_hours": gold["top_three_projects"][1]["rework_hours"],
                "third_project_hours": gold["top_three_projects"][2]["rework_hours"],
                "top_three_hours": gold["top_three_hours"],
                "loaded_rate": gold["rate"],
                "top_three_cost": gold["top_three_cost"],
                "top_three_concentration": gold["top_three_concentration"],
            }, abs_tol=.00005),
        ]
    elif task_id == "task_011":
        criteria = [
            *_numeric_criteria(mapping, "work_orders", "Completed work orders and callbacks are correct", {
                "completed_work_orders": gold["completed_work_orders"],
                "callbacks": gold["callbacks"],
            }),
            *_field_numeric_criteria(mapping, "productivity", "First-time-fix rate and billable hours are correct", {
                "first_time_fix_rate": (gold["first_time_fix_rate"], .00005),
                "billable_hours": (gold["billable_hours"], .02),
            }),
            *_numeric_criteria(mapping, "value", "Labor-plus-material proxy and average per work order are correct", {
                "revenue_proxy": gold["revenue_proxy"],
                "average_revenue_per_work_order": gold["average_revenue_per_work_order"],
            }),
        ]
    elif task_id == "task_013":
        criteria = [
            *_numeric_criteria(mapping, "advance", "The invoice-level borrowing-base calculation is correct", {
                "invoice_row_count": gold["row_count"],
                "gross_open_ar": gold["gross_open_ar"],
                "pre_concentration_eligible": gold["pre_concentration_eligible"],
                "customer_cap": gold["customer_cap"],
                "post_concentration_eligible": gold["post_concentration_eligible"],
                "advance_rate": gold["advance_rate"],
                "borrowing_base": gold["borrowing_base"],
            }, abs_tol=.00005),
        ]
    elif task_id == "task_014":
        criteria = [
            *_numeric_criteria(mapping, "debt_ebitda", "Funded debt and posted/pro-forma EBITDA are correct", {
                "funded_debt": gold["funded_debt"],
                "posted_ebitda": gold["ltm_adjusted_ebitda_posted"],
                "pro_forma_ebitda": gold["ltm_adjusted_ebitda_pro_forma"],
            }),
            *_numeric_criteria(mapping, "leverage", "Posted/pro-forma leverage is correct", {
                "posted_leverage": gold["leverage_posted"],
                "pro_forma_leverage": gold["leverage_pro_forma"],
                "maximum_leverage": gold["max_leverage"],
            }, abs_tol=.0001, rel_tol=1e-6),
            *_numeric_criteria(mapping, "fccr", "Posted/pro-forma FCCR is correct", {
                "posted_fccr": gold["fccr_posted"],
                "pro_forma_fccr": gold["fccr_pro_forma"],
                "minimum_fccr": gold["min_fccr"],
            }, abs_tol=.0001, rel_tol=1e-6),
            *_numeric_criteria(mapping, "net_worth", "Posted/pro-forma tangible net worth is correct", {
                "posted_tangible_net_worth": gold["tangible_net_worth_posted"],
                "pro_forma_tangible_net_worth": gold["tangible_net_worth_pro_forma"],
                "minimum_tangible_net_worth": gold["min_tangible_net_worth"],
            }),
            *_string_criteria(mapping, "compliance", "Both reported bases are correctly identified as compliant", {
                "posted_compliance": "compliant",
                "pro_forma_compliance": "compliant",
            }),
        ]
    elif task_id == "task_016":
        criteria = [
            *_numeric_criteria(mapping, "cash_control", "Accounting cash reconciles to the cash workpaper", {
                "system_cash": gold["system_cash"], "cash_workpaper": gold["cash_workpaper"], "cash_delta": gold["cash_delta"],
            }),
            *_numeric_criteria(mapping, "ar_control", "Accounting AR reconciles to the AR workpaper", {
                "system_ar": gold["system_ar"], "ar_workpaper": gold["ar_workpaper"], "ar_delta": gold["ar_delta"],
            }),
            *_numeric_criteria(mapping, "ap_control", "Accounting AP reconciles to the AP workpaper", {
                "system_ap": gold["system_ap"], "ap_workpaper": gold["ap_workpaper"], "ap_delta": gold["ap_delta"],
            }),
            *_numeric_criteria(mapping, "debt_control", "Accounting debt reconciles to the debt schedule", {
                "system_debt": gold["system_debt"], "debt_schedule": gold["debt_schedule"], "debt_delta": gold["debt_delta"],
            }),
            *_numeric_criteria(mapping, "close_view", "The proposed WIP adjustment and pro-forma operating income are correct", {
                "proposed_wip_adjustment": gold["proposed_wip_adjustment"],
                "pro_forma_operating_income": gold["pro_forma_operating_income"],
            }),
        ]
    elif task_id == "task_017":
        criteria = [
            *_numeric_criteria(mapping, "ltm_denominators", "LTM revenue and cost-of-revenue denominators are correct", {
                "ltm_revenue": gold["ltm_revenue"], "ltm_cost_of_revenue": gold["ltm_cost_of_revenue"],
            }),
            *_numeric_criteria(mapping, "ar_population", "Gross AR, retainage and over-90 populations are correct", {
                "gross_open_ar": gold["gross_open_ar"], "retainage": gold["retainage"], "ar_over_90": gold["ar_over_90"],
            }),
            *_numeric_criteria(mapping, "ap_population", "Gross AP, held, overdue and net working-capital populations are correct", {
                "gross_open_ap": gold["gross_open_ap"], "held_ap": gold["held_ap"], "overdue_ap": gold["overdue_ap"],
                "net_ar_less_ap": gold["net_ar_less_ap"],
            }),
            *_numeric_criteria(mapping, "days", "DSO and DPO are correct", {"dso": gold["dso"], "dpo": gold["dpo"]}, abs_tol=.0001),
            *_numeric_criteria(mapping, "cash_release", "The DSO and DPO cash-release sensitivities and total are correct", {
                "five_day_dso_cash_release": gold["five_day_dso_cash_release"],
                "three_day_dpo_cash_release": gold["three_day_dpo_cash_release"],
                "total_cash_release": gold["total_cash_release"],
            }),
        ]
    elif task_id == "task_018":
        criteria = [
            *_numeric_criteria(mapping, "coverage", "Backlog and weighted-pipeline coverage calculation is correct", {
                "signed_backlog": gold["signed_backlog"],
                "gross_pipeline": gold["gross_pipeline"],
                "weighted_pipeline": gold["weighted_pipeline"],
                "h2_plan_revenue": gold["h2_plan_revenue"],
                "signed_backlog_coverage": gold["signed_backlog_coverage"],
                "total_coverage": gold["total_coverage"],
                "capacity_review_count": gold["capacity_review_count"],
                "capacity_review_gross": gold["capacity_review_gross"],
                "capacity_review_weighted": gold["capacity_review_weighted"],
                "top_three_weighted": gold["top_three_weighted"],
                "top_three_concentration": gold["top_three_concentration"],
            }, abs_tol=.00005),
            *_list_criteria(mapping, "top_three_opportunities", "Top-three weighted opportunities are correctly ranked", "top_three_opportunities", gold["top_three_opportunities"]),
        ]
    elif task_id == "task_019":
        criteria = [
            *_field_numeric_criteria(mapping, "approved_portfolio", "Approved spend, savings and portfolio payback are correct", {
                "approved_spend": (gold["approved_spend"], .02),
                "approved_annual_savings": (gold["approved_annual_savings"], .02),
                "portfolio_simple_payback_years": (gold["portfolio_simple_payback_years"], .0001),
            }),
            *_numeric_criteria(mapping, "financing", "Equipment-line basis, proceeds, cash need and remaining facility are correct", {
                "equipment_line_eligible_basis": gold["equipment_line_eligible_basis"],
                "equipment_line_proceeds": gold["equipment_line_proceeds"], "cash_required": gold["cash_required"],
                "facility_remaining": gold["facility_remaining"],
            }),
            *_numeric_criteria(mapping, "aop", "AOP envelope, posted additions and remaining headroom are correct", {
                "aop_capex_envelope": gold["aop_capex_envelope"],
                "fy26_posted_additions": gold["fy26_posted_additions"],
                "remaining_aop_before_requests": gold["remaining_aop_before_requests"],
                "remaining_aop_after_approved_requests": gold["remaining_aop_after_approved_requests"],
            }),
            *_numeric_criteria(mapping, "liquidity", "Downside revolver use and remaining commitment are correct", {
                "downside_maximum_revolver": gold["downside_maximum_revolver"],
                "revolver_capacity_after_downside": gold["revolver_capacity_after_downside"],
            }),
            *_numeric_criteria(mapping, "leverage", "Pro-forma funded debt and leverage are correct", {
                "pro_forma_funded_debt": gold["pro_forma_funded_debt"],
                "pro_forma_leverage_after_capex_financing": gold["pro_forma_leverage_after_capex_financing"],
            }, abs_tol=.0001),
            *_list_criteria(mapping, "release_now", "Release-now requests are correct", "release_now", gold["release_now"]),
            *_list_criteria(mapping, "conditional", "Conditional requests are correct", "conditional", gold["conditional"]),
            *_list_criteria(mapping, "hold_or_defer", "Hold/defer requests are correct", "hold_or_defer", gold["hold_or_defer"]),
        ]
    elif task_id == "task_020":
        criteria = [
            *_numeric_criteria(mapping, "headcount", "Active headcount, plan target and gap are correct", {
                "active_headcount": gold["active_headcount"], "plan_year_end_headcount": gold["plan_year_end_headcount"],
                "headcount_gap": gold["headcount_gap"],
            }),
            *_numeric_criteria(mapping, "payroll", "Gross pay, employer taxes, benefits and burdened payroll are correct", {
                "gross_pay": gold["gross_pay"], "employer_taxes": gold["employer_taxes"], "benefits": gold["benefits"],
                "total_burdened_payroll": gold["total_burdened_payroll"],
            }),
            *_numeric_criteria(mapping, "hourly_population", "Hourly population, hours and overtime mix are correct", {
                "hourly_employee_count": gold["hourly_employee_count"], "regular_hours": gold["regular_hours"],
                "overtime_hours": gold["overtime_hours"], "overtime_percent": gold["overtime_percent"],
            }, abs_tol=.00005),
            *_numeric_criteria(mapping, "actual_rate", "Hourly payroll cost and actual cost per hour are correct", {
                "hourly_payroll_cost": gold["hourly_payroll_cost"],
                "actual_hourly_payroll_cost_per_hour": gold["actual_hourly_payroll_cost_per_hour"],
            }, abs_tol=.01),
            *_numeric_criteria(mapping, "rate_bridge", "Approved estimating rate and rate difference are correct", {
                "approved_loaded_rate": gold["approved_loaded_rate"], "rate_difference": gold["rate_difference"],
            }, abs_tol=.01),
        ]
    elif task_id == "task_021":
        criteria = [
            *_numeric_criteria(mapping, "fixed_assets", "Gross fixed assets, accumulated depreciation and NBV are correct", {
                "gross_fixed_assets": gold["gross_fixed_assets"], "accumulated_depreciation": gold["accumulated_depreciation"],
                "net_book_value": gold["net_book_value"],
            }),
            *_numeric_criteria(mapping, "debt", "Funded debt classification is correct", {
                "funded_debt": gold["funded_debt"], "current_debt": gold["current_debt"], "long_term_debt": gold["long_term_debt"],
            }),
            *_numeric_criteria(mapping, "net_debt", "Cash and historical net debt are correct", {"cash": gold["cash"], "net_debt": gold["net_debt"]}),
            *_numeric_criteria(mapping, "funding", "Approved capex, line proceeds and cash requirement are correct", {
                "approved_capex": gold["approved_capex"], "equipment_line_proceeds": gold["equipment_line_proceeds"],
                "cash_required": gold["cash_required"],
            }),
            *_numeric_criteria(mapping, "pro_forma", "Pro-forma assets, cash, debt and net debt are correct", {
                "pro_forma_gross_fixed_assets": gold["pro_forma_gross_fixed_assets"], "pro_forma_cash": gold["pro_forma_cash"],
                "pro_forma_funded_debt": gold["pro_forma_funded_debt"], "pro_forma_net_debt": gold["pro_forma_net_debt"],
            }),
        ]
    elif task_id == "task_022":
        criteria = [
            *_field_numeric_criteria(mapping, "current_population", "Current work-order population, technicians, callbacks and first-time-fix rate are correct", {
                "completed_work_orders": (gold["completed_work_orders"], .005),
                "unique_technicians": (gold["unique_technicians"], .005),
                "callbacks": (gold["callbacks"], .005),
                "first_time_fix_rate": (gold["first_time_fix_rate"], .00005),
            }),
            *_field_numeric_criteria(mapping, "current_output", "Current billable hours and operating-value proxy are correct", {
                "billable_hours": (gold["billable_hours"], .005),
                "revenue_proxy": (gold["revenue_proxy"], .02),
            }),
            *_field_numeric_criteria(mapping, "per_technician", "Per-technician productivity is correct", {
                "work_orders_per_technician": (gold["work_orders_per_technician"], .005),
                "hours_per_technician": (gold["hours_per_technician"], .005),
                "proxy_per_technician": (gold["proxy_per_technician"], .02),
            }),
            *_field_numeric_criteria(mapping, "incremental", "The three-technician 90-percent ramp outputs are correct", {
                "incremental_work_orders": (gold["incremental_work_orders"], .005),
                "incremental_billable_hours": (gold["incremental_billable_hours"], .005),
                "incremental_revenue_proxy": (gold["incremental_revenue_proxy"], .02),
            }),
            *_field_numeric_criteria(mapping, "pro_forma", "Pro-forma quarterly output, van capex and ramp factor are correct", {
                "pro_forma_work_orders": (gold["pro_forma_work_orders"], .005),
                "pro_forma_billable_hours": (gold["pro_forma_billable_hours"], .005),
                "pro_forma_revenue_proxy": (gold["pro_forma_revenue_proxy"], .02),
                "service_van_capex": (gold["service_van_capex"], .02),
                "ramp_factor": (gold["ramp_factor"], .0001),
            }),
        ]
    else:
        raise KeyError(task_id)
    result = _result(criteria)
    semantic_specs: list[tuple[str, dict[str, Any]]] = []
    if task_id == "task_005":
        semantic_specs = [
            ("liquidity__minimum_cash_week", {"minimum_cash_week": gold["minimum_cash_week"]})
        ]
    elif task_id == "task_006":
        semantic_specs = [
            (f"timing__{key}", {key: gold[key]})
            for key in ("minimum_cash_week", "first_draw_week")
        ]
    elif task_id == "task_007":
        semantic_specs = [
            ("concentration__largest_customer", {"largest_customer": gold["largest_customer"]})
        ]
    elif task_id in {"task_010", "task_018"}:
        key = "top_three_projects" if task_id == "task_010" else "top_three_opportunities"
        values = (
            [row["project_id"] for row in gold["top_three_projects"]]
            if task_id == "task_010" else gold["top_three_opportunities"]
        )
        for index, item in enumerate(values, start=1):
            item_id = _normalize(item).replace(" ", "_")
            semantic_specs.append((f"{key}__item_{index:02d}_{item_id}", {key: item}))
        semantic_specs.append((f"{key}__complete_set", {key: values}))
    elif task_id == "task_014":
        semantic_specs = [
            (f"compliance__{key}", {key: "compliant"})
            for key in ("posted_compliance", "pro_forma_compliance")
        ]
    elif task_id == "task_019":
        for key in ("release_now", "conditional", "hold_or_defer"):
            for index, item in enumerate(gold[key], start=1):
                item_id = _normalize(item).replace(" ", "_")
                semantic_specs.append(
                    (f"{key}__item_{index:02d}_{item_id}", {key: item})
                )
            semantic_specs.append((f"{key}__complete_set", {key: gold[key]}))
    if semantic_specs:
        result = _attach_semantic_review(
            result,
            task_id=task_id,
            evidence=str(answer or "")[:60_000],
            artifact_type="submitted answer",
            specs=[
                {
                    "criterion_id": criterion_id,
                    "expected_facts": expected,
                    "hard_gate_met": bool(mapping),
                    "hard_gate_evidence": "answer parsed into a non-empty field mapping" if mapping else "answer did not parse",
                }
                for criterion_id, expected in semantic_specs
            ],
        )
    return result


def _attach_semantic_review(
    result: dict[str, Any],
    *,
    task_id: str,
    evidence: str,
    artifact_type: str,
    specs: list[dict[str, Any]],
) -> dict[str, Any]:
    # Assign the objective rubric policy before routing any criterion. Semantic
    # verification is a grading method, not an importance category: a title
    # stays weight 1, provenance stays weight 3, and only a true central
    # decision receives the weight-10 failure cap.
    result = attach_default_policy(result)
    by_id = {str(row["id"]): row for row in result.get("criteria", [])}
    reviews = []
    for spec in specs:
        criterion_id = str(spec["criterion_id"])
        criterion = by_id[criterion_id]
        # A criterion routed to the semantic verifier is semantic by
        # construction. Mark it explicitly so the released catalog, weights,
        # failure caps, and runtime result cannot drift from the review plan.
        criterion["semantic"] = True
        if criterion.get("category") == "decision":
            criterion["weight"] = 10
            criterion["failure_cap"] = 0.49
        reviews.append({
            "criterion_id": criterion_id,
            "requirement": semantic_requirement(
                criterion_id=criterion_id,
                description=str(criterion["description"]),
                expected_facts=spec.get("expected_facts"),
                artifact_type=artifact_type,
            ),
            "hard_gate_met": bool(spec["hard_gate_met"]),
            "hard_gate_evidence": str(spec["hard_gate_evidence"]),
            "legacy_lexical_match": bool(criterion["value"]),
        })
    result["semantic_review"] = {
        "version": 2,
        "mode": "deterministic_hard_gates_plus_bounded_semantic_judge",
        "task_id": task_id,
        "artifact": None,
        "evidence": evidence[:60_000],
        "criteria": reviews,
        "policy": "Semantic wording cannot rescue a failed objective numeric, formula, or structure gate.",
    }
    return result


def _legacy_artifact_evidence(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pptx":
        presentation = Presentation(path)
        chunks = []
        for index, slide in enumerate(presentation.slides, start=1):
            chunks.append(f"[Slide {index}]")
            for shape in slide.shapes:
                if getattr(shape, "text", ""):
                    chunks.append(shape.text)
                if getattr(shape, "has_table", False):
                    chunks.append("TABLE:")
                    chunks.extend(
                        " | ".join(cell.text for cell in row.cells)
                        for row in shape.table.rows
                    )
        return "\n".join(chunks)[:60_000]
    if suffix == ".docx":
        return _document_text(Document(path))[:60_000]
    if suffix == ".xlsx":
        workbook = load_workbook(path, data_only=False, read_only=False)
        chunks = []
        for sheet in workbook.worksheets:
            chunks.append(f"[Sheet: {sheet.title}]")
            for row in sheet.iter_rows():
                values = [f"{cell.coordinate}={cell.value}" for cell in row if cell.value is not None]
                if values:
                    chunks.append(" | ".join(values))
        return "\n".join(chunks)[:60_000]
    return ""


CELL_REFERENCE_PATTERN = re.compile(
    r"(?:'[^']+'!|(?:[A-Za-z0-9_ ]+!)?)\$?[A-Z]{1,3}\$?[0-9]+"
)
STRUCTURED_REFERENCE_PATTERN = re.compile(r"\b[A-Za-z_][A-Za-z0-9_.]*\[[^\]]+\]")


def _formula_count(ws, ranges: list[str], *, require_reference: bool = False) -> int:
    count = 0
    for cell_range in ranges:
        for row in ws[cell_range]:
            for cell in row:
                is_formula = isinstance(cell.value, str) and cell.value.startswith("=")
                has_reference = bool(
                    is_formula
                    and (CELL_REFERENCE_PATTERN.search(cell.value) or STRUCTURED_REFERENCE_PATTERN.search(cell.value))
                )
                if is_formula and (not require_reference or has_reference):
                    count += 1
    return count


def _formula_text(cell: Any) -> str:
    return str(getattr(cell, "value", "") or "").lower().replace("$", "").replace("'", "").replace(" ", "")


def _formula_has_refs(cell: Any, *references: str) -> bool:
    formula = _formula_text(cell)
    return formula.startswith("=") and all(
        reference.lower().replace("$", "").replace("'", "").replace(" ", "") in formula
        for reference in references
    )


def _formula_has_large_literal(cell: Any) -> bool:
    formula = _formula_text(cell)
    if not formula.startswith("="):
        return False
    without_references = re.sub(r"(?:[a-z0-9_ ]+!)?[a-z]{1,3}[0-9]+", "", formula)
    return any(abs(float(token)) >= 10_000 for token in re.findall(r"(?<![a-z])\d+(?:\.\d+)?", without_references))


def _recalculated_data_workbook(path: Path):
    """Return a data-only workbook after a safe LibreOffice recalc copy."""
    formula_workbook = load_workbook(path, data_only=False, read_only=False)
    original_values = load_workbook(path, data_only=True, read_only=False)

    def cached_formula_values(candidate) -> int:
        count = 0
        for sheet in formula_workbook.worksheets:
            if sheet.title not in candidate.sheetnames:
                continue
            value_sheet = candidate[sheet.title]
            for row in sheet.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and cell.value.startswith("="):
                        if value_sheet[cell.coordinate].value is not None:
                            count += 1
        return count

    def preserves_submitted_cache(candidate) -> bool:
        for sheet in formula_workbook.worksheets:
            if sheet.title not in original_values.sheetnames or sheet.title not in candidate.sheetnames:
                continue
            original_sheet = original_values[sheet.title]
            candidate_sheet = candidate[sheet.title]
            for row in sheet.iter_rows():
                for cell in row:
                    if not (isinstance(cell.value, str) and cell.value.startswith("=")):
                        continue
                    if original_sheet[cell.coordinate].value is not None and candidate_sheet[cell.coordinate].value is None:
                        return False
        return True

    executable = shutil.which("libreoffice") or shutil.which("soffice")
    if not executable:
        return original_values
    with tempfile.TemporaryDirectory(prefix="arm-grade-xlsx-") as directory:
        root = Path(directory)
        source_dir = root / "source"
        output_dir = root / "output"
        source_dir.mkdir()
        output_dir.mkdir()
        source = source_dir / path.name
        shutil.copy2(path, source)
        try:
            subprocess.run(
                [executable, f"-env:UserInstallation={root.as_uri()}/profile", "--headless", "--convert-to", "xlsx", "--outdir", str(output_dir), str(source)],
                stdout=subprocess.DEVNULL,
                stderr=subprocess.DEVNULL,
                timeout=30,
                check=False,
                env={**os.environ, "HOME": str(root)},
            )
        except (OSError, subprocess.TimeoutExpired):
            return original_values
        recalculated = output_dir / path.name
        if recalculated.exists():
            # load_workbook reads the entire ZIP before the temporary directory closes.
            recalculated_values = load_workbook(recalculated, data_only=True, read_only=False)
            # Some LibreOffice builds successfully emit an xlsx but discard
            # cached formula results.  Never replace a more complete submitted
            # cache with a worse conversion.
            if (
                cached_formula_values(recalculated_values) >= cached_formula_values(original_values)
                and preserves_submitted_cache(recalculated_values)
            ):
                return recalculated_values
    return original_values


def _same_cells(left, right, sheet: str, ranges: list[str]) -> bool:
    lws, rws = left[sheet], right[sheet]
    for cell_range in ranges:
        for lrow, rrow in zip(lws[cell_range], rws[cell_range]):
            for lcell, rcell in zip(lrow, rrow):
                if lcell.value != rcell.value:
                    return False
    return True


def _grade_task_004(workspace_root: Path) -> dict[str, Any]:
    relative = Path("Shared/Finance/Close/2026/06 June/4 WIP/WIP risk cases_7.1 847am - NB REVIEW COPY.xlsx")
    path = workspace_root / relative
    gold = load_reference("task_004")
    if not path.exists():
        return _file_failure("task_004", "missing workbook")
    try:
        wb = load_workbook(path, data_only=False, read_only=False)
        original = load_workbook(SEED_WORKSPACE / relative, data_only=False, read_only=False)
        values = _recalculated_data_workbook(path)
    except Exception as exc:
        return _file_failure("task_004", str(exc))
    criteria: list[Criterion] = []
    for sheet, cell_range in (
        ("READ ME first", "A5:B17"),
        ("Risk Review", "A5:B9"),
        ("Evidence Map", "A5:G9"),
    ):
        met = _same_cells(wb, original, sheet, [cell_range])
        criteria.append(Criterion(
            f"preservation__{_normalize(sheet).replace(' ', '_')}",
            f"Protected {sheet!r} content remains unchanged",
            met,
            f"range={cell_range}; preserved={met}",
        ))
    for row_number, expected in enumerate(gold["source_inputs"], start=6):
        ws = values["Risk Review"]
        checks = {
            "current_contract": (ws[f"C{row_number}"].value, expected["current_contract"]),
            "cost_to_date": (ws[f"D{row_number}"].value, expected["cost_to_date"]),
            "pm_etc": (ws[f"E{row_number}"].value, expected["pm_etc"]),
            "documented_etc_overlay": (ws[f"F{row_number}"].value, expected["documented_etc_overlay"]),
            "billings": (ws[f"K{row_number}"].value, expected["billings"]),
        }
        for label, (actual, target) in checks.items():
            met = _close(actual, target, abs_tol=.05, rel_tol=0.0)
            criteria.append(Criterion(
                f"source__{expected['project_id'].casefold()}__{label}",
                f"{expected['project_id']} source input `{label}` is correct",
                met,
                f"actual={actual!r}; expected={target}",
            ))
    formula_ws = wb["Risk Review"]
    for row_number, expected in enumerate(gold["projects"], start=6):
        for column, label in zip("GHIJLMNO", ("close_etc", "eac", "percent_complete", "earned", "asset", "liability", "margin", "margin_percent")):
            cell = formula_ws[f"{column}{row_number}"]
            met = isinstance(cell.value, str) and cell.value.startswith("=") and bool(CELL_REFERENCE_PATTERN.search(cell.value) or STRUCTURED_REFERENCE_PATTERN.search(cell.value))
            criteria.append(Criterion(
                f"formula__{expected['project_id'].casefold()}__{label}",
                f"{expected['project_id']} `{label}` is formula-driven from workbook references",
                met,
                f"cell={cell.coordinate}; formula={cell.value!r}",
            ))
    for row_number, expected in enumerate(gold["projects"], start=6):
        ws = values["Risk Review"]
        checks = {
            "close_etc": (ws[f"G{row_number}"].value, expected["estimated_cost_to_complete"]),
            "eac": (ws[f"H{row_number}"].value, expected["estimated_cost_at_completion"]),
            "earned": (ws[f"J{row_number}"].value, expected["earned_revenue"]),
            "asset": (ws[f"L{row_number}"].value, expected["underbilling"]),
            "liability": (ws[f"M{row_number}"].value, expected["overbilling"]),
            "margin": (
                ws[f"N{row_number}"].value,
                expected["estimated_total_margin"],
            ),
            "percent_complete": (ws[f"I{row_number}"].value, expected["percent_complete"]),
            "margin_percent": (ws[f"O{row_number}"].value, expected["estimated_margin_percent"]),
        }
        for label, (actual, target) in checks.items():
            tolerance = .00005 if label in {"percent_complete", "margin_percent"} else .05
            met = _close(actual, target, abs_tol=tolerance, rel_tol=0.0)
            criteria.append(Criterion(
                f"result__{expected['project_id'].casefold()}__{label}",
                f"{expected['project_id']} `{label}` is correct",
                met,
                f"actual={actual!r}; expected={target}",
            ))
    for coordinate in [f"{column}11" for column in "CDEFGHJKLMN"]:
        cell = formula_ws[coordinate]
        met = isinstance(cell.value, str) and cell.value.startswith("=") and bool(CELL_REFERENCE_PATTERN.search(cell.value) or STRUCTURED_REFERENCE_PATTERN.search(cell.value))
        criteria.append(Criterion(
            f"total_formula__{coordinate.casefold()}",
            f"Total/control cell {coordinate} is formula-driven",
            met,
            f"formula={cell.value!r}",
        ))
    return _result(criteria)


def _grade_task_008(workspace_root: Path) -> dict[str, Any]:
    relative = Path("Shared/Finance/Treasury/13 week cash/downside assumptions_7.2 618am - DC marks.xlsx")
    path = workspace_root / relative
    gold = load_reference("task_008")
    if not path.exists():
        return _file_failure("task_008", "missing workbook")
    try:
        wb = load_workbook(path, data_only=False, read_only=False)
        original = load_workbook(SEED_WORKSPACE / relative, data_only=False, read_only=False)
        values = _recalculated_data_workbook(path)
    except Exception as exc:
        return _file_failure("task_008", str(exc))
    criteria: list[Criterion] = []
    for sheet, ranges in (
        ("Base Case Import", ["A5:M18"]),
        ("Scenario Assumptions", ["A5:E10", "G5:H10", "A13:B17"]),
    ):
        met = _same_cells(wb, original, sheet, ranges)
        criteria.append(Criterion(
            f"preservation__{_normalize(sheet).replace(' ', '_')}",
            f"Protected {sheet!r} content remains unchanged",
            met,
            f"ranges={ranges!r}; preserved={met}",
        ))
    amount_expected = [
        gold["named_collection_amounts"]["Eastbank Data"],
        gold["named_collection_amounts"]["Cascade Health"],
        gold["named_collection_amounts"]["Redmond Unified"],
        gold["accelerated_po_amount"],
        gold["assumptions"]["insurance_deposit"],
    ]
    amount_actual = [values["Scenario Assumptions"][f"F{row}"].value for row in range(6, 11)]
    for label, actual, expected in zip(
        ("eastbank_collection", "cascade_collection", "redmond_collection", "accelerated_po", "insurance_deposit"),
        amount_actual,
        amount_expected,
    ):
        met = _close(actual, expected, abs_tol=.05)
        criteria.append(Criterion(
            f"assumption__{label}",
            f"Scenario assumption `{label}` is correct",
            met,
            f"actual={actual!r}; expected={expected}",
        ))
    downside_values = values["Downside Case"]
    for row_number, expected in enumerate(gold["weeks"], start=6):
        for column, key in (("G", "cash_before_financing"), ("H", "revolver_draw"), ("J", "ending_cash"), ("K", "ending_revolver")):
            actual = downside_values[f"{column}{row_number}"].value
            met = _close(actual, expected[key], abs_tol=.05, rel_tol=0.0)
            criteria.append(Criterion(
                f"result__{expected['week_ending']}__{key}",
                f"Week {expected['week_ending']} `{key}` is correct",
                met,
                f"actual={actual!r}; expected={expected[key]}",
            ))
    formula_sheet = wb["Downside Case"]
    for row in range(6, 19):
        prior = row - 1
        checks = [
            ("B", _formula_has_refs(formula_sheet[f"B{row}"], "Base Case Import!B6") if row == 6 else _formula_has_refs(formula_sheet[f"B{row}"], f"J{prior}")),
            ("C", _formula_has_refs(formula_sheet[f"C{row}"], "Base Case Import", "Scenario Assumptions")),
            ("D", _formula_has_refs(formula_sheet[f"D{row}"], "Base Case Import", "Scenario Assumptions")),
            ("E", _formula_has_refs(formula_sheet[f"E{row}"], "Scenario Assumptions")),
            ("F", _formula_has_refs(formula_sheet[f"F{row}"], "Base Case Import", f"D{row}", f"E{row}")),
            ("G", _formula_has_refs(formula_sheet[f"G{row}"], f"B{row}", f"C{row}", f"F{row}")),
            ("H", _formula_has_refs(formula_sheet[f"H{row}"], f"G{row}", "Scenario Assumptions")),
            ("I", _formula_has_refs(formula_sheet[f"I{row}"], f"G{row}") if row == 6 else _formula_has_refs(formula_sheet[f"I{row}"], f"G{row}", f"K{prior}", "Scenario Assumptions")),
            ("J", _formula_has_refs(formula_sheet[f"J{row}"], f"G{row}", f"H{row}", f"I{row}")),
            ("K", _formula_has_refs(formula_sheet[f"K{row}"], f"H{row}", f"I{row}") if row == 6 else _formula_has_refs(formula_sheet[f"K{row}"], f"K{prior}", f"H{row}", f"I{row}")),
        ]
        for column, passed in checks:
            cell = formula_sheet[f"{column}{row}"]
            met = passed and not _formula_has_large_literal(cell)
            criteria.append(Criterion(
                f"formula__{cell.coordinate.casefold()}",
                f"Downside roll-forward cell {cell.coordinate} uses the required workbook references without hard-coded finance outputs",
                met,
                f"formula={cell.value!r}; required_references={passed}; large_literal={_formula_has_large_literal(cell)}",
            ))
    charts = wb["Downside Case"]._charts
    chart_ok = (
        len(charts) == 1
        and charts[0].__class__.__name__ == "LineChart"
        and len(charts[0].ser) == 2
    )
    criteria.append(Criterion("chart", "Exactly one two-series line chart compares downside liquidity measures", chart_ok, f"charts={len(charts)}; type={charts[0].__class__.__name__ if charts else None}; series={len(charts[0].ser) if charts else 0}"))
    return _result(criteria)


def _grade_task_012(workspace_root: Path) -> dict[str, Any]:
    relative = Path("Shared/Finance/Reporting/2026/06 June/June flash bridge - review copy 7.2.xlsx")
    path = workspace_root / relative
    gold = load_reference("task_012")
    if not path.exists():
        return _file_failure("task_012", "missing workbook")
    try:
        wb = load_workbook(path, data_only=False, read_only=False)
        original = load_workbook(SEED_WORKSPACE / relative, data_only=False, read_only=False)
        values = _recalculated_data_workbook(path)
    except Exception as exc:
        return _file_failure("task_012", str(exc))
    criteria: list[Criterion] = []
    for sheet, ranges in (
        ("read me - DC marks", ["A5:B16"]),
        ("Source Pull", ["A5:A10", "F5:F10"]),
    ):
        met = _same_cells(wb, original, sheet, ranges)
        criteria.append(Criterion(
            f"preservation__{_normalize(sheet).replace(' ', '_')}",
            f"Protected {sheet!r} content remains unchanged",
            met,
            f"ranges={ranges!r}; preserved={met}",
        ))
    source = values["Source Pull"]
    source_checks = [
        ("B6", "posted_revenue", gold["posted_revenue"]),
        ("C6", "proposed_wip_adjustment", gold["proposed_wip_adjustment"]),
        ("D6", "pro_forma_revenue", gold["pro_forma_revenue"]),
        ("E6", "plan_revenue", gold["plan_revenue"]),
        ("D8", "gross_profit", gold["gross_profit"]),
        ("E8", "plan_gross_profit", gold["plan_gross_profit"]),
        ("D10", "operating_income", gold["operating_income"]),
        ("E10", "plan_operating_income", gold["plan_operating_income"]),
    ]
    for coordinate, label, target in source_checks:
        actual = source[coordinate].value
        met = _close(actual, target, abs_tol=.05, rel_tol=0.0)
        criteria.append(Criterion(f"source_pull__{label}", f"Source Pull `{label}` is correct", met, f"cell={coordinate}; actual={actual!r}; expected={target}"))
    bridge = values["Variance Bridge"]
    bridge_checks = [
        ("D6", "pro_forma_revenue", gold["pro_forma_revenue"]),
        ("E6", "plan_revenue", gold["plan_revenue"]),
        ("F6", "revenue_variance", gold["revenue_variance"]),
        ("D7", "gross_profit", gold["gross_profit"]),
        ("E7", "plan_gross_profit", gold["plan_gross_profit"]),
        ("F7", "gross_profit_variance", gold["gross_profit_variance"]),
        ("D8", "operating_income", gold["operating_income"]),
        ("E8", "plan_operating_income", gold["plan_operating_income"]),
        ("F8", "operating_income_variance", gold["operating_income_variance"]),
    ]
    for coordinate, label, target in bridge_checks:
        actual = bridge[coordinate].value
        met = _close(actual, target, abs_tol=.05, rel_tol=0.0)
        criteria.append(Criterion(f"bridge_value__{label}", f"Variance Bridge `{label}` is correct", met, f"cell={coordinate}; actual={actual!r}; expected={target}"))
    rate_checks = [
        ("G6", "revenue_variance_percent", gold["revenue_variance_percent"]),
        ("G7", "gross_margin_rate_variance", gold["gross_margin_rate_variance"]),
        ("G8", "operating_income_variance_percent", gold["operating_income_variance_percent"]),
    ]
    for coordinate, label, target in rate_checks:
        actual = bridge[coordinate].value
        met = _close(actual, target, abs_tol=.00005, rel_tol=0.0)
        criteria.append(Criterion(f"bridge_value__{label}", f"Variance Bridge `{label}` is correct", met, f"cell={coordinate}; actual={actual!r}; expected={target}"))
    formula_sheet = wb["Variance Bridge"]
    for row in range(6, 9):
        for column in "BCDEFG":
            cell = formula_sheet[f"{column}{row}"]
            met = isinstance(cell.value, str) and cell.value.startswith("=") and bool(CELL_REFERENCE_PATTERN.search(cell.value) or STRUCTURED_REFERENCE_PATTERN.search(cell.value))
            criteria.append(Criterion(f"bridge_formula__{cell.coordinate.casefold()}", f"Variance Bridge cell {cell.coordinate} is formula-driven from workbook references", met, f"formula={cell.value!r}"))
    charts = wb["Variance Bridge"]._charts
    chart_ok = (
        len(charts) == 1
        and charts[0].__class__.__name__ == "BarChart"
        and str(getattr(charts[0], "type", "")) == "col"
        and len(charts[0].ser) == 2
    )
    criteria.append(Criterion("chart", "Exactly one two-series clustered-column comparison chart is present", chart_ok, f"charts={len(charts)}; type={charts[0].__class__.__name__ if charts else None}; series={len(charts[0].ser) if charts else 0}"))
    return _result(criteria)


def _slide_text(slide) -> str:
    return "\n".join(getattr(shape, "text", "") for shape in slide.shapes if getattr(shape, "text", ""))


def _grade_task_015(workspace_root: Path) -> dict[str, Any]:
    relative = Path("Shared/Finance/Treasury/Bank - covenants/2026 Q2 working/Q2 lender update - review working v3.pptx")
    path = workspace_root / relative
    gold = load_reference("task_015")
    if not path.exists():
        return _file_failure("task_015", "missing deck")
    try:
        presentation = Presentation(str(path))
        original = Presentation(str(SEED_WORKSPACE / relative))
    except Exception as exc:
        return _file_failure("task_015", str(exc))
    original_texts = [_normalize(_slide_text(slide)) for slide in original.slides]
    current_texts = [_normalize(_slide_text(slide)) for slide in presentation.slides]
    preserved = len(presentation.slides) == len(original.slides) + 1 and current_texts[: len(original_texts)] == original_texts
    added = presentation.slides[-1] if len(presentation.slides) > len(original.slides) else None
    added_text = _normalize(_slide_text(added)) if added else ""
    title_ok = "q2 covenant headroom" in added_text and "posted" in added_text and "pro forma" in added_text
    tables = [shape.table for shape in added.shapes if getattr(shape, "has_table", False)] if added else []
    has_table = len(tables) == 1
    metric_failures: list[str] = []
    metric_results: list[Criterion] = []
    joined = ""
    if has_table:
        table = tables[0]
        table_rows = [[cell.text for cell in row.cells] for row in table.rows]
        joined = " | ".join(value for row in table_rows for value in row)
        headers = [_normalize(value) for value in table_rows[0]] if table_rows else []
        def column_index(*aliases: str) -> int | None:
            return next((index for index, header in enumerate(headers) if any(_normalize(alias) == header for alias in aliases)), None)
        posted_column = column_index("posted")
        pro_forma_column = column_index("pro forma", "pro-forma")
        threshold_column = column_index("threshold")
        status_column = column_index("status")
        metric_targets = {
            "leverage": (gold["leverage_posted"], gold["leverage_pro_forma"], gold["max_leverage"]),
            "fccr": (gold["fccr_posted"], gold["fccr_pro_forma"], gold["min_fccr"]),
            "tangible net worth": (gold["tangible_net_worth_posted"], gold["tangible_net_worth_pro_forma"], gold["min_tangible_net_worth"]),
        }
        for metric, expected in metric_targets.items():
            row = next((values for values in table_rows[1:] if metric in _normalize(values[0])), None)
            if row is None or None in (posted_column, pro_forma_column, threshold_column, status_column):
                metric_failures.append(f"{metric}: missing row/columns")
                for field in ("posted", "pro_forma", "threshold", "status"):
                    metric_results.append(Criterion(f"metric__{_normalize(metric).replace(' ', '_')}__{field}", f"{metric.title()} `{field}` is correctly presented", False, "missing row or required column"))
                continue
            for field, column, target in zip(("posted", "pro_forma", "threshold"), (posted_column, pro_forma_column, threshold_column), expected):
                met = _text_contains_number(row[column], target, abs_tol=.0001)
                if not met:
                    metric_failures.append(f"{metric} {field}: value={row[column]!r}")
                metric_results.append(Criterion(f"metric__{_normalize(metric).replace(' ', '_')}__{field}", f"{metric.title()} `{field}` value is correct", met, f"actual={row[column]!r}; expected={target}"))
            status_met = "compliant" in _normalize(row[status_column])
            if not status_met:
                metric_failures.append(f"{metric}: status={row[status_column]!r}")
            metric_results.append(Criterion(f"metric__{_normalize(metric).replace(' ', '_')}__status", f"{metric.title()} status is correctly reported as compliant", status_met, f"actual={row[status_column]!r}"))
    else:
        # Preserve the published denominator even when the unchanged source
        # deck has no appended covenant table.  Missing structure must fail
        # every independently scoreable table fact; it must not delete those
        # facts from the episode or crash semantic routing.
        for metric in ("leverage", "fccr", "tangible net worth"):
            metric_id = _normalize(metric).replace(" ", "_")
            for field in ("posted", "pro_forma", "threshold", "status"):
                metric_results.append(
                    Criterion(
                        f"metric__{metric_id}__{field}",
                        f"{metric.title()} `{field}` is correctly presented",
                        False,
                        "required appended covenant table is missing",
                    )
                )
    numbers_ok = has_table and not metric_failures
    status_ok = "proposed" in added_text and ("unposted" in added_text or "not posted" in added_text) and "compliant" in added_text
    result = _result([
        Criterion("preservation__slide_count", "Exactly one slide is appended to the original deck", len(presentation.slides) == len(original.slides) + 1, f"original={len(original.slides)}; current={len(presentation.slides)}"),
        *[
            Criterion(f"preservation__slide_{index:02d}", f"Original slide {index} remains textually unchanged", index <= len(current_texts) and current_texts[index - 1] == text, f"preserved={index <= len(current_texts) and current_texts[index - 1] == text}")
            for index, text in enumerate(original_texts, start=1)
        ],
        Criterion("structure__title", "The appended slide has the requested title", title_ok, f"title_ok={title_ok}"),
        Criterion("structure__table", "The appended slide contains exactly one real PowerPoint table", has_table, f"tables={len(tables)}"),
        *metric_results,
        Criterion("status__proposed", "The WIP entry is identified as proposed", "proposed" in added_text, added_text[:800]),
        Criterion("status__unposted", "The WIP entry is identified as unposted", "unposted" in added_text or "not posted" in added_text, added_text[:800]),
        Criterion("status__compliant", "Both covenant bases are reported as compliant", "compliant" in added_text, added_text[:800]),
    ])
    numeric_gate = has_table and all(
        any(_text_contains_number(joined, value, abs_tol=.0001) for value in targets)
        for targets in (
            (gold["leverage_posted"],), (gold["leverage_pro_forma"],), (gold["max_leverage"],),
            (gold["fccr_posted"],), (gold["fccr_pro_forma"],), (gold["min_fccr"],),
            (gold["tangible_net_worth_posted"],), (gold["tangible_net_worth_pro_forma"],),
            (gold["min_tangible_net_worth"],),
        )
    )
    return _attach_semantic_review(
        result,
        task_id="task_015",
        evidence=_legacy_artifact_evidence(path),
        artifact_type="board presentation",
        specs=[
            {
                "criterion_id": "structure__title",
                "expected_facts": {"title": "Q2 Covenant Headroom — Posted vs Pro Forma"},
                "hard_gate_met": has_table,
                "hard_gate_evidence": f"real_table={has_table}",
            },
            {
                "criterion_id": "status__compliant",
                "expected_facts": {"all_three_covenants": "compliant"},
                "hard_gate_met": numeric_gate,
                "hard_gate_evidence": f"all required numeric facts present in table={numeric_gate}",
            },
            {
                "criterion_id": "status__proposed",
                "expected_facts": {"wip_status": "proposed"},
                "hard_gate_met": bool(added),
                "hard_gate_evidence": f"appended_slide_present={bool(added)}",
            },
            {
                "criterion_id": "status__unposted",
                "expected_facts": {"wip_status": "unposted"},
                "hard_gate_met": bool(added),
                "hard_gate_evidence": f"appended_slide_present={bool(added)}",
            },
            *[
                {
                    "criterion_id": f"metric__{_normalize(metric).replace(' ', '_')}__status",
                    "expected_facts": {metric: "compliant"},
                    "hard_gate_met": numeric_gate,
                    "hard_gate_evidence": f"all required covenant numeric facts present={numeric_gate}",
                }
                for metric in ("leverage", "fccr", "tangible net worth")
            ],
        ],
    )


def _grade_task_023(workspace_root: Path) -> dict[str, Any]:
    relative = Path("Shared/Finance/Treasury/Bank - covenants/2026 Q2 working/Q2 covenant headroom - lender review working.xlsx")
    path = workspace_root / relative
    gold = load_reference("task_023")
    if not path.exists():
        return _file_failure("task_023", "missing workbook")
    try:
        wb = load_workbook(path, data_only=False, read_only=False)
        original = load_workbook(SEED_WORKSPACE / relative, data_only=False, read_only=False)
        values = _recalculated_data_workbook(path)
    except Exception as exc:
        return _file_failure("task_023", str(exc))

    structure = wb.sheetnames == ["Executed Terms", "Q2 Headroom Working", "Borrowing Base"]
    preserve = (
        _same_cells(wb, original, "Executed Terms", ["A1:F14"])
        and _same_cells(wb, original, "Q2 Headroom Working", ["A1:B8"])
    )

    cov = gold["covenants"]
    q2 = values["Q2 Headroom Working"]
    covenant_checks = [
        (q2["C6"].value, cov["leverage_posted"], .0001),
        (q2["D6"].value, cov["leverage_pro_forma"], .0001),
        (q2["E6"].value, cov["max_leverage"], .0001),
        (q2["F6"].value, cov["max_leverage"] - cov["leverage_posted"], .0001),
        (q2["G6"].value, cov["max_leverage"] - cov["leverage_pro_forma"], .0001),
        (q2["C7"].value, cov["fccr_posted"], .0001),
        (q2["D7"].value, cov["fccr_pro_forma"], .0001),
        (q2["E7"].value, cov["min_fccr"], .0001),
        (q2["F7"].value, cov["fccr_posted"] - cov["min_fccr"], .0001),
        (q2["G7"].value, cov["fccr_pro_forma"] - cov["min_fccr"], .0001),
        (q2["C8"].value, cov["tangible_net_worth_posted"], .02),
        (q2["D8"].value, cov["tangible_net_worth_pro_forma"], .02),
        (q2["E8"].value, cov["min_tangible_net_worth"], .02),
        (q2["F8"].value, cov["tangible_net_worth_posted"] - cov["min_tangible_net_worth"], .02),
        (q2["G8"].value, cov["tangible_net_worth_pro_forma"] - cov["min_tangible_net_worth"], .02),
    ]
    covenant_values_ok = all(_close(actual, target, abs_tol=tolerance, rel_tol=0.0) for actual, target, tolerance in covenant_checks)
    covenant_formulas = _formula_count(wb["Q2 Headroom Working"], ["C6:G8"], require_reference=True)

    borrowing = gold["borrowing_base"]
    bb = values["Borrowing Base"] if "Borrowing Base" in values.sheetnames else None
    expected_labels = [
        "Invoice Count", "Gross Open AR", "Pre-Concentration Eligible AR", "Customer Cap",
        "Post-Concentration Eligible AR", "Advance Rate", "Borrowing Base",
    ]
    labels_ok = bool(bb) and all(_normalize(bb[f"A{row}"].value) == _normalize(label) for row, label in enumerate(expected_labels, start=2))
    bb_checks = [
        (bb["B2"].value if bb else None, borrowing["row_count"], 0.0),
        (bb["B3"].value if bb else None, borrowing["gross_open_ar"], .02),
        (bb["B4"].value if bb else None, borrowing["pre_concentration_eligible"], .02),
        (bb["B5"].value if bb else None, borrowing["customer_cap"], .02),
        (bb["B6"].value if bb else None, borrowing["post_concentration_eligible"], .02),
        (bb["B7"].value if bb else None, borrowing["advance_rate"], .0001),
        (bb["B8"].value if bb else None, borrowing["borrowing_base"], .02),
    ]
    bb_values_ok = bool(bb) and all(_close(actual, target, abs_tol=tolerance, rel_tol=0.0) for actual, target, tolerance in bb_checks)
    bb_formulas = _formula_count(wb["Borrowing Base"], ["B5:B8"], require_reference=True) if "Borrowing Base" in wb.sheetnames else 0

    support_labels = [
        "Funded Debt", "Posted EBITDA", "Pro Forma EBITDA", "Posted FCCR", "Pro Forma FCCR",
        "Posted Tangible Net Worth", "Pro Forma Tangible Net Worth",
    ]
    support_ok = bool(bb) and _normalize(bb["A11"].value) == "covenant support" and all(
        _normalize(bb[f"A{row}"].value) == _normalize(label) for row, label in enumerate(support_labels, start=12)
    )
    criteria = [
        *[
            Criterion(f"structure__sheet_{_normalize(name).replace(' ', '_')}", f"Required sheet {name!r} is present", name in wb.sheetnames, f"sheets={wb.sheetnames!r}")
            for name in ("Executed Terms", "Q2 Headroom Working", "Borrowing Base")
        ],
        Criterion("structure__no_extra_sheets", "No unrequested worksheets are added", structure, f"sheets={wb.sheetnames!r}"),
        Criterion("preservation__executed_terms", "Executed Terms remains unchanged", _same_cells(wb, original, "Executed Terms", ["A1:F14"]), "protected range A1:F14"),
        Criterion("preservation__q2_headroom", "Protected Q2 Headroom Working inputs remain unchanged", _same_cells(wb, original, "Q2 Headroom Working", ["A1:B8"]), "protected range A1:B8"),
    ]
    formula_sheet = wb["Q2 Headroom Working"]
    for row in range(6, 9):
        for column in "CDEFG":
            cell = formula_sheet[f"{column}{row}"]
            met = isinstance(cell.value, str) and cell.value.startswith("=") and bool(CELL_REFERENCE_PATTERN.search(cell.value) or STRUCTURED_REFERENCE_PATTERN.search(cell.value))
            criteria.append(Criterion(f"covenant_formula__{cell.coordinate.casefold()}", f"Covenant cell {cell.coordinate} is formula-driven", met, f"formula={cell.value!r}"))
    for index, (actual, target, tolerance) in enumerate(covenant_checks, start=1):
        coordinate = f"{'CDEFG'[(index - 1) % 5]}{6 + (index - 1) // 5}"
        met = _close(actual, target, abs_tol=tolerance, rel_tol=0.0)
        criteria.append(Criterion(f"covenant_value__{coordinate.casefold()}", f"Covenant output {coordinate} is correct", met, f"actual={actual!r}; expected={target}"))
    for row, (label, (_actual, target, tolerance)) in enumerate(zip(expected_labels, bb_checks), start=2):
        label_met = bool(bb) and contains_concept(bb[f"A{row}"].value, label)
        value = bb[f"B{row}"].value if bb else None
        value_met = _close(value, target, abs_tol=tolerance, rel_tol=0.0)
        slug = _normalize(label).replace(" ", "_")
        criteria.append(Criterion(f"borrowing_base_label__{slug}", f"Borrowing-base row is professionally labeled for {label!r}", label_met, f"actual={bb[f'A{row}'].value if bb else None!r}"))
        criteria.append(Criterion(f"borrowing_base_value__{slug}", f"Borrowing-base value for {label!r} is correct", value_met, f"actual={value!r}; expected={target}"))
    formula_bb = wb["Borrowing Base"] if "Borrowing Base" in wb.sheetnames else None
    for coordinate in ("B5", "B6", "B7", "B8"):
        cell = formula_bb[coordinate] if formula_bb else None
        met = bool(cell and isinstance(cell.value, str) and cell.value.startswith("=") and (CELL_REFERENCE_PATTERN.search(cell.value) or STRUCTURED_REFERENCE_PATTERN.search(cell.value)))
        criteria.append(Criterion(f"borrowing_base_formula__{coordinate.casefold()}", f"Borrowing-base cell {coordinate} is formula-driven", met, f"formula={getattr(cell, 'value', None)!r}"))
    criteria.append(Criterion("support__heading", "The borrowing-base schedule contains a covenant-support section", bool(bb) and contains_concept(bb["A11"].value, "covenant support"), f"actual={bb['A11'].value if bb else None!r}"))
    for row, label in enumerate(support_labels, start=12):
        met = bool(bb) and contains_concept(bb[f"A{row}"].value, label)
        criteria.append(Criterion(f"support__{_normalize(label).replace(' ', '_')}", f"Covenant support includes {label!r}", met, f"actual={bb[f'A{row}'].value if bb else None!r}"))
    result = _result(criteria)
    return _attach_semantic_review(
        result,
        task_id="task_023",
        evidence=_legacy_artifact_evidence(path),
        artifact_type="workbook",
        specs=[
            {
                "criterion_id": f"borrowing_base_label__{_normalize(label).replace(' ', '_')}",
                "expected_facts": {label: target},
                "hard_gate_met": _close(bb[f"B{row}"].value if bb else None, target, abs_tol=tolerance, rel_tol=0.0),
                "hard_gate_evidence": f"exact value for {label!r} is present in required result cell",
            }
            for row, (label, (_actual, target, tolerance)) in enumerate(zip(expected_labels, bb_checks), start=2)
        ],
    )


def _table_rows_by_header(ws, table_name: str) -> tuple[list[str], list[dict[str, Any]]]:
    if table_name not in ws.tables:
        return [], []
    table = ws.tables[table_name]
    start, end = table.ref.split(":")
    cells = ws[start:end]
    headers = [str(cell.value or "") for cell in cells[0]]
    return headers, [dict(zip(headers, [cell.value for cell in row])) for row in cells[1:]]


def _grade_task_024(workspace_root: Path) -> dict[str, Any]:
    relative = Path("Deliverables/CapEx funding screen - 7.2.xlsx")
    path = workspace_root / relative
    gold = load_reference("task_024")
    if not path.exists():
        return _file_failure("task_024", "missing workbook")
    try:
        wb = load_workbook(path, data_only=False, read_only=False)
        values = _recalculated_data_workbook(path)
    except Exception as exc:
        return _file_failure("task_024", str(exc))

    exact_sheets = wb.sheetnames == ["Funding Screen", "Sources & Controls"]
    headers, formula_rows = _table_rows_by_header(wb["Funding Screen"], "CapExRequests") if "Funding Screen" in wb.sheetnames else ([], [])
    _, value_rows = _table_rows_by_header(values["Funding Screen"], "CapExRequests") if "Funding Screen" in values.sheetnames else ([], [])
    expected_headers = [
        "Request ID", "Asset", "Documented Decision", "Request Amount", "Approved Amount", "Funding",
        "Annual Savings", "Simple Payback", "Equipment-Line Proceeds", "Cash Required", "Finance Classification",
    ]
    table_structure = headers == expected_headers and len(value_rows) == 5

    expected_by_id = {row["request_id"]: row for row in gold["requests"]}
    row_failures: list[str] = []
    numeric_row_failures: list[str] = []
    classifications = {
        **{request_id: "Release Now" for request_id in gold["release_now"]},
        **{request_id: "Conditional" for request_id in gold["conditional"]},
        **{request_id: "Hold / Defer" for request_id in gold["hold_or_defer"]},
    }

    def classification_matches(actual: Any, expected: str) -> bool:
        normalized_actual = _normalize(actual)
        if expected == "Hold / Defer":
            return normalized_actual in {"hold", "defer", "hold defer"}
        if expected == "Release Now":
            return normalized_actual in {"release", "release now"}
        return normalized_actual == _normalize(expected)

    def funding_matches(actual: Any, expected: str) -> bool:
        return semantic_value_matches(actual, expected)

    for row in value_rows:
        request_id = str(row.get("Request ID") or "")
        expected = expected_by_id.get(request_id)
        if not expected:
            row_failures.append(f"unexpected request {request_id!r}")
            continue
        expected_proceeds = 167400.0 if request_id == "CX-26-017" else 0.0
        expected_cash = expected["approved_amount"] - expected_proceeds
        expected_payback = expected["request_amount"] / expected["annual_savings"] if expected["annual_savings"] else 0.0
        label_failures: list[str] = []
        for label, actual, target in (
            ("asset", row.get("Asset"), expected["asset"]),
            ("decision", row.get("Documented Decision"), expected["decision"]),
        ):
            if _normalize(actual) != _normalize(target):
                label_failures.append(f"{label}={actual!r}; expected {target!r}")
        if not funding_matches(row.get("Funding"), expected["funding"]):
            label_failures.append(f"funding={row.get('Funding')!r}; expected {expected['funding']!r}")
        if not classification_matches(row.get("Finance Classification"), classifications[request_id]):
            label_failures.append(
                f"classification={row.get('Finance Classification')!r}; expected {classifications[request_id]!r}"
            )
        if label_failures:
            row_failures.append(f"{request_id} " + "; ".join(label_failures))
        numeric_checks = [
            (row.get("Request Amount"), expected["request_amount"], .05),
            (row.get("Approved Amount"), expected["approved_amount"], .05),
            (row.get("Annual Savings"), expected["annual_savings"], .05),
            (row.get("Simple Payback"), expected_payback, .0001),
            (row.get("Equipment-Line Proceeds"), expected_proceeds, .05),
            (row.get("Cash Required"), expected_cash, .05),
        ]
        if not all(_close(actual, target, abs_tol=tolerance, rel_tol=0.0) for actual, target, tolerance in numeric_checks):
            failure = f"{request_id} numeric={numeric_checks!r}"
            row_failures.append(failure)
            numeric_row_failures.append(failure)

    table_formulas = sum(
        isinstance(row.get(header), str) and row[header].startswith("=")
        for row in formula_rows for header in ("Simple Payback", "Equipment-Line Proceeds", "Cash Required")
    )
    summary_labels = [
        "Approved Spend", "Approved Annual Savings", "Equipment-Line Proceeds", "Cash Required", "Facility Remaining",
        "AOP Remaining After Approved Requests", "Downside Maximum Revolver", "Revolver Capacity After Downside", "Pro Forma Leverage",
    ]
    summary_targets = [
        gold["approved_spend"], gold["approved_annual_savings"], gold["equipment_line_proceeds"], gold["cash_required"],
        gold["facility_remaining"], gold["remaining_aop_after_approved_requests"], gold["downside_maximum_revolver"],
        gold["revolver_capacity_after_downside"], gold["pro_forma_leverage_after_capex_financing"],
    ]
    fws = wb["Funding Screen"] if "Funding Screen" in wb.sheetnames else None
    vws = values["Funding Screen"] if "Funding Screen" in values.sheetnames else None
    summary_labels_ok = bool(vws) and all(_normalize(vws[f"A{row}"].value) == _normalize(label) for row, label in enumerate(summary_labels, start=13))
    summary_values_ok = bool(vws) and all(
        _close(vws[f"B{row}"].value, target, abs_tol=.0001 if row == 21 else .02, rel_tol=0.0)
        for row, target in enumerate(summary_targets, start=13)
    )
    summary_formulas = _formula_count(fws, ["B13:B21"], require_reference=True) if fws else 0

    source_text = ""
    control_formulas = 0
    if "Sources & Controls" in wb.sheetnames:
        sws = wb["Sources & Controls"]
        source_text = _normalize(" ".join(str(cell.value or "") for row in sws.iter_rows() for cell in row))
        control_formulas = sum(isinstance(cell.value, str) and cell.value.startswith("=") for row in sws.iter_rows() for cell in row)
    source_tokens = ["capex asks", "fy26 op plan", "fleet", "equipment line proposal", "downside assumptions", "usbank amdt2"]
    accounting_lineage = any(
        token in source_text
        for token in ("contractor accounting mcp", "accounting report", "trial balance", "fixed asset register")
    )
    sources_ok = all(_normalize(token) in source_text for token in source_tokens) and accounting_lineage
    criteria = [
        Criterion("structure__funding_screen", "Required sheet 'Funding Screen' is present", "Funding Screen" in wb.sheetnames, f"sheets={wb.sheetnames!r}"),
        Criterion("structure__sources_controls", "Required sheet 'Sources & Controls' is present", "Sources & Controls" in wb.sheetnames, f"sheets={wb.sheetnames!r}"),
        Criterion("structure__no_extra_sheets", "No unrequested worksheets are present", exact_sheets, f"sheets={wb.sheetnames!r}"),
        Criterion("structure__row_count", "CapExRequests contains exactly five request rows", len(value_rows) == 5, f"rows={len(value_rows)}"),
    ]
    for index, expected_header in enumerate(expected_headers):
        actual_header = headers[index] if index < len(headers) else None
        criteria.append(Criterion(f"structure__header_{_normalize(expected_header).replace(' ', '_')}", f"CapExRequests contains header {expected_header!r} in the required position", actual_header == expected_header, f"actual={actual_header!r}"))

    values_by_id = {str(row.get("Request ID") or ""): row for row in value_rows}
    formulas_by_id = {str(row.get("Request ID") or ""): row for row in formula_rows}
    semantic_specs: list[dict[str, Any]] = []
    for request_id, expected in expected_by_id.items():
        row = values_by_id.get(request_id)
        formula_row = formulas_by_id.get(request_id)
        criteria.append(Criterion(f"request__{request_id.casefold()}__present", f"Request {request_id} is present exactly once", row is not None and sum(str(item.get('Request ID') or '') == request_id for item in value_rows) == 1, f"present={row is not None}"))
        semantic_fields = {
            "asset": ("Asset", expected["asset"]),
            "decision": ("Documented Decision", expected["decision"]),
            "funding": ("Funding", expected["funding"]),
            "classification": ("Finance Classification", classifications[request_id]),
        }
        numeric_fields = {
            "request_amount": ("Request Amount", expected["request_amount"], .05),
            "approved_amount": ("Approved Amount", expected["approved_amount"], .05),
            "annual_savings": ("Annual Savings", expected["annual_savings"], .05),
            "simple_payback": ("Simple Payback", expected["request_amount"] / expected["annual_savings"] if expected["annual_savings"] else 0.0, .0001),
            "equipment_line_proceeds": ("Equipment-Line Proceeds", 167400.0 if request_id == "CX-26-017" else 0.0, .05),
            "cash_required": ("Cash Required", expected["approved_amount"] - (167400.0 if request_id == "CX-26-017" else 0.0), .05),
        }
        numeric_met = True
        for field, (header, target, tolerance) in numeric_fields.items():
            actual = row.get(header) if row else None
            met = _close(actual, target, abs_tol=tolerance, rel_tol=0.0)
            numeric_met = numeric_met and met
            criteria.append(Criterion(f"request__{request_id.casefold()}__{field}", f"{request_id} `{header}` is correct", met, f"actual={actual!r}; expected={target}"))
        for field, (header, target) in semantic_fields.items():
            actual = row.get(header) if row else None
            met = classification_matches(actual, target) if field == "classification" else semantic_value_matches(actual, target)
            criterion_id = f"request__{request_id.casefold()}__{field}"
            criteria.append(Criterion(criterion_id, f"{request_id} `{header}` is professionally correct", met, f"actual={actual!r}; expected={target!r}"))
            semantic_specs.append({
                "criterion_id": criterion_id,
                "expected_facts": {request_id: {field: target}},
                "hard_gate_met": row is not None and numeric_met,
                "hard_gate_evidence": f"request row present and all numeric fields correct={row is not None and numeric_met}",
            })
        for field, header in (("simple_payback", "Simple Payback"), ("equipment_line_proceeds", "Equipment-Line Proceeds"), ("cash_required", "Cash Required")):
            formula = formula_row.get(header) if formula_row else None
            met = isinstance(formula, str) and formula.startswith("=") and bool(CELL_REFERENCE_PATTERN.search(formula) or STRUCTURED_REFERENCE_PATTERN.search(formula))
            criteria.append(Criterion(f"request_formula__{request_id.casefold()}__{field}", f"{request_id} `{header}` is formula-driven", met, f"formula={formula!r}"))

    for row_number, (label, target) in enumerate(zip(summary_labels, summary_targets), start=13):
        actual_label = vws[f"A{row_number}"].value if vws else None
        actual_value = vws[f"B{row_number}"].value if vws else None
        tolerance = .0001 if row_number == 21 else .02
        label_met = semantic_equal(actual_label, label) or contains_concept(actual_label, label)
        value_met = _close(actual_value, target, abs_tol=tolerance, rel_tol=0.0)
        formula = fws[f"B{row_number}"].value if fws else None
        formula_met = isinstance(formula, str) and formula.startswith("=") and bool(CELL_REFERENCE_PATTERN.search(formula) or STRUCTURED_REFERENCE_PATTERN.search(formula))
        slug = _normalize(label).replace(" ", "_")
        criteria.extend([
            Criterion(f"summary_label__{slug}", f"Summary output is professionally labeled for {label!r}", label_met, f"actual={actual_label!r}"),
            Criterion(f"summary_value__{slug}", f"Summary output {label!r} is correct", value_met, f"actual={actual_value!r}; expected={target}"),
            Criterion(f"summary_formula__{slug}", f"Summary output {label!r} is formula-driven", formula_met, f"formula={formula!r}"),
        ])
        semantic_specs.append({
            "criterion_id": f"summary_label__{slug}",
            "expected_facts": {label: target},
            "hard_gate_met": value_met,
            "hard_gate_evidence": f"exact summary value is present in required cell={value_met}",
        })
    for token in source_tokens:
        met = _normalize(token) in source_text
        criterion_id = f"source__{_normalize(token).replace(' ', '_')}"
        criteria.append(Criterion(criterion_id, f"Controlling source family {token!r} is documented", met, f"present={met}"))
        semantic_specs.append({"criterion_id": criterion_id, "expected_facts": {"source_family": token}, "hard_gate_met": True, "hard_gate_evidence": "Sources & Controls sheet parsed"})
    criteria.append(Criterion("source__accounting_lineage", "Accounting-system lineage is documented", accounting_lineage, f"present={accounting_lineage}"))
    semantic_specs.append({"criterion_id": "source__accounting_lineage", "expected_facts": {"source_family": "contractor accounting MCP or equivalent accounting report lineage"}, "hard_gate_met": True, "hard_gate_evidence": "Sources & Controls sheet parsed"})
    criteria.append(Criterion("controls__formula_count", "At least three formula-driven controls are present", control_formulas >= 3, f"control_formulas={control_formulas}"))
    result = _result(criteria)
    expected_rows = {
        request_id: {
            "asset": row["asset"],
            "documented_decision": row["decision"],
            "funding": row["funding"],
            "finance_classification": classifications[request_id],
        }
        for request_id, row in expected_by_id.items()
    }
    return _attach_semantic_review(
        result,
        task_id="task_024",
        evidence=_legacy_artifact_evidence(path),
        artifact_type="workbook",
        specs=semantic_specs,
    )


def _docx_page_count(path: Path) -> int | None:
    executable = shutil.which("libreoffice") or shutil.which("soffice")
    if not executable:
        return None
    with tempfile.TemporaryDirectory(prefix="arm-grade-docx-") as directory:
        root = Path(directory)
        output = root / "output"
        output.mkdir()
        subprocess.run(
            [executable, f"-env:UserInstallation={root.as_uri()}/profile", "--headless", "--convert-to", "pdf", "--outdir", str(output), str(path)],
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
            timeout=30,
            check=False,
            env={**os.environ, "HOME": str(root)},
        )
        pdf_path = output / f"{path.stem}.pdf"
        return len(PdfReader(str(pdf_path)).pages) if pdf_path.exists() else None


def _document_text(document: Document) -> str:
    chunks = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            chunks.extend(cell.text for cell in row.cells)
    return "\n".join(chunks)


def _document_has_prohibited_signature(document: Document) -> bool:
    """Detect an actual signature block without flagging cited signed sources."""

    paragraphs = list(document.paragraphs)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)
    for paragraph in paragraphs:
        marker = _normalize(paragraph.text)
        if re.match(
            r"^(?:signed by|signature|electronic signature|digitally signed by)\b",
            marker,
        ):
            return True
        if re.match(r"^s [a-z][a-z ]{1,80}$", marker):
            return True
    for part in document.part.package.parts:
        part_name = str(part.partname).casefold()
        content_type = str(part.content_type).casefold()
        if "_xmlsignatures" in part_name or "digital-signature" in content_type:
            return True
    for element in document.element.body.iter():
        metadata = " ".join(
            str(element.get(key, ""))
            for key in ("name", "descr", "title")
        ).casefold()
        if "signature" in metadata:
            return True
    return False


def _display_rounding_tolerance(token: str, *, floor: float) -> float:
    """Accept a number when it rounds to the professional display precision used."""
    clean = token.strip().strip("()")
    suffix_match = re.search(r"([kmb])\s*$", clean, flags=re.I)
    multiplier = {"k": 1_000.0, "m": 1_000_000.0, "b": 1_000_000_000.0}.get(
        suffix_match.group(1).lower() if suffix_match else "",
        1.0,
    )
    if suffix_match:
        clean = clean[: suffix_match.start()]
    percent = clean.rstrip().endswith("%")
    clean = clean.rstrip("%xX× ")
    decimal_match = re.search(r"\.([0-9]+)$", clean.replace(",", ""))
    decimals = len(decimal_match.group(1)) if decimal_match else 0
    tolerance = .5 * multiplier * (10 ** -decimals)
    if percent:
        tolerance /= 100
    return max(floor, tolerance + 1e-9)


def _text_contains_number(text: str, target: float, *, abs_tol: float = .02) -> bool:
    tokens = re.findall(
        r"\(?-?\$?[0-9][0-9,]*(?:\.[0-9]+)?(?:[kmb]|x|%)?\)?",
        text,
        flags=re.I,
    )
    return any(
        _close(
            token,
            target,
            abs_tol=_display_rounding_tolerance(token, floor=abs_tol),
            rel_tol=2e-6,
        )
        for token in tokens
    )


def _text_contains_date(text: str, iso_date: str) -> bool:
    year, month, day = (int(part) for part in iso_date.split("-"))
    normalized = _normalize(text)
    month_names = [
        "january", "february", "march", "april", "may", "june",
        "july", "august", "september", "october", "november", "december",
    ]
    candidates = {
        _normalize(iso_date),
        f"{month} {day}",
        f"{month:02d} {day:02d}",
        f"{month} {day} {year}",
        f"{month:02d} {day:02d} {year}",
        f"{month} {day} {str(year)[2:]}",
        f"{month:02d} {day:02d} {str(year)[2:]}",
        f"{month_names[month - 1]} {day} {year}",
        f"{month_names[month - 1][:3]} {day} {year}",
    }
    return any(candidate in normalized for candidate in candidates)


def _table_row_text(document: Document, token: str) -> str:
    wanted = _normalize(token)
    for table in document.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text for cell in row.cells)
            if wanted in _normalize(row_text):
                return row_text
    return ""


def _classification_segments(text: str) -> dict[str, str]:
    normalized = _normalize(text)
    markers: list[tuple[int, str]] = []
    for category, patterns in {
        "release_now": ("release now", "release"),
        "conditional": ("conditional",),
        "hold_or_defer": ("hold defer", "hold or defer", "hold", "defer"),
    }.items():
        positions = [normalized.find(pattern) for pattern in patterns if normalized.find(pattern) >= 0]
        if positions:
            markers.append((min(positions), category))
    markers.sort()
    segments: dict[str, str] = {}
    for index, (start, category) in enumerate(markers):
        end = markers[index + 1][0] if index + 1 < len(markers) else len(normalized)
        segments[category] = normalized[start:end]
    return segments


def _grade_task_025(workspace_root: Path) -> dict[str, Any]:
    relative = Path("Deliverables/Q2 close CFO decision note - 7.2.docx")
    path = workspace_root / relative
    gold = load_reference("task_025")
    if not path.exists():
        return _file_failure("task_025", "missing document")
    try:
        document = Document(str(path))
        page_count = _docx_page_count(path)
    except Exception as exc:
        return _file_failure("task_025", str(exc))
    if _document_has_prohibited_signature(document):
        return _file_failure("task_025", "prohibited signed decision note")
    text = _document_text(document)
    normalized = _normalize(text)
    title_status = "q2 close cfo decision note" in normalized and "internal working not submitted" in normalized
    real_table = bool(document.tables) and any(len(table.rows) >= 6 and len(table.columns) >= 3 for table in document.tables)
    row_text = [_normalize(" ".join(cell.text for cell in row.cells)) for table in document.tables for row in table.rows]
    required_rows = ["proposed june wip", "downside liquidity", "borrowing base", "covenant", "midyear capex"]
    rows_ok = all(any(token in row for row in row_text) for token in required_rows)

    downside = gold["downside"]
    borrowing = gold["borrowing_base"]
    covenants = gold["covenants"]
    capex = gold["capex"]
    wip_row = _table_row_text(document, "proposed june wip")
    liquidity_row = _table_row_text(document, "downside liquidity")
    borrowing_row = _table_row_text(document, "borrowing base")
    covenant_row = _table_row_text(document, "covenant")
    capex_row = _table_row_text(document, "midyear capex")
    wip_normalized = _normalize(wip_row)
    covenant_normalized = _normalize(covenant_row)
    wip_ok = (
        _text_contains_number(wip_row, gold["wip_adjustment"])
        and "proposed" in wip_normalized
        and ("unposted" in wip_normalized or "not posted" in wip_normalized)
    )
    liquidity_ok = (
        all(
            _text_contains_number(liquidity_row, number)
            for number in [downside["minimum_cash_before_financing"], downside["maximum_revolver"]]
        )
        and _text_contains_date(liquidity_row, downside["first_draw_week"])
    )
    borrowing_ok = _text_contains_number(borrowing_row, borrowing["borrowing_base"])
    covenant_ok = (
        all(
            _text_contains_number(covenant_row, number, abs_tol=.0001)
            for number in [covenants["leverage_posted"], covenants["leverage_pro_forma"]]
        )
        and "compliant" in covenant_normalized
    )
    requests_by_id = {row["request_id"]: row for row in capex["requests"]}
    request_ids = capex["release_now"] + capex["conditional"] + capex["hold_or_defer"]
    segments = _classification_segments(capex_row)
    requests_named = all(
        _normalize(request_id) in segments.get(category, "")
        or _normalize(requests_by_id[request_id]["asset"]) in segments.get(category, "")
        for category, ids in (
            ("release_now", capex["release_now"]),
            ("conditional", capex["conditional"]),
            ("hold_or_defer", capex["hold_or_defer"]),
        )
        for request_id in ids
    )
    classification_labels = set(segments) == {"release_now", "conditional", "hold_or_defer"}
    capex_ok = (
        _text_contains_number(capex_row, capex["cash_required"])
        and requests_named
        and classification_labels
    )
    source_groups = {
        "accounting": ("contractor accounting mcp", "accounting mcp"),
        "wip": ("pm etc updates", "revenue recognition wip policy", "co log master", "wip policy"),
        "liquidity": ("downside assumptions", "13wk cash", "13 week cash", "ar notes"),
        "borrowing_base": ("usbank ar eligibility", "ar eligibility exhibit", "eligibility exhibit"),
        "covenants": ("usbank amdt2", "executed amendment", "debt sched"),
        "capex": ("capex asks roi", "capex asks", "capital committee decision memo", "aop approval deck"),
    }
    matched_source_groups = {
        group for group, aliases in source_groups.items() if any(alias in normalized for alias in aliases)
    }
    criteria = [
        Criterion("identity__title", "The note has the requested Q2 Close CFO Decision Note title", "q2 close cfo decision note" in normalized, normalized[:500]),
        Criterion("identity__status", "The note is clearly labeled Internal Working — Not Submitted", "internal working not submitted" in normalized, normalized[:500]),
        Criterion(
            "table__real_word_table",
            "The deliverable contains a real Word decision table",
            real_table,
            f"real_table={real_table}",
            category="structure",
            weight=1,
            semantic=False,
        ),
        *[
            Criterion(f"table__row_{_normalize(row).replace(' ', '_')}", f"The decision table contains the required {row!r} workstream row", any(row in candidate for candidate in row_text), f"rows={row_text!r}")
            for row in required_rows
        ],
        Criterion("wip__amount", "The proposed WIP adjustment amount is correct", _text_contains_number(wip_row, gold["wip_adjustment"]), wip_row),
        Criterion("wip__proposed", "The WIP adjustment is identified as proposed", "proposed" in wip_normalized, wip_row),
        Criterion("wip__unposted", "The WIP adjustment is identified as unposted", "unposted" in wip_normalized or "not posted" in wip_normalized, wip_row),
        Criterion("liquidity__minimum_cash", "The downside minimum cash is correct", _text_contains_number(liquidity_row, downside["minimum_cash_before_financing"]), liquidity_row),
        Criterion("liquidity__first_draw_week", "The first revolver-draw week is correct", _text_contains_date(liquidity_row, downside["first_draw_week"]), liquidity_row),
        Criterion("liquidity__maximum_revolver", "The downside maximum revolver is correct", _text_contains_number(liquidity_row, downside["maximum_revolver"]), liquidity_row),
        Criterion("borrowing_base__amount", "The borrowing-base result is correct", borrowing_ok, borrowing_row),
        Criterion("covenants__posted_leverage", "Posted leverage is correct", _text_contains_number(covenant_row, covenants["leverage_posted"], abs_tol=.0001), covenant_row),
        Criterion("covenants__pro_forma_leverage", "Pro-forma leverage is correct", _text_contains_number(covenant_row, covenants["leverage_pro_forma"], abs_tol=.0001), covenant_row),
        Criterion("covenants__compliance", "The covenant conclusion is compliant", "compliant" in covenant_normalized, covenant_row),
        Criterion("capex__cash_required", "The capex cash requirement is correct", _text_contains_number(capex_row, capex["cash_required"]), capex_row),
        Criterion("capex__classification_labels", "Release Now, Conditional, and Hold/Defer categories are clearly labeled", classification_labels, f"segments={segments!r}"),
    ]
    for category, ids in (
        ("release_now", capex["release_now"]),
        ("conditional", capex["conditional"]),
        ("hold_or_defer", capex["hold_or_defer"]),
    ):
        for request_id in ids:
            met = _normalize(request_id) in segments.get(category, "") or _normalize(requests_by_id[request_id]["asset"]) in segments.get(category, "")
            criteria.append(Criterion(f"capex__{request_id.casefold()}__{category}", f"Request {request_id} is correctly classified as {category.replace('_', ' ')}", met, f"segment={segments.get(category, '')!r}"))
    for group in source_groups:
        criteria.append(Criterion(f"source__{group}", f"The controlling {group.replace('_', ' ')} source family is named", group in matched_source_groups, f"matched_groups={sorted(matched_source_groups)!r}"))
    criteria.append(Criterion("length", "The rendered note is no more than two pages", page_count is not None and page_count <= 2, f"pages={page_count}"))
    result = _result(criteria)
    return _attach_semantic_review(
        result,
        task_id="task_025",
        evidence=_legacy_artifact_evidence(path),
        artifact_type="decision memorandum",
        specs=[
            {
                "criterion_id": "identity__title",
                "expected_facts": {"title": "Q2 Close CFO Decision Note", "status": "Internal Working — Not Submitted"},
                "hard_gate_met": True,
                "hard_gate_evidence": "document parsed",
            },
            {
                "criterion_id": "wip__proposed",
                "expected_facts": {"proposed_wip_adjustment": gold["wip_adjustment"], "status": "proposed"},
                "hard_gate_met": _text_contains_number(wip_row, gold["wip_adjustment"]),
                "hard_gate_evidence": "exact WIP amount present in decision row",
            },
            {
                "criterion_id": "liquidity__first_draw_week",
                "expected_facts": {"first_draw_week": downside["first_draw_week"]},
                "hard_gate_met": all(_text_contains_number(liquidity_row, number) for number in [downside["minimum_cash_before_financing"], downside["maximum_revolver"]]),
                "hard_gate_evidence": "exact minimum-cash and maximum-revolver facts present in liquidity row",
            },
            {
                "criterion_id": "covenants__compliance",
                "expected_facts": {"status": "compliant"},
                "hard_gate_met": all(_text_contains_number(covenant_row, number, abs_tol=.0001) for number in [covenants["leverage_posted"], covenants["leverage_pro_forma"]]),
                "hard_gate_evidence": "exact posted and pro-forma leverage facts present in covenant row",
            },
            {
                "criterion_id": "capex__classification_labels",
                "expected_facts": {"cash_required": capex["cash_required"], "release_now": capex["release_now"], "conditional": capex["conditional"], "hold_or_defer": capex["hold_or_defer"]},
                "hard_gate_met": classification_labels,
                "hard_gate_evidence": (
                    "all three classification sections are present"
                    if classification_labels
                    else f"classification sections present={sorted(segments)!r}"
                ),
            },
            {
                "criterion_id": "source__accounting",
                "expected_facts": {"required_source_workstreams": sorted(source_groups)},
                "hard_gate_met": True,
                "hard_gate_evidence": "document parsed; source naming is semantic",
            },
            {
                "criterion_id": "identity__status",
                "expected_facts": {"status": "Internal Working — Not Submitted"},
                "hard_gate_met": True,
                "hard_gate_evidence": "document parsed",
            },
            {
                "criterion_id": "wip__unposted",
                "expected_facts": {"status": "unposted"},
                "hard_gate_met": _text_contains_number(wip_row, gold["wip_adjustment"]),
                "hard_gate_evidence": "exact WIP amount present in decision row",
            },
            *[
                {
                    "criterion_id": f"capex__{request_id.casefold()}__{category}",
                    "expected_facts": {"request_id": request_id, "classification": category.replace("_", " ")},
                    "hard_gate_met": (
                        _normalize(request_id) in _normalize(capex_row)
                        or _normalize(requests_by_id[request_id]["asset"]) in _normalize(capex_row)
                    ),
                    "hard_gate_evidence": "request id or asset name is present in the capex decision row",
                }
                for category, ids in (
                    ("release_now", capex["release_now"]),
                    ("conditional", capex["conditional"]),
                    ("hold_or_defer", capex["hold_or_defer"]),
                )
                for request_id in ids
            ],
            *[
                {
                    "criterion_id": f"source__{group}",
                    "expected_facts": {"source_family": group},
                    "hard_gate_met": True,
                    "hard_gate_evidence": "document parsed; source naming is semantic",
                }
                for group in source_groups
                if group != "accounting"
            ],
        ],
    )


FILE_GRADERS: dict[str, Callable[[Path], dict[str, Any]]] = {
    "task_004": _grade_task_004,
    "task_008": _grade_task_008,
    "task_012": _grade_task_012,
    "task_015": _grade_task_015,
    "task_023": _grade_task_023,
    "task_024": _grade_task_024,
    "task_025": _grade_task_025,
}


def grade_task(task_id: str, answer: Any, workspace_root: str | Path) -> dict[str, Any]:
    if 26 <= int(task_id[-3:]) <= 100:
        from evaluator.corporate_finance import grade_corporate_finance_task
        result = grade_corporate_finance_task(task_id, answer, workspace_root)
    elif task_id in FILE_GRADERS:
        result = FILE_GRADERS[task_id](Path(workspace_root))
        result = _canonicalize_legacy_file_policy(task_id, result)
    else:
        result = _grade_numeric(task_id, answer)
    return apply_reward_policy(task_id, attach_default_policy(result))
