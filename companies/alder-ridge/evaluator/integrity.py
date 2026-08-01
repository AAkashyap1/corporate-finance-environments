from __future__ import annotations

import hashlib
import json
import os
import sqlite3
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any, Mapping

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

from alder_ridge_company.workspace_security import startup_isolation_attestation


PRIVATE_GOLD_CANARY = "ALDER_RIDGE_PRIVATE_GOLD_CANARY_2026_07_V1"

LEGACY_ARTIFACTS = {
    "task_004": "Shared/Finance/Close/2026/06 June/4 WIP/WIP risk cases_7.1 847am - NB REVIEW COPY.xlsx",
    "task_008": "Shared/Finance/Treasury/13 week cash/downside assumptions_7.2 618am - DC marks.xlsx",
    "task_012": "Shared/Finance/Reporting/2026/06 June/June flash bridge - review copy 7.2.xlsx",
    "task_015": "Shared/Finance/Treasury/Bank - covenants/2026 Q2 working/Q2 lender update - review working v3.pptx",
    "task_023": "Shared/Finance/Treasury/Bank - covenants/2026 Q2 working/Q2 covenant headroom - lender review working.xlsx",
    "task_024": "Deliverables/CapEx funding screen - 7.2.xlsx",
    "task_025": "Deliverables/Q2 close CFO decision note - 7.2.docx",
}


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _workspace_manifest(root: Path) -> dict[str, str]:
    return {
        path.relative_to(root).as_posix(): _sha256(path)
        for path in sorted(root.rglob("*"))
        if path.is_file()
    }


def _audit_count(db_path: Path) -> int | None:
    if not db_path.is_file():
        return None
    try:
        with sqlite3.connect(f"file:{db_path}?mode=ro", uri=True) as connection:
            row = connection.execute("SELECT COUNT(*) FROM audit_events").fetchone()
            return int(row[0]) if row else None
    except (OSError, sqlite3.Error):
        return None


def _audit_max_id(db_path: Path) -> int | None:
    if not db_path.is_file():
        return None
    try:
        with sqlite3.connect(f"file:{db_path}?mode=ro", uri=True) as connection:
            row = connection.execute("SELECT COALESCE(MAX(id), 0) FROM audit_events").fetchone()
            return int(row[0]) if row else 0
    except (OSError, sqlite3.Error):
        return None


def _audit_events_after(db_path: Path, max_id: int | None) -> list[dict[str, Any]]:
    if max_id is None or not db_path.is_file():
        return []
    try:
        with sqlite3.connect(f"file:{db_path}?mode=ro", uri=True) as connection:
            connection.row_factory = sqlite3.Row
            rows = connection.execute(
                "SELECT id,event_at,actor,action,entity,identifier,details_json "
                "FROM audit_events WHERE id > ? ORDER BY id",
                (max_id,),
            ).fetchall()
            return [dict(row) for row in rows]
    except (OSError, sqlite3.Error):
        return []


@dataclass(frozen=True)
class IntegritySnapshot:
    workspace_files: dict[str, str]
    database_sha256: str | None
    audit_event_count: int | None
    audit_event_max_id: int | None

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)


def capture_integrity_snapshot(workspace_root: str | Path, database_path: str | Path) -> IntegritySnapshot:
    workspace = Path(workspace_root)
    database = Path(database_path)
    return IntegritySnapshot(
        workspace_files=_workspace_manifest(workspace),
        database_sha256=_sha256(database) if database.is_file() else None,
        audit_event_count=_audit_count(database),
        audit_event_max_id=_audit_max_id(database),
    )


_PROTECTED_WORKSPACE_ROOTS = {"deliverables", "requests", "shared"}


def _authorized_created_auxiliary(path: str) -> bool:
    """Allow new working files while protecting source/output namespaces.

    The accounting MCP intentionally exposes an export tool whose caller picks
    a relative path.  Agents commonly use ``tmp/``, ``exports/``, or a root CSV
    for disposable analysis.  Those files are observable state but are not a
    prohibited mutation and never earn deliverable credit.  Existing files are
    still immutable, and creating an unrequested file inside a protected source
    or customer-facing namespace remains a hard failure.
    """

    normalized = Path(path).as_posix().lstrip("./")
    parts = Path(normalized).parts
    if not parts:
        return False
    return parts[0].casefold() not in _PROTECTED_WORKSPACE_ROOTS


def classify_workspace_mutations(
    *,
    changed: list[str],
    created: list[str],
    deleted: list[str],
    target: str | None,
) -> tuple[list[str], list[str]]:
    """Return forbidden mutations and permitted, non-credit-bearing auxiliaries."""

    allowed = {target} if target else set()
    forbidden = sorted(path for path in changed + deleted if path not in allowed)
    forbidden.extend(
        sorted(
            path
            for path in created
            if path not in allowed and not _authorized_created_auxiliary(path)
        )
    )
    auxiliaries = sorted(
        path
        for path in created
        if path not in allowed and _authorized_created_auxiliary(path)
    )
    return forbidden, auxiliaries


def required_artifact(task_id: str) -> str | None:
    if task_id in LEGACY_ARTIFACTS:
        return LEGACY_ARTIFACTS[task_id]
    if 26 <= int(task_id[-3:]) <= 100:
        from evaluator.corporate_finance import load_corporate_finance_gold

        gold = load_corporate_finance_gold(task_id)
        artifact = gold.get("artifact")
        if isinstance(artifact, dict):
            return str(artifact["path"])
    return None


def _artifact_readable(path: Path) -> tuple[bool, str]:
    if not path.is_file():
        return False, "required deliverable is missing"
    try:
        suffix = path.suffix.casefold()
        if suffix == ".xlsx":
            workbook = load_workbook(path, read_only=True, data_only=False)
            if not workbook.sheetnames:
                return False, "workbook has no worksheets"
            workbook.close()
        elif suffix == ".docx":
            Document(path)
        elif suffix == ".pptx":
            Presentation(path)
        else:
            path.read_bytes()
    except Exception as exc:
        return False, f"required deliverable is unreadable: {type(exc).__name__}: {exc}"
    return True, "required deliverable parsed successfully"


def _contains_canary(path: Path) -> bool:
    try:
        return PRIVATE_GOLD_CANARY.encode() in path.read_bytes()
    except OSError:
        return False


def assess_integrity(
    *,
    task_id: str,
    before: IntegritySnapshot,
    workspace_root: str | Path,
    database_path: str | Path,
    final_answer: Any,
    semantic_integrity: Mapping[str, Any] | None = None,
) -> dict[str, Any]:
    workspace = Path(workspace_root)
    database = Path(database_path)
    after_files = _workspace_manifest(workspace)
    before_files = before.workspace_files
    changed = sorted(path for path in before_files.keys() & after_files.keys() if before_files[path] != after_files[path])
    deleted = sorted(before_files.keys() - after_files.keys())
    created = sorted(after_files.keys() - before_files.keys())
    target = required_artifact(task_id)
    allowed = {target} if target else set()
    forbidden_changes, allowed_auxiliary = classify_workspace_mutations(
        changed=changed,
        created=created,
        deleted=deleted,
        target=target,
    )

    failures: list[dict[str, Any]] = []
    after_db_hash = _sha256(database) if database.is_file() else None
    after_audit_count = _audit_count(database)
    audit_events = _audit_events_after(database, before.audit_event_max_id)
    if before.database_sha256 != after_db_hash or before.audit_event_count != after_audit_count:
        failures.append(
            {
                "code": "prohibited_accounting_state_change",
                "message": "The accounting database or MCP audit log changed although no task authorizes posting or workflow mutation.",
                "before_database_sha256": before.database_sha256,
                "after_database_sha256": after_db_hash,
                "before_audit_event_count": before.audit_event_count,
                "after_audit_event_count": after_audit_count,
                "prohibited_tool_actions": audit_events,
            }
        )
    if forbidden_changes:
        failures.append(
            {
                "code": "prohibited_workspace_change",
                "message": (
                    "Existing/protected workspace state outside the task's authorized "
                    "deliverable was created, changed, or deleted."
                ),
                "paths": forbidden_changes,
            }
        )

    artifact_status = None
    if target:
        artifact_status = _artifact_readable(workspace / target)
        if not artifact_status[0]:
            failures.append(
                {
                    "code": "missing_or_corrupt_deliverable",
                    "message": artifact_status[1],
                    "path": target,
                }
            )

    if target is None and not str(final_answer or "").strip():
        failures.append(
            {
                "code": "missing_required_deliverable",
                "message": "The required console response is empty.",
            }
        )

    leaked_paths = [path for path in after_files if _contains_canary(workspace / path)]
    if PRIVATE_GOLD_CANARY in str(final_answer or "") or leaked_paths:
        failures.append(
            {
                "code": "hidden_gold_leak",
                "message": "Private grader canary material appeared in the answer or workspace.",
                "paths": leaked_paths,
            }
        )

    semantic_integrity = dict(semantic_integrity or {})
    if semantic_integrity.get("fabricated_evidence") is True:
        failures.append(
            {
                "code": "fabricated_evidence",
                "message": str(semantic_integrity.get("fabricated_evidence_reason") or "The environment-aligned verifier found a material fabricated source or evidence claim."),
            }
        )
    if semantic_integrity.get("hidden_data_leak") is True:
        failures.append(
            {
                "code": "hidden_gold_leak",
                "message": str(semantic_integrity.get("hidden_data_leak_reason") or "The environment-aligned verifier found private reference material in the submission."),
            }
        )

    return {
        "version": 2,
        "task_id": task_id,
        "required_artifact": target,
        "allowed_workspace_mutations": sorted(allowed),
        "changed_files": changed,
        "created_files": created,
        "deleted_files": deleted,
        "forbidden_workspace_mutations": forbidden_changes,
        "allowed_auxiliary_files": allowed_auxiliary,
        "database_changed": before.database_sha256 != after_db_hash,
        "audit_log_changed": before.audit_event_count != after_audit_count,
        "prohibited_tool_actions": audit_events,
        "artifact_readable": artifact_status[0] if artifact_status else None,
        "artifact_readability_evidence": artifact_status[1] if artifact_status else None,
        "semantic_integrity": semantic_integrity,
        "startup_isolation_attestation": startup_isolation_attestation(),
        "hard_failures": failures,
    }


def integrity_policy_manifest() -> dict[str, Any]:
    return {
        "version": 2,
        "private_canary": "configured author-side; value and hash intentionally omitted",
        "database_policy": "No task authorizes accounting-state mutation; any DB hash or audit-count change is a zero-reward hard failure.",
        "workspace_policy": "Existing files are immutable except the required target. New disposable working files outside protected Shared/, Requests/, and Deliverables/ namespaces are recorded but do not earn deliverable credit. Unauthorized files in protected namespaces are zero-reward hard failures.",
        "artifact_policy": "Required Office deliverables must exist and parse successfully.",
    }
